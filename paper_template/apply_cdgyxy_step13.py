from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCX = Path(r"C:\Users\Quant\Desktop\Paper2026_1\word_output\基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx")

# Widths are sized to the contents while keeping every table within the text area.
TABLE_WIDTHS = (
    (2.40, 0.80, 1.70, 1.70),
    (1.25, 1.25, 2.05, 2.05),
    (2.40, 0.60, 1.80, 1.80),
)


def set_run_font(run, east_asia: str, latin: str, size_pt: float, bold: bool) -> None:
    run.font.name = latin
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def set_border(parent, edge: str, value: str, size: int = 0) -> None:
    borders = parent.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        parent.append(borders)
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), value)
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), "000000" if value != "nil" else "auto")


def set_cell_border(cell, edge: str, value: str, size: int = 0) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), value)
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), "000000" if value != "nil" else "auto")


def clear_cell_fill(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shading = tcpr.find(qn("w:shd"))
    if shading is not None:
        tcpr.remove(shading)


def set_cell_margins(cell, top: int = 55, start: int = 70, bottom: int = 55, end: int = 70) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header_repeat(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        trpr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    cant_split = trpr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        trpr.append(cant_split)


def format_table(table, widths: tuple[float, ...]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tblpr = table._tbl.tblPr
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        set_border(tblpr, edge, "nil")

    mark_header_repeat(table.rows[0])

    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for column_index, cell in enumerate(row.cells):
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                set_cell_border(cell, edge, "nil")
            if row_index == 0:
                set_cell_border(cell, "top", "single", 12)
                set_cell_border(cell, "bottom", "single", 6)
            if row_index == len(table.rows) - 1:
                set_cell_border(cell, "bottom", "single", 12)
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            clear_cell_fill(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if row_index > 0 and column_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, "宋体", "Times New Roman", 9, row_index == 0)


def main() -> None:
    doc = Document(DOCX)
    if len(doc.tables) != 3:
        raise RuntimeError(f"应有3张表，实际找到{len(doc.tables)}张")

    captions = [p for p in doc.paragraphs if re.match(r"^表\d+\s", p.text.strip())]
    if len(captions) != 3:
        raise RuntimeError(f"应有3个表题，实际找到{len(captions)}个")
    for caption in captions:
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(2)
        caption.paragraph_format.keep_with_next = True
        for run in caption.runs:
            set_run_font(run, "黑体", "Times New Roman", 9, True)

    for table, widths in zip(doc.tables, TABLE_WIDTHS):
        format_table(table, widths)

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
