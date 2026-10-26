#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Integrate exactly cropped original fig2-G/fig3-H/fig4-H and their table-only
facts into the 2026-10-26 Changjiang Information & Communications manuscript.

This script does NOT regenerate plot data, bars, axes, or labels.
"""
from __future__ import annotations

from pathlib import Path
from hashlib import sha256
import json
import re
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu

HERE = Path(__file__).resolve().parent
FIGS = HERE / "figure_previews_20261026"
SOURCE = HERE / "自适应多尺度加权的多变量网络流量预测_czxxtx20261026_结构精简稿.docx"
OUTPUT = HERE / "自适应多尺度加权的多变量网络流量预测_czxxtx20261026_图文整合稿.docx"
REPORT = HERE / "图2G图3H图4H_图文整合核验_20261026.json"

# The clip heights are the exact SVG-crop heights already approved; in points.
# Original PNGs, not invented charts, are cropped, leaving the plotted area untouched.
SPECS = [
    ("fig2_G_compact_bars.png", "fig2_G_compact_bars_no_table.png", 356.4, 270.0),
    ("fig3_H_grouped_bars.png", "fig3_H_grouped_bars_no_table.png", 360.0, 274.5),
    ("fig4_H_grouped_bars.png", "fig4_H_grouped_bars_no_table.png", 381.6, 274.0),
]
EXPECTED_VALUES = {
    "2.2": ["4 656", "130 162", "384 382", "8 339", "6.41%", "2.17%"],
    "2.4": ["17.47%", "20.84%", "3.87%", "6.65%"],
    "2.5": ["23.16%", "26.11%", "26.02%", "23.77%",
            "11.16%", "12.39%", "12.34%", "11.19%", "8/8"],
}
INSERT_TEXT = {
    "2.2": ("在模型规模方面，DLinear在Abilene和GÉANT上的参数量均为4 656，"
            "LightTS分别为130 162和384 382，本文方法在两个数据集上均为8 339。"
            "相对LightTS，本文方法的参数量分别仅占6.41%和2.17%。"),
    "2.4": ("与固定等权融合相比，本文方法在Abilene高需求目标点上的MAE和平均低估误差"
            "（AUE）配对平均降幅分别为17.47%和20.84%；在GÉANT上分别为3.87%和6.65%。"
            "这组结果补充了图3相对DLinear的比较，反映窗口级自适应加权的作用。"),
    "2.5": ("在GÉANT一周预测任务中，相对DLinear、固定等权、全局权重和恒定输入路由，"
            "本文方法的MSE配对平均降幅分别为23.16%、26.11%、26.02%和23.77%，"
            "MAE分别为11.16%、12.39%、12.34%和11.19%。"
            "八个共同随机种子下，针对上述四种对照的MSE和MAE配对改善次数均为8/8。"),
}

def digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()

def make_exact_crops():
    results = []
    for original, output, original_height, crop_height in SPECS:
        source = FIGS / original
        destination = FIGS / output
        assert source.is_file(), source
        with Image.open(source) as picture:
            rgb = picture.convert("RGB")
            width, height = rgb.size
            crop_px = round(height * crop_height / original_height)
            assert 0 < crop_px < height
            assert crop_px >= round(height * 0.70), "unexpected crop"
            clipped = rgb.crop((0, 0, width, crop_px))
            clipped.save(destination, format="PNG", optimize=True)
            # Ensure pixels in the entire retained region are unchanged.
            from PIL import ImageChops
            assert ImageChops.difference(clipped, rgb.crop((0, 0, width, crop_px))).getbbox() is None
            results.append({
                "original": original, "new": output,
                "input_size": [width, height], "cropped_size": [width, crop_px],
                "source_sha256": digest(source), "cropped_sha256": digest(destination),
                "crop_top_region_pixel_identical": True
            })
    return results

def get_section_label(paragraph):
    m = re.match(r"^\s*([0123])[.．]\s*([1-5])\s*", paragraph.text)
    if m:
        return m.group(1) + "." + m.group(2)
    return None

def section_ranges(paragraphs):
    starts={}
    for i, p in enumerate(paragraphs):
        sec=get_section_label(p)
        if sec in EXPECTED_VALUES:
            starts[sec] = i
    assert set(starts) == set(EXPECTED_VALUES), f"missing target section(s): {starts}"
    ranges={}
    for key,start in starts.items():
        finish=len(paragraphs)
        for j in range(start+1,len(paragraphs)):
            s=paragraphs[j].text.strip()
            if re.match(r"^([0-3])[.．]\s*[1-9]\s*|^[23]\s*[　 ]*[.．]?\s*(结语|实验|参考文献)|^3\s*[　 ]*结语|^参考文献",s):
                finish=j;break
        assert finish>start+1,(key,start,finish)
        ranges[key]=(start,finish)
    return ranges

def add_after(paragraph, content):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, paragraph._parent)
    p.style = paragraph.style
    if paragraph.paragraph_format.first_line_indent is not None:
        p.paragraph_format.first_line_indent=paragraph.paragraph_format.first_line_indent
    run = p.add_run(content)
    # Use the same text appearance as an adjacent body paragraph.
    for r in paragraph.runs:
        if r.text.strip() and not r._r.xpath('.//w:drawing'):
            try:
                run.style = r.style
                run.bold = r.bold
                run.italic = r.italic
                if r.font.name:run.font.name = r.font.name
                if r.font.size:run.font.size = r.font.size
            except Exception:
                pass
            break
    return p

def insert_missing_facts(doc):
    paras=list(doc.paragraphs)
    ranges=section_ranges(paras)
    outcome={}
    # Work backwards so insertions do not invalidate earlier indexes.
    for key,(start,end) in sorted(ranges.items(),key=lambda kv:kv[1][0],reverse=True):
        existing="\n".join(p.text for p in paras[start:end])
        missing=[x for x in EXPECTED_VALUES[key] if x not in existing]
        if not missing:
            outcome[key] = {"inserted":False,"reason":"all figure-table values already in section"}
            continue
        anchor_idx=end-1
        while anchor_idx>start and not paras[anchor_idx].text.strip():
            anchor_idx-=1
        add_after(paras[anchor_idx],INSERT_TEXT[key])
        outcome[key]={"inserted":True,"missing_before":missing,
                      "text":INSERT_TEXT[key],"after_paragraph_index":anchor_idx}
    return outcome

def replace_figures(doc):
    # Four original figures: fig 1 remains unchanged; replace images 2–4 only.
    shapes=list(doc.inline_shapes)
    if len(shapes) != 4:
        print("FIGURE INVENTORY: shape count =", len(shapes), flush=True)
        for i,s in enumerate(shapes):
            bs=s._inline.xpath('.//a:blip')
            ids=[b.get(qn("r:embed")) for b in bs]
            parts=[str(doc.part.related_parts[x].partname) for x in ids if x in doc.part.related_parts]
            print("SHAPE",i, "size",s.width,s.height, "rel",ids, "parts",parts,flush=True)
        for i,p in enumerate(doc.paragraphs):
            inline=p._p.xpath('.//wp:inline')
            if inline:
                print("INLINE PARA",i, "blips",len(p._p.xpath('.//a:blip')),
                      "text",repr(p.text[:100]),"previous",repr(doc.paragraphs[i-1].text[:85] if i else ""),flush=True)
        raise AssertionError(f"expected four inline figures, found {len(shapes)}")
    outcome=[]
    for figure_index, (_, cropped, _, _) in enumerate(SPECS, start=2):
        shape=shapes[figure_index-1]
        blips=shape._inline.xpath('.//a:blip')
        assert len(blips)==1,(figure_index,len(blips))
        old_rel=blips[0].get(qn("r:embed"))
        assert old_rel, f"image {figure_index} has no relationship"
        old_part=doc.part.related_parts[old_rel]
        # Keep the original horizontal size; preserve cropped picture aspect ratio.
        new_path=FIGS/cropped
        relation_id,image=doc.part.get_or_add_image(str(new_path))
        blips[0].set(qn("r:embed"),relation_id)
        shape.height = Emu(round(shape.width*image.px_height/image.px_width))
        outcome.append({
            "figure": figure_index, "old_partname":str(old_part.partname),
            "new_png": cropped, "new_partname":str(doc.part.related_parts[relation_id].partname),
            "new_dimensions_emu":[shape.width,shape.height],
        })
    return outcome

def main():
    assert SOURCE.exists(),SOURCE
    audit={"source_document":SOURCE.name,"source_sha256":digest(SOURCE)}
    audit["cropping"]=make_exact_crops()
    source_doc=Document(SOURCE)
    orig={
        "tables":len(source_doc.tables),
        "inline_images":len(source_doc.inline_shapes),
        "nonempty_paragraphs":sum(bool(p.text.strip()) for p in source_doc.paragraphs),
    }
    original_table_text=[[[c.text for c in row.cells] for row in t.rows] for t in source_doc.tables]
    original_numbered_formulas=[
        p.text for p in source_doc.paragraphs if re.search(r"[（(]\s*\d+\s*[）)]",p.text)
    ]
    audit["new_text"]=insert_missing_facts(source_doc)
    audit["figure_replacements"]=replace_figures(source_doc)
    source_doc.save(OUTPUT)
    check=Document(OUTPUT)
    assert len(check.tables)==orig["tables"],"table count changed"
    assert len(check.inline_shapes)==orig["inline_images"],"figure count changed"
    assert [[[c.text for c in row.cells] for row in t.rows] for t in check.tables]==original_table_text,"table text changed"
    assert [p.text for p in check.paragraphs if re.search(r"[（(]\s*\d+\s*[）)]",p.text)]==original_numbered_formulas,"formula-number labels changed"
    ranges=section_ranges(list(check.paragraphs))
    for section,req in EXPECTED_VALUES.items():
        start,end=ranges[section]
        txt="\n".join(p.text for p in check.paragraphs[start:end])
        assert all(v in txt for v in req),(section,[v for v in req if v not in txt])
    audit["after"]={
        "source_and_output_distinct":str(SOURCE)!=str(OUTPUT),
        "tables_unchanged":True,
        "inline_images":len(check.inline_shapes),
        "numbered_formula_paragraphs_unchanged":True,
        "required_values_in_corresponding_sections":True,
        "output_sha256":digest(OUTPUT),
        "new_document":OUTPUT.name,
        "source_counts":orig
    }
    REPORT.write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(audit,ensure_ascii=False,indent=2))

if __name__=="__main__":
    main()
