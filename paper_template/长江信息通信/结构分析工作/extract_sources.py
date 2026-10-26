from pathlib import Path
import json
import re
import hashlib
from zipfile import ZipFile
from lxml import etree
import pdfplumber
from docx import Document

BASE = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}

def chars(s):
    return len(re.sub(r'\s+', '', s))

def main():
    inventory = []
    for i, path in enumerate(sorted(BASE.glob('*.pdf')), 1):
        slug = f'sample_{i:02}'
        pages = []
        with pdfplumber.open(path) as pdf:
            for n, page in enumerate(pdf.pages, 1):
                text = page.extract_text(x_tolerance=2, y_tolerance=3) or ''
                pages.append({'page': n, 'width': float(page.width), 'height': float(page.height),
                              'text': text, 'nonspace_chars': chars(text)})
        (OUT / f'{slug}.txt').write_text('\n\n'.join(f'===== PDF页 {p["page"]} =====\n{p["text"]}' for p in pages), encoding='utf-8')
        item = {'slug': slug, 'filename': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'pages': pages}
        inventory.append(item)
        print(slug, path.name, len(pages), sum(p['nonspace_chars'] for p in pages))
    path = next(p for p in BASE.glob('*czxxtx20261024.docx') if not p.name.startswith('~$'))
    doc = Document(path)
    paras = []
    for i, p in enumerate(doc.paragraphs):
        math = p._p.xpath('.//m:t')
        text = p.text
        if math:
            text += ' [公式文本: ' + ''.join(t.text or '' for t in math) + ']'
        if text.strip() or p._p.xpath('.//w:drawing'):
            paras.append({'index': i, 'style': p.style.name, 'text': text, 'chars': chars(text),
                          'images': len(p._p.xpath('.//w:drawing'))})
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    with ZipFile(path) as z:
        app = z.read('docProps/app.xml').decode('utf-8') if 'docProps/app.xml' in z.namelist() else ''
        media = [n for n in z.namelist() if n.startswith('word/media/')]
    manuscript = {'filename': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                  'paragraphs': paras, 'tables': tables, 'media': media, 'app_xml': app}
    (OUT / 'manuscript.txt').write_text('\n\n'.join(f'[P{p["index"]:03} {p["style"]}] {p["text"]}' for p in paras) + '\n\n' + '\n\n'.join(f'[表格{i+1}]\n' + '\n'.join(' | '.join(row) for row in t) for i,t in enumerate(tables)), encoding='utf-8')
    (OUT / 'source_inventory.json').write_text(json.dumps({'samples': inventory, 'manuscript': manuscript}, ensure_ascii=False, indent=2), encoding='utf-8')
    print('manuscript', len(paras), len(tables), len(media), sum(p['chars'] for p in paras))

if __name__ == '__main__':
    main()
