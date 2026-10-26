#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Inspect unmodified original DOCX and the original binary Word97 journal template.
Does not depend on apt-get / LibreOffice; does not change manuscript."""
from pathlib import Path
import importlib.util, json, re, hashlib
from collections import Counter
import olefile
BASE=Path(__file__).resolve().parent
imp=BASE/"核对格式_20261026.py"
spec=importlib.util.spec_from_file_location("format_analyzer",imp)
mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
paper=BASE/"自适应多尺度加权的多变量网络流量预测_czxxtx20261026_结构精简稿.docx"
binary=BASE/"论文格式模板 (1).doc"
out=BASE/"结构精简稿_格式比对_20261026.json"
def word97_text(path):
    ole=olefile.OleFileIO(str(path))
    w=ole.openstream("WordDocument").read()
    u16=lambda x:int.from_bytes(w[x:x+2],"little")
    u32=lambda x:int.from_bytes(w[x:x+4],"little")
    flags=u16(10)
    tblname="1Table" if (flags&0x200) else "0Table"
    table=ole.openstream(tblname).read()
    pos=32
    csw=u16(pos);pos+=2+csw*2
    cslw=u16(pos);pos+=2+cslw*4
    count=u16(pos);pos+=2
    assert count>33
    start=u32(pos+33*8);length=u32(pos+33*8+4)
    clx=table[start:start+length]
    if clx[0]==1:
        j=0
        while clx[j]==1:
            n=int.from_bytes(clx[j+1:j+3],"little")
            j+=3+n
        clx=clx[j:]
    assert clx[0]==2,"expected piece table"
    datalen=int.from_bytes(clx[1:5],"little")
    count=(datalen-4)//12
    cps=[int.from_bytes(clx[5+i*4:9+i*4],"little") for i in range(count+1)]
    pcdStart=5+(count+1)*4
    texts=[]
    for i in range(count):
        pcd=clx[pcdStart+i*8:pcdStart+i*8+8]
        raw=int.from_bytes(pcd[2:6],"little")
        short=bool(raw&0x40000000)
        offset=raw&0x3fffffff
        if short:
            content=w[offset//2:offset//2+cps[i+1]-cps[i]]
            piece=content.decode("cp1252",errors="replace")
        else:
            content=w[offset:offset+(cps[i+1]-cps[i])*2]
            piece=content.decode("utf-16-le",errors="replace")
        texts.append(piece)
    raw="".join(texts)
    raw=re.sub(r"\x13.*?\x14", "", raw,flags=re.S)
    raw=raw.replace("\x15","")
    return {"chars":len(raw),"paragraphs":[x.strip() for x in raw.split("\r") if x.strip()],
            "preview":raw[:18000],"ole_table":tblname}
def main():
    a=mod.inspect_doc("original-structure-concise",paper)
    b=word97_text(binary)
    ps=[x["text"].strip() for x in a["prefix_paragraphs"]]
    orig_abstract=next((x for x in ps if x.startswith("摘要")),None)
    orig_en_abstract=next((x for x in ps if x.startswith("Abstract")),None)
    orig_keywords=next((x for x in ps if x.startswith("关键词")),None)
    orig_en_keywords=next((x for x in ps if x.startswith("Key words")),None)
    from docx import Document
    from docx.oxml.ns import qn
    from collections import Counter
    doc=Document(paper)
    styles={}
    for st in doc.styles:
        if st.name.startswith("Journal") or st.name in ("Title","Normal"):
            fonts=[]
            for el in st._element.xpath(".//w:rFonts"):
                fonts.append({"eastAsia":el.get(qn("w:eastAsia")),"ascii":el.get(qn("w:ascii")),"hAnsi":el.get(qn("w:hAnsi")),"cs":el.get(qn("w:cs"))})
            styles[st.name]={"font":st.font.name,"size_pt":st.font.size.pt if st.font.size else None,
                              "bold":st.font.bold,"font_names":fonts[:5],
                              "paragraph_line":st.paragraph_format.line_spacing,
                              "first_line_pt":st.paragraph_format.first_line_indent.pt if st.paragraph_format.first_line_indent else None,
                              "before_pt":st.paragraph_format.space_before.pt if st.paragraph_format.space_before else None,
                              "after_pt":st.paragraph_format.space_after.pt if st.paragraph_format.space_after else None}
    def defaults():
        opts=[]
        for el in doc.styles.element.xpath(".//w:docDefaults//w:rFonts"):
            opts.append({"ascii":el.get(qn("w:ascii")),"eastAsia":el.get(qn("w:eastAsia")),
                         "hAnsi":el.get(qn("w:hAnsi"))})
        return opts
    border_reports=[]
    for tbl in doc.tables:
        counts=Counter()
        for row in tbl.rows:
            for cell in row.cells:
                for el in cell._tc.xpath(".//w:tcBorders/*"):
                    counts[(el.tag.split("}")[-1],el.get(qn("w:val")))]+=1
        border_reports.append({str(key):v for key,v in counts.items()})
    a["full_main_paragraphs"]=[{"index":i,"text":p.text,"style":p.style.name if p.style else None}
                              for i,p in enumerate(doc.paragraphs)]
    a["computed_style_rules"]=styles
    a["document_default_fonts"]=defaults()
    a["table_cell_borders"]=border_reports
    result={
        "source_document":paper.name,
        "source_sha256":hashlib.sha256(paper.read_bytes()).hexdigest(),
        "template_sha256":hashlib.sha256(binary.read_bytes()).hexdigest(),
        "template_binary_original":b,
        "paper":a,
        "quick_findings":{
            "title":ps[0] if ps else None,
            "title_len":len(ps[0]) if ps else 0,
            "abstract":orig_abstract,"english_abstract":orig_en_abstract,
            "keywords":orig_keywords,"english_keywords":orig_en_keywords,
            "figure_captions":a["captions"],
            "table_count":len(a["tables"]),
            "bibliography_count":len(a["bibliography"]),
            "image_colors":a["xmlmetrics"]["images"],
            "word_equations_omath":a["xmlmetrics"]["oMath"],
            "explicit_fonts":a["xmlmetrics"]["fonts_explicit"],
        }
    }
    out.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    print("SOURCE:",json.dumps({k:v for k,v in a.items() if k not in("prefix_paragraphs","last_paragraphs","last_100_text","tables","bibliography")},ensure_ascii=False)[:25000])
    print("SOURCE PREFIX",json.dumps(a["prefix_paragraphs"][:27],ensure_ascii=False)[:12500])
    print("SOURCE TABLES",json.dumps(a["tables"],ensure_ascii=False)[:9000])
    print("SOURCE REFS",json.dumps(a["bibliography"],ensure_ascii=False)[:10000])
    print("TEMPLATE TEXT",json.dumps(b["paragraphs"][:35],ensure_ascii=False)[:10000])
    print("DONE",out)
if __name__=="__main__":main()
