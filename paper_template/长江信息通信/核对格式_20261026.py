#!/usr/bin/env python3
# -*- coding:utf-8 -*-
"""Read-only inspection of original structural concise manuscript vs DOC template."""
from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import json, re, subprocess, os
from lxml import etree
from PIL import Image
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
HERE=Path(__file__).resolve().parent
PAPER=HERE/"自适应多尺度加权的多变量网络流量预测_czxxtx20261026_结构精简稿.docx"
TEMPLATE=HERE/"论文格式模板 (1).doc"
OUT=HERE/"结构精简稿_期刊格式逐项核验_20261026.json"
TEMP=Path("/tmp/czxxtx-format-audit")
TEMP.mkdir(exist_ok=True,parents=True)
NAMESPACE={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
           "m":"http://schemas.openxmlformats.org/officeDocument/2006/math"}
def run(cmd):
    x=subprocess.run(cmd,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=110)
    print("COMMAND",cmd[0],x.returncode,x.stdout[-1800:],flush=True)
    return x
def dimensions(doc):
    return [{"width_cm":round(s.page_width.cm,3),"height_cm":round(s.page_height.cm,3),
             "top_cm":round(s.top_margin.cm,3),"bottom_cm":round(s.bottom_margin.cm,3),
             "left_cm":round(s.left_margin.cm,3),"right_cm":round(s.right_margin.cm,3),
             "header_cm":round(s.header_distance.cm,3),"footer_cm":round(s.footer_distance.cm,3)}
            for s in doc.sections]
def paragraph_info(p):
    return {"text":p.text[:1000],"style":p.style.name if p.style else "?",
            "align":str(p.alignment),"runs":len(p.runs),
            "runfonts":list({r.font.name for r in p.runs if r.font.name}),
            "run_sizes":[round(r.font.size.pt,1) for r in p.runs if r.font.size is not None][:15],
            "space_before":p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None,
            "space_after":p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None,
            "line_spacing":p.paragraph_format.line_spacing}
def xmlmetrics(docx_path):
    with ZipFile(docx_path) as z:
        blob=z.read("word/document.xml")
        root=etree.fromstring(blob)
        alltext=root.xpath("//w:t/text()",namespaces=NAMESPACE)
        pnodes=root.xpath("//w:p",namespaces=NAMESPACE)
        xmlfonts=Counter()
        for e in root.xpath("//w:rFonts",namespaces=NAMESPACE):
            a=e.get(qn("w:eastAsia"));b=e.get(qn("w:ascii"));c=e.get(qn("w:hAnsi"))
            if a:xmlfonts["eastAsia:"+a]+=1
            if b:xmlfonts["ascii:"+b]+=1
            if c:xmlfonts["hAnsi:"+c]+=1
        images=[]
        for file in z.namelist():
            if file.startswith("word/media/") and file.lower().endswith((".png",".jpeg",".jpg",".gif",".tiff",".tif")):
                try:
                    im=Image.open(BytesIO(z.read(file))).convert("RGB").resize((320,200))
                    values=list(im.getdata())
                    colored=sum(max(rgb)-min(rgb)>25 for rgb in values)
                    vibr=sum(max(rgb)-min(rgb)>50 for rgb in values)
                    images.append({"filename":file,"original_size":Image.open(BytesIO(z.read(file))).size,
                                   "color_pixels_pct":round(colored/len(values)*100,2),
                                   "strong_color_pixels_pct":round(vibr/len(values)*100,2),
                                   "is_grayscale_all_pixels":colored==0})
                except Exception as e:images.append({"filename":file,"error":str(e)})
        xmlstr=blob.decode("utf-8")
        return {"oMath":len(root.xpath("//m:oMath",namespaces=NAMESPACE)),
                "oMathPara":len(root.xpath("//m:oMathPara",namespaces=NAMESPACE)),
                "OLE_object_count":len(root.xpath("//*[local-name()='OLEObject']")),
                "drawing_count":len(root.xpath("//*[local-name()='drawing']")),
                "p_count_including_tables":len(pnodes),
                "char_count_including_tables":len("".join(alltext)),
                "fonts_explicit":xmlfonts.most_common(30),
                "images":images,
                "common_font_names":root.xpath("//w:rFonts/@w:eastAsia",namespaces=NAMESPACE)[:20],
                "formula_number_patterns":re.findall(r"[（(]\d{1,2}[）)]", "".join(alltext))[:30],
                "references_english_initials":sum(1 for t in alltext if re.match(r"\[\d+\].*[A-Za-z]",t))}

def inspect_doc(label,path):
    d=Document(path)
    ps=[paragraph_info(p) for p in d.paragraphs]
    mainps=[p.text.strip() for p in d.paragraphs if p.text.strip()]
    sections=dimensions(d)
    hfreq=Counter((p.style.name if p.style else "") for p in d.paragraphs)
    specs=[]
    for i,t in enumerate(d.tables,1):
        cols=[len(row.cells) for row in t.rows]
        c=t._element.xpath(".//w:tblBorders/*")
        borders=[{"side":etree.QName(el).localname,"val":el.get(qn("w:val"))} for el in c]
        specs.append({"number":i,"rows":len(t.rows),"columns_each_row":cols,
                      "cell_texts":[[x.text[:120] for x in row.cells] for row in list(t.rows)[:8]],
                      "borders":borders,"style":t.style.name if t.style else None})
    metrics=xmlmetrics(path)
    form_text=[s for s in mainps if re.match(r"^(图\s*\d|表\s*\d)",s)]
    refs=[s for s in mainps if re.match(r"^\[\d+\]",s)]
    author=[s for s in mainps if s.startswith("作者简介")]
    summary={
       "label":label,"path":str(path.relative_to(HERE)) if path.is_relative_to(HERE) else str(path),
       "pagesetup":sections,"paragraph_count":len(d.paragraphs),
       "styles":hfreq.most_common(15),
       "nonblank_paragraphs":len(mainps),
       "prefix_paragraphs":ps[:27],
       "body_headings":[{"idx":i,"text":p.text[:160],"style":p.style.name}
             for i,p in enumerate(d.paragraphs) if re.match(r"^\s*[0-9]\s*[.．、\s]|^\s*[0-9]\s*[\s　]*(引言|结语|方法|实验)|^\s*[0-9]\.[0-9]",p.text)],
       "captions":form_text,"tables":specs,"bibliography":refs,
       "author_intro":author,
       "last_paragraphs":ps[-21:],
       "last_100_text":"\n".join(mainps[-14:])[:5000],
       "xmlmetrics":metrics}
    return summary
def main():
    assert PAPER.exists() and TEMPLATE.exists()
    # Read original binary Word 97-2003 .doc using LibreOffice only for inspection.
    convert=run(["libreoffice","-env:UserInstallation=file:///tmp/czxxtx_audit_template",
      "--headless","--convert-to","docx","--outdir",str(TEMP),str(TEMPLATE)])
    tpl=TEMP/(TEMPLATE.stem+".docx")
    assert convert.returncode==0 and tpl.exists(), "template conversion failed"
    source=inspect_doc("original_20261026",PAPER)
    template=inspect_doc("publisher_original_doc_converted_for_readonly_audit",tpl)
    r={"source_sha":__import__("hashlib").sha256(PAPER.read_bytes()).hexdigest(),
       "template_sha":__import__("hashlib").sha256(TEMPLATE.read_bytes()).hexdigest(),
       "source":source,"template":template}
    OUT.write_text(json.dumps(r,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    print("======= SOURCE OVERVIEW =======",flush=True)
    print(json.dumps({k:v for k,v in source.items() if k not in ("prefix_paragraphs","last_paragraphs","last_100_text","tables","bibliography","xmlmetrics")},ensure_ascii=False,indent=2),flush=True)
    print("SOURCE FORMATS",json.dumps(source["xmlmetrics"],ensure_ascii=False,indent=2),flush=True)
    print("SOURCE FIRST PARAGRAPHS",json.dumps(source["prefix_paragraphs"][:20],ensure_ascii=False),flush=True)
    print("SOURCE TABLES",json.dumps(source["tables"],ensure_ascii=False)[:9000],flush=True)
    print("SOURCE REFERENCES",json.dumps(source["bibliography"],ensure_ascii=False)[:12000],flush=True)
    print("======= TEMPLATE OVERVIEW =======",flush=True)
    print(json.dumps({k:v for k,v in template.items() if k not in ("prefix_paragraphs","last_paragraphs","last_100_text","tables","bibliography","xmlmetrics")},ensure_ascii=False,indent=2),flush=True)
    print("TEMPLATE FIRST PARAGRAPHS",json.dumps(template["prefix_paragraphs"][:27],ensure_ascii=False),flush=True)
    print("TEMPLATE STYLES",json.dumps(template["xmlmetrics"]["fonts_explicit"],ensure_ascii=False),flush=True)
    print("TEMPLATE TABLES",json.dumps(template["tables"][:4],ensure_ascii=False)[:9000],flush=True)
    print("AUDIT DONE",OUT,flush=True)
if __name__=="__main__":
    main()
