#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read-only audit of the latest submission manuscript; only writes JSON audit."""
from pathlib import Path
from collections import Counter
import hashlib, json, re, zipfile, subprocess
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from io import BytesIO
B=Path(__file__).resolve().parent.parent
F=B/"自适应多尺度加权的多变量网络流量预测_czxxtx20261027.docx"
J=B/"格式核对工作_20261027/最新投稿稿_整体核查记录.json"
D=Document(F)
W={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
   'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
paras=[{'index':i,'style':p.style.name,'text':p.text,
        'runs':len(p.runs)} for i,p in enumerate(D.paragraphs)]
def full_text(p):
    chunks=[]
    for el in p._p.xpath('.//w:t | .//m:t'):
        if el.text:chunks.append(el.text)
    return ''.join(chunks)
allparas=[{'index':i,'style':p.style.name,'text':p.text, 'with_equations':full_text(p)} for i,p in enumerate(D.paragraphs)]
full='\n'.join(x['text'] for x in allparas)
items=[x for x in allparas if re.match(r'^\s*\[\d+\]',x['text'])]
byref=[]
for ref in items:
    m=re.search(r'^\s*\[(\d+)\]',ref['text'])
    if not m:continue
    num=int(m.group(1))
    para=ref['text']
    # Capture bilingual continuation paragraphs before next numbered reference.
    idx=ref['index'];extras=[]
    for nxt in allparas[idx+1:]:
        if re.match(r'^\s*\[\d+\]',nxt['text']) or nxt['text'].startswith('作者简介'):break
        if nxt['text'].strip(): extras.append(nxt['text'])
    byref.append({'n':num,'reference':para,'continued':extras[:2]})
cites=[]
for p in allparas:
    if re.match(r'^\s*\[\d+\]',p['text']) or p['style']=='JournalReference':continue
    for match in re.finditer(r'(?<!\[)\[(\d{1,2})\](?!\])',p['text']):
        cites.append({'n':int(match.group(1)),'paragraph_index':p['index'],'context':p['text'][:1000]})
first={}
for c in cites:
    first.setdefault(c['n'],c['context'])
intro=next((p for p in allparas if ('TimeMixer' in p['text'] and 'LightTS' in p['text'])),{})
t=[]
for i,table in enumerate(D.tables,1):
    borders=[]
    for row in table.rows:
        for cell in row.cells:
            for el in cell._tc.xpath('.//w:tcBorders/*'):
                if el.get(qn('w:val'))!='nil':borders.append({'edge':el.tag.rsplit('}',1)[-1],'val':el.get(qn('w:val'))})
    vals=[[c.text for c in row.cells] for row in table.rows]
    t.append({'number':i,'rows':len(table.rows),'cols':len(table.columns),'borders':dict(Counter(x['edge'] for x in borders)),
              'top':vals[:3],'bottom':vals[-2:]})
pages=[{'width_cm':round(s.page_width.cm,3),'height_cm':round(s.page_height.cm,3),
        'top_cm':round(s.top_margin.cm,3),'bottom_cm':round(s.bottom_margin.cm,3),
        'left_cm':round(s.left_margin.cm,3),'right_cm':round(s.right_margin.cm,3)}
       for s in D.sections]
styles={}
for key in ['Title','JournalBody','JournalHeading1','JournalHeading2','JournalAbstract','JournalEquation',
            'JournalReference','JournalFigure','JournalCaption','JournalMetadata']:
    try:
        s=D.styles[key]; styles[key]={'font_pt':s.font.size.pt if s.font.size else None,
        'bold':s.font.bold,'line_spacing':s.paragraph_format.line_spacing,
        'first_indent_pt':s.paragraph_format.first_line_indent.pt if s.paragraph_format.first_line_indent else None}
    except:pass
im=[];floaty=0
with zipfile.ZipFile(F) as z:
    body=etree.fromstring(z.read('word/document.xml'))
    formula=[x for x in body.xpath('//w:p[.//m:oMath]',namespaces=W)]
    for name in z.namelist():
        if name.startswith('word/media/') and name.lower().endswith(('.png','.jpg','.jpeg','.tif','.tiff')):
            dat=z.read(name)
            img=Image.open(BytesIO(dat))
            sample=img.convert('RGB').resize((250,100))
            pix=list(sample.getdata());color=sum(max(p)-min(p)>5 for p in pix)/len(pix)*100
            im.append({'name':name,'px':list(img.size),'mode':img.mode,'pct_non_gray':round(color,2),'bytes':len(dat)})
    omath=len(body.xpath('//m:oMath',namespaces=W))
    drawings=len(body.xpath('//*[local-name()="drawing"]'))
    anchors=len(body.xpath('//*[local-name()="anchor"]'))
    cols=len(body.xpath('//*[local-name()="sectPr"]'))
    embedded_title=[x.text for x in allparas if re.match(r'^\s*图\s*[1-9]',x['text'])]
    formats=[]
    for e in body.xpath('//w:rFonts',namespaces=W):
        formats.append({'eastAsia':e.get(qn('w:eastAsia')),'ascii':e.get(qn('w:ascii'))})
    colors=list(set(x['eastAsia'] for x in formats if x['eastAsia']))
head=[p for p in allparas[:15]]
end=[p for p in allparas[-22:]]
title=allparas[0]['text']
zha=next((p['text'] for p in allparas if p['text'].startswith('摘要：')),None)
ena=next((p['text'] for p in allparas if p['text'].startswith('Abstract:')),None)
zhkey=next((p['text'] for p in allparas if p['text'].startswith('关键词：')),None)
enkey=next((p['text'] for p in allparas if p['text'].startswith('Key words:')),None)
numtxt=full.replace(' ', '').replace('\u00a0','')
chars_han=len(re.findall(r'[\u3400-\u9fff]',zha or ''))
english_words=len((ena or '').replace('Abstract:','').split())
report={
 'document_name':F.name,'sha256':hashlib.sha256(F.read_bytes()).hexdigest(),'bytes':F.stat().st_size,
 'title':title,'title_hanzi':len(re.findall(r'[\u3400-\u9fff]',title)),
 'head':head,'tail':end,'abstract':zha,'abstract_hanzi':chars_han,
 'abstract_chars_all_without_spaces':len(re.sub(r'\s+','',zha or '')),
 'english_abstract':ena,'english_abstract_words':english_words,
 'keywords':zhkey,'english_keywords':enkey,
 'styles':styles,'pages':pages,'paragraph_count':len(allparas),
 'headings':[x for x in allparas if re.match(r'^\s*(?:[0-9]+[.．][0-9]+|[0-9]+\s+)(?:[^\n]+)$',x['text']) and x['style'].startswith('JournalHeading')],
 'figure_captions':embedded_title, 'tables':t,'image_embedded':im,
 'omath':omath,'equation_paragraphs':[{'index':i,'math':full_text(D.paragraphs[i])[:1800]}
       for i,p in enumerate(D.paragraphs) if p._p.xpath('.//m:oMath')],
 'drawings':drawings,'anchors':anchors,'sectPr_count':cols,'explicit_eastAsia_fonts':colors,
 'references':byref,'citations_in_body':cites,'first_citation_context':first,
 'intro_sentence':intro,'all_paragraphs':allparas,
 'numeric_unspaced_long_numbers':list(dict.fromkeys(re.findall(r'(?<!\d)\d{5,}(?!\d)',full)))[:60],
 'reference_order':[i['n'] for i in byref],
 'reference_citation_order':list(dict.fromkeys(c['n'] for c in cites))
}
J.parent.mkdir(parents=True,exist_ok=True)
J.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('CURRENT_REVIEW_SHA256',report['sha256'])
print('HEADINGS',json.dumps(report['headings'],ensure_ascii=False))
print('ABSTRACT_SIZE',chars_han,report['abstract_chars_all_without_spaces'],english_words)
print('TABLE_IMAGE_FORMULA',len(t),len(im),omath)
print('REFERENCE_ORDER',report['reference_order'],report['reference_citation_order'])
print('INTRO',json.dumps(intro,ensure_ascii=False))
print('REFERENCE_LIST',json.dumps(byref,ensure_ascii=False))
print('PAGES',json.dumps(pages,ensure_ascii=False))
print('NUMERIC',report['numeric_unspaced_long_numbers'])
print('OUTPUT',J.as_posix())
