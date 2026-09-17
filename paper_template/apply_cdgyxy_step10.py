from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"C:\Users\Quant\Desktop\Paper2026_1")
DOCX = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx"
QA_DIR = ROOT / "word_output" / "qa_cdgyxy_step10"
BACKUP = QA_DIR / "before_step10.docx"


def set_run_font(run, east_asia: str, size_pt: float, *, bold: bool | None = None) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if east_asia == "黑体":
        rfonts.set(qn("w:ascii"), "Arial")
        rfonts.set(qn("w:hAnsi"), "Arial")
    else:
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")
    color.attrib.pop(qn("w:themeColor"), None)


def replace_text(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def paragraph_has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def count_media(path: Path) -> int:
    with zipfile.ZipFile(path) as zf:
        return sum(name.startswith("word/media/") and not name.endswith("/") for name in zf.namelist())


QA_DIR.mkdir(parents=True, exist_ok=True)
if not BACKUP.exists():
    shutil.copy2(DOCX, BACKUP)

before = Document(DOCX)
before_paragraph_count = len(before.paragraphs)
before_table_count = len(before.tables)
before_media_count = count_media(DOCX)

doc = Document(DOCX)

# Requirement 10: the Introduction is unnumbered; subsequent sections restart at 1.
heading_text = {
    11: "引言",
    17: "1 相关工作",
    18: "1.1 网络流量预测",
    21: "1.2 轻量多变量时间序列预测",
    24: "1.3 频域建模与自适应多尺度融合",
    27: "2 自适应多尺度网络流量预测方法",
    31: "2.1 问题定义与数据预处理",
    40: "2.2 多尺度序列构造",
    44: "2.3 DLinear尺度专家",
    51: "2.4 样本级自适应尺度加权",
    62: "2.5 辅助频域门控",
    72: "2.6 训练目标与复杂度",
    76: "3 实验与结果分析",
    77: "3.1 数据集与运行环境",
    80: "3.2 评价指标与统计方法",
    85: "3.3 基础模型与组件消融",
    92: "3.4 自适应权重消融与路由行为",
    102: "3.5 频率保留比例敏感性",
    109: "3.6 代表性预测与应用含义",
    114: "3.7 讨论与局限性",
    118: "4 结论",
}
for index, text in heading_text.items():
    replace_text(doc.paragraphs[index], text)

# Expand technical abbreviations at their first occurrence in the manuscript body.
term_text = {
    14: "分解线性模型（Decomposition-Linear，DLinear）通过移动平均分解趋势项和余项，再分别进行线性预测，在多个长期序列基准上形成了具有竞争力的简单基线[7]。iTransformer代表了显式学习变量关系的另一条路线[8]；频率增强分解Transformer（Frequency Enhanced Decomposed Transformer，FEDformer）、频率插值时间序列分析基线（Frequency Interpolation Time Series Analysis Baseline，FITS）与频域感知器分别利用频域增强、频率插值或频域多层感知机（multilayer perceptron，MLP）提取周期结构[9-11]。多尺度方面，TimeMixer通过不同下采样层级混合信息[12]，Pathformer根据输入选择尺度路径[13]，TimesNet依据频谱周期构造二维变化表示[14]，OneNet则通过在线组合应对概念漂移[15]。这些研究为“多尺度互补”和“输入相关组合”提供了方法基础，但其中多数结构并非针对骨干网OD矩阵设计。",
    15: "频域处理也存在明确边界。低频成分常对应慢变趋势和主要周期，但网络流量中的突发和状态切换可能分布在中高频段。过强低通虽然能降低局部噪声，也可能删除与预测有关的变化。因此，本文不预设频域模块必然有效，而是把频率保留比例作为敏感性因素进行消融。与面向蜂窝网格、结合频域MLP和多尺度注意力的时空多层感知机（a frequency-domain multilayer perceptron with multiscale attention network for spatial-temporal traffic prediction，STMLP）[16]不同，本文面向Abilene的144维OD矩阵，采用三个DLinear专家和由四个统计量驱动的轻量路由器，不使用二维卷积、复杂注意力或网络拓扑。",
    20: "另一类研究突出周期性和多尺度性。多尺度深层回声状态网络用于捕获物联网流量在不同尺度上的非线性依赖[5]；层次化平稳过程用于描述长期5G业务中的非平稳变化[6]；国内运营级互联网协议（Internet Protocol，IP）网络研究验证了OD矩阵直接预测对规划和运维的实际意义[17]，CycleLLH则通过周期整合和轻量预测头利用网络流量周期[19]。这些结果支持在骨干网流量中同时观察细粒度变化和较慢趋势。本文进一步把尺度贡献设计成随输入窗口变化的权重，并以固定等权组检验路由机制本身。",
    23: "国内多变量时序研究从不同角度扩展了尺度建模。基于离散傅里叶变换（discrete Fourier transform，DFT）的双分支Transformer把频率成分划分后分别预测[20]；多尺度特征融合与双注意力模型使用不同感受野提取局部模式[21]；统计特征搜索方法说明低维统计描述能够辅助结构选择[22]。本文的四个统计量不参与模型搜索，也不直接输出预测，而是作为路由信号控制三个尺度专家的组合，因而计算链路更短。",
    25: "傅里叶变换可把时域序列分解为不同频率的幅值与相位。FEDformer将分解和频域增强结合，并提示仅保留低频可能损失重要事件信息[9]；FITS使用复数线性层进行频域插值，以很少参数完成预测[10]；面向时间序列预测的频域多层感知机（Frequency-domain MLPs for Time Series forecasting，FreTS）通过频域MLP学习时间维和变量维依赖[11]。本文的频域模块仅执行标准化、实数输入快速傅里叶变换（real-input fast Fourier transform，RFFT）、低频保留和可学习门控，并非原始FITS复数插值层的复现，其目的主要是构造可控实验变量。",
    71: "门控参数在同一尺度内按频率共享给所有OD变量。逆实数输入快速傅里叶变换（inverse real-input fast Fourier transform，IRFFT）将信号恢复到时域，再还原窗口均值与标准差。实验比较r∈{0.25,0.50,0.75}。该设计把频域处理限定为可控低通和衰减，不将其等同于FITS[10]的复数频率插值。",
    79: "实验运行环境为Windows 64位操作系统、Python 3.11.9、PyTorch 2.5.1+cu118、统一计算设备架构（compute unified device architecture，CUDA）11.8、NumPy 2.4.6、pandas 2.3.3和SciPy 1.17.1。模型采用自适应矩估计（adaptive moment estimation，Adam）优化器，学习率0.001，批量大小16，最多训练30轮；验证集MSE连续6轮未改善时提前停止，并保留验证误差最低的参数。每个配置使用随机种子42、43和44独立训练。实验档案共含48次正式训练，每次保存评价指标、最佳检查点和预测数组，多尺度模型另保存测试窗口的路由权重。",
    84: "MSE、MAE和均方根误差（root mean squared error，RMSE）均越小越好。由于RMSE是MSE的单调变换，正文表格主要列出MSE和MAE。每个配置以三个随机种子的均值±样本标准差报告。统计推断的独立重复单位为随机种子，即每个配置n=3，不能把813个测试窗口当作813次独立重复。固定权重与自适应权重使用相同种子进行双侧配对t检验，并报告配对差的95%置信区间（confidence interval，CI）及Cohen's d_z。在r=0.50下的两项预设权重比较使用Holm法校正。由于n=3无法可靠检验正态性，显著性结果仅作为均值、标准差和效应方向的补充。",
    93: "为分离多尺度结构和权重生成方式的影响，表2在尺度集合与专家结构完全相同的条件下比较固定等权和自适应权重。无频域时，自适应模型相对固定等权模型的MSE降低3.95%，MAE降低11.38%。按种子计算的MSE平均配对差为0.007651，95% CI为[0.005238, 0.010064]，t(2)=13.642，Holm校正后p=0.01066，Cohen's d_z=7.876。该结果在三个种子上方向一致，但效应量和置信区间建立在很少的算法重复上，应谨慎解释。",
}
for index, text in term_text.items():
    replace_text(doc.paragraphs[index], text)

# Apply the prescribed heading hierarchy without altering reference formatting (point 17).
intro_index = 11
level1_indices = {17, 27, 76, 118}
level2_indices = {18, 21, 24, 31, 40, 44, 51, 62, 72, 77, 80, 85, 92, 102, 109, 114}

for index in {intro_index} | level1_indices | level2_indices:
    paragraph = doc.paragraphs[index]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        if run.text:
            if index == intro_index:
                set_run_font(run, "黑体", 14, bold=True)
            elif index in level1_indices:
                set_run_font(run, "黑体", 12, bold=True)
            else:
                set_run_font(run, "黑体", 10.5, bold=True)

# Body text: Chinese Songti, size 5 (10.5 pt). Exclude captions and references,
# whose dedicated requirements are handled in later points.
for index in range(12, 122):
    if index in level1_indices or index in level2_indices:
        continue
    paragraph = doc.paragraphs[index]
    text = paragraph.text.strip()
    if not text or text.startswith("图") or text.startswith("表") or paragraph_has_drawing(paragraph):
        continue
    for run in paragraph.runs:
        if run.text:
            set_run_font(run, "宋体", 10.5)

doc.save(DOCX)

after = Document(DOCX)
assert len(after.paragraphs) == before_paragraph_count
assert len(after.tables) == before_table_count
assert count_media(DOCX) == before_media_count
assert after.paragraphs[11].text == "引言"
assert after.paragraphs[17].text == "1 相关工作"
assert after.paragraphs[118].text == "4 结论"
assert "Frequency Enhanced Decomposed Transformer，FEDformer" in after.paragraphs[14].text
assert "discrete Fourier transform，DFT" in after.paragraphs[23].text
assert "real-input fast Fourier transform，RFFT" in after.paragraphs[25].text
assert "root mean squared error，RMSE" in after.paragraphs[84].text
assert "confidence interval，CI" in after.paragraphs[84].text

body_text = "".join(p.text for p in after.paragraphs[11:122])
han_count = len(re.findall(r"[\u4e00-\u9fff]", body_text))
print(f"UPDATED={DOCX}")
print(f"PARAGRAPHS={len(after.paragraphs)} TABLES={len(after.tables)} MEDIA={count_media(DOCX)}")
print(f"BODY_HAN_CHARACTERS={han_count}")

