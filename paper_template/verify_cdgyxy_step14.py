from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(r"C:\Users\Quant\Desktop\Paper2026_1\word_output\基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx")
FIGURE_DIR = Path(r"C:\Users\Quant\Desktop\Paper2026_1\abilene_forecasting\paper_figures\scipilot_figure_set")


def border(cell, edge: str) -> tuple[str | None, str | None]:
    node = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
    item = None if node is None else node.find(qn(f"w:{edge}"))
    return (None, None) if item is None else (item.get(qn("w:val")), item.get(qn("w:sz")))


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> None:
    doc = Document(DOCX)
    assert len(doc.tables) == 3
    for table_index, table in enumerate(doc.tables, 1):
        last = len(table.rows) - 1
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                expected_top = ("single", "12") if row_index == 0 else ("nil", "0")
                expected_bottom = (
                    ("single", "12") if row_index == last
                    else ("single", "6") if row_index == 0
                    else ("nil", "0")
                )
                assert border(cell, "top") == expected_top, (table_index, row_index, "top", border(cell, "top"))
                assert border(cell, "bottom") == expected_bottom, (table_index, row_index, "bottom", border(cell, "bottom"))
                for edge in ("left", "right", "insideH", "insideV"):
                    assert border(cell, edge) == ("nil", "0"), (table_index, row_index, edge, border(cell, edge))

    equations = [
        p for p in doc.paragraphs
        if any(shape.get("equationxml") for shape in p._p.xpath(".//*[local-name()='shape']"))
    ]
    assert len(equations) == 20
    numbers = [int(re.search(r"\((\d+)\)", p.text).group(1)) for p in equations]
    assert numbers == list(range(1, 21)), numbers

    with zipfile.ZipFile(DOCX) as archive:
        assert sha256(archive.read("word/media/image2.png")) == sha256((FIGURE_DIR / "S1_小样本透明化消融.png").read_bytes())
        assert sha256(archive.read("word/media/image5.png")) == sha256((FIGURE_DIR / "S3_频域敏感性原始轨迹.png").read_bytes())

    required = [sentence for sentence in (
        "q为移动平均窗口长度",
        "W_T^{(s)}、W_R^{(s)}",
        "a_n为第n个样本的尺度分数",
        "F^{(s)}_k表示第k个频率系数",
        "L_MSE为训练损失",
    )]
    body = "\n".join(p.text for p in doc.paragraphs)
    assert all(sentence in body for sentence in required)
    print("PASS: 3 tables, 20 formulas, 2 corrected figures, and formula explanations verified")


if __name__ == "__main__":
    main()
