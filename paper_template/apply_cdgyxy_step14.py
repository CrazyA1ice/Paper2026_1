from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path(
    r"C:\Users\Quant\Desktop\Paper2026_1\word_output\基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx"
)

EXPLANATIONS = {
    "趋势项T^{(s)}和余项R^{(s)}": (
        "其中，q为移动平均窗口长度，m=(q-1)/2=12为端点复制的半窗口长度，"
        "T^{(s)}和R^{(s)}分别表示尺度s的趋势项和余项。"
    ),
    "同一尺度内的144个变量共享时间映射权重": (
        "式中，W_T^{(s)}、W_R^{(s)}和b_T^{(s)}、b_R^{(s)}分别为趋势支路与余项支路的"
        "线性映射权重和偏置，Ŷ^{(s)}为尺度s的预测结果。"
    ),
    "Softmax将分数转换为非负、和为1的权重": (
        "式中，a_n为第n个样本的尺度分数，W_1、W_2和b_1、b_2为两层感知器的权重和偏置，"
        "ReLU为线性整流函数。"
    ),
    "门控参数在同一尺度内按频率共享": (
        "式中，X̃^{(s)}表示标准化后的尺度序列，F^{(s)}_k表示第k个频率系数，K_s为保留的低频数，"
        "g^{(s)}_k为门控系数，F̂^{(s)}_k为门控后的频率系数。"
    ),
    "使用Adam优化器更新参数": (
        "式中，L_MSE为训练损失，N为批量样本数，H为预测长度，C为变量数，"
        "Y和Ŷ分别表示真实值和预测值。"
    ),
}


def set_run_font(run) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def main() -> None:
    doc = Document(DOCX)

    # The template stores editable Word equation objects as VML shapes carrying
    # an equationxml payload (rather than as modern OMML nodes).
    equation_paragraphs = [
        p for p in doc.paragraphs
        if any(shape.get("equationxml") for shape in p._p.xpath(".//*[local-name()='shape']"))
    ]
    if len(equation_paragraphs) != 20:
        raise RuntimeError(f"应有20个公式，实际找到{len(equation_paragraphs)}个")
    numbers = []
    for paragraph in equation_paragraphs:
        matches = re.findall(r"\((\d+)\)", paragraph.text)
        if len(matches) != 1:
            raise RuntimeError(f"公式编号无法唯一识别：{paragraph.text!r}")
        numbers.append(int(matches[0]))
        if paragraph.style.name != "Equation":
            raise RuntimeError(f"公式({matches[0]})未使用Equation样式")
        tabs = paragraph._p.xpath("./w:pPr/w:tabs/w:tab")
        tab_values = {(tab.get(qn("w:val")), tab.get(qn("w:pos"))) for tab in tabs}
        if ("center", "4819") not in tab_values or ("right", "9524") not in tab_values:
            raise RuntimeError(f"公式({matches[0]})缺少居中或右对齐制表位")
    if numbers != list(range(1, 21)):
        raise RuntimeError(f"公式编号不连续：{numbers}")

    for prefix, sentence in EXPLANATIONS.items():
        matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(matches) != 1:
            raise RuntimeError(f"无法唯一定位公式说明段：{prefix}")
        paragraph = matches[0]
        if sentence not in paragraph.text:
            spacer = "" if paragraph.text.endswith((" ", "　")) else " "
            run = paragraph.add_run(spacer + sentence)
            set_run_font(run)

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
