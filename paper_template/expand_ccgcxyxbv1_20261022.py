# -*- coding: utf-8 -*-
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
from hashlib import sha256
import re, json

root=Path("paper_template/长春工程学院学报(自然科学版)")
src=root/"自适应多尺度加权的多变量网络流量预测_ccgcxyxbV1.docx"
out=root/"自适应多尺度加权的多变量网络流量预测_ccgcxyxbV2_扩充稿.docx"
report=root/"ccgcxyxbV1至V2_扩充修改对比报告.md"
old=Path("word_output/2026.10.08v75.docx")

# (section, following section, restored topic, newly written content)
blocks=[
("0 引言","1自适应多尺度","OD研究对象和任务边界","从预测任务来看，OD流量矩阵以源节点和目的节点的组合为变量，记录了各节点对之间随时间变化的通信需求。与只考察全网总流量不同，多变量预测需要在同一时间窗口内同时给出多个OD变量的未来轨迹。本文不在预测阶段显式建立节点拓扑关系，而是将计算重点放在各变量共享的时间变化模式上，使模型能够在较多OD变量的条件下保持相对简洁的参数结构。"),
("1.1 数据表示与预处理","1.2","预处理合理性与对照公平性","上述处理同时服务于模型训练和结果比较。时间顺序划分保证验证集、测试集位于训练阶段之后；标准化参数只由训练部分确定，避免后续时间段的信息参与预处理参数估计。每个OD变量分别进行变换与标准化，可以减小不同流量量级对损失计算的影响。构造样本时，输入窗口和预测目标均限制在各自数据子集内，因而同一数据划分能够供DLinear、LightTS及各融合模型使用。"),
("1.2 多尺度构造与线性专家","1.3","尺度专家设计逻辑","三个尺度并非对应三个不同的预测时段，而是在同一段历史上形成粒度不同的输入表示。细粒度分支直接保留逐小时观测，较粗粒度分支通过局部平均削弱相邻小时之间的短时起伏；各分支随后独立完成趋势项和余项的线性映射。这样设置的目的，是使候选预测来自不同时间分辨率，同时保持输出维度一致。尺度数和专家结构在全部融合消融模型中保持不变，有利于单独考察权重分配方式的作用。"),
("1.3 窗口统计特征与自适应加权","1.4","低维窗口路由机制","与直接把全部OD时间序列输入另一套预测网络不同，路由分支仅接收四个由当前窗口计算得到的统计量。这些量依次从整体位置、离散程度、末端状态和相邻时刻变化四个方面描述窗口，但不对某个单独OD变量生成专属权重。Softmax将三个尺度分数归一化，使最终输出保持为各专家候选预测的加权组合。对于不同历史窗口，专家网络保持不变，融合系数则由各自窗口的特征确定。"),
("2.1 数据与实验设置","2.2","重复实验与配对统计口径","基线与本文方法使用相同数据集、时间划分、输入窗口和预测目标，所有模型均在各自验证集上选择检查点，再使用测试集评价。随机种子42～49用于重复训练，并按种子保留预测记录和误差指标。表中“均值±标准差”反映不同初始化下的描述性变化；相对降幅先在同一种子内配对计算，再对八次结果取平均，以避免将两个模型的非配对汇总值直接作为单次改善比例。"),
("2.2 预测精度比较","2.3","模型对照与参数量解读","与单尺度DLinear相比，本文方法通过另外两个聚合尺度提供补充候选预测；与LightTS相比，主要差别在于以三个分解线性专家和低维路由完成任务，而非采用外部基线的网络结构。表2和图2将绝对误差、逐种子配对降幅及参数数量分别呈现，可以同时观察预测结果和模型规模。需要指出，参数比例以相应数据集上的LightTS配置为分母，并非声称本文方法是所有对照中参数最少的模型。"),
("2.3 融合方式消融","2.4","受控消融与窗口信息作用","消融实验将预测专家、输入尺度和训练任务固定，仅改变融合系数的生成方式。固定等权不学习尺度偏好，全局可学习权重虽可在训练中调整，却对所有窗口使用同一组系数；恒定输入路由则保留路由网络的结构与参数数量，但不向其提供窗口统计信息。与这些设置相比较，本文方法的融合权重直接依赖当前输入窗口。两套数据集的平均MSE、MAE结果均支持这种窗口级融合方式在所设24 h任务中的作用。"),
("2.4 高需求预测表现","2.5","高需求评价指标解释","高需求点由训练集各OD变量的分位数单独确定，再将同一阈值应用于该变量的测试目标，因此不同变量的高需求判定不依赖统一的绝对流量门槛。MAE用于衡量这些目标点上的总体预测偏差；AUE只累计预测低于真实值的差额，并对全部入选目标点取平均。两项指标从不同角度反映高需求时段的预测质量。Abilene上两项指标均获得明确的平均改善，可为后续需求估计提供补充依据。"),
("2.5 一周预测表现","3结语","一周预测证据完整性","一周任务与24 h任务采用相同的96 h输入和数据预处理，区别在于输出序列延长至168个小时，并为这一输出长度重新训练相应模型。表4报告每种方法八次重复的绝对误差，图4则进一步展示逐种子配对降幅，二者分别回答平均误差大小和重复结果方向的问题。GÉANT的这一设置覆盖307个共同测试窗口，评价时包含全部预测小时及全部保留的OD变量，避免仅用少量典型曲线代替整体结果。")
]
normalize=lambda s:re.sub(r"\s+","",s).replace("．",".")
def hashfile(p):return sha256(p.read_bytes()).hexdigest()
d=Document(src)
original=[p.text for p in d.paragraphs if p.text.strip()]
for section,next_title,title,text in blocks:
    paragraphs=list(d.paragraphs)
    found=[p for p in paragraphs if normalize(p.text).startswith(normalize(next_title)) and len(normalize(p.text))<66]
    assert len(found)==1,(section,next_title,[p.text for p in found])
    anchor=found[0]
    offset=paragraphs.index(anchor)
    source=next((p for p in reversed(paragraphs[:offset]) if len(p.text)>40 and not re.match(r"^\s*\d[\d.]*\s",p.text)),None)
    assert source is not None,section
    node=OxmlElement("w:p")
    if source._p.pPr is not None:node.insert(0,deepcopy(source._p.pPr))
    anchor._p.addprevious(node)
    para=Paragraph(node,anchor._parent);para.style=source.style
    run=para.add_run(text)
    for existing in source.runs:
        if existing.text and existing._r.rPr is not None:
            run._r.insert(0,deepcopy(existing._r.rPr));break
    para.paragraph_format.keep_with_next=False
d.save(out)
after=Document(out)
paragraphs=[p.text for p in after.paragraphs if p.text.strip()]
it=iter(paragraphs)
for item in original:
    for found in it:
        if item==found:break
    else:raise AssertionError("original paragraph changed: "+item[:50])
for _,_,_,text in blocks:assert paragraphs.count(text)==1
before_count=sum(len(p.text) for p in Document(src).paragraphs)
after_count=sum(len(p.text) for p in after.paragraphs)
gain=after_count-before_count
assert gain==sum(len(x[3]) for x in blocks)==1580,gain
assert len(Document(src).tables)==len(after.tables)
with ZipFile(src) as z0,ZipFile(out) as z1:
    common=set(z0.namelist()) & set(z1.namelist())
    changed=sorted(k for k in common if z0.read(k)!=z1.read(k))
    assert "word/document.xml" in changed, changed
    # python-docx may normalize unrelated OOXML serialization (namespace, rels);
    # strict guards below verify retained original content and objects.
    for key in common:
        if key.startswith("word/media/"):assert z0.read(key)==z1.read(key)
    for token in (b"<m:oMath",b"<w:tbl",b"<w:drawing"):
        assert z0.read("word/document.xml").count(token)==z1.read("word/document.xml").count(token)
    assert len(Document(src).inline_shapes)==len(after.inline_shapes)
rows=[
"# V1至V2扩充修改对比报告",
"",
"## 1. 修改对象与结果",
"",
"原稿："+str(src),
"",
"扩充稿："+str(out),
"",
"参考历史版本："+str(old),
"",
"本次新增9段，共1580个字符（含标点和英数，按Word正文段落文本字符计）。原有段落保留且顺序不变；原稿没有被覆盖。",
"",
"## 2. 相对于较长历史版本恢复的内容",
"",
"对照历史长稿V75与本次面向长春工程学院学报压缩的稿件，原有的预测任务边界、数据处理依据、尺度专家设计动机、窗口路由作用和实验控制变量解释等论证被明显精简。本次根据当前正文和既有实验证据重新组织，而非直接复制旧稿。",
"",
"| 位置 | 恢复内容 | 字符数 |",
"|---|---|---:|"]
for section,_,why,text in blocks:rows.append("| "+section+" | "+why+" | "+str(len(text))+" |")
rows+=["","## 3. 新增正文逐段对比",""]
for section,_,why,text in blocks:
    rows += ["### "+section+"："+why, "", "**修改前：** 此处结束后直接进入下一节，缺少以下解释性文字。", "",
             "**修改后新增：** "+text,""]
rows+=["## 4. 未恢复的内容及选取原则","",
"没有恢复早期频域探索和与本次正向论文主张无关的扩展结果；没有添加在线部署、显著性、所有种子或所有预测长度普遍优于对照等缺乏当前证据支持的表述。所有原始实验记录保持不变。",
"",
"## 5. 保护性核验","",
"- 原稿SHA256："+hashfile(src),
"- 扩充稿SHA256："+hashfile(out),
"- 原稿正文段落字符数："+str(before_count)+"；扩充稿："+str(after_count)+"；净新增："+str(gain),
"- 原有非空正文段落逐一保留，新增9段。",
"- 表格数量、嵌入图片数量、数学对象标签数均一致，媒体资源字节未变。",
"- 修改过的DOCX内部组件："+", ".join(changed),
"- 字数说明：段落文本字符统计不包含表格、文本框、公式内字符，和Word全文字数口径不同；期刊“6000字左右”应以其指定口径再核对。",
""]
report.write_text("\n".join(rows),encoding="utf-8")
qa={"status":"PASS","added_chars":gain,"new_paragraphs":len(blocks),
"before_chars":before_count,"after_chars":after_count,"changed_zip_parts":changed,
"source_sha256":hashfile(src),"output_sha256":hashfile(out)}
(root/"ccgcxyxbV1至V2_扩充QA.json").write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(qa,ensure_ascii=False))
