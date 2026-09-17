from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_SciPilot图表整合最新版.docx"
OUTPUT = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


source = Document(SOURCE)
output = Document(OUTPUT)
assert [p.text for p in source.paragraphs] == [p.text for p in output.paragraphs]
assert len(source.tables) == len(output.tables)
assert len(source.inline_shapes) == len(output.inline_shapes)

section = output.sections[0]
assert round(section.top_margin.cm, 1) == 3.1
assert round(section.bottom_margin.cm, 1) == 2.5
assert round(section.left_margin.cm, 1) == 2.0
assert round(section.right_margin.cm, 1) == 2.0

with ZipFile(OUTPUT) as archive:
    document_root = etree.fromstring(archive.read("word/document.xml"))
    paragraphs = document_root.xpath(".//w:p", namespaces=NS)
    assert paragraphs
    assert all(
        (spacing := paragraph.find("w:pPr/w:spacing", namespaces=NS)) is not None
        and spacing.get(w("line")) == "240"
        and spacing.get(w("lineRule")) == "auto"
        for paragraph in paragraphs
    )
    section_props = document_root.xpath(".//w:sectPr", namespaces=NS)
    assert all(section.find("w:cols", namespaces=NS).get(w("num")) == "1" for section in section_props)
    assert all(section.find("w:pgSz", namespaces=NS).get(w("orient")) is None for section in section_props)

    styles_root = etree.fromstring(archive.read("word/styles.xml"))
    default_fonts = styles_root.find("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", namespaces=NS)
    normal_styles = styles_root.xpath(
        "./w:style[@w:type='paragraph' and (@w:default='1' or w:name/@w:val='Normal' or w:name/@w:val='正文')]",
        namespaces=NS,
    )
    assert normal_styles
    normal_fonts = normal_styles[0].find("w:rPr/w:rFonts", namespaces=NS)
    assert default_fonts is not None and default_fonts.get(w("eastAsia")) == "宋体"
    assert normal_fonts is not None and normal_fonts.get(w("eastAsia")) == "宋体"

print(f"段落={len(output.paragraphs)}，表格={len(output.tables)}，图片={len(output.inline_shapes)}")
print(
    "页边距(cm)："
    f"上={section.top_margin.cm:.2f}，下={section.bottom_margin.cm:.2f}，"
    f"左={section.left_margin.cm:.2f}，右={section.right_margin.cm:.2f}"
)
print(f"document.xml 段落单倍行距={len(paragraphs)}/{len(paragraphs)}，纵向单栏=通过，默认中文字体=宋体")
print("正文文本、表格数量、图片数量与源稿一致")
