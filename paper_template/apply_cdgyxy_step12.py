from __future__ import annotations

import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path(r"C:\Users\Quant\Desktop\Paper2026_1\word_output\基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx")
FIGURE_DIR = Path(r"C:\Users\Quant\Desktop\Paper2026_1\abilene_forecasting\paper_figures\scipilot_figure_set")

FIGURE_MEDIA = {
    "word/media/image2.png": FIGURE_DIR / "S1_小样本透明化消融.png",
    "word/media/image3.png": FIGURE_DIR / "S2_配对权重消融.png",
    "word/media/image4.png": FIGURE_DIR / "S4a_路由权重诊断.png",
    "word/media/image5.png": FIGURE_DIR / "S3_频域敏感性原始轨迹.png",
    "word/media/image6.png": FIGURE_DIR / "S4b_代表性预测诊断.png",
}

TEXT_APPEND = {
    86: "随机种子分布及均值见图2。",
    93: "同种子配对变化见图3。",
    98: "路由权重的样本变化与总体分布见图4。",
    103: "各随机种子的敏感性轨迹见图5。",
}


def set_run_font(run, east_asia: str, latin: str, size_pt: float, bold: bool) -> None:
    run.font.name = latin
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def update_document(source: Path, intermediate: Path) -> None:
    doc = Document(source)

    for index, suffix in TEXT_APPEND.items():
        paragraph = doc.paragraphs[index]
        for run in paragraph.runs:
            if run.text == " " + suffix:
                run.text = suffix
        if suffix not in paragraph.text:
            run = paragraph.add_run(suffix)
            set_run_font(run, "宋体", "Times New Roman", 10.5, False)

    captions = [p for p in doc.paragraphs if re.match(r"^图\d+\s", p.text.strip())]
    if len(captions) != 6:
        raise RuntimeError(f"应有6个图题，实际找到{len(captions)}个")

    for caption in captions:
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(0)
        caption.paragraph_format.keep_with_next = False
        for run in caption.runs:
            set_run_font(run, "黑体", "Times New Roman", 9, True)

    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:drawing|.//w:pict"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = True

    doc.save(intermediate)


def replace_media(source: Path, destination: Path) -> None:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=destination.parent) as handle:
        rebuilt = Path(handle.name)
    try:
        with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as zout:
            names = set(zin.namelist())
            missing = set(FIGURE_MEDIA) - names
            if missing:
                raise RuntimeError(f"DOCX中缺少预期图片：{sorted(missing)}")
            for item in zin.infolist():
                data = FIGURE_MEDIA[item.filename].read_bytes() if item.filename in FIGURE_MEDIA else zin.read(item.filename)
                zout.writestr(item, data)
        os.replace(rebuilt, destination)
    finally:
        if rebuilt.exists():
            rebuilt.unlink()


def main() -> None:
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    for figure in FIGURE_MEDIA.values():
        if not figure.exists():
            raise FileNotFoundError(figure)

    with tempfile.TemporaryDirectory(dir=DOCX.parent) as temp_dir:
        intermediate = Path(temp_dir) / "step12_text.docx"
        update_document(DOCX, intermediate)
        replace_media(intermediate, DOCX)

    print(DOCX)


if __name__ == "__main__":
    main()
