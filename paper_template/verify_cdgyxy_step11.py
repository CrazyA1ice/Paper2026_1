from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Quant\Desktop\Paper2026_1")
DOCX = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx"
PAGES = ROOT / "word_output" / "qa_cdgyxy_step11" / "pages"


doc = Document(DOCX)
assert len(doc.paragraphs) == 145
assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 6
assert doc.sections[0].different_first_page_header_footer

footer_lines = [p.text for p in doc.sections[0].first_page_footer.paragraphs if p.text]
assert footer_lines == [
    "收稿日期：[待填写]",
    "基金项目：[基金来源]（[基金项目编号]）；如无基金项目请填写“无”。",
    "第一作者简介：[姓名]（[出生年]—），[性别]，[职称]，[学位]，研究方向：[请填写]。",
    "通信作者简介：[姓名]（[出生年]—），[性别]，[职称]，[学位]，研究方向：[请填写]，电子邮箱：[请填写]。",
]
for paragraph in doc.sections[0].first_page_footer.paragraphs:
    if not paragraph.text:
        continue
    for run in (r for r in paragraph.runs if r.text):
        rpr = run._element.rPr
        assert rpr is not None and rpr.rFonts is not None
        assert rpr.rFonts.get(qn("w:eastAsia")) == "宋体"
        assert rpr.sz.get(qn("w:val")) == "16"

main_text = "\n".join(p.text for p in doc.paragraphs)
assert not any(
    p.text.startswith(("收稿日期：", "基金项目：", "作者简介：", "联系方式："))
    for p in doc.paragraphs
)

# All numbered references remain cited after compression.
cited = set()
for match in re.finditer(r"\[(\d+)(?:-(\d+))?(?:,([\d,-]+))?\]", "\n".join(p.text for p in doc.paragraphs[11:122])):
    start = int(match.group(1))
    end = int(match.group(2) or start)
    cited.update(range(start, end + 1))
    if match.group(3):
        for part in match.group(3).split(","):
            if "-" in part:
                left, right = map(int, part.split("-"))
                cited.update(range(left, right + 1))
            elif part:
                cited.add(int(part))
assert cited == set(range(1, 23)), sorted(cited)

for token in (
    "0.186099", "0.168203", "4.06%", "10.52%", "3.95%",
    "48次训练", "三个随机种子", "频域门控未超过无频域主模型",
):
    assert token in main_text

with zipfile.ZipFile(DOCX) as archive:
    names = archive.namelist()
    media = [name for name in names if name.startswith("word/media/") and not name.endswith("/")]
    assert len(media) == 6
    document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "w:titlePg" in document_xml
    assert 'w:type="first"' in document_xml

pages = sorted(PAGES.glob("page-*.png"))
assert len(pages) == 11

print("首页专用页脚=4行；正文末尾旧题录信息=已移除")
print("正文引用覆盖=[1]-[22]；核心实验数值与结论边界=保留")
print(f"段落={len(doc.paragraphs)}，表格={len(doc.tables)}，图片={len(doc.inline_shapes)}，渲染页={len(pages)}")

