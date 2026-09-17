from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"C:\Users\Quant\Desktop\Paper2026_1")
DOCX = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx"
QA_DIR = ROOT / "word_output" / "qa_cdgyxy_step11"
BACKUP = QA_DIR / "before_compression_and_step11.docx"
AUDIT = QA_DIR / "compression_audit.json"


def count_units(text: str) -> int:
    return len(re.findall(r"[\u3400-\u9fff]|[A-Za-z0-9]+(?:[-'’][A-Za-z0-9]+)*", text))


def set_run_font(run, east_asia: str = "宋体", size_pt: float = 10.5) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size_pt)


def replace_text(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
    else:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    for run in paragraph.runs:
        if run.text:
            set_run_font(run)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_single_spacing(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")


def count_media(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return sum(name.startswith("word/media/") and not name.endswith("/") for name in archive.namelist())


QA_DIR.mkdir(parents=True, exist_ok=True)
if not BACKUP.exists():
    shutil.copy2(DOCX, BACKUP)

doc = Document(DOCX)
before_paragraphs = len(doc.paragraphs)
before_tables = len(doc.tables)
before_images = len(doc.inline_shapes)
before_media = count_media(DOCX)

# Compress repeated positioning and interpretation while retaining every citation,
# acronym definition, experimental fact, limitation, and conclusion boundary.
replacements = {
    12: "骨干网络中的源—目的节点对（origin-destination，OD）流量刻画入口与出口之间的传输需求，短期预测可服务容量规划、拥塞预警、流量工程和资源调度。现有方法涵盖统计、机器学习和深度学习[1-2]，并从特征交互、拓扑传播、多尺度动态及非平稳性等方面改善预测[3-6]。这些研究表明，网络流量具有多变量关联和多时间尺度变化。",
    13: "细粒度观测保留突发变化，粗尺度平滑噪声并突出趋势。固定尺度或固定融合比例难以随当前流量状态调整预测依据；图卷积和复杂Transformer又增加参数、训练成本及拓扑要求。因此，本研究检验轻量线性主干配合样本级尺度选择，能否在不引入复杂拓扑建模时获得稳定收益。",
    14: "分解线性模型（Decomposition-Linear，DLinear）以移动平均分离趋势和余项并进行线性预测，是轻量基线[7]；iTransformer显式学习变量关系[8]。频率增强分解Transformer（Frequency Enhanced Decomposed Transformer，FEDformer）、频率插值时间序列分析基线（Frequency Interpolation Time Series Analysis Baseline，FITS）及频域多层感知机（multilayer perceptron，MLP）分别利用频域增强、插值或MLP提取周期结构[9-11]。TimeMixer、Pathformer、TimesNet和OneNet通过多尺度混合、自适应路径、频谱周期或在线组合处理尺度变化[12-15]。这些工作支持多尺度互补与输入相关组合，但并非专为骨干网OD矩阵设计。",
    15: "低频通常对应趋势和主要周期，但突发与状态切换可能位于中高频，过强低通会删除预测信息。因此，频率保留比例仅作为敏感性因素。与面向蜂窝网格、结合频域MLP和多尺度注意力的时空多层感知机（a frequency-domain multilayer perceptron with multiscale attention network for spatial-temporal traffic prediction，STMLP）[16]不同，本研究面向Abilene的144维OD矩阵，采用DLinear专家和四统计量路由器，不使用二维卷积、复杂注意力或网络拓扑。",
    16: "主要工作包括：1）构建尺度1、2和4的三分支DLinear专家；2）以窗口均值、波动、最近状态和变化强度生成样本级非负归一化权重，并设置固定等权对照；3）在Abilene上完成48次训练，覆盖基础模型、组件消融、频率敏感性和路由分布；4）如实报告频域门控未超过无频域主模型。结论仅适用于既定数据划分、预测任务和三个随机种子。",
    19: "网络流量预测对象包括链路负载、节点业务和OD矩阵。传统统计模型解释性较好，深度模型更能表示高维非线性关系，但需权衡精度、成本与可解释性[1-2]。特征交互模型[3]、图卷积与多头注意力[4]以及动态扩散卷积交互图网络[18]可建模时空依赖，适合拓扑明确的场景，但结构和先验要求较高。",
    20: "周期性与多尺度性是另一研究重点。多尺度回声状态网络捕获物联网流量的跨尺度依赖[5]，层次化平稳过程描述长期5G业务的非平稳变化[6]。运营级互联网协议（Internet Protocol，IP）网络研究验证了OD预测的运维价值[17]，CycleLLH利用周期整合和轻量预测头[19]。本研究据此设计随窗口变化的尺度权重，并用固定等权组检验路由贡献。",
    22: "DLinear把趋势和余项分别映射为多步预测[7]，参数少且训练稳定，适合作为轻量主干。iTransformer把变量作为标记学习变量相关性[8]。本研究使用共享时间映射处理144个OD变量，聚焦时间尺度融合，不把结果解释为已充分建模OD拓扑。",
    23: "国内研究以基于离散傅里叶变换（discrete Fourier transform，DFT）的双分支Transformer[20]、多尺度特征融合与双注意力[21]和统计特征搜索[22]扩展尺度建模。本研究的四个统计量仅作为路由信号，控制三个尺度专家组合，不参与模型搜索或直接预测。",
    25: "傅里叶变换表征不同频率的幅值与相位。FEDformer结合分解与频域增强[9]；FITS用复数线性层进行频域插值[10]；面向时间序列预测的频域多层感知机（Frequency-domain MLPs for Time Series forecasting，FreTS）学习时间维和变量维依赖[11]。本研究仅执行标准化、实数输入快速傅里叶变换（real-input fast Fourier transform，RFFT）、低频保留和可学习门控，用于构造可控变量，不复现FITS的复数插值层。",
    26: "TimeMixer、Pathformer、TimesNet和OneNet分别通过下采样混合、自适应路径、频谱周期和在线组合利用多尺度信息[12-15]。本文仅保留三条线性专家及两层路由器。最近邻工作STMLP将频域MLP和多尺度注意力用于二维蜂窝网格[16]，而本文面向公开骨干网OD矩阵，以统计量驱动样本级尺度加权，并用固定等权对照直接检验动态权重收益。",
    115: "三层证据支持主要结论：主模型优于单尺度DLinear；在尺度和专家相同条件下，自适应权重优于固定等权；路由权重随窗口变化。因此，收益主要来自样本级尺度组合，但实验未分别检验四个统计量的因果贡献。",
    116: "较高频率保留比例优于较低比例，提示中高频信息对Abilene短期预测有用；但含门控模型未超过无频域主模型，因此频域门控不能作为精度创新。该边界与低频截断风险的相关研究一致[9-10,20]。",
    117: "结论受三点限制：仅使用一个骨干网数据集，跨网络迁移性尚未验证[17-19]；每个配置只有三个随机种子，统计推断应以均值、标准差、置信区间和方向一致性为主；模型未使用OD拓扑或显式变量注意力，不能据此否定相关路线[4,8,18]。后续需在第二个公开数据集和12、48、96等预测长度上复核固定/自适应消融；频域模块可尝试残差旁路或分频带选择。",
}

audit_rows = []
for index, new_text in replacements.items():
    old_text = doc.paragraphs[index].text
    audit_rows.append(
        {
            "paragraph": index,
            "before_units": count_units(old_text),
            "after_units": count_units(new_text),
            "delta": count_units(new_text) - count_units(old_text),
        }
    )
    replace_text(doc.paragraphs[index], new_text)

# Remove the four metadata paragraphs previously placed after the references.
# Their content is relocated to the first-page footer below.
for index in (148, 147, 146, 145):
    remove_paragraph(doc.paragraphs[index])

section = doc.sections[0]
section.different_first_page_header_footer = True
section.footer_distance = Cm(0.8)
footer = section.first_page_footer
footer.is_linked_to_previous = False

while len(footer.paragraphs) > 1:
    remove_paragraph(footer.paragraphs[-1])
first = footer.paragraphs[0]
clear_paragraph(first)

footer_lines = [
    "收稿日期：[待填写]",
    "基金项目：[基金来源]（[基金项目编号]）；如无基金项目请填写“无”。",
    "第一作者简介：[姓名]（[出生年]—），[性别]，[职称]，[学位]，研究方向：[请填写]。",
    "通信作者简介：[姓名]（[出生年]—），[性别]，[职称]，[学位]，研究方向：[请填写]，电子邮箱：[请填写]。",
]
footer_paragraphs = [first]
for _ in footer_lines[1:]:
    footer_paragraphs.append(footer.add_paragraph())

for paragraph, text in zip(footer_paragraphs, footer_lines):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_single_spacing(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 8)

doc.save(DOCX)

after = Document(DOCX)
assert len(after.paragraphs) == before_paragraphs - 4
assert len(after.tables) == before_tables
assert len(after.inline_shapes) == before_images
assert count_media(DOCX) == before_media
assert after.sections[0].different_first_page_header_footer
assert [p.text for p in after.sections[0].first_page_footer.paragraphs] == footer_lines
assert not any(p.text.startswith(("收稿日期：", "基金项目：", "作者简介：", "联系方式：")) for p in after.paragraphs)

AUDIT.write_text(
    json.dumps(
        {
            "revised_paragraphs": audit_rows,
            "estimated_unit_delta": sum(row["delta"] for row in audit_rows),
            "body_metadata_paragraphs_removed": 4,
            "preserved": {
                "tables": before_tables,
                "inline_shapes": before_images,
                "media_files": before_media,
            },
        },
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)

print(f"UPDATED={DOCX}")
print(f"PARAGRAPHS={len(after.paragraphs)} TABLES={len(after.tables)} IMAGES={len(after.inline_shapes)}")
print(f"ESTIMATED_REVISED_PARAGRAPH_DELTA={sum(row['delta'] for row in audit_rows)}")
print(f"FIRST_PAGE_FOOTER_LINES={len(footer_lines)}")

