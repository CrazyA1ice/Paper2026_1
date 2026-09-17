from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Quant\Desktop\Paper2026_1")
DOCX = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_成都工业学院学报版.docx"
PAGES = ROOT / "word_output" / "qa_cdgyxy_step10" / "pages"


def assert_run_format(run, font: str, size_half_points: int, bold: bool | None = None) -> None:
    rpr = run._element.rPr
    assert rpr is not None
    assert rpr.rFonts is not None
    assert rpr.rFonts.get(qn("w:eastAsia")) == font
    assert int(rpr.sz.get(qn("w:val"))) == size_half_points
    if bold is True:
        assert run.bold is True


doc = Document(DOCX)
expected_headings = {
    11: "引言",
    17: "1 相关工作",
    18: "1.1 网络流量预测",
    21: "1.2 轻量多变量时间序列预测",
    24: "1.3 频域建模与自适应多尺度融合",
    27: "2 自适应多尺度网络流量预测方法",
    31: "2.1 问题定义与数据预处理",
    40: "2.2 多尺度序列构造",
    44: "2.3 DLinear尺度专家",
    51: "2.4 样本级自适应尺度加权",
    62: "2.5 辅助频域门控",
    72: "2.6 训练目标与复杂度",
    76: "3 实验与结果分析",
    77: "3.1 数据集与运行环境",
    80: "3.2 评价指标与统计方法",
    85: "3.3 基础模型与组件消融",
    92: "3.4 自适应权重消融与路由行为",
    102: "3.5 频率保留比例敏感性",
    109: "3.6 代表性预测与应用含义",
    114: "3.7 讨论与局限性",
    118: "4 结论",
}
for index, text in expected_headings.items():
    assert doc.paragraphs[index].text == text

level1 = {17, 27, 76, 118}
level2 = set(expected_headings) - {11} - level1
for run in (r for r in doc.paragraphs[11].runs if r.text):
    assert_run_format(run, "黑体", 28, True)
for index in level1:
    for run in (r for r in doc.paragraphs[index].runs if r.text):
        assert_run_format(run, "黑体", 24, True)
for index in level2:
    for run in (r for r in doc.paragraphs[index].runs if r.text):
        assert_run_format(run, "黑体", 21, True)

body_checked = 0
for index in range(12, 122):
    if index in level1 or index in level2:
        continue
    paragraph = doc.paragraphs[index]
    text = paragraph.text.strip()
    if not text or text.startswith("图") or text.startswith("表"):
        continue
    if paragraph._p.xpath(".//w:drawing | .//w:pict"):
        continue
    for run in (r for r in paragraph.runs if r.text):
        assert_run_format(run, "宋体", 21)
    body_checked += 1

required_definitions = {
    "DLinear": "Decomposition-Linear，DLinear",
    "FEDformer": "Frequency Enhanced Decomposed Transformer，FEDformer",
    "FITS": "Frequency Interpolation Time Series Analysis Baseline，FITS",
    "MLP": "multilayer perceptron，MLP",
    "STMLP": "spatial-temporal traffic prediction，STMLP",
    "IP": "Internet Protocol，IP",
    "DFT": "discrete Fourier transform，DFT",
    "FreTS": "Frequency-domain MLPs for Time Series forecasting，FreTS",
    "RFFT": "real-input fast Fourier transform，RFFT",
    "IRFFT": "inverse real-input fast Fourier transform，IRFFT",
    "CUDA": "compute unified device architecture，CUDA",
    "Adam": "adaptive moment estimation，Adam",
    "RMSE": "root mean squared error，RMSE",
    "CI": "confidence interval，CI",
}
body_text = "\n".join(p.text for p in doc.paragraphs[11:122])
for abbreviation, definition in required_definitions.items():
    assert definition in body_text, abbreviation

body_units = len(re.findall(r"[\u3400-\u9fff]|[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*", body_text))
assert body_units < 8000
assert len(doc.paragraphs) == 149
assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 6
with zipfile.ZipFile(DOCX) as archive:
    media_count = sum(name.startswith("word/media/") and not name.endswith("/") for name in archive.namelist())
assert media_count == 6

page_files = sorted(PAGES.glob("page-*.png"))
assert len(page_files) == 12

print(f"标题层级={len(expected_headings)}处；引言不编号；一级标题从1开始")
print(f"正文宋体五号检查={body_checked}段；一级/二级标题字号与字体=通过")
print(f"首次术语全称检查={len(required_definitions)}项；正文估算字数单位={body_units}（<8000）")
print(f"段落={len(doc.paragraphs)}，表格={len(doc.tables)}，图片={len(doc.inline_shapes)}，渲染页={len(page_files)}")

