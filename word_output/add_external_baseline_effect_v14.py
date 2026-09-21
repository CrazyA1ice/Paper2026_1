from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.21v13.docx"
TARGET = ROOT / "word_output" / "2026.09.21v14.docx"
FIGURE = (
    ROOT / "abilene_forecasting" / "paper_figures" / "exports" / "color"
    / "fig2_external_baseline_reduction_color.png"
)


def find_one(document: Document, startswith: str):
    matches = [p for p in document.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph starting with {startswith!r}, found {len(matches)}"
        )
    return matches[0]


def set_paragraph_text(paragraph, text: str) -> None:
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def clone_paragraph(template, text: str):
    clone = deepcopy(template._p)
    for child in list(clone):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            clone.remove(child)
    run = OxmlElement("w:r")
    if template.runs and template.runs[0]._r.rPr is not None:
        run.append(deepcopy(template.runs[0]._r.rPr))
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    clone.append(run)
    return clone


def keep_with_next(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def remove_range(start_paragraph, end_paragraph) -> None:
    """Remove all body elements from start paragraph up to but not including end."""
    current = start_paragraph._p
    stop = end_paragraph._p
    parent = current.getparent()
    while current is not None and current is not stop:
        nxt = current.getnext()
        parent.remove(current)
        current = nxt
    if current is None:
        raise RuntimeError("End paragraph was not reached while removing section")


def main() -> None:
    if TARGET.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {TARGET}")
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    doc = Document(SOURCE)

    # 1) Remove the weak representative-window subsection and old Figure 5.
    old_26 = find_one(doc, "2.6 代表性窗口预测结果")
    conclusion_heading = find_one(doc, "3 结论")
    remove_range(old_26, conclusion_heading)

    # 2) Strengthen section 2.3 with seed-level consistency and explicit figure reading.
    p23_analysis = find_one(doc, "两数据集的平均结果均表明")
    set_paragraph_text(
        p23_analysis,
        "逐随机种子结果进一步表明，Abilene上本文方法相对DLinear的MSE和MAE在随机种子42—44下均为3/3次降低；"
        "GÉANT相对DLinear时，两项指标均为2/3次降低，seed43出现退化，与表2所示较大的跨种子波动一致。"
        "相对LightTS时，两个数据集的MSE和MAE均在3/3个随机种子下取得更低误差。"
        "为同时呈现平均降幅和跨种子一致性，图2给出了本文方法相对两类基线的逐种子误差降低率；"
        "圆点表示单个随机种子，黑色菱形表示由3个共同随机种子均值计算的总体降幅，正值表示本文方法误差更低。"
        "其中，相对LightTS的MSE总体降幅在Abilene和GÉANT上分别为29.89%和46.83%，"
        "且三个随机种子全部保持正向改善。本文方法仅含8339个可训练参数，明显少于当前适配下的LightTS；"
        "但参数量仅反映模型规模，不等同于训练或推理速度。外部基线结果用于说明当前统一实验协议下的相对表现，"
        "不据此推断本文方法在其他任务上普遍优于LightTS。"
    )

    # 3) Insert new Figure 2 before section 2.4.
    heading24 = find_one(doc, "2.4 八随机种子核心消融与稳定性")
    caption_template = find_one(doc, "图2 固定等权与自适应权重的八随机种子配对MSE及配对差值分布")

    pic_para = doc.add_paragraph()
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_run = pic_para.add_run()
    pic_run.add_picture(str(FIGURE), width=Inches(5.65))
    keep_with_next(pic_para)
    drawing = pic_run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("name", "图2 外部基线相对误差降低率")
        drawing[0].set(
            "descr",
            "Abilene和GÉANT上，本文方法相对DLinear与LightTS在随机种子42—44下的MSE和MAE相对降低率；"
            "圆点为单个随机种子，黑色菱形为3种子均值降幅。"
        )

    new_caption = clone_paragraph(
        caption_template,
        "图2 本文方法相对DLinear和LightTS的三随机种子误差降低率"
    )
    heading24._p.addprevious(pic_para._p)
    heading24._p.addprevious(new_caption)

    # 4) Renumber the existing figures and all local references.
    old_fig2_caption = find_one(doc, "图2 固定等权与自适应权重的八随机种子配对MSE及配对差值分布")
    set_paragraph_text(
        old_fig2_caption,
        "图3 固定等权与自适应权重的八随机种子配对MSE及配对差值分布"
    )
    p24 = find_one(doc, "同种子配对结果及配对差值分布见图2。")
    set_paragraph_text(
        p24,
        p24.text.replace("见图2", "见图3").replace("图2上排", "图3上排")
    )

    old_fig3_caption = find_one(doc, "图3 八随机种子下样本级自适应路由权重分布")
    set_paragraph_text(
        old_fig3_caption,
        "图4 八随机种子下样本级自适应路由权重分布"
    )
    p25_intro = find_one(doc, "为从总体统计视角分析样本级路由器")
    set_paragraph_text(p25_intro, p25_intro.text.replace("图3汇总", "图4汇总"))
    p25_dist = find_one(doc, "从图3可以看出")
    set_paragraph_text(p25_dist, p25_dist.text.replace("从图3", "从图4"))

    old_fig4_caption = find_one(doc, "图4 两数据集样本级路由权重变化")
    set_paragraph_text(
        old_fig4_caption,
        old_fig4_caption.text.replace("图4 ", "图5 ", 1)
    )
    p25_dyn_intro = find_one(doc, "为进一步观察权重在具体测试样本上的动态变化")
    set_paragraph_text(p25_dyn_intro, p25_dyn_intro.text.replace("图4", "图5"))
    p25_dyn = find_one(doc, "图4中三个尺度权重均随测试窗口发生变化")
    set_paragraph_text(
        p25_dyn,
        p25_dyn.text.replace("图4中", "图5中").replace("与图3", "与图4")
    )
    p25_summary = find_one(doc, "综合表2、表3和图2—图4可见")
    set_paragraph_text(
        p25_summary,
        p25_summary.text.replace("图2—图4", "图2—图5")
    )

    # 5) Remove representative-window language from the conclusion.
    conclusion = find_one(doc, "两数据集实验表明")
    set_paragraph_text(
        conclusion,
        "两数据集实验表明，样本级自适应多尺度加权可在较小参数增量下改善固定等权融合，"
        "并降低相对DLinear和LightTS的平均预测误差；跨随机种子比较和路由权重分析进一步表明，"
        "模型能够依据输入状态动态调整尺度组合。GÉANT上的跨种子波动仍较大。后续将扩展更多网络、"
        "预测长度及拓扑约束，进一步检验稳定性与泛化能力。"
    )

    # 6) Structural guardrails.
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "2.6 代表性窗口预测结果" not in full_text
    assert "代表性窗口" not in full_text
    assert "图2 本文方法相对DLinear和LightTS的三随机种子误差降低率" in full_text
    assert "图3 固定等权与自适应权重的八随机种子配对MSE及配对差值分布" in full_text
    assert "图4 八随机种子下样本级自适应路由权重分布" in full_text
    assert "图5 两数据集样本级路由权重变化" in full_text
    assert "相对LightTS时，两个数据集的MSE和MAE均在3/3个随机种子下取得更低误差" in full_text
    assert len(doc.tables) == 3
    assert len(doc.inline_shapes) == 5

    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
