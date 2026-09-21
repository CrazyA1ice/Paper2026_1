from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.21v7.docx"
TARGET = ROOT / "word_output" / "2026.09.21v13.docx"
FIGURE = (
    ROOT
    / "abilene_forecasting"
    / "paper_figures"
    / "exports"
    / "fig5_prediction_curves"
    / "fig5_representative_forecasts.png"
)


def find_one(document: Document, startswith: str):
    matches = [p for p in document.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph starting with {startswith!r}, found {len(matches)}"
        )
    return matches[0]


def set_paragraph_text(paragraph, text: str) -> None:
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def clone_paragraph(template, text: str):
    clone = deepcopy(template._p)
    for child in list(clone):
        if child.tag.endswith("}r"):
            clone.remove(child)
    run = OxmlElement("w:r")
    if template.runs and template.runs[0]._r.rPr is not None:
        run.append(deepcopy(template.runs[0]._r.rPr))
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    clone.append(run)
    return clone


def keep_with_next(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def main() -> None:
    if TARGET.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {TARGET}")
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    document = Document(SOURCE)
    conclusion_heading = find_one(document, "3 结论")
    heading_template = find_one(document, "2.5 路由行为与适用范围")
    body_template = find_one(document, "为从总体统计视角分析样本级路由器")
    caption_template = find_one(document, "图4 两数据集样本级路由权重变化")

    heading = clone_paragraph(heading_template, "2.6 代表性窗口预测结果")
    intro = clone_paragraph(
        body_template,
        "图5固定随机种子42，在两个数据集中选择本文方法窗口MAE最接近测试集中位数的窗口，并选择真实值波动接近非恒定OD通道中位水平的通道。该规则选择典型而非最佳案例，不按相对DLinear优势筛选；流量采用标准化值。",
    )

    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_run = picture_paragraph.add_run()
    picture_run.add_picture(str(FIGURE), width=Inches(5.2))
    keep_with_next(picture_paragraph)
    drawing = picture_run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("name", "图5 Abilene与GÉANT代表性预测窗口")
        drawing[0].set(
            "descr",
            "Abilene和GÉANT代表性测试窗口中真实值、DLinear与本文方法的24步标准化流量预测曲线。",
        )

    caption = clone_paragraph(
        caption_template,
        "图5 Abilene与GÉANT代表性预测窗口的真实值与预测值对比（随机种子42）",
    )
    result = clone_paragraph(
        body_template,
        "Abilene代表窗口中，本文方法MAE为0.0680，低于DLinear的0.1057；GÉANT中两者分别为0.2094和0.2087，差异很小。两个面板显示主要趋势能够被跟随，但突发峰值存在平滑现象。图5仅作直观展示，不替代多随机种子总体比较。",
    )

    anchor = conclusion_heading._p
    for element in [heading, intro, picture_paragraph._p, caption, result]:
        anchor.addprevious(element)

    conclusion = find_one(document, "在Abilene和GÉANT上的实验表明")
    set_paragraph_text(
        conclusion,
        "两数据集实验表明，样本级自适应多尺度加权可在较小参数增量下改善固定等权融合，并降低相对DLinear和LightTS的平均预测误差。代表性窗口显示模型能够跟随主要趋势，但对突发峰值仍存在平滑偏差；GÉANT跨种子波动较大。后续将扩展更多网络、预测长度及拓扑约束，进一步检验稳定性与泛化能力。",
    )

    while document.paragraphs and not document.paragraphs[-1].text.strip():
        paragraph = document.paragraphs[-1]._p
        paragraph.getparent().remove(paragraph)

    document.save(TARGET)

    check = Document(TARGET)
    full_text = "\n".join(p.text for p in check.paragraphs)
    assert "2.6 代表性窗口预测结果" in full_text
    assert "图5 Abilene与GÉANT代表性预测窗口的真实值与预测值对比" in full_text
    assert "本文方法MAE为0.0680" in full_text
    assert "对突发峰值仍存在平滑偏差" in full_text
    assert len(check.tables) == 3
    assert len(check.inline_shapes) == 5
    assert len(re.findall(r"(?m)^\[\d+\]", full_text.split("参考文献", 1)[1])) == 11
    print(TARGET)


if __name__ == "__main__":
    main()
