"""Replace the selected Figure 1 and update only its introductory explanation."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
from io import BytesIO

from lxml import etree
from PIL import Image
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'word_output/2026.09.30v54.docx'
TARGET = ROOT / 'word_output/2026.10.02v55.docx'
FIGURE = ROOT / 'paper_figures/fig1_fusion_options_20261002/B_weighted_branches.png'
REVIEW = ROOT / 'word_output/v55_review'
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}

with ZipFile(SOURCE) as z:
    original_xml = z.read('word/document.xml')
    root = etree.fromstring(original_xml)
    paragraphs = root.find('w:body',NS).findall('w:p',NS)
    para = paragraphs[19]
    nodes = para.xpath('.//w:t',namespaces=NS)
    old = ''.join(n.text or '' for n in nodes)
    assert old.startswith('图1给出模型的整体流程。')
    new = (
        '图1给出模型的整体流程。历史OD窗口经对数变换和标准化后，构造尺度1、2和4的输入序列，'
        '分别交给独立DLinear专家。各专家将趋势项和余项分别映射到24步预测，再将两部分结果相加。'
        '窗口路由网络根据原始分辨率输入的4项统计特征生成Softmax权重，对3个尺度的候选预测加权求和。'
        '图中省略样本下标n，曲边连接表示DLinear专家内部的展开；实线表示预测计算流程，'
        '虚线表示窗口特征和融合权重的传递。图中的矩阵、曲线和色块用于说明计算过程。'
        '固定等权、全局可学习权重与恒定输入路由对照采用相同的尺度与专家结构，并分别训练，'
        '用以比较等权、全局非等权与输入相关融合。'
    )
    nodes[0].text = new
    for node in nodes[1:]: node.text = ''
    drawing = paragraphs[32].find('.//w:drawing',NS)
    assert drawing is not None
    assert drawing.find('.//a:blip',NS).get('{'+NS['r']+'}embed') == 'rId6'
    image_bytes = FIGURE.read_bytes()
    with Image.open(BytesIO(image_bytes)) as image: px_width,px_height = image.size
    extent = drawing.find('.//wp:extent',NS)
    cx = int(extent.get('cx'))
    cy = round(cx*px_height/px_width)
    extent.set('cy',str(cy))
    drawing.find('.//a:xfrm/a:ext',NS).set('cy',str(cy))
    drawing.find('.//wp:docPr',NS).set('descr',
        '历史OD输入经三个独立DLinear尺度专家预测，由窗口统计量路由产生权重并加权求和；曲边区域展开专家内部的趋势与余项预测。')
    new_xml = etree.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    with ZipFile(TARGET,'w') as output:
        for item in z.infolist():
            payload = (new_xml if item.filename=='word/document.xml' else image_bytes
                       if item.filename=='word/media/image1.png' else z.read(item.filename))
            output.writestr(item,payload)

before,after = Document(SOURCE),Document(TARGET)
assert len(before.paragraphs)==len(after.paragraphs)==124
assert len(before.tables)==len(after.tables)==3
assert len(before.inline_shapes)==len(after.inline_shapes)==5
for i,(p,q) in enumerate(zip(before.paragraphs,after.paragraphs)):
    if i not in (19,32): assert p._p.xml==q._p.xml,i
for p,q in zip(before.tables,after.tables): assert p._tbl.xml==q._tbl.xml
assert [etree.tostring(x) for x in before._element.xpath('.//m:oMath')] == [etree.tostring(x) for x in after._element.xpath('.//m:oMath')]
with ZipFile(SOURCE) as a,ZipFile(TARGET) as b:
    assert a.namelist()==b.namelist()
    changed_parts = [name for name in a.namelist() if a.read(name)!=b.read(name)]
    assert changed_parts == ['word/document.xml','word/media/image1.png'],changed_parts
REVIEW.mkdir(exist_ok=True)
(REVIEW/'figure1_change_record.json').write_text(json.dumps({
    'source':str(SOURCE),'output':str(TARGET),'figure':str(FIGURE),
    'paragraph_index':19,'before':old,'after':new,
    'figure_width_mm':cx/36000,'figure_height_mm':cy/36000,
    'figure_sha256':hashlib.sha256(image_bytes).hexdigest(),
    'changed_docx_parts':changed_parts,'tables_unchanged':True,
    'equations_unchanged':True,'other_figures_unchanged':True,
},ensure_ascii=False,indent=2),encoding='utf-8')
print(TARGET)
print('Protected tables, equations, other paragraphs and Figure 2–5 verified unchanged.')
