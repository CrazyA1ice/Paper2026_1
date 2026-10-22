from pathlib import Path
import json
import sys
from collections import Counter
from lxml import etree
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'review_scratch/figureqa_dependencies'))
import pymupdf

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
def segments(paragraph):
    result, nodes = [], []
    for child in paragraph:
        if child.tag in (f'{{{M}}}oMath', f'{{{M}}}oMathPara'):
            if nodes:
                result.append(nodes)
                nodes = []
            result.append(child)
        else:
            nodes.extend(child.findall(f'.//{{{W}}}t'))
    if nodes:
        result.append(nodes)
    return result

document = Document(ROOT / 'word_output/2026.10.03v62.docx')
pdf = pymupdf.open('C:/Users/Quant/Downloads/检测报告_2026.10.03v62_.pdf')
report_chars, report_colors, report_pages = [], [], []
for i, page in enumerate(pdf):
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            for span in line['spans']:
                part = ''.join(span['text'].split())
                report_chars.append(part)
                report_colors.extend([span['color']] * len(part))
                report_pages.extend([i+1] * len(part))
report = ''.join(report_chars)
records, readable = [], []
for i, paragraph in enumerate(document.paragraphs):
    parts, math_index = [], 0
    for group in segments(paragraph._p):
        if isinstance(group, list):
            parts.append({'text': ''.join(n.text or '' for n in group), 'nodes': len(group)})
        else:
            parts.append({'math': math_index})
            math_index += 1
    normalized = ''.join(paragraph.text.split())
    match = report.find(normalized) if normalized else -1
    if match < 0 and len(normalized) > 40:
        normalized = normalized[:40]
        match = report.find(normalized)
    counts = Counter(report_colors[match:match+len(normalized)]) if match >= 0 else Counter()
    risk = '高风险' if counts[0xE53935] else '中风险' if counts[0xF9A825] else '提示风险' if counts[0x7B1FA2] else '未标记/不参与'
    record = {'i': i, 'risk': risk, 'parts': parts, 'colors': {hex(k): v for k, v in counts.items()},
              'report_page': report_pages[match] if match >= 0 else None}
    records.append(record)
    readable.append(f'\nP{i} {risk} 报告页{record["report_page"]}\n' + '\n'.join(
        f'S{j} ({part.get("nodes", 0)} nodes): {part["text"]}' if 'text' in part else f'S{j}: <M{part["math"]}>'
        for j, part in enumerate(parts)))
(QA / 'source_segments.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
(QA / 'source_segments.txt').write_text('\n'.join(readable), encoding='utf-8')
(QA / 'scan_before.txt').write_text('\n\n'.join(p.text for p in document.paragraphs[12:117] if p.text and not p.text.startswith(('注：', '图', '表')) and len(p.text)>35), encoding='utf-8')
print(json.dumps([{'i': r['i'], 'risk': r['risk'], 'page':r['report_page']} for r in records if r['i']<118 and r['risk']!='未标记/不参与'], ensure_ascii=False))
