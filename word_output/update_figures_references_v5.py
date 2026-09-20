from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.20v4.docx"
TARGET = ROOT / "word_output" / "2026.09.20v5.docx"
FIGURE_2 = (
    ROOT
    / "abilene_forecasting"
    / "paper_figures"
    / "exports"
    / "color"
    / "fig2_core_weight_pairing_v2_color.png"
)


def find_one(document: Document, startswith: str):
    matches = [p for p in document.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {startswith!r}, found {len(matches)}")
    return matches[0]


def replace_exact_run(paragraph, old: str, new: str) -> None:
    matches = [run for run in paragraph.runs if run.text == old]
    if not matches:
        raise RuntimeError(f"Expected at least one run {old!r} in {paragraph.text!r}")
    matches[0].text = new


def set_paragraph_text_like(paragraph, text: str) -> None:
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def clone_reference(template_paragraph, text: str):
    clone = deepcopy(template_paragraph._p)
    for child in list(clone):
        if child.tag.endswith("}r"):
            clone.remove(child)
    run = OxmlElement("w:r")
    if template_paragraph.runs and template_paragraph.runs[0]._r.rPr is not None:
        run.append(deepcopy(template_paragraph.runs[0]._r.rPr))
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    clone.append(run)
    return clone


def replace_figure(document: Document) -> None:
    if len(document.inline_shapes) != 3:
        raise RuntimeError(f"Expected 3 inline figures, found {len(document.inline_shapes)}")
    shape = document.inline_shapes[1]
    blip = shape._inline.xpath(".//a:blip")[0]
    embed = blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    image_part = document.part.related_parts[embed]
    image_part._blob = FIGURE_2.read_bytes()
    with Image.open(FIGURE_2) as image:
        px_w, px_h = image.size
    shape.width = Inches(6.30)
    shape.height = Inches(6.30 * px_h / px_w)


def update_citations(document: Document) -> None:
    intro_network = find_one(document, "随着云计算")
    replace_exact_run(intro_network, "[6]", "[10]")

    intro_models = find_one(document, "网络流量在不同时间尺度上")
    replace_exact_run(intro_models, "[7]", "[11]")
    replace_exact_run(intro_models, "[8]", "[6]")

    data_para = find_one(document, "实验选用")
    replace_exact_run(data_para, "CESNET TS-Zoo", "CESNET TS-Zoo[7]")
    replace_exact_run(data_para, "Abilene", "Abilene[8]")
    replace_exact_run(data_para, "GÉANT", "GÉANT[9]")


def reorder_references(document: Document) -> None:
    heading = find_one(document, "参考文献")
    body = heading._p.getparent()
    paragraphs = document.paragraphs
    heading_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is heading._p)
    refs = {}
    for paragraph in paragraphs[heading_index + 1 :]:
        match = re.match(r"^\[(\d+)\]", paragraph.text)
        if match:
            refs[int(match.group(1))] = paragraph
    if set(refs) != set(range(1, 9)):
        raise RuntimeError(f"Unexpected reference numbers: {sorted(refs)}")

    existing_order = [1, 2, 3, 4, 5, 8, 6, 7]
    new_number = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 8: 6, 6: 10, 7: 11}
    elements = []
    for old in existing_order[:6]:
        paragraph = refs[old]
        set_paragraph_text_like(
            paragraph,
            re.sub(r"^\[\d+\]", f"[{new_number[old]}]", paragraph.text, count=1),
        )
        elements.append(paragraph._p)

    template = refs[8]
    additions = [
        "[7] KUREŠ M, KOUMAR J, HYNEK K. CESNET TS-Zoo: a library for reproducible analysis of network traffic time series[C]//2025 21st International Conference on Network and Service Management. 2025. DOI: 10.23919/CNSM67658.2025.11297513.",
        "[8] ZHANG Y, ROUGHAN M, DUFFIELD N, et al. Fast accurate computation of large-scale IP traffic matrices from link loads[J]. ACM SIGMETRICS Performance Evaluation Review, 2003, 31(1): 206-217. DOI: 10.1145/885651.781053.",
        "[9] UHLIG S, QUOITIN B, LEPROPRE J, et al. Providing public intradomain traffic matrices to the research community[J]. ACM SIGCOMM Computer Communication Review, 2006, 36(1): 83-86. DOI: 10.1145/1111322.1111341.",
    ]
    elements.extend(clone_reference(template, text) for text in additions)

    for old in existing_order[6:]:
        paragraph = refs[old]
        set_paragraph_text_like(
            paragraph,
            re.sub(r"^\[\d+\]", f"[{new_number[old]}]", paragraph.text, count=1),
        )
        elements.append(paragraph._p)

    for paragraph in refs.values():
        body.remove(paragraph._p)
    anchor = heading._p
    for element in elements:
        anchor.addnext(element)
        anchor = element


def validate(document: Document) -> None:
    text = "\n".join(p.text for p in document.paragraphs)
    body_text, reference_text = text.split("\n参考文献\n", 1)
    citations = {int(x) for x in re.findall(r"\[(\d+)\]", body_text)}
    references = {int(x) for x in re.findall(r"(?m)^\[(\d+)\]", reference_text)}
    assert citations == references == set(range(1, 12)), (citations, references)
    assert "CESNET TS-Zoo[7]中的Abilene[8]和GÉANT[9]" in body_text
    assert "韦烜等[10]" in body_text
    assert "CycleLLH从周期性整合角度改进网络流量预测[11]" in body_text
    assert "LightTS采用连续采样和间隔采样配合多层感知机结构，以较低计算开销建模多变量时间序列[6]" in body_text
    reference_lines = [line for line in reference_text.splitlines() if re.match(r"^\[\d+\]", line)]
    assert all(not re.search(r"[\u4e00-\u9fff]", line) for line in reference_lines[:9])
    assert all(re.search(r"[\u4e00-\u9fff]", line) for line in reference_lines[9:])
    assert len(document.tables) == 3
    assert len(document.inline_shapes) == 3


def main() -> None:
    if TARGET.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {TARGET}")
    document = Document(SOURCE)
    replace_figure(document)
    update_citations(document)
    caption = find_one(document, "图2 固定等权与自适应权重")
    set_paragraph_text_like(
        caption,
        "图2 固定等权与自适应权重的八随机种子配对MSE（橙色方形与蓝色圆形分别表示固定等权和自适应权重；绿色连线表示误差降低，红色连线表示误差升高；竖向虚线和点线表示两模型8种子均值）",
    )
    reorder_references(document)
    validate(document)
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
