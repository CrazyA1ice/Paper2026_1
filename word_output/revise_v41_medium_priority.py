"""Improve three figure/table reading transitions in the v40 manuscript."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.27v40.docx"
TARGET = ROOT / "word_output/2026.09.27v41.docx"

document = Document(SOURCE)
paragraphs = document.paragraphs
assert len(paragraphs) == 114 and len(document.inline_shapes) == 5
assert paragraphs[20]._p.xpath(".//w:drawing")
assert paragraphs[21].text == "图1 自适应多尺度网络流量预测方法框架"
assert paragraphs[32].text.startswith("经过上述处理后")

# Let the task definition and preprocessing fill the large gap before Fig. 1.
# The model overview at paragraph 19 still introduces the figure before it appears.
paragraphs[32]._p.addnext(paragraphs[20]._p)
paragraphs[20]._p.addnext(paragraphs[21]._p)
# Keep the next subsection's opening sentence with Eq. (5).
paragraphs[34].paragraph_format.keep_with_next = True
paragraphs[34].paragraph_format.keep_together = True

# Give the reader the table before the numerical interpretation. The DLinear
# and equal-weight percentages already appear with Fig. 3, so avoid repeating them.
assert paragraphs[84].text.startswith("表3汇总随机种子42～49的8次核心消融。")
paragraphs[84].runs[0].text = (
    "表3汇总随机种子42～49的8次核心消融，列出五种模型的参数量、MSE和MAE。"
)
# The closing signpost repeats the next subsection and otherwise leaves one
# isolated line above Fig. 2 after pagination.
assert paragraphs[79].runs[0].text.endswith(
    "图2因此只承担外部结构基线比较，DLinear和固定等权模型的内部比较放在表3和图3中。"
)
paragraphs[79].runs[0].text = paragraphs[79].runs[0].text.removesuffix(
    "图2因此只承担外部结构基线比较，DLinear和固定等权模型的内部比较放在表3和图3中。"
)
detail = (
    "全局可学习权重与恒定输入路由均对所有窗口使用一组可学习权重，后者与本文方法同为8339个标称参数。"
    "相对这两种对照，本文方法在Abilene上的逐种子平均MSE改进均约1.80%，在GÉANT上分别为4.31%和4.97%，"
    "各有6/8和7/8个种子胜出。相对恒定输入路由的95%置信区间分别为[0.49%,3.11%]和[−1.62%,11.56%]，"
    "相对全局权重则为[0.01%,3.58%]和[−2.51%,11.13%]；GÉANT两项区间均跨零。"
)
new_paragraph = paragraphs[86].insert_paragraph_before(detail)
new_paragraph.style = paragraphs[84].style
# Avoid splitting the Fig. 3 explanation across two pages.
paragraphs[86].paragraph_format.keep_together = True

# Finish the model-behaviour evidence before discussing inference cost.
assert paragraphs[93].text.startswith("模型规模之外，还需考察前向开销。")
paragraphs[93].runs[0].text = paragraphs[93].runs[0].text.replace(
    "模型规模之外，还需考察前向开销。",
    "除预测误差和路由行为外，还需考察前向开销。",
    1,
)
paragraphs[99]._p.addnext(paragraphs[93]._p)
# The moved cost paragraph shifts the reference block; keep entries intact.
for paragraph in paragraphs[103:114]:
    paragraph.paragraph_format.keep_together = True
paragraphs[102].paragraph_format.page_break_before = True

document.save(TARGET)
check = Document(TARGET)
assert len(check.inline_shapes) == 5 and len(check.tables) == 3
def position(prefix):
    matches = [i for i, p in enumerate(check.paragraphs) if p.text.startswith(prefix)]
    assert len(matches) == 1, prefix
    return matches[0]

assert position("图1 自适应") > position("经过上述处理后")
assert position("除预测误差和路由行为外") > position("进一步以输入窗口")
print(TARGET)
