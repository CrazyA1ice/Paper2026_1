from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"word_output"/"2026.09.21v16.docx"
TARGET=ROOT/"word_output"/"2026.09.21v17.docx"

doc=Document(SOURCE)

def find(prefix):
    xs=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(xs)!=1:
        raise RuntimeError(f"{prefix!r}: {len(xs)} matches")
    return xs[0]

def set_text(p,text):
    first=p.runs[0] if p.runs else p.add_run()
    first.text=text
    for r in p.runs[1:]:
        r.text=""

def repl(prefix,text):
    p=find(prefix)
    set_text(p,text)
    return p

def clone_after(template,text):
    new_p=deepcopy(template._p)
    for child in list(new_p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            new_p.remove(child)
    run=OxmlElement("w:r")
    if template.runs and template.runs[0]._r.rPr is not None:
        run.append(deepcopy(template.runs[0]._r.rPr))
    node=OxmlElement("w:t")
    node.text=text
    run.append(node)
    new_p.append(run)
    template._p.addnext(new_p)
    return new_p

# 1) Abstracts: restore journal-like completeness while remaining within the 300-400 Chinese-character requirement.
repl(
"摘要：",
"摘要：针对骨干网络源—目的节点对流量具有明显时变性和多时间尺度特征、固定尺度模型难以根据当前状态调整信息利用方式的问题，提出一种样本级自适应多尺度预测方法。采用尺度1、2和4的非重叠平均池化构造多分辨率序列，并以DLinear作为各尺度预测专家；进一步提取窗口均值、标准差、末时刻均值和平均绝对一阶差分，经轻量路由网络生成Softmax尺度权重，实现多尺度预测结果的动态融合。在Abilene和GÉANT两个骨干网流量数据集上进行实验，并与DLinear、LightTS及固定等权多尺度模型比较。共同3个随机种子下，所提方法的MSE相对DLinear分别降低4.06%和4.43%，相对LightTS分别降低29.89%和46.83%；8个随机种子的核心消融中，相对固定等权融合分别降低1.67%和4.68%。结果表明，该方法能够以较小参数增量有效利用不同时间尺度的互补信息，提高多变量网络流量预测精度。"
)
repl(
"Abstract:",
"Abstract: Backbone origin-destination (OD) traffic exhibits pronounced temporal variation and multiscale characteristics, while fixed-scale models cannot adapt their information use to the current traffic state. A sample-wise adaptive multiscale forecasting method is therefore proposed. Non-overlapping average pooling at scales 1, 2, and 4 is used to construct multiresolution sequences, and DLinear is employed as the forecasting expert at each scale. Four window statistics, including the mean, standard deviation, latest-step mean, and mean absolute first difference, are further fed into a lightweight routing network to generate Softmax scale weights for dynamic fusion. Experiments are conducted on the Abilene and GÉANT backbone traffic datasets against DLinear, LightTS, and fixed equal-weight multiscale fusion. Under three common random seeds, the proposed method reduces MSE by 4.06% and 4.43% over DLinear and by 29.89% and 46.83% over LightTS on the two datasets, respectively. In the eight-seed core ablation, MSE is reduced by 1.67% and 4.68% over fixed equal weighting. These results indicate that sample-wise scale weighting can effectively exploit complementary temporal information at different resolutions with only a small parameter increase."
)

# 2) Introduction: restore background and related work to the density seen in published journal samples.
p12=repl(
"随着云计算、视频业务和数据中心互联的发展",
"随着云计算、视频业务、数据中心互联及智能网络应用的持续发展，骨干网络承载的业务类型与流量规模不断增加，OD流量在小时级尺度上表现出明显的时变性、突发性和周期性。准确预测未来一段时间的OD流量，可为容量规划、拥塞预警、资源调度和流量工程提供先验依据，因此网络流量预测一直是网络管理与智能运维中的重要研究问题。"
)
p13=repl(
"不同时间尺度包含互补流量信息",
"网络流量预测方法已由传统统计模型逐步发展到机器学习和深度学习方法。Lohrasbinasab等[1]系统梳理了网络流量预测从统计方法向机器学习方法的演进，指出数据驱动模型在复杂流量模式建模方面具有更强的表达能力。Qin等[2]利用图神经网络联合建模通信网络的空间关联与时间依赖；韦烜等[10]围绕大型IP网络流量矩阵的分析与预测开展研究，说明OD维度间关联和时间演化对预测性能具有重要影响。针对网络流量中的周期特征，唐文杰等[11]提出CycleLLH模型，通过周期性信息整合增强流量变化规律的建模能力。"
)
new_intro=(
"除拓扑关联与周期规律外，多时间尺度建模也是提升预测性能的重要路径。Zhou等[3]提出基于深度回声状态网络的多尺度网络流量预测方法，通过不同时间分辨率捕获互补变化模式。DLinear[4]采用时间序列分解与线性映射构建简洁预测结构，表明在长期预测任务中较低复杂度模型同样具有较强竞争力；TimeMixer[5]通过可分解的多尺度混合机制联合利用不同分辨率信息；LightTS[6]则以采样导向的MLP结构实现快速多变量时间序列预测。这些研究表明，多尺度信息与轻量预测结构具有良好的结合潜力，但多数方法仍采用预设的尺度组合或统一的融合方式。"
)
clone_after(p13,new_intro)
repl(
"现有方法多侧重拓扑关联",
"在实际骨干网络中，不同输入窗口的流量水平、波动幅度和短期变化状态并不相同，固定尺度或固定融合比例难以始终匹配当前预测样本。基于此，提出一种样本级自适应多尺度网络流量预测方法：首先在尺度1、2和4上构造多分辨率序列，以DLinear作为尺度专家；随后利用窗口统计特征驱动轻量路由器，为每个样本生成动态尺度权重并完成预测融合。实验在Abilene和GÉANT两个公开骨干网数据集上展开，并通过外部基线比较、固定等权消融、8随机种子重复实验以及路由权重分析验证方法的预测效果与自适应融合机制。"
)

# 3) Method: add design motivation and enough explanation without returning to audit-style prose.
repl(
"本文方法由多尺度序列构造",
"本文方法由数据预处理、多尺度序列构造、DLinear尺度专家和样本级自适应融合组成，整体流程如图1所示。核心思想是先在不同时间分辨率上形成具有互补信息的候选预测，再由轻量路由器根据当前输入窗口的统计状态动态调整各尺度贡献，使模型同时利用细粒度局部变化与较粗粒度趋势信息。"
)
repl(
"其中μctr和σctr仅由训练集计算",
"其中μctr和σctr仅由训练集计算，并统一用于验证集和测试集，以避免测试信息参与归一化过程。对原始非负流量采用log(1+x)变换，可压缩不同OD流量之间较大的数值动态范围；随后按变量标准化，使各维流量在相近尺度上参与模型训练。数据按时间顺序以6∶2∶2划分，各集合内部独立构造滑动窗口。"
)
repl(
"三个尺度的输入长度依次为96、48和24",
"三个尺度的输入长度依次为96、48和24，分别对应原始分辨率、2点平均和4点平均。尺度1完整保留短时波动，尺度2在平滑局部扰动的同时保持较高时间分辨率，尺度4进一步突出较慢变化趋势。采用非重叠平均池化无需引入额外可训练参数，也便于不同尺度专家保持统一的输入构造方式。"
)
repl(
"同一尺度内各变量共享时间映射参数",
"同一尺度内各变量共享时间映射参数，不同尺度使用独立专家参数。DLinear先通过移动平均将输入分为趋势项与余项，再分别进行线性映射，这种结构参数规模较小，适合作为多尺度框架中的基础专家；不同尺度专家独立学习，可避免将不同时间分辨率下的变化模式强制映射到同一组参数中。"
)
repl(
"固定等权无法反映不同输入窗口的状态差异",
"固定等权融合默认所有预测窗口始终采用相同的尺度比例，难以反映流量状态随时间变化的特点。为此，从原始输入Xn提取四维统计向量zn=[μn,σn,ℓn,dn]，其中"
)
repl(
"μn、σn、ℓn和dn分别描述窗口总体水平",
"μn、σn、ℓn和dn分别描述窗口总体水平、离散程度、最近状态和平均变化强度，能够以较低维度概括当前输入的整体状态。四维向量经隐藏维数为16的两层感知器映射为三个尺度分数，在不引入复杂特征提取网络的情况下，为尺度选择提供样本相关信息："
)
repl(
"每个样本生成一组共享于全部OD变量的尺度权重",
"每个样本生成一组共享于全部OD变量的尺度权重，使同一时间窗口内的各OD流量采用一致的时间尺度组合，并避免路由参数量随变量数显著增加。为检验动态权重的作用，固定等权消融将αn,s统一设为1/3，其尺度集合、专家结构及训练过程均保持不变。"
)
repl(
"模型采用Adam优化并按验证集MSE选择参数",
"模型采用Adam优化，并根据验证集MSE选择最优参数。DLinear、固定等权多尺度和本文模型分别含4656、8208和8339个可训练参数，其中自适应路由器在固定等权多尺度模型基础上仅增加131个参数。由此，模型性能提升主要通过尺度组织与动态融合获得，而无需显著扩大参数规模。"
)

# 4) Experiments: restore context and result interpretation, keeping the main evidence prominent.
repl(
"实验采用CESNET TS-Zoo",
"实验采用CESNET TS-Zoo[7]中的Abilene[8]和GÉANT[9] 1 h聚合流量矩阵。Abilene包含较长时间序列并保留144维有效OD流量，GÉANT去除恒定通道后保留524维，变量规模明显更高。两个数据集在网络规模和流量分布上具有差异，可用于从不同骨干网场景考察模型的适用性。两者均按时间顺序以6∶2∶2划分，输入长度L=96、预测长度H=24，并采用第1.1节的预处理方式，具体规模见表1。"
)
repl(
"所有模型采用统一数据划分和训练接口",
"所有模型采用统一的数据划分、输入长度和评价流程，批量大小为16，并按验证集MSE选择最优模型。DLinear、固定等权多尺度和本文模型最多训练30轮，LightTS最多训练50轮。外部基线比较使用共同随机种子42—44，以保证各模型重复次数一致；对核心模型进一步扩展到42—49共8个随机种子，用于观察多尺度机制在重复实验中的总体表现。"
)
repl(
"对比模型包括DLinear、LightTS",
"对比模型包括单尺度DLinear、轻量多变量预测模型LightTS、固定等权多尺度模型和本文方法。DLinear用于衡量多尺度结构相对单尺度线性预测的增益，LightTS作为外部轻量时间序列基线；固定等权模型与本文方法采用完全相同的尺度集合{1,2,4}和DLinear专家，仅将三个尺度权重固定为1/3，从而直接检验样本级自适应加权的贡献。"
)
repl(
"上述指标均在标准化空间计算",
"上述指标均在标准化空间计算，数值越小表示预测误差越低。其中MSE对较大预测偏差更敏感，MAE反映平均绝对偏差，RMSE与原始误差量纲形式一致。正文主要结合MSE和MAE分析模型差异；表2报告3个共同随机种子的均值±标准差，表3报告8个随机种子的均值±标准差。"
)
repl(
"表2给出了随机种子42—44的外部基线结果",
"表2给出了随机种子42—44的外部基线结果。本文方法在Abilene和GÉANT上的MSE分别为0.186和0.434，均低于DLinear和LightTS；相对DLinear分别降低4.06%和4.43%，相对LightTS分别降低29.89%和46.83%。MAE同样保持较低水平，说明自适应多尺度结构在两个不同规模的骨干网数据集上均能提升平均预测精度。"
)
repl(
"逐随机种子结果见图2",
"逐随机种子结果见图2。Abilene上，本文方法相对DLinear和LightTS的MSE、MAE均在3个随机种子下保持较低误差；GÉANT上，相对LightTS同样表现出一致优势，相对DLinear的平均误差仍更低。结合表2与图2可以看出，多尺度动态融合不仅改善了单尺度DLinear的预测结果，在统一实验设置下相较LightTS也具有明显的误差优势。"
)
repl(
"表3使用随机种子42—49比较DLinear",
"为进一步验证核心结构，表3将DLinear、固定等权多尺度和本文方法扩展到随机种子42—49共8次重复实验。相较DLinear，本文方法在Abilene和GÉANT上的平均MSE分别降低2.77%和2.37%；相较固定等权多尺度分别降低1.67%和4.68%。固定等权模型已引入相同的多尺度专家，因此该结果进一步说明，在多尺度表征基础上根据输入样本动态调整尺度权重能够继续带来预测增益。"
)
repl(
"同种子配对结果见图3",
"同种子配对结果见图3。Abilene中本文方法在6/8个随机种子上取得更低MSE，GÉANT为7/8；两个数据集的平均ΔMSE均为正，与表3的均值结果一致。相较单纯比较总体均值，配对结果进一步表明自适应加权的收益在多数重复实验中能够得到体现。"
)
repl(
"图4汇总了随机种子42—49下全部测试窗口",
"为观察路由器对不同时间尺度的实际分配情况，图4汇总随机种子42—49下全部测试窗口的三尺度权重分布。固定等权情况下三个尺度权重均为1/3，而自适应模型允许权重随输入样本变化，因此其分布可以反映模型学习到的总体尺度偏好。"
)
repl(
"图4显示，两数据集的权重均未固定在1/3",
"图4显示，两数据集的路由权重均形成了明显的非等权分布。Abilene的α1、α2和α4平均为0.257、0.319和0.424，GÉANT分别为0.283、0.318和0.399，两个数据集均表现为尺度4权重最高、尺度1最低。说明在当前24步预测任务中，较粗时间尺度提供的平滑趋势信息具有较高贡献，同时尺度1和尺度2仍保留稳定权重，从而形成多时间尺度的联合预测。"
)
repl(
"综合表2、表3和图2—图4",
"综合表2、表3和图2—图4可以看出，本文方法在两个骨干网数据集上均取得更低的平均预测误差；固定等权消融进一步验证了动态尺度权重的有效性，路由权重分布则从模型内部行为角度表明三个尺度并非简单平均，而是形成了与输入样本相关的差异化贡献。"
)

# 5) Conclusion: short, affirmative, and journal-like.
repl(
"两数据集实验表明",
"两数据集实验表明，样本级自适应多尺度加权能够有效融合不同时间分辨率的信息，在较小参数增量下提升多变量网络流量预测精度；相对DLinear、LightTS和固定等权多尺度模型均取得更低的平均误差。路由权重结果进一步验证了动态尺度分配机制的作用。后续将面向更多网络场景和预测长度扩展验证。"
)

# Guardrails.
full="\n".join(p.text for p in doc.paragraphs)
abstract=find("摘要：").text.replace("摘要：","")
conclusion=find("两数据集实验表明").text
assert 350 <= len(abstract) <= 400, len(abstract)
assert len(conclusion) <= 150, len(conclusion)
assert len(doc.tables)==3
assert len(doc.inline_shapes)==4
for required in [
    "图1 自适应多尺度网络流量预测方法框架",
    "图2 Abilene与GÉANT上三随机种子外部基线配对结果",
    "图3 固定等权与自适应权重的八随机种子配对MSE及配对差值分布",
    "图4 八随机种子下样本级自适应路由权重分布",
]:
    assert required in full
assert "Lohrasbinasab等[1]" in full
assert "TimeMixer[5]" in full
assert "LightTS[6]" in full

doc.save(TARGET)
print(TARGET)
print("abstract chars",len(abstract))
print("conclusion chars",len(conclusion))
