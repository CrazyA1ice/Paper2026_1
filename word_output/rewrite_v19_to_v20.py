from __future__ import annotations
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.21v19.docx"
TARGET = ROOT / "word_output" / "2026.09.21v20.docx"

doc = Document(SOURCE)

def find(prefix):
    xs = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(xs) != 1:
        raise RuntimeError(f"{prefix!r}: {len(xs)} matches")
    return xs[0]

def set_text(p, text):
    first = p.runs[0] if p.runs else p.add_run()
    first.text = text
    for r in p.runs[1:]:
        r.text = ""

def repl(prefix, text):
    p = find(prefix)
    set_text(p, text)

replacements = {
"摘要：": "摘要：骨干网络源—目的节点对流量在时间上并非稳定变化，短时波动、突发变化和较慢趋势往往同时存在。固定尺度模型使用单一时间分辨率，难以随当前流量状态调整不同时间尺度信息的利用方式。为此，提出一种样本级自适应多尺度预测方法。采用尺度1、2和4的非重叠平均池化得到多分辨率序列，并分别输入DLinear专家；同时提取输入窗口的均值、标准差、末时刻均值和平均绝对一阶差分，通过轻量路由网络计算Softmax尺度权重，完成3个尺度预测结果的融合。在Abilene和GÉANT两个骨干网流量数据集上进行实验。3个共同随机种子下，所提方法的MSE相对DLinear分别降低4.06%和4.43%，相对LightTS分别降低29.89%和46.83%；8个随机种子的消融实验中，相对固定等权融合分别降低1.67%和4.68%。在仅增加少量参数的情况下，不同时间尺度的信息可以按照输入样本重新分配权重，两个数据集上的平均预测误差均得到改善。",

"Abstract:": "Abstract: Backbone origin-destination (OD) traffic does not evolve at a single temporal pace. Short-term fluctuations, abrupt changes, and slower trends may appear in the same traffic sequence, whereas a fixed-scale model always uses one temporal resolution. A sample-wise adaptive multiscale forecasting method is therefore proposed. Non-overlapping average pooling at scales 1, 2, and 4 is applied to obtain multiresolution sequences, which are processed by separate DLinear experts. The mean, standard deviation, latest-step mean, and mean absolute first difference of each input window are also extracted and fed into a lightweight routing network. Softmax weights produced by the router are used to combine the three forecasts. Experiments are carried out on the Abilene and GÉANT backbone traffic datasets. With three common random seeds, the proposed method reduces MSE by 4.06% and 4.43% compared with DLinear and by 29.89% and 46.83% compared with LightTS. In the eight-seed ablation experiment, the MSE reductions over fixed equal weighting are 1.67% and 4.68%, respectively. The results on both datasets show that the relative contribution of different temporal scales can be adjusted for each input sample with only a small increase in model parameters.",

"随着云计算、视频业务、数据中心互联及智能网络应用的发展": "随着云计算、视频业务、数据中心互联及智能网络应用的发展，骨干网络承载的数据量不断增加，业务流量的变化也更加复杂。源—目的节点对（origin-destination，OD）流量既会受到日常周期的影响，也可能在较短时间内出现明显波动。对于网络管理而言，若能提前获得后一时段的OD流量估计，可为容量规划、拥塞预警、资源调度和流量工程提供数据依据。因此，OD流量预测不仅要求模型刻画长期变化，还要保留对局部波动的响应能力。",

"网络流量预测的研究经历了从统计方法到机器学习、深度学习方法的发展": "早期网络流量预测主要采用统计模型，随后逐渐转向机器学习和深度学习方法。Lohrasbinasab等[1]对这一发展过程进行了综述，指出数据驱动方法在复杂流量模式下具有更强的建模能力。Qin等[2]把图神经网络用于通信网络流量预测，将空间关联和时间依赖同时纳入模型；韦烜等[10]针对大型IP网络流量矩阵开展分析与预测，强调了OD变量之间关联关系的影响。对于周期性较明显的流量，唐文杰等[11]提出CycleLLH模型，将周期信息纳入预测过程。已有研究分别从变量关联、时间依赖和周期变化等方面改进了预测模型。",

"网络流量在不同时间尺度上呈现的特征并不完全相同": "除了上述特征外，网络流量还具有明显的多时间尺度特性。原始分辨率序列能够反映较细的波动，但其中也包含较多局部扰动；时间尺度变粗后，短时波动被部分平滑，整体趋势会更加突出。Zhou等[3]采用深度回声状态网络进行多尺度网络流量预测。DLinear[4]通过时间序列分解和线性映射完成预测，模型结构相对简单；TimeMixer[5]利用多尺度混合机制处理不同分辨率序列；LightTS[6]则通过采样导向的MLP结构进行快速多变量时间序列预测。由此，多尺度信息与轻量预测结构之间具有进一步结合的空间。",

"实际骨干网络的流量状态会不断变化": "从实际流量序列看，不同输入窗口的波动程度并不一致。有些时段以局部变化为主，有些时段则表现出更明显的平滑趋势，此时3个时间尺度的作用很难长期保持相同。固定尺度或固定融合比例不能随这种状态变化调整。本文据此在尺度1、2和4上构造多分辨率序列，由DLinear分别完成预测，再根据当前输入窗口的统计特征生成样本级尺度权重。实验采用Abilene和GÉANT两个公开骨干网数据集，并设置基线比较、固定等权消融、8个随机种子重复实验和路由权重分析。",

"本文方法的整体结构如图1所示": "图1给出了本文方法的整体流程。原始OD流量首先完成对数变换和标准化处理，随后在尺度1、2和4上形成3组输入序列。每组序列对应一个DLinear专家，各自输出长度为H的预测结果。与此同时，轻量路由网络根据当前输入窗口的统计量给出3个尺度权重，最终预测由3个专家输出加权得到。",

"其中μctr和σctr由训练集计算": "其中μctr和σctr只在训练集上计算，验证集和测试集直接使用训练阶段得到的参数。由于不同OD变量的流量量级差异较大，原始非负流量先进行log(1+x)变换，再按变量标准化。数据仍按时间顺序以6∶2∶2划分为训练集、验证集和测试集，各集合内部独立生成滑动窗口。",

"本文选取尺度1、2和4构造多分辨率序列": "尺度集合取S={1,2,4}。尺度1对应原始96步输入，不做时间压缩；尺度2经过2点非重叠平均后得到48步序列；尺度4得到24步序列。这样处理后，尺度1保留更多短时变化，尺度2和尺度4逐渐削弱局部扰动。由于采用平均池化，多尺度构造本身不增加可训练参数。",

"各时间尺度均采用独立的DLinear专家": "3个尺度分别配置DLinear专家[4]。对每个尺度的输入，先用长度q=25的中心移动平均得到趋势项，再由原序列与趋势项之差得到余项，二者沿时间维分别进行线性映射。不同尺度使用独立参数，同一尺度内的OD变量共享时间映射参数。这样设置后，每个专家只需学习本尺度下的时间变化，而不必共用一套时间映射。",

"得到3个尺度的预测结果后，还需要确定各尺度在当前样本中的占比": "3个DLinear专家输出后，融合权重由当前输入样本决定。若直接取1/3等权，流量平稳时和波动明显时采用的尺度比例完全相同。本文从原始输入Xn提取四维统计向量zn=[μn,σn,ℓn,dn]，其中",

"μn、σn、ℓn和dn分别表示输入窗口的总体水平": "μn、σn、ℓn和dn分别表示输入窗口的总体水平、离散程度、最近状态和平均变化强度。4个统计量输入隐藏维数为16的两层感知器，输出3个尺度分数。该路由器只处理低维统计量，没有再叠加额外的时序编码模块。",

"经Softmax归一化后，每个样本得到一组和为1的尺度权重": "3个尺度分数经过Softmax后转化为非负且和为1的权重。每个输入样本对应一组权重，并在该窗口内由全部OD变量共享。固定等权消融模型直接令3个权重均为1/3，除此之外，尺度集合、DLinear专家和训练过程与本文方法一致。",

"模型采用Adam进行优化，并根据验证集MSE选择参数": "训练采用Adam优化，以验证集MSE选择模型。DLinear、固定等权多尺度模型和本文方法的可训练参数量分别为4656、8208和8339。与固定等权模型相比，加入样本级路由后增加131个参数，参数规模基本保持在同一量级。",

"实验选用CESNET TS-Zoo[7]中的Abilene[8]和GÉANT[9] 1 h聚合流量矩阵": "实验数据来自CESNET TS-Zoo[7]，选取其中的Abilene[8]和GÉANT[9] 1 h聚合流量矩阵。Abilene保留144维有效OD流量；GÉANT去除恒定通道后保留524维，OD变量数量更大。两个数据集分别按时间顺序以6∶2∶2划分，输入长度L=96，预测长度H=24，数据预处理与第1.1节相同。表1给出了具体数据规模和预测任务设置。",

"各模型使用相同的数据划分、输入长度和评价指标": "所有对比模型使用相同的数据划分、输入长度和评价指标，批量大小均为16。DLinear、固定等权多尺度模型和本文方法最多训练30轮，LightTS最多训练50轮，各模型均以验证集MSE选择结果。外部基线比较使用随机种子42—44；核心模型另外在42—49共8个随机种子下重复训练。",

"对比模型包括DLinear、LightTS和固定等权多尺度模型": "实验设置3类对比模型：单尺度DLinear、LightTS以及固定等权多尺度模型。前两者用于观察本文方法与现有轻量预测结构的差异。固定等权多尺度模型使用尺度集合{1,2,4}和相同的DLinear专家，但不经过路由网络，3个尺度始终按1/3融合。",

"MSE、MAE和RMSE均在标准化空间中计算": "评价指标为MSE、MAE和RMSE，均在标准化空间中计算。MSE按式（14）计算，MAE和RMSE分别见式（15）和式（16）。正文主要讨论MSE和MAE；表2统计随机种子42—44的均值±标准差，表3统计随机种子42—49的均值±标准差。",

"表2列出了随机种子42—44下的预测结果": "由表2可知，本文方法在Abilene上的MSE为0.186，在GÉANT上为0.434。相对于DLinear，两者分别降低4.06%和4.43%；相对于LightTS，降幅分别为29.89%和46.83%。MAE的变化方向与MSE一致。两个数据集的结果都显示，加入多尺度输入和样本级融合后，平均预测误差低于对应基线。",

"图2给出了3个共同随机种子的逐次结果": "图2进一步列出了3个随机种子的单次结果。Abilene的3次实验中，本文方法的MSE和MAE均低于DLinear和LightTS。GÉANT中，本文方法相对LightTS仍保持较低误差；与DLinear相比有个别随机种子波动，但3次结果取平均后仍低于DLinear。",

"表3将DLinear、固定等权多尺度模型和本文方法的重复次数扩展到8个随机种子": "表3把核心模型的重复实验扩展到8个随机种子。相对DLinear，本文方法在Abilene和GÉANT上的平均MSE分别降低2.77%和2.37%。固定等权模型已经包含相同的3个DLinear尺度专家，因此更能直接反映路由权重的影响；在该模型基础上，本文方法的MSE继续降低1.67%和4.68%。",

"图3比较了固定等权模型与本文方法在相同随机种子下的MSE": "图3按相同随机种子比较固定等权模型与本文方法。Abilene的8个随机种子中有6个随机种子取得更低MSE，GÉANT中为7个。两个数据集的平均ΔMSE均为正，与表3给出的均值差异相符。",

"图4给出了随机种子42—49下全部测试窗口的尺度权重分布": "图4汇总了随机种子42—49下全部测试窗口的尺度权重。与固定的1/3权重不同，路由网络输出会随输入窗口变化，因此这里统计的是各尺度在测试集上的整体分布。",

"从图4可以看出，两个数据集的3个尺度并未得到相同权重": "Abilene中α1、α2和α4的平均权重分别为0.257、0.319和0.424；GÉANT中分别为0.283、0.318和0.399。两个数据集都表现为尺度4的平均权重最高，尺度1最低。对于本文设置的24步预测任务，较粗尺度的趋势信息占比较大，但另外两个尺度并未被路由器舍弃。",

"结合表2、表3和图2—图4": "表2和图2反映了本文方法与外部基线的误差差异，表3与图3则比较了固定等权和自适应加权。两组结果的变化方向基本一致。结合图4可见，3个尺度在测试样本中并非保持相同占比，路由器会改变各尺度参与最终预测的程度。",

"本文研究了骨干网络OD流量的多尺度预测问题": "本文针对骨干网络OD流量的多时间尺度变化，构建了样本级自适应多尺度预测方法。尺度1、2和4的序列分别由DLinear专家预测，轻量路由网络根据输入窗口统计量给出融合权重。Abilene和GÉANT上的结果显示，相比DLinear、LightTS和固定等权多尺度模型，本文方法的平均预测误差更低。路由权重在3个尺度间呈现不同分布，说明固定等权并非唯一的融合方式。整个路由模块仅增加131个参数，可在较小模型增量下实现多尺度预测结果的动态融合。"
}

for prefix, text in replacements.items():
    repl(prefix, text)

full = "\n".join(p.text for p in doc.paragraphs)

# 核心概念与实验数据保护
must_keep = [
    "源—目的节点对", "DLinear", "Softmax", "Abilene", "GÉANT", "LightTS",
    "CycleLLH", "TimeMixer", "CESNET TS-Zoo",
    "4.06%", "4.43%", "29.89%", "46.83%", "1.67%", "4.68%",
    "0.186", "0.434", "2.77%", "2.37%",
    "0.257", "0.319", "0.424", "0.283", "0.318", "0.399",
    "4656", "8208", "8339", "131",
    "L=96", "H=24", "q=25", "6∶2∶2", "{1,2,4}"
]
for token in must_keep:
    assert token in full, token

assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 4
for caption in [
    "图1 自适应多尺度网络流量预测方法框架",
    "图2 Abilene与GÉANT上三随机种子外部基线配对结果",
    "图3 固定等权与自适应权重的八随机种子配对MSE及配对差值分布",
    "图4 八随机种子下样本级自适应路由权重分布",
]:
    assert caption in full

doc.save(TARGET)
print(TARGET)
print("tables", len(doc.tables), "figures", len(doc.inline_shapes))
print("abstract_chars", len(find("摘要：").text.replace("摘要：", "")))
