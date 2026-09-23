from pathlib import Path
import csv
from collections import defaultdict

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.23v24.docx"
TARGET = ROOT / "word_output" / "2026.09.23v25.docx"
CORE = ROOT / "abilene_forecasting" / "results" / "core_v2_summary.csv"

OUT_COLOR = ROOT / "abilene_forecasting" / "paper_figures" / "exports" / "color"
OUT_GRAY = ROOT / "abilene_forecasting" / "paper_figures" / "exports" / "grayscale"
OUT_DATA = ROOT / "abilene_forecasting" / "paper_figures" / "source_data_v25"
for d in (OUT_COLOR, OUT_GRAY, OUT_DATA):
    d.mkdir(parents=True, exist_ok=True)

COLOR_FIG = OUT_COLOR / "fig3_lightts_improvement_heatstrip.png"
GRAY_FIG = OUT_GRAY / "fig3_lightts_improvement_heatstrip.png"
AUDIT_CSV = OUT_DATA / "fig3_lightts_improvement_seed42_44.csv"
SUMMARY_CSV = OUT_DATA / "fig3_lightts_improvement_summary.csv"

# ---------- Load and strictly verify repository data ----------
rows = list(csv.DictReader(CORE.open(encoding="utf-8")))
wanted = {"lightts", "dlinear_scale"}
by = defaultdict(dict)
for r in rows:
    if r["model"] in wanted and int(r["seed"]) in {42, 43, 44}:
        by[(r["dataset"], int(r["seed"]))][r["model"]] = r

datasets = ["abilene", "geant"]
metrics = [("MSE", "mse"), ("MAE", "mae"), ("RMSE", "rmse")]
seeds = [42, 43, 44]

records = []
matrices = {}
params = {}

for ds in datasets:
    mat = np.zeros((3, 4), dtype=float)
    for j, seed in enumerate(seeds):
        pair = by[(ds, seed)]
        assert set(pair) == wanted, (ds, seed, pair.keys())
        light = pair["lightts"]
        ours = pair["dlinear_scale"]
        p_ours = int(ours["parameters"])
        p_light = int(light["parameters"])
        params.setdefault(ds, (p_ours, p_light))
        assert params[ds] == (p_ours, p_light)
        for i, (label, key) in enumerate(metrics):
            lv = float(light[key])
            ov = float(ours[key])
            improvement = (lv - ov) / lv * 100.0
            assert improvement > 0, (ds, seed, label, improvement)
            mat[i, j] = improvement
            records.append({
                "dataset": ds,
                "seed": seed,
                "metric": label,
                "lightts": lv,
                "adaptive": ov,
                "relative_improvement_pct": improvement,
                "adaptive_parameters": p_ours,
                "lightts_parameters": p_light,
            })
    mat[:, 3] = mat[:, :3].mean(axis=1)
    matrices[ds] = mat

assert params["abilene"] == (8339, 130162), params["abilene"]
assert params["geant"] == (8339, 384382), params["geant"]

# Guardrails for all 18 plotted per-seed values and six means.
expected_means = {
    ("abilene", "MSE"): 29.803334304754742,
    ("abilene", "MAE"): 17.5741855339389,
    ("abilene", "RMSE"): 16.231227254261075,
    ("geant", "MSE"): 46.88658964188419,
    ("geant", "MAE"): 29.152506963018922,
    ("geant", "RMSE"): 27.17615137962596,
}
for ds in datasets:
    for i, (label, _) in enumerate(metrics):
        got = float(matrices[ds][i, 3])
        assert abs(got - expected_means[(ds, label)]) < 1e-10, (ds, label, got)

with AUDIT_CSV.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(records[0].keys()))
    w.writeheader()
    w.writerows(records)

summary_rows = []
for ds in datasets:
    p_ours, p_light = params[ds]
    for i, (label, _) in enumerate(metrics):
        vals = matrices[ds][i, :3]
        summary_rows.append({
            "dataset": ds,
            "metric": label,
            "seed42_pct": vals[0],
            "seed43_pct": vals[1],
            "seed44_pct": vals[2],
            "mean_pct": vals.mean(),
            "min_pct": vals.min(),
            "max_pct": vals.max(),
            "adaptive_parameters": p_ours,
            "lightts_parameters": p_light,
            "adaptive_parameter_ratio_vs_lightts_pct": p_ours / p_light * 100.0,
            "parameter_reduction_pct": (p_light - p_ours) / p_light * 100.0,
        })

with SUMMARY_CSV.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(summary_rows[0].keys()))
    w.writeheader()
    w.writerows(summary_rows)

# ---------- Figure 3B: separated heatstrip + parameter mini-panels ----------
matplotlib.rcParams["font.family"] = ["Noto Serif CJK SC", "Noto Sans CJK SC", "DejaVu Serif"]
matplotlib.rcParams["axes.unicode_minus"] = False
matplotlib.rcParams["pdf.fonttype"] = 42
matplotlib.rcParams["ps.fonttype"] = 42

def draw(path, grayscale=False):
    if grayscale:
        cmap = plt.get_cmap("Greys")
        ours_color, light_color = "#333333", "#B3B3B3"
    else:
        cmap = LinearSegmentedColormap.from_list(
            "paper_teal_blue", ["#EEF7E8", "#A5DDD2", "#4DB3C2", "#1578A8", "#0A4C87"]
        )
        ours_color, light_color = "#2C5B75", "#A8B8C4"

    fig = plt.figure(figsize=(9.4, 4.6), dpi=300)
    gs = fig.add_gridspec(
        2, 4,
        height_ratios=[16, 1.15],
        width_ratios=[4.25, 1.35, 4.25, 1.35],
        left=0.065, right=0.985, bottom=0.15, top=0.93,
        wspace=0.28, hspace=0.28
    )

    heat_axes = [fig.add_subplot(gs[0, 0]), fig.add_subplot(gs[0, 2])]
    param_axes = [fig.add_subplot(gs[0, 1]), fig.add_subplot(gs[0, 3])]
    cax = fig.add_subplot(gs[1, :])

    vmin, vmax = 0.0, 55.0
    im = None
    for ax, pax, ds, panel in zip(heat_axes, param_axes, datasets, ["a", "b"]):
        data = matrices[ds]
        im = ax.imshow(data, cmap=cmap, vmin=vmin, vmax=vmax, aspect="auto", interpolation="nearest")
        ax.set_xticks(range(4), ["seed42", "seed43", "seed44", "均值"], fontsize=8.4)
        ax.set_yticks(range(3), ["MSE", "MAE", "RMSE"], fontsize=9.8)
        ax.tick_params(axis="both", length=0)
        ax.set_title(f"({panel}) " + ("Abilene" if ds == "abilene" else "GÉANT"),
                     loc="left", fontsize=13.2, pad=7)

        # Crisp cell boundaries and a stronger separator before the mean column.
        ax.set_xticks(np.arange(-0.5, 4, 1), minor=True)
        ax.set_yticks(np.arange(-0.5, 3, 1), minor=True)
        ax.grid(which="minor", color="white", linewidth=1.05)
        ax.tick_params(which="minor", bottom=False, left=False)
        ax.axvline(2.5, color="white", linewidth=2.4)

        for i in range(3):
            for j in range(4):
                val = data[i, j]
                rgba = im.cmap(im.norm(val))
                luminance = 0.2126 * rgba[0] + 0.7152 * rgba[1] + 0.0722 * rgba[2]
                txt_color = "#0B2942" if luminance > 0.60 else "white"
                weight = "bold" if j == 3 else "normal"
                ax.text(j, i, f"{val:.1f}%", ha="center", va="center",
                        fontsize=8.7, color=txt_color, fontweight=weight)

        for s in ax.spines.values():
            s.set_visible(False)

        # Parameter panel is physically separate from the heatmap to prevent overlaps.
        p_ours, p_light = params[ds]
        ratios = [p_ours / p_light * 100.0, 100.0]
        bars = pax.barh([0, 1], ratios, height=0.42, color=[ours_color, light_color], alpha=0.92)
        pax.set_xlim(0, 115)
        pax.set_ylim(-0.65, 1.65)
        pax.set_yticks([0, 1], ["本文方法", "LightTS"], fontsize=7.7)
        pax.set_xticks([0, 50, 100])
        pax.tick_params(axis="x", direction="in", labelsize=7.2)
        pax.set_xlabel("参数量比例（%）", fontsize=7.6)
        pax.grid(axis="x", color="#DDDDDD", linestyle=":", linewidth=0.55)
        for bar, count, ratio in zip(bars, [p_ours, p_light], ratios):
            x = bar.get_width()
            # Place labels above/right with reserved margin; never overlay bars or heatmap.
            xpos = min(x + 2.0, 111.0)
            pax.text(xpos, bar.get_y() + bar.get_height()/2,
                     f"{count:,}", va="center", ha="left", fontsize=7.2, color="#222222")
        for s in ["top", "right"]:
            pax.spines[s].set_visible(False)
        pax.spines["left"].set_linewidth(0.65)
        pax.spines["bottom"].set_linewidth(0.65)

    cb = fig.colorbar(im, cax=cax, orientation="horizontal")
    cb.set_label("相对 LightTS 的误差降低率（%）", fontsize=8.5, labelpad=2)
    cb.ax.tick_params(labelsize=7.2, direction="in")
    cb.outline.set_linewidth(0.55)

    fig.savefig(path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(fig)

draw(COLOR_FIG, grayscale=False)
draw(GRAY_FIG, grayscale=True)

# ---------- Replace Figure 3 in v24 and revise related text ----------
doc = Document(SOURCE)

def find_exact(text):
    found = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(found) != 1:
        raise RuntimeError(f"Expected one exact paragraph {text!r}, found {len(found)}")
    return found[0]

def find_prefix(prefix):
    found = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(found) != 1:
        raise RuntimeError(f"Expected one prefix {prefix!r}, found {len(found)}")
    return found[0]

def replace_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)

old_caption = "图3 Abilene与GÉANT上三随机种子外部基线配对结果"
new_caption = "图3 Abilene与GÉANT上本文方法相对LightTS的三随机种子多指标性能提升"
caption = find_exact(old_caption)

prev = caption._element.getprevious()
pic_p = None
for _ in range(5):
    if prev is None:
        break
    if prev.tag.endswith("}p") and ("<w:drawing" in prev.xml or "<pic:pic" in prev.xml):
        pic_p = Paragraph(prev, caption._parent)
        break
    prev = prev.getprevious()
if pic_p is None:
    raise RuntimeError("Could not find Figure 3 image paragraph")
for run in list(pic_p.runs):
    run._element.getparent().remove(run._element)
pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic_p.add_run().add_picture(str(COLOR_FIG), width=Inches(6.25))
replace_text(caption, new_caption)

fig3_para = find_prefix("图3进一步给出了随机种子42—44的单次实验结果")
replace_text(
    fig3_para,
    "为避免与图5中DLinear和固定等权多尺度模型的八随机种子核心消融重复，图3仅保留独立外部基线LightTS，并使用两种方法共同具备的随机种子42—44进行比较。"
    "对每个随机种子分别按（LightTS误差−本文方法误差）/LightTS误差计算相对降低率后，Abilene上的MSE、MAE和RMSE三随机种子平均降低率分别为29.80%、17.57%和16.23%；"
    "GÉANT上分别为46.89%、29.15%和27.18%。三个共同随机种子在两个数据集的三项指标上均保持正向降低。"
    "图3右侧同时给出模型参数量比例：本文方法参数量为8339，LightTS在Abilene和GÉANT上分别为130162和384382，本文方法分别约为其6.41%和2.17%。"
    "因此，图3用于集中展示本文方法相对外部结构基线的多指标误差优势及参数规模差异，而DLinear相关的内部基线比较由表3和图5承担。"
)

# Keep Table 2's DLinear comparison, but clarify Figure 3's narrower role.
baseline_para = find_prefix("表2给出了3个共同随机种子下的外部基线结果")
replace_text(
    baseline_para,
    "表2给出了3个共同随机种子下的总体基线结果，保留DLinear和LightTS两类参照。本文方法在Abilene上的MSE为0.186，在GÉANT上为0.434；"
    "相对于DLinear，两项MSE分别降低4.06%和4.43%，相对于LightTS分别降低29.89%和46.83%，MAE的变化方向与MSE一致。"
    "其中，DLinear作为与本文尺度专家同源的单尺度内部基线，主要用于后续核心消融；LightTS作为独立外部结构基线，则由图3进一步展开逐随机种子、多指标及参数规模比较。"
)

# Revise experiment-chain summary so Figure 3 and Figure 5 have distinct roles.
chain = find_prefix("综合表2、表3以及图3、图5和图6")
replace_text(
    chain,
    "综合表2、表3以及图3、图5和图6形成互补的实验链。表2给出DLinear、LightTS与本文方法的总体误差统计；"
    "图3进一步聚焦独立外部基线LightTS，以共同随机种子42—44的MSE、MAE和RMSE相对降低率及参数规模差异展示外部竞争力；"
    "表3和图5则使用随机种子42—49，从绝对误差统计和相对改进分布两个角度比较DLinear、固定等权多尺度模型与本文方法，分析多尺度结构和样本级融合方式的贡献；"
    "图6最后从内部尺度权重分布说明路由器在测试样本上形成了非等权融合。这样，外部模型比较、内部消融和机制解释分别由不同图表承担，减少了重复信息。"
)

# Guardrails.
full = "\n".join(p.text for p in doc.paragraphs)
assert old_caption not in full
assert new_caption in full
for token in ["29.80%", "17.57%", "16.23%", "46.89%", "29.15%", "27.18%", "6.41%", "2.17%", "8339", "130162", "384382"]:
    assert token in full, token
assert "图3进一步给出了随机种子42—44的单次实验结果" not in full
assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 4

doc.save(TARGET)
print("saved", TARGET)
print("figure", COLOR_FIG)
print("audit", AUDIT_CSV, SUMMARY_CSV)
print("tables", len(doc.tables), "figures", len(doc.inline_shapes), "paragraphs", len(doc.paragraphs))
