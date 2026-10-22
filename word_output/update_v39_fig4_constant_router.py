"""Replace Figure 4 in v38 with the constant-input router comparison."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.27v38.docx"
TARGET = ROOT / "word_output/2026.09.27v39.docx"
FIGURE = (ROOT / "paper_figures/fig4_calibration_trajectories_v39" /
          "fig4_calibration_plus_constant_router.png")

document = Document(SOURCE)
paragraphs = document.paragraphs
assert len(paragraphs) == 114 and len(document.inline_shapes) == 5
assert paragraphs[91].text == "图4 Abilene与GÉANT全测试点预测校准及典型窗口预测轨迹"
assert len(paragraphs[90]._p.xpath(".//w:drawing")) == 1

for run in list(paragraphs[90].runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
paragraphs[90].add_run().add_picture(str(FIGURE), width=Inches(5.85))

description = paragraphs[92]
assert len(description.runs) == 1
old = description.text
assert "三种模型在高真实值区域均有低估" in old
assert "低于DLinear和固定等权模型" in old
updated = old.replace(
    "曲线表示各区间的8种子平均预测值，浅色带为种子间±1个标准差，虚线为理想预测。三种模型在高真实值区域均有低估，",
    "四种模型的曲线表示各区间的8种子平均预测值，浅色带为种子间±1个标准差，灰色虚线为理想预测。四种模型在高真实值区域均有低估，",
).replace(
    "低于DLinear和固定等权模型",
    "低于DLinear、固定等权与恒定输入路由模型",
)
assert updated != old
description.runs[0].text = updated

document.save(TARGET)
check = Document(TARGET)
assert len(check.paragraphs) == 114 and len(check.inline_shapes) == 5
assert check.paragraphs[92].text == updated
assert len(check.paragraphs[90]._p.xpath(".//w:drawing")) == 1
print(TARGET)
