"""Add the paired-seed figure and finalize the real-data router figure."""

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.25v28_fig6_realdata_preview.docx"
TARGET = ROOT / "word_output/2026.09.25v29.docx"
PAIRED_FIGURE = ROOT / "paper_figures/fig3_global_vs_adaptive_paired.png"


def replace_text(paragraph, new_text):
    assert paragraph.runs
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def clone_after(anchor, template):
    element = deepcopy(template._p)
    anchor.addnext(element)
    return Paragraph(element, template._parent)


doc = Document(SOURCE)
paragraphs = doc.paragraphs
assert "全局可学习权重模型" in paragraphs[84].text
assert paragraphs[85].text.startswith("表3")
assert paragraphs[93].text.startswith("图6")
assert len(doc.tables) == 3

# Renumber the existing sequence so the added paired comparison follows Table 3.
number_map = {"3": "2", "5": "4", "6": "5"}
for paragraph in paragraphs:
    if re.search(r"图[356]", paragraph.text):
        replace_text(paragraph, re.sub(r"图([356])", lambda m: "图" + number_map[m[1]], paragraph.text))

replace_text(
    paragraphs[84],
    paragraphs[84].text + "图3展示八个种子的配对改进率及其跨种子置信区间；正值表示样本级路由的MSE更低，菱形与误差线分别表示均值及t分布近似95%置信区间。",
)
replace_text(
    paragraphs[89],
    "固定等权模型只能检验等权融合是否足够；全局可学习权重模型进一步控制了“权重不等”这一因素。图3显示，样本级路由在Abilene的6个种子、GÉANT的7个种子上优于全局权重模型，但GÉANT种子43的MSE升高13.22%，使跨种子95%置信区间跨越零。图4另汇总相对DLinear和固定等权模型的多指标结果。",
)

table = doc.tables[2]
bridge = clone_after(table._tbl, paragraphs[86])
replace_text(bridge, "图3逐随机种子展示本文方法相对全局可学习权重模型的MSE变化。")
image = clone_after(bridge._p, paragraphs[87])
for run in list(image.runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
image.add_run().add_picture(str(PAIRED_FIGURE), width=Inches(5.8))
image.paragraph_format.keep_with_next = True

caption = clone_after(image._p, paragraphs[88])
replace_text(caption, "图3 样本级路由相对全局可学习权重模型的八随机种子配对MSE改进率")
caption.paragraph_format.keep_with_next = False

doc.save(TARGET)
print(TARGET)
