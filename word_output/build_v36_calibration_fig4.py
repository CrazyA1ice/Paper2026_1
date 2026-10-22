"""Replace Figure 4 with all-test calibration and revise its surrounding claims."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.26v35.docx"
TARGET = ROOT / "word_output/2026.09.27v36.docx"
FIGURE = ROOT / "paper_figures/forecast_truth_previews/B_全测试点预测校准_8种子.png"


def set_text(paragraph, value):
    assert paragraph.runs
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


doc = Document(SOURCE)
p = doc.paragraphs
assert len(p) == 114 and len(doc.inline_shapes) == 5 and len(doc.tables) == 3
assert p[91].text.startswith("图4 样本级路由相对全局可学习权重模型")
assert len(p[90]._p.xpath(".//w:drawing")) == 1

set_text(
    p[89],
    "固定等权模型只能检验等权融合是否足够；全局可学习权重模型进一步控制了“权重不等”这一因素。"
    "表3显示，样本级路由在Abilene的6个种子、GÉANT的7个种子上取得较低MSE，"
    "但GÉANT种子43的MSE比全局权重模型高13.22%，使跨种子95%置信区间跨越零。"
    "图4进一步考察不同真实流量水平下各模型的平均预测偏差。",
)
for run in list(p[90].runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
p[90].add_run().add_picture(str(FIGURE), width=Inches(5.85))
set_text(p[91], "图4 Abilene与GÉANT全测试点的八随机种子预测校准比较")
set_text(
    p[92],
    "图4使用随机种子42～49的全部测试窗口、24个预测步和有效OD通道，将每个数据集的真实值"
    "按分位点划为最多40个区间，并对重复边界合并。横轴为区间内真实值均值，纵轴为对应的预测值均值；"
    "各模型实线表示8个随机种子的平均预测值，浅色带表示种子间±1个标准差，虚线为预测均值与真实值相等。"
    "三种模型在较高真实值区域均存在低估；本文方法的平均预测曲线在高值端相对两类对照更接近虚线，"
    "但偏差仍然明显。该图展示的是按真实值分组后的平均偏差，不能代表逐点误差分布；"
    "相互重叠的测试窗口也不能视为独立重复，总体预测精度仍以表3的配对结果为准。",
)
set_text(
    p[101],
    "本文面向骨干网络OD流量中的短时波动和较慢趋势，以尺度1、2和4的DLinear专家构建多尺度预测，"
    "并通过轻量路由网络计算输入相关融合权重。在预测长度24的Abilene和GÉANT任务上，"
    "本文方法相对DLinear、LightTS和固定等权模型取得较低平均误差。加入全局可学习权重对照后，"
    "样本级路由的逐种子平均MSE改进率分别为1.80%和4.31%，但GÉANT的95%置信区间跨越零。"
    "按真实值分组的预测校准表明，各模型对较高流量取值仍有低估；本文方法在高值端的平均偏差较小，"
    "但不能由此推断其在每个测试点均更准确。逐种子权重及波动分组分析也未能充分证明路由器"
    "稳定地匹配每个窗口的尺度需求。后续需在更多数据和预测设置下检验这种输入相关融合的稳定收益。",
)

doc.save(TARGET)
check = Document(TARGET)
all_text = "\n".join(paragraph.text for paragraph in check.paragraphs)
assert len(check.paragraphs) == 114 and len(check.inline_shapes) == 5
assert "图4 Abilene与GÉANT全测试点的八随机种子预测校准比较" in all_text
assert "预测步分解显示" not in all_text
assert "H=12" not in all_text and "H=48" not in all_text
print(TARGET)
