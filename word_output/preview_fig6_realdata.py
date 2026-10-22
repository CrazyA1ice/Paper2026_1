"""Place the real-data Figure 6 in a copy of the v28 manuscript."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.23v28.docx"
TARGET = ROOT / "word_output" / "2026.09.25v28_fig6_realdata_preview.docx"
FIGURE = ROOT / "paper_figures" / "fig6_router_seedwise_realdata_compact.png"


def set_text(paragraph, expected, value):
    assert expected in paragraph.text
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


document = Document(SOURCE)
paragraphs = document.paragraphs
set_text(
    paragraphs[91],
    "图6分别显示",
    "图6保留分布轮廓与散点。轮廓为8个种子均值的描述性平滑分布，彩色点对应种子42～49；细线表示各自窗口内权重的均值±1个标准差，并限于0～1。空心菱形及黑色误差线为跨种子均值及95%置信区间，水平虚线为固定等权1/3。",
)
set_text(
    paragraphs[93],
    "图6 各随机种子",
    "图6 八随机种子下路由权重均值分布与窗口内标准差",
)
picture_paragraph = paragraphs[92]
assert len(picture_paragraph._p.xpath(".//w:drawing")) == 1
for run in list(picture_paragraph.runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
picture_paragraph.add_run().add_picture(str(FIGURE), width=Inches(5.8))
picture_paragraph.paragraph_format.keep_with_next = True
paragraphs[93].paragraph_format.keep_with_next = False
document.save(TARGET)
print(TARGET)
