from pathlib import Path
import csv, math, re
from collections import defaultdict

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy.stats import gaussian_kde, t

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.22v23.docx"
TARGET = ROOT / "word_output" / "2026.09.23v24.docx"
CORE = ROOT / "abilene_forecasting" / "results" / "core_v2_summary.csv"
PAIR = ROOT / "abilene_forecasting" / "paper_figures" / "source_data_v2" / "fig2_weight_pairing_8seeds.csv"

OUT_COLOR = ROOT / "abilene_forecasting" / "paper_figures" / "exports" / "color"
OUT_GRAY = ROOT / "abilene_forecasting" / "paper_figures" / "exports" / "grayscale"
OUT_DATA = ROOT / "abilene_forecasting" / "paper_figures" / "source_data_v24"
for d in (OUT_COLOR, OUT_GRAY, OUT_DATA):
    d.mkdir(parents=True, exist_ok=True)

COLOR_FIG = OUT_COLOR / "fig5_relative_improvement_raincloud.png"
GRAY_FIG = OUT_GRAY / "fig5_relative_improvement_raincloud.png"
AUDIT_CSV = OUT_DATA / "fig5_relative_improvement_8seeds.csv"
AUDIT_SUMMARY = OUT_DATA / "fig5_relative_improvement_summary.csv"

# ---------- Load and verify repository data ----------
core_rows = list(csv.DictReader(CORE.open(encoding="utf-8")))
pair_rows = list(csv.DictReader(PAIR.open(encoding="utf-8")))

core = defaultdict(dict)
for r in core_rows:
    if r["model"] in {"dlinear", "dlinear_scale"}:
        core[(r["dataset"], int(r["seed"]))][r["model"]] = r

pair = {(r["dataset"], int(r["seed"])): r for r in pair_rows}
datasets = ["abilene", "geant"]
metrics = [("MSE", "mse"), ("MAE", "mae"), ("RMSE", "rmse")]

records = []
values = defaultdict(lambda: defaultdict(dict))

for ds in datasets:
    for seed in range(42, 50):
        assert (ds, seed) in core and set(core[(ds, seed)]) == {"dlinear", "dlinear_scale"}
        assert (ds, seed) in pair
        for label, key in metrics:
            b = float(core[(ds, seed)]["dlinear"][key])
            a = float(core[(ds, seed)]["dlinear_scale"][key])
            imp_d = (b - a) / b * 100.0
            values[ds]["vs_dlinear"].setdefault(label, []).append(imp_d)
            records.append({
                "dataset": ds, "seed": seed, "baseline": "DLinear",
                "metric": label, "baseline_value": b, "adaptive_value": a,
                "relative_improvement_pct": imp_d
            })

            pr = pair[(ds, seed)]
            fixed = float(pr[f"fixed_{key}"])
            adap = float(pr[f"adaptive_{key}"])
            # Cross-check the adaptive values from the paired source against core_v2_summary.
            assert abs(adap - a) < 1e-9, (ds, seed, key, adap, a)
            imp_f = (fixed - adap) / fixed * 100.0
            if key in {"mse", "mae"}:
                supplied = float(pr[f"{key}_improve_pct"])
                assert abs(imp_f - supplied) < 1e-8, (ds, seed, key, imp_f, supplied)
            values[ds]["vs_fixed"].setdefault(label, []).append(imp_f)
            records.append({
                "dataset": ds, "seed": seed, "baseline": "固定等权",
                "metric": label, "baseline_value": fixed, "adaptive_value": adap,
                "relative_improvement_pct": imp_f
            })

# Persist exact plotted data.
with AUDIT_CSV.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(records[0].keys()))
    w.writeheader()
    w.writerows(records)

summary_rows = []
for ds in datasets:
    for comp, baseline in [("vs_dlinear", "DLinear"), ("vs_fixed", "固定等权")]:
        for label, _ in metrics:
            arr = np.array(values[ds][comp][label], dtype=float)
            mean = float(arr.mean())
            sd = float(arr.std(ddof=1))
            ci = float(t.ppf(0.975, len(arr)-1) * sd / math.sqrt(len(arr)))
            wins = int(np.sum(arr > 0))
            summary_rows.append({
                "dataset": ds, "baseline": baseline, "metric": label,
                "n_seeds": len(arr), "mean_relative_improvement_pct": mean,
                "sd_pct": sd, "ci95_halfwidth_pct": ci, "wins": wins
            })

with AUDIT_SUMMARY.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(summary_rows[0].keys()))
    w.writeheader()
    w.writerows(summary_rows)

summary = {(r["dataset"], r["baseline"], r["metric"]): r for r in summary_rows}

# Guardrail values used in manuscript prose.
expected = {
    ("abilene", "DLinear", "MSE"): (2.7677, 8),
    ("abilene", "DLinear", "MAE"): (8.0114, 7),
    ("abilene", "DLinear", "RMSE"): (1.3953, 8),
    ("geant", "DLinear", "MSE"): (2.3483, 5),
    ("geant", "DLinear", "MAE"): (1.3243, 4),
    ("geant", "DLinear", "RMSE"): (1.2205, 5),
}
for k, (m, wins) in expected.items():
    got = summary[k]
    assert abs(float(got["mean_relative_improvement_pct"]) - m) < 5e-4, (k, got)
    assert int(got["wins"]) == wins, (k, got)

# ---------- Publication-style Figure 5 ----------
# CJK font installed by workflow. Do not bake a figure title into the image.
matplotlib.rcParams["font.family"] = ["Noto Serif CJK SC", "Noto Sans CJK SC", "DejaVu Serif"]
matplotlib.rcParams["axes.unicode_minus"] = False
matplotlib.rcParams["pdf.fonttype"] = 42
matplotlib.rcParams["ps.fonttype"] = 42

blue = "#2B6CB0"
orange = "#E87722"
blue_fill = "#BFD7F2"
orange_fill = "#F6CFB1"

def draw_figure(path, grayscale=False):
    if grayscale:
        c1, c2 = "#333333", "#777777"
        f1, f2 = "#D0D0D0", "#E2E2E2"
    else:
        c1, c2 = blue, orange
        f1, f2 = blue_fill, orange_fill

    fig, axes = plt.subplots(1, 2, figsize=(9.2, 5.3), dpi=300)
    fig.subplots_adjust(left=0.105, right=0.985, bottom=0.17, top=0.92, wspace=0.34)

    rng = np.random.default_rng(20260923)
    y_top = [6.0, 5.0, 4.0]
    y_bot = [2.75, 1.75, 0.75]

    for ax, ds, panel in zip(axes, datasets, ["a", "b"]):
        ax.set_xlim(-15, 17)
        ax.set_ylim(0.25, 6.7)
        ax.axvline(0, color="#666666", ls="--", lw=0.9, zorder=0)
        ax.axhline(3.38, color="#B3B3B3", ls="--", lw=0.8, zorder=0)
        ax.grid(axis="x", color="#D9D9D9", ls=":", lw=0.55)
        ax.set_xticks(np.arange(-15, 20, 5))
        ax.tick_params(axis="both", direction="in", length=3.5, width=0.8, labelsize=8.8)
        ax.set_yticks(y_top + y_bot)
        ax.set_yticklabels(["MSE", "MAE", "RMSE", "MSE", "MAE", "RMSE"], fontsize=9.4)
        ax.set_xlabel("相对改进率（%）", fontsize=9.8)
        ax.set_title(f"({panel}) " + ("Abilene" if ds == "abilene" else "GÉANT"),
                     loc="left", fontsize=13.5, pad=6)

        # Compact group labels outside the axes; these are necessary legend-like identifiers.
        ax.text(-0.17, 0.77, "相对 DLinear", transform=ax.transAxes, rotation=90,
                ha="center", va="center", color=c1, fontsize=8.8, fontweight="bold")
        ax.text(-0.17, 0.245, "相对固定等权", transform=ax.transAxes, rotation=90,
                ha="center", va="center", color=c2, fontsize=8.8, fontweight="bold")

        for comp, color, fill, ys in [
            ("vs_dlinear", c1, f1, y_top),
            ("vs_fixed", c2, f2, y_bot),
        ]:
            baseline = "DLinear" if comp == "vs_dlinear" else "固定等权"
            for (metric_label, _), y in zip(metrics, ys):
                arr = np.array(values[ds][comp][metric_label], dtype=float)

                # Half-raincloud KDE: density sits above the row, leaving scatter below.
                xs = np.linspace(max(-15, arr.min() - 3.0), min(17, arr.max() + 3.0), 350)
                if np.ptp(arr) > 1e-8:
                    kde = gaussian_kde(arr)
                    den = kde(xs)
                    den = den / den.max() * 0.27
                    ax.fill_between(xs, y, y + den, color=fill, alpha=0.78, lw=0, zorder=1)
                    ax.plot(xs, y + den, color=color, alpha=0.45, lw=0.7, zorder=2)

                # All 8 seeds.
                jitter = rng.uniform(-0.15, -0.035, len(arr))
                ax.scatter(arr, y + jitter, s=15, color=color, alpha=0.62,
                           edgecolors="white", linewidths=0.3, zorder=3)

                # Mean and 95% t confidence interval computed from the same 8 per-seed improvements.
                row = summary[(ds, baseline, metric_label)]
                mean = float(row["mean_relative_improvement_pct"])
                ci = float(row["ci95_halfwidth_pct"])
                ax.errorbar(mean, y + 0.055, xerr=ci, fmt="o", ms=6.2, color=color,
                            ecolor=color, elinewidth=1.4, capsize=3.2, zorder=5)
                ax.text(mean, y + 0.31, f"{mean:.2f}", color=color, fontsize=8.1,
                        ha="center", va="center", fontweight="bold")

        for s in ax.spines.values():
            s.set_linewidth(0.8)

    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    handles = [
        Line2D([0], [0], marker="o", color="none", markerfacecolor="#888888",
               markeredgecolor="white", markersize=5, label="单次实验结果"),
        Patch(facecolor="#CFCFCF", edgecolor="none", alpha=0.8, label="分布密度"),
        Line2D([0], [0], marker="o", color="black", linewidth=1.2,
               markersize=5.5, label="均值 ± 95%置信区间"),
        Line2D([0], [0], color=c1, linewidth=4, alpha=0.78, label="相对 DLinear"),
        Line2D([0], [0], color=c2, linewidth=4, alpha=0.78, label="相对固定等权"),
    ]
    fig.legend(handles=handles, loc="lower center", ncol=5, frameon=True,
               bbox_to_anchor=(0.5, 0.02), fontsize=7.9,
               handlelength=1.8, columnspacing=1.3, borderpad=0.45)
    fig.savefig(path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(fig)

draw_figure(COLOR_FIG, grayscale=False)
draw_figure(GRAY_FIG, grayscale=True)

# ---------- Replace Figure 5 and related manuscript prose ----------
doc = Document(SOURCE)

def find_exact(text):
    m = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(m) != 1:
        raise RuntimeError(f"Expected one exact paragraph {text!r}, found {len(m)}")
    return m[0]

def find_prefix(prefix):
    m = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(m) != 1:
        raise RuntimeError(f"Expected one prefix {prefix!r}, found {len(m)}")
    return m[0]

def replace_paragraph_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)

old_caption = "图5 固定等权与自适应权重的八随机种子配对MSE及配对差值分布"
new_caption = "图5 Abilene与GÉANT上本文方法相对两类基线的八随机种子多指标相对改进分布"
caption = find_exact(old_caption)

# Replace the preceding image paragraph while preserving location/alignment.
prev = caption._element.getprevious()
pic_p = None
for _ in range(5):
    if prev is None:
        break
    if prev.tag.endswith("}p") and ("<w:drawing" in prev.xml or "<pic:pic" in prev.xml):
        pic_p = Paragraph(prev, caption._parent)
        break
    prev = prev.getprevious()
if pic_p is None:
    raise RuntimeError("Could not locate image paragraph preceding Figure 5 caption")
for r in list(pic_p.runs):
    r._element.getparent().remove(r._element)
pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic_p.add_run().add_picture(str(COLOR_FIG), width=Inches(6.25))
replace_paragraph_text(caption, new_caption)

p_fig5 = find_prefix("图5按照相同随机种子对固定等权模型和本文方法进行配对比较")
replace_paragraph_text(
    p_fig5,
    "图5进一步从随机种子42—49的逐次结果出发，计算本文方法相对DLinear和固定等权多尺度模型的相对误差改进率，并同时给出MSE、MAE和RMSE三个指标。"
    "图中每个散点对应一个随机种子，半小提琴表示8次结果的分布，圆点与误差线表示逐随机种子相对改进率的均值及95%置信区间，0%虚线表示两种方法误差相同。"
    "相对DLinear时，Abilene的MSE、MAE和RMSE逐随机种子平均改进率分别为2.77%、8.01%和1.40%，其中MSE与RMSE在8次实验中均取得改善；"
    "GÉANT对应的平均改进率分别为2.35%、1.32%和1.22%，但分布跨越0%的情况更多，说明跨随机种子波动更明显。"
    "相对固定等权模型时，Abilene三项指标的逐随机种子平均改进率分别为1.65%、8.92%和0.84%，GÉANT分别为4.54%、3.05%和2.37%。"
    "因此，在两种基线和三个误差指标下，本文方法的平均相对改进方向均为正，但不同数据集上的稳定性存在差异。"
)

p_summary = find_prefix("从DLinear、固定等权多尺度模型到本文方法")
replace_paragraph_text(
    p_summary,
    "从DLinear、固定等权多尺度模型到本文方法，模型结构逐步增加了两个因素：先加入多个时间尺度，再把固定权重替换为样本级权重。"
    "表3给出三类核心模型的绝对误差均值与标准差，图5则从8个共同随机种子的相对改进分布补充展示本文方法相对单尺度基线和固定等权多尺度基线的变化，并将比较扩展到MSE、MAE和RMSE三个指标。"
    "尤其是在专家结构和尺度集合完全相同的固定等权对照下，两数据集三项指标的平均相对改进均为正，说明性能变化不能仅用增加多尺度专家解释，尺度融合方式本身也会影响结果。"
    "这也是后续分析路由权重分布的原因：如果路由器始终输出接近1/3的权重，那么自适应模块与固定等权模型实际上不会形成明显区别。"
)

# v23-specific experiment-chain summary.
chain_matches = [p for p in doc.paragraphs if p.text.strip().startswith("综合表2、表3以及图3、图5和图6")]
if len(chain_matches) == 1:
    replace_paragraph_text(
        chain_matches[0],
        "综合表2、表3以及图3、图5和图6可以形成一条较清晰的实验链。"
        "表2和图3用于比较本文方法与外部轻量基线的总体误差水平；"
        "表3和图5进一步从绝对误差统计与八随机种子多指标相对改进分布两个角度，比较单尺度DLinear、固定等权多尺度模型与本文方法，考察多尺度结构和样本级融合方式带来的变化；"
        "图6则从模型内部的尺度分配结果说明路由器在测试样本上形成了非等权的融合方式。"
        "各部分分别对应外部基线比较、核心结构消融与稳定性、内部权重行为，与第1节提出的多尺度专家和样本级路由设计相互对应。"
    )
elif len(chain_matches) != 0:
    raise RuntimeError(f"Unexpected chain summary count: {len(chain_matches)}")

# Guardrails.
full = "\n".join(p.text for p in doc.paragraphs)
assert old_caption not in full
assert new_caption in full
assert "配对差值分布" not in full
assert "平均ΔMSE" not in full
for tok in ["2.77%", "8.01%", "1.40%", "2.35%", "1.32%", "1.22%",
            "1.65%", "8.92%", "0.84%", "4.54%", "3.05%", "2.37%"]:
    assert tok in full, tok
assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 4

doc.save(TARGET)
print("saved", TARGET)
print("figure", COLOR_FIG)
print("audit", AUDIT_CSV, AUDIT_SUMMARY)
print("tables", len(doc.tables), "figures", len(doc.inline_shapes), "paragraphs", len(doc.paragraphs))
