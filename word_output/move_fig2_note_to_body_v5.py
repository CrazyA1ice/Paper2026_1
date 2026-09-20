from __future__ import annotations

from pathlib import Path
from docx import Document


ROOT=Path(__file__).resolve().parents[1]
DOCX=ROOT/"word_output"/"2026.09.20v5.docx"


def find_one(doc,prefix):
    xs=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(xs)!=1:
        raise RuntimeError(f"{prefix!r}: {len(xs)} matches")
    return xs[0]


def set_text_like(p,text):
    first=p.runs[0] if p.runs else p.add_run()
    first.text=text
    for r in p.runs[1:]:
        r.text=""


def remove_para(p):
    el=p._element
    el.getparent().remove(el)
    p._p=p._element=None


doc=Document(DOCX)

analysis=find_one(doc,"同种子配对结果及配对差值分布见图2。")
set_text_like(
    analysis,
    "同种子配对结果及配对差值分布见图2。图2上排灰线连接同一随机种子下固定等权与自适应权重模型的MSE，"
    "黑色菱形及误差线表示8个随机种子的均值±标准差；下排以ΔMSE=固定等权−自适应权重表示效应方向，"
    "圆点表示各随机种子的配对差值，半小提琴表示差值分布，黑色圆点及误差线表示平均ΔMSE及其95%置信区间，"
    "其中ΔMSE>0表示自适应权重模型的MSE更低。Abilene中，自适应模型在8个种子中的6个取得更低MSE，"
    "平均ΔMSE为0.00320，95%置信区间为[-0.00020,0.00660]，配对t检验p=0.061；"
    "GÉANT中为7个，平均ΔMSE为0.02094，95%置信区间为[-0.00958,0.05146]，p=0.149。"
    "两数据集的平均差值均为正，说明样本级自适应加权在平均意义上降低了MSE；但相对固定等权的95%置信区间均跨越0，"
    "因此不能表述为两数据集均达到统计显著。结合表3可见，Abilene上的跨种子波动较小，而GÉANT仍存在更明显的稳定性提升空间。"
)

notes=[p for p in doc.paragraphs if p.text.startswith("注：a为Abilene，b为GÉANT。")]
for p in notes:
    remove_para(p)

text="\n".join(p.text for p in doc.paragraphs)
assert "注：a为Abilene，b为GÉANT。" not in text
assert "半小提琴表示差值分布" in text
assert "ΔMSE>0表示自适应权重模型的MSE更低" in text

tmp=DOCX.with_suffix(".tmp.docx")
doc.save(tmp)
tmp.replace(DOCX)
print(DOCX)
