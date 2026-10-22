"""Replace Figure 4 with calibration and truth-selected forecast examples."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.27v36.docx"
TARGET = ROOT / "word_output/2026.09.27v37.docx"
FIGURE = ROOT / "paper_figures/fig4_calibration_trajectories/fig4_calibration_plus_trajectories_preview.png"


def set_text(paragraph, value):
    assert paragraph.runs
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


doc = Document(SOURCE)
p = doc.paragraphs
assert len(p) == 114 and len(doc.inline_shapes) == 5 and len(doc.tables) == 3
assert p[91].text == "图4 Abilene与GÉANT全测试点的八随机种子预测校准比较"
assert len(p[90]._p.xpath(".//w:drawing")) == 1

set_text(p[89], p[89].text.replace(
    "图4进一步考察不同真实流量水平下各模型的平均预测偏差。",
    "图4进一步结合全测试点的平均预测偏差与按真实值波动选取的预测轨迹，观察整体趋势和具体误差。",
))
for run in list(p[90].runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
p[90].add_run().add_picture(str(FIGURE), width=Inches(5.85))
set_text(p[91], "图4 Abilene与GÉANT全测试点预测校准及典型窗口预测轨迹")
set_text(
    p[92],
    "图4(a、b)汇总随机种子42～49的全部测试窗口、24个预测步及有效OD通道："
    "按真实值分位点划为最多40个区间，曲线表示各区间的8种子平均预测值，浅色带为种子间±1个标准差，"
    "虚线为理想预测。三种模型在高真实值区域均有低估，本文方法的平均偏差较小，但分组均值不能代表逐点误差。"
    "图4(c～f)固定随机种子42，仅根据真实序列选择标准差接近第75百分位的OD通道，"
    "再选择波动程度接近中位数和第90百分位的窗口，"
    "未依据模型误差挑选。四例中本文方法有三例的24步均方误差低于DLinear和固定等权模型；"
    "Abilene的典型波动窗口则较差。轨迹展示局部行为，总体精度以表3的配对结果为准；"
    "重叠测试窗口也不视为独立重复。",
)

doc.save(TARGET)
check = Document(TARGET)
all_text = "\n".join(paragraph.text for paragraph in check.paragraphs)
assert len(check.paragraphs) == 114 and len(check.inline_shapes) == 5 and len(check.tables) == 3
assert "图4 Abilene与GÉANT全测试点预测校准及典型窗口预测轨迹" in all_text
assert "Abilene的典型波动窗口则较差" in all_text
assert "H=12" not in all_text and "H=48" not in all_text
print(TARGET)
