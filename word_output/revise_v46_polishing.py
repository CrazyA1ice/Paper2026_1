from pathlib import Path
from zipfile import ZipFile
import hashlib

from docx import Document


root = Path(__file__).resolve().parent
source = root / "2026.09.27v45.docx"
target = root / "2026.09.27v46.docx"
report = root / "2026.09.27v46_语言修改对照.md"

edits = {
    15: (
        "因此，模型在构造多个时间尺度的同时，根据当前输入样本调整各尺度在最终预测中的权重。",
        "因此，本文在构造多个时间尺度的同时，依据当前输入窗口调整各尺度的融合权重。",
    ),
    20: (
        "DLinear结构较简洁，可在控制模型规模的同时承担各尺度的预测任务，使实验重点落在多尺度输入和样本级融合本身。",
        "DLinear结构较简洁，可在控制模型规模的同时承担各尺度的预测任务，也便于比较多尺度输入和样本级融合的作用。",
    ),
    64: (
        "依次检验多尺度结构、非等权融合及输入相关权重的作用",
        "分别比较多尺度结构、非等权融合与输入相关权重的表现",
    ),
    77: (
        "DLinear与本文尺度专家采用同源结构，后续核心消融继续以它作为单尺度内部基线；",
        "DLinear与本文尺度专家采用同源结构，后续核心消融也以DLinear作为单尺度内部基线；",
    ),
    90: (
        "这一比较仍无法单独证明逐窗口权重分配有效。",
        "这些对照显示完整方法与两类常权重方案存在误差差异，但尚不能将收益唯一归因于逐窗口权重变化。",
    ),
    93: (
        "四种模型在高真实值区域均出现低估。本文方法的平均偏差较小，但分组均值不代表逐点误差。",
        "四种模型在高真实值区间均有低估；本文方法的分组平均偏差较小，但这一结果不能替代逐点误差比较。",
    ),
    100: (
        "对8个随机种子的预测长度24检查点，在NVIDIA GTX 1060上以批量16对预加载的测试张量计时，",
        "使用预测长度为24的8个随机种子检查点，在NVIDIA GTX 1060上以批量大小16对预加载的测试张量计时，",
    ),
    102: (
        "在预测长度24的Abilene和GÉANT任务上，本文方法相对DLinear、LightTS和固定等权模型取得较低平均误差。",
        "在Abilene和GÉANT预测长度为24的任务中，本文方法的平均误差低于DLinear、LightTS和固定等权模型。",
    ),
}

before = Document(source)
after = Document(source)
changes = []
for index, (old, new) in edits.items():
    paragraph = after.paragraphs[index]
    assert paragraph.text.count(old) == 1, index
    assert not paragraph._p.xpath(".//m:oMath | .//w:drawing"), index
    previous = paragraph.text
    assert paragraph.runs[0].text.count(old) == 1, index
    paragraph.runs[0].text = paragraph.runs[0].text.replace(old, new)
    changes.append((index, previous, paragraph.text))

after.save(target)
verify = Document(target)
assert len(before.paragraphs) == len(verify.paragraphs)
for index, (left, right) in enumerate(zip(before.paragraphs, verify.paragraphs)):
    if index not in edits:
        assert left.text == right.text, index
assert len(before.tables) == len(verify.tables) == 3
for left, right in zip(before.tables, verify.tables):
    assert [[c.text for c in r.cells] for r in left.rows] == [[c.text for c in r.cells] for r in right.rows]
assert len(before.inline_shapes) == len(verify.inline_shapes) == 5
with ZipFile(source) as left, ZipFile(target) as right:
    media_hashes = lambda z: {
        n: hashlib.sha256(z.read(n)).hexdigest()
        for n in z.namelist() if n.startswith("word/media/")
    }
    assert media_hashes(left) == media_hashes(right)
    assert left.read("word/document.xml").count(b"<m:oMath>") == right.read("word/document.xml").count(b"<m:oMath>")

lines = ["# v45 → v46 语言修改对照", "", "仅调整以下8处表述；数据、表格、图像、公式与参考文献未修改。", ""]
for index, old, new in changes:
    lines.extend([f"## 段落 {index}", "", f"原文：{old}", "", f"修改：{new}", ""])
report.write_text("\n".join(lines), encoding="utf-8")
print(target)
print(report)
print("changed paragraphs:", len(changes))
