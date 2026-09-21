from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "word_output" / "2026.09.20v5.docx"
CANDIDATE_DIR = ROOT / "abilene_forecasting" / "paper_figures" / "candidates_v5"
EXPORT_DIR = ROOT / "abilene_forecasting" / "paper_figures" / "exports"
FIG3_COLOR = CANDIDATE_DIR / "color" / "candidate_A_router_weight_distribution_color.png"
FIG3_GRAY = CANDIDATE_DIR / "grayscale" / "candidate_A_router_weight_distribution_grayscale.png"
FORMAL_COLOR = EXPORT_DIR / "color" / "fig3_router_weight_distribution_color.png"
FORMAL_GRAY = EXPORT_DIR / "grayscale" / "fig3_router_weight_distribution_grayscale.png"


def find_one(document: Document, startswith: str):
    matches = [p for p in document.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {startswith!r}, found {len(matches)}")
    return matches[0]


def set_paragraph_text_like(paragraph, text: str) -> None:
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def copy_paragraph_format(source: Paragraph, target: Paragraph) -> None:
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))
    target.alignment = source.alignment


def set_text_with_template(target: Paragraph, template: Paragraph, text: str) -> None:
    copy_paragraph_format(template, target)
    run = target.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))


def main() -> None:
    if not FIG3_COLOR.exists() or not FIG3_GRAY.exists():
        raise FileNotFoundError("Candidate A figure files are missing")

    FORMAL_COLOR.parent.mkdir(parents=True, exist_ok=True)
    FORMAL_GRAY.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FIG3_COLOR, FORMAL_COLOR)
    shutil.copy2(FIG3_GRAY, FORMAL_GRAY)

    document = Document(DOCX)
    if len(document.inline_shapes) != 3:
        raise RuntimeError(f"Expected 3 inline figures before insertion, found {len(document.inline_shapes)}")

    heading = find_one(document, "2.5 路由行为与适用范围")
    old_intro = find_one(document, "为观察样本级路由器是否退化为固定平均融合")
    old_caption = find_one(document, "图3 两数据集样本级路由权重变化")
    old_analysis = find_one(document, "图3中三个尺度权重均随测试窗口发生变化")
    summary = find_one(document, "综合表2、表3和图2可见")

    # Use existing Section 2.5 body/caption formatting as templates.
    body_template = old_intro
    caption_template = old_caption

    set_paragraph_text_like(
        old_intro,
        "为从总体统计视角分析样本级路由器对不同时间尺度的权重分配，"
        "图3汇总了Abilene和GÉANT数据集在随机种子42—49下全部测试窗口的三尺度路由权重。"
        "图中小提琴轮廓表示样本级权重分布，彩色圆点表示各随机种子的平均权重，"
        "黑色菱形及误差线表示8个随机种子均值及标准差，灰色虚线为固定等权基准1/3。"
    )

    # New Figure 3.
    picture_p = insert_after(old_intro)
    copy_paragraph_format(find_one(document, "图1 自适应多尺度网络流量预测方法框架")._p.getprevious() if False else old_intro, picture_p)
    picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_p.add_run().add_picture(str(FORMAL_COLOR), width=Inches(6.30))

    caption_p = insert_after(picture_p)
    set_text_with_template(
        caption_p,
        caption_template,
        "图3 八随机种子下样本级自适应路由权重分布",
    )

    analysis_p = insert_after(caption_p)
    set_text_with_template(
        analysis_p,
        body_template,
        "从图3可以看出，两数据集的三尺度权重均呈非退化分布，并未长期固定在1/3。"
        "按8个随机种子的平均权重计算，Abilene的α1、α2和α4分别为0.257、0.319和0.424，"
        "GÉANT分别为0.283、0.318和0.399；两个数据集上尺度4的平均权重均最高，尺度1最低，"
        "尺度2整体接近等权基准。与此同时，不同随机种子的平均权重仍存在明显差异，"
        "其中Abilene的尺度2和尺度4离散程度更大。该结果说明路由器并非学习到单一固定尺度，"
        "而是在总体尺度偏好基础上保留了随训练初始化和输入样本变化的自适应调节能力。"
    )

    transition_p = insert_after(analysis_p)
    set_text_with_template(
        transition_p,
        body_template,
        "为进一步观察权重在具体测试样本上的动态变化，图4固定选取随机种子42，"
        "展示两个数据集测试集前150个窗口的三尺度权重。该选择规则在两数据集间保持一致，"
        "仅用于机制可视化，不参与模型优劣的统计比较。"
    )

    # Existing Figure 3 becomes Figure 4.
    set_paragraph_text_like(
        old_caption,
        "图4 两数据集样本级路由权重变化（随机种子42，灰色点线为等权基准1/3）",
    )
    set_paragraph_text_like(
        old_analysis,
        "图4中三个尺度权重均随测试窗口发生变化，并未长期固定在1/3附近，"
        "与图3的总体分布结果相互印证，说明路由器能够依据输入窗口统计状态动态调整尺度组合。"
        "以随机种子42为例，两数据集均表现出较明显的尺度偏好变化；但不同随机种子的平均权重并不完全一致，"
        "因此不能将某一尺度解释为所有网络场景下持续占优的固定尺度。"
    )
    set_paragraph_text_like(
        summary,
        "综合表2、表3和图2—图4可见，样本级自适应尺度加权的平均收益已在Abilene和GÉANT两个骨干网数据集上得到验证，"
        "路由权重的总体分布与局部动态也表明模型确实在不同时间尺度之间进行自适应调节；"
        "但GÉANT上的跨种子方差更大，稳定性仍有提升空间。当前实验仅考察L=96、H=24的单一预测设置，"
        "未显式建模网络拓扑，外部模型也仅包含LightTS一种轻量基线；同时，“轻量”主要指参数规模，而非训练时间最短。"
        "因此，后续仍需在更多网络、预测长度和拓扑约束条件下进一步检验泛化性。"
    )

    text = "\n".join(p.text for p in document.paragraphs)
    assert "图3 八随机种子下样本级自适应路由权重分布" in text
    assert "图4 两数据集样本级路由权重变化" in text
    assert "Abilene的α1、α2和α4分别为0.257、0.319和0.424" in text
    assert "GÉANT分别为0.283、0.318和0.399" in text
    assert "图2—图4" in text

    tmp = DOCX.with_suffix(".tmp.docx")
    document.save(tmp)
    tmp.replace(DOCX)

    # Re-open to validate relationships and inline shapes after save.
    check = Document(DOCX)
    if len(check.inline_shapes) != 4:
        raise RuntimeError(f"Expected 4 inline figures after insertion, found {len(check.inline_shapes)}")
    print(DOCX)


if __name__ == "__main__":
    main()
