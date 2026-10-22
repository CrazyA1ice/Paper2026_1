"""Replace Figures 3 and 5 and align their manuscript descriptions."""

from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.25v31.docx"
TARGET = ROOT / "word_output" / "2026.09.25v32.docx"
FIG3 = ROOT / "paper_figures" / "fig3_classic_companion" / "fig3_relative_improvement_classic.png"
FIG5 = ROOT / "paper_figures" / "fig5_classic_restored" / "fig5_router_weights_classic.png"


def replace_text(paragraph, original, revised):
    assert original in paragraph.text
    text = paragraph.text.replace(original, revised)
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_text(paragraph, old_start, revised):
    assert paragraph.text.startswith(old_start)
    paragraph.runs[0].text = revised
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_picture(paragraph, image):
    assert len(paragraph.runs) == 1
    assert "graphic" in paragraph._p.xml
    paragraph.runs[0].clear()
    paragraph.runs[0].add_picture(str(image), width=Inches(6.25))


def main():
    document = Document(SOURCE)
    paragraphs = document.paragraphs
    assert len(paragraphs) == 113

    replace_text(
        paragraphs[86],
        "半小提琴表示8次结果的分布，圆点和误差线表示逐随机种子改进率的均值及95%置信区间",
        "半小提琴以平滑轮廓展示8次结果的分布，黑色菱形和横向误差线表示8次改进率的均值及95%置信区间",
    )
    replace_picture(paragraphs[87], FIG3)

    set_text(
        paragraphs[94],
        "图5保留分布轮廓与散点。",
        "图5展示随机种子42～49下三个尺度的路由权重。蓝、绿、橙色小提琴分别对应尺度1、2和4在全部测试窗口及种子上的汇总分布；彩色散点为各随机种子的窗口平均权重，短横线为汇总分布中位数。黑色菱形和误差线表示8个种子均值的均值±1个标准差，水平虚线为固定等权1/3。小提琴合并了种子与窗口两个层次，其宽度不能直接解释为同一种子内的窗口响应强度。",
    )
    replace_picture(paragraphs[95], FIG5)
    set_text(
        paragraphs[96],
        "图5 八随机种子下路由权重均值分布与窗口内标准差",
        "图5 Abilene与GÉANT的逐窗口路由权重汇总分布及八种子均值",
    )
    replace_text(
        paragraphs[97],
        "因此，合并所有种子的权重分布不能单独证明路由器有效识别每个窗口的尺度需求。",
        "图5的汇总分布同时包含种子间差异与同一种子内的窗口差异，因此不能单独证明路由器有效识别每个窗口的尺度需求。",
    )

    document.save(TARGET)
    check = Document(TARGET)
    assert len(check.inline_shapes) == 5
    assert "均值±1个标准差" in check.paragraphs[94].text
    assert "横向误差线" in check.paragraphs[86].text
    print(TARGET)


if __name__ == "__main__":
    main()
