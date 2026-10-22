"""Make the two high-priority notation and baseline fixes in v39."""

from copy import deepcopy
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.27v39.docx"
TARGET = ROOT / "word_output/2026.09.27v40.docx"
PANDOC = Path(r"D:\software\anaconda\Library\bin\pandoc.exe")

EQUATIONS = {
    25: r"\mathbf{X}_n=[\mathbf{x}_n,\ldots,\mathbf{x}_{n+L-1}]\in\mathbb{R}^{L\times C}.",
    26: r"\mathbf{Y}_n=[\mathbf{x}_{n+L},\ldots,\mathbf{x}_{n+L+H-1}]\in\mathbb{R}^{H\times C}.",
    28: r"u_{t,c}=\log(1+v_{t,c}).",
    30: r"x_{t,c}=\frac{u_{t,c}-\mu_c^{\mathrm{tr}}}{\sigma_c^{\mathrm{tr}}}.",
    35: r"\mathbf{X}^{(s)}_{n,t}=\frac{1}{s}\sum_{i=0}^{s-1}\mathbf{X}_{n,st+i},\quad t=0,\ldots,\frac{L}{s}-1.",
    40: r"\mathbf{T}^{(s)}_{n,t}=\frac{1}{q}\sum_{i=-m}^{m}\widetilde{\mathbf{X}}^{(s)}_{n,t+i},\quad q=2m+1=25.",
    41: r"\mathbf{R}^{(s)}_n=\mathbf{X}^{(s)}_n-\mathbf{T}^{(s)}_n.",
    43: r"\widehat{\mathbf{Y}}^{(s)}_n=\mathbf{W}^{(s)}_{T}\mathbf{T}^{(s)}_n+\mathbf{b}^{(s)}_{T}+\mathbf{W}^{(s)}_{R}\mathbf{R}^{(s)}_n+\mathbf{b}^{(s)}_{R}.",
    48: r"\mu_n=\frac{1}{LC}\sum_{t=0}^{L-1}\sum_{c=1}^{C}X_{n,t,c},\quad \sigma_n=\sqrt{\frac{1}{LC}\sum_{t=0}^{L-1}\sum_{c=1}^{C}(X_{n,t,c}-\mu_n)^2}.",
    49: r"\ell_n=\frac{1}{C}\sum_{c=1}^{C}X_{n,L-1,c},\quad d_n=\frac{1}{(L-1)C}\sum_{t=1}^{L-1}\sum_{c=1}^{C}|X_{n,t,c}-X_{n,t-1,c}|.",
    51: r"\mathbf{a}_n=\mathbf{W}_2\operatorname{ReLU}(\mathbf{W}_1\mathbf{z}_n+\mathbf{b}_1)+\mathbf{b}_2.",
    54: r"\alpha_{n,s}=\frac{\exp(a_{n,s})}{\sum_{j\in S}\exp(a_{n,j})},\quad \alpha_{n,s}\geq 0,\quad \sum_{s\in S}\alpha_{n,s}=1.",
    56: r"\widehat{\mathbf{Y}}_n=\sum_{s\in S}\alpha_{n,s}\widehat{\mathbf{Y}}^{(s)}_n.",
    60: r"\mathcal{L}_{\mathrm{MSE}}=\frac{1}{NHC}\sum_{n=1}^{N}\sum_{h=1}^{H}\sum_{c=1}^{C}(Y_{n,h,c}-\widehat{Y}_{n,h,c})^2.",
    73: r"\mathrm{MAE}=\frac{1}{N_{\mathrm{te}}HC}\sum_{n=1}^{N_{\mathrm{te}}}\sum_{h=1}^{H}\sum_{c=1}^{C}|Y_{n,h,c}-\widehat{Y}_{n,h,c}|.",
    74: r"\mathrm{RMSE}=\sqrt{\mathrm{MSE}}.",
}

INLINE = {
    24: {
        "xt∈RC": r"\mathbf{x}_t\in\mathbb{R}^{C}",
        "Xn": r"\mathbf{X}_n",
        "Yn": r"\mathbf{Y}_n",
    },
    27: {
        "Ŷn=fθ(Xn)": r"\widehat{\mathbf{Y}}_n=f_{\theta}(\mathbf{X}_n)",
        "vt,c": r"v_{t,c}",
    },
    31: {
        "μctr": r"\mu_c^{\mathrm{tr}}",
        "σctr": r"\sigma_c^{\mathrm{tr}}",
    },
    39: {"X̃(s)": r"\widetilde{\mathbf{X}}^{(s)}"},
    42: {
        "T(s)": r"\mathbf{T}^{(s)}",
        "R(s)": r"\mathbf{R}^{(s)}",
    },
    44: {
        "T(s)": r"\mathbf{T}^{(s)}",
        "R(s)": r"\mathbf{R}^{(s)}",
    },
    47: {
        "zn=[μn,σn,ℓn,dn]": r"\mathbf{z}_n=[\mu_n,\sigma_n,\ell_n,d_n]",
        "Xn": r"\mathbf{X}_n",
    },
    50: {
        "μn": r"\mu_n",
        "σn": r"\sigma_n",
        "ℓn": r"\ell_n",
        "dn": r"d_n",
    },
    57: {"Ŷn": r"\widehat{\mathbf{Y}}_n"},
}


def replace_text(paragraph, old, new):
    assert old in paragraph.runs[0].text
    paragraph.runs[0].text = paragraph.runs[0].text.replace(old, new)


document = Document(SOURCE)
paragraphs = document.paragraphs
assert len(paragraphs) == 114 and len(document.inline_shapes) == 5

# Give the replicated endpoint sequence an explicit name before Eq. (6).
replace_text(
    paragraphs[39],
    "边界位置采用端点复制方式补齐，从而保持分解前后的时间长度一致：",
    "边界位置采用端点复制方式补齐，记补齐后的序列为X̃(s)；越界索引取最近端点值，从而保持分解前后的时间长度一致：",
)

# List every control and distinguish independent architecture from shared parameters.
replace_text(
    paragraphs[19],
    "对照保留相同的尺度和专家",
    "对照采用相同的尺度与专家结构，并分别训练",
)
replace_text(
    paragraphs[45],
    "单尺度预测、固定等权、全局可学习固定权重和样本级路由四种设置",
    "单尺度预测、固定等权、全局可学习固定权重、恒定输入路由和样本级路由五种设置",
)
replace_text(
    paragraphs[62],
    "单尺度DLinear、固定等权多尺度、全局可学习固定权重模型和本文方法",
    "单尺度DLinear、固定等权多尺度、全局可学习固定权重模型、恒定输入路由模型和本文方法",
)
replace_text(
    paragraphs[64],
    "DLinear和LightTS提供相同预测任务下的外部误差参照；固定等权、全局可学习常权重、恒定输入路由与样本级路由模型共享三个尺度专家，",
    "LightTS提供独立结构的外部基线，DLinear提供单尺度参照；固定等权、全局可学习常权重、恒定输入路由与样本级路由模型采用相同的三个尺度专家结构，并分别训练，",
)

latex_items = list(EQUATIONS.values()) + [
    latex for mapping in INLINE.values() for latex in mapping.values()
]
unique_latex = list(dict.fromkeys(latex_items))
with tempfile.TemporaryDirectory() as temporary:
    temp = Path(temporary)
    markdown = temp / "math.md"
    math_docx = temp / "math.docx"
    markdown.write_text("\n\n".join(f"$${latex}$$" for latex in unique_latex), encoding="utf-8")
    subprocess.run([str(PANDOC), str(markdown), "-o", str(math_docx)], check=True)
    source_math = Document(math_docx)
    assert len(source_math.paragraphs) == len(unique_latex)
    math = {}
    for latex, paragraph in zip(unique_latex, source_math.paragraphs):
        nodes = paragraph._p.xpath(".//m:oMath")
        assert len(nodes) == 1, latex
        math[latex] = nodes[0]

    for index, latex in EQUATIONS.items():
        old_nodes = paragraphs[index]._p.xpath(".//m:oMath")
        assert len(old_nodes) == 1, index
        old_nodes[0].getparent().replace(old_nodes[0], deepcopy(math[latex]))

    for index, replacements in INLINE.items():
        paragraph = paragraphs[index]
        original = paragraph.text
        for token in replacements:
            assert token in original, (index, token)
        parts = re.split("(" + "|".join(re.escape(s) for s in sorted(replacements, key=len, reverse=True)) + ")", original)
        paragraph.clear()
        for part in parts:
            if not part:
                continue
            if part in replacements:
                paragraph._p.append(deepcopy(math[replacements[part]]))
            else:
                paragraph.add_run(part)

# In the source, the explanatory lead-in followed rather than preceded Eq. (11).
paragraphs[51]._p.addprevious(paragraphs[52]._p)
# Keep the Softmax explanation with Eq. (12) across the page boundary.
paragraphs[53].paragraph_format.keep_with_next = True
paragraphs[53].paragraph_format.keep_together = True

document.save(TARGET)
check = Document(TARGET)
assert len(check.paragraphs) == 114 and len(check.inline_shapes) == 5
for number in range(1, 17):
    matches = [p for p in check.paragraphs if p.text.strip().endswith(f"({number})")]
    assert len(matches) == 1 and len(matches[0]._p.xpath(".//m:oMath")) == 1, number
assert check.paragraphs[51].text.endswith("其尺度分数计算为")
assert check.paragraphs[52].text.strip().endswith("(11)")
print(TARGET)
