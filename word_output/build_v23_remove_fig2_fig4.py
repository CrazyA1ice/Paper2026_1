from pathlib import Path
import re

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.22v22.docx"
TARGET = ROOT / "word_output" / "2026.09.22v23.docx"

doc = Document(SOURCE)


def paragraph_texts():
    return [p.text for p in doc.paragraphs]


def find_exact(text: str):
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph {text!r}, got {len(matches)}")
    return matches[0]


def find_prefix(prefix: str):
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one prefix {prefix!r}, got {len(matches)}")
    return matches[0]


def remove_paragraph(p):
    el = p._element
    parent = el.getparent()
    if parent is None:
        raise RuntimeError("Paragraph has no parent")
    parent.remove(el)


def remove_picture_before_caption(caption_text: str):
    cap = find_exact(caption_text)
    prev = cap._element.getprevious()
    searched = 0
    while prev is not None and searched < 4:
        if prev.tag.endswith("}p"):
            xml = prev.xml
            text = "".join(prev.itertext()).strip()
            if "<w:drawing" in xml or "<pic:pic" in xml or "<a:blip" in xml:
                prev.getparent().remove(prev)
                return
            if text:
                break
            searched += 1
        prev = prev.getprevious()
    raise RuntimeError(f"Could not find picture paragraph before caption: {caption_text}")


# --- Remove current Figure 2 and the two paragraphs written specifically for it.
fig2_caption = "图2 Abilene单个OD变量的多尺度序列构造与平均池化示意"
remove_picture_before_caption(fig2_caption)
remove_paragraph(find_exact(fig2_caption))
remove_paragraph(find_prefix("为直观说明式（5）的尺度变换"))
remove_paragraph(find_prefix("尺度1保留原始96步输入；尺度2对相邻2点取平均后得到48步序列；尺度4"))

# --- Remove current Figure 4 and the two paragraphs written specifically for it.
fig4_caption = "图4 Abilene与GÉANT代表性测试窗口的24步预测及逐步误差差值"
remove_picture_before_caption(fig4_caption)
remove_paragraph(find_exact(fig4_caption))
remove_paragraph(find_prefix("平均指标能够反映整体误差水平，但不容易展示24步预测过程中预测轨迹与真实序列之间的局部差异"))
remove_paragraph(find_prefix("图4(a)、(b)展示真实值、DLinear和本文方法在24个预测步上的标准化预测轨迹"))

# --- Rewrite the experiment-chain summary so it no longer depends on either deleted figure.
summary = find_prefix("综合表2、表3和图3—图6")
new_summary = (
    "综合表2、表3以及图3、图5和图6可以形成一条较清晰的实验链。"
    "表2和图3用于比较本文方法与外部轻量基线的总体误差水平；"
    "表3和图5进一步控制专家结构，考察固定等权与样本级权重之间的差异；"
    "图6则从模型内部的尺度分配结果说明路由器在测试样本上形成了非等权的融合方式。"
    "各部分分别对应总体性能、核心结构消融和内部权重行为，"
    "与第1节提出的多尺度专家和样本级路由设计相互对应。"
)
if summary.runs:
    summary.runs[0].text = new_summary
    for run in summary.runs[1:]:
        run.text = ""
else:
    summary.add_run(new_summary)

# Guardrails: the deleted figures and their direct textual references must be completely absent.
full = "\n".join(paragraph_texts())
for forbidden in [
    "图2 Abilene单个OD变量的多尺度序列构造与平均池化示意",
    "图4 Abilene与GÉANT代表性测试窗口的24步预测及逐步误差差值",
    "为直观说明式（5）的尺度变换",
    "逐步误差差值ΔAE",
    "图4(a)、(b)",
]:
    assert forbidden not in full, forbidden

# Keep figure numbers 2 and 4 intentionally vacant for the replacement figures.
assert re.search(r"图2(?!\d)", full) is None, "A Figure 2 reference remains"
assert re.search(r"图4(?!\d)", full) is None, "A Figure 4 reference remains"

# Unrelated manuscript structure should remain intact.
assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 4
for kept in [
    "图1 自适应多尺度网络流量预测方法框架",
    "图3 Abilene与GÉANT上三随机种子外部基线配对结果",
    "图5 固定等权与自适应权重的八随机种子配对MSE及配对差值分布",
    "图6 八随机种子下样本级自适应路由权重分布",
]:
    assert kept in full, kept

doc.save(TARGET)
print(f"Saved {TARGET}")
print("tables", len(doc.tables), "figures", len(doc.inline_shapes), "paragraphs", len(doc.paragraphs))
