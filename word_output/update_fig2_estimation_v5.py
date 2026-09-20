from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "word_output" / "2026.09.20v5.docx"
FIGURE_2 = (
    ROOT
    / "abilene_forecasting"
    / "paper_figures"
    / "exports"
    / "color"
    / "fig2_core_weight_estimation_color.png"
)


def find_one(document: Document, startswith: str):
    matches = [p for p in document.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph starting with {startswith!r}, found {len(matches)}"
        )
    return matches[0]


def set_paragraph_text_like(paragraph, text: str) -> None:
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_note_after(paragraph, text: str):
    # Remove an older Figure 2 note if the script is rerun.
    next_el = paragraph._p.getnext()
    if next_el is not None:
        next_text = "".join(
            node.text or "" for node in next_el.iter() if node.tag.endswith("}t")
        ).strip()
        if next_text.startswith("注：上排为8个随机种子"):
            next_el.getparent().remove(next_el)

    note = paragraph._parent.add_paragraph()
    note._p.getparent().remove(note._p)
    paragraph._p.addnext(note._p)
    run = note.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(7.5)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_before = Pt(0)
    note.paragraph_format.space_after = Pt(3)
    note.paragraph_format.line_spacing = 1.0
    return note


def replace_figure(document: Document) -> None:
    if len(document.inline_shapes) != 3:
        raise RuntimeError(f"Expected 3 inline figures, found {len(document.inline_shapes)}")
    shape = document.inline_shapes[1]
    blip = shape._inline.xpath(".//a:blip")[0]
    embed = blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    image_part = document.part.related_parts[embed]
    image_part._blob = FIGURE_2.read_bytes()

    with Image.open(FIGURE_2) as image:
        px_w, px_h = image.size
    shape.width = Inches(6.30)
    shape.height = Inches(6.30 * px_h / px_w)


def main() -> None:
    document = Document(DOCX)
    replace_figure(document)

    analysis = find_one(document, "同种子配对结果")
    set_paragraph_text_like(
        analysis,
        "同种子配对结果及配对差值分布见图2。图2上排保留随机种子42—49的逐一配对MSE，"
        "下排以ΔMSE=固定等权−自适应权重表示效应方向。Abilene中，自适应模型在8个种子中的6个取得更低MSE，"
        "平均ΔMSE为0.00320，95%置信区间为[-0.00020,0.00660]，配对t检验p=0.061；"
        "GÉANT中为7个，平均ΔMSE为0.02094，95%置信区间为[-0.00958,0.05146]，p=0.149。"
        "两数据集的平均差值均为正，说明样本级自适应加权在平均意义上降低了MSE；但相对固定等权的95%置信区间均跨越0，"
        "因此不能表述为两数据集均达到统计显著。结合表3可见，Abilene上的跨种子波动较小，而GÉANT仍存在更明显的稳定性提升空间。"
    )

    caption = find_one(document, "图2 ")
    set_paragraph_text_like(
        caption,
        "图2 固定等权与自适应权重的八随机种子配对MSE及配对差值分布",
    )

    insert_note_after(
        caption,
        "注：a为Abilene，b为GÉANT。上排灰线连接同一随机种子下的固定等权与自适应权重结果，"
        "黑色菱形及误差线表示均值±标准差；下排圆点为各随机种子的ΔMSE，半小提琴表示配对差值的核密度分布，"
        "黑色圆点及误差线表示平均ΔMSE及其95%置信区间。ΔMSE>0表示自适应权重模型的MSE更低。"
    )

    # Guardrails.
    text = "\n".join(p.text for p in document.paragraphs)
    assert "图2 固定等权与自适应权重的八随机种子配对MSE及配对差值分布" in text
    assert "平均ΔMSE为0.00320" in text
    assert "平均ΔMSE为0.02094" in text
    assert "ΔMSE>0表示自适应权重模型的MSE更低" in text
    assert len(document.inline_shapes) == 3

    tmp = DOCX.with_suffix(".tmp.docx")
    document.save(tmp)
    tmp.replace(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
