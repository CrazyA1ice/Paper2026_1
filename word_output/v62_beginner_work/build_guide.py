from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import re, csv, json, hashlib
import numpy as np
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\Quant\Desktop\Paper2026_1')
WORK=ROOT/'word_output/v62_beginner_work'
SOURCE=ROOT/'word_output/2026.10.03v62.docx'
OUT=ROOT/'word_output/2026.10.03v62_零基础逐图逐式逐表说明.docx'
blocks=json.loads((WORK/'source_blocks.json').read_text(encoding='utf-8'))
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
sec.footer_distance=Inches(.3)
for name,size,font in [('Normal',12,'SimSun'),('Title',23,'Microsoft YaHei'),('Heading 1',17,'Microsoft YaHei'),('Heading 2',14,'Microsoft YaHei'),('Heading 3',12,'Microsoft YaHei')]:
 s=doc.styles[name];s.font.name=font;s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font)
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.3
 if name.startswith('Heading'):
  s.paragraph_format.space_before=Pt(13);s.paragraph_format.keep_with_next=True
doc.styles['Normal'].paragraph_format.widow_control=True
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='从零读懂网络流量预测论文'
doc.core_properties.subject='2026.10.03v62论文零基础逐图逐式逐表说明'
doc.core_properties.author=''
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('第 ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld);foot.add_run(' 页')
for r in foot.runs:r.font.size=Pt(9)

def para(text,style=None):
 p=doc.add_paragraph(text,style);p.paragraph_format.keep_with_next=False if style is None else True
 return p
def heading(text,level=2):return doc.add_heading(text,level)
def table(headers,rows,widths=None,size=10.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=widths or [6.9/len(headers)]*len(headers)
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   marg=OxmlElement('w:tcMar')
   for edge,value in [('top',70),('bottom',70),('left',70),('right',70)]:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');marg.append(el)
   pr.append(marg)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DEE8EE' if ri==0 else ('F5F7F8' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
    p.paragraph_format.line_spacing=1.15
    p.paragraph_format.keep_with_next=(ri==0 or (len(rows)<=12 and ri<len(t.rows)-1))
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(str(cell.text))<22 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.font.size=Pt(size);r.bold=ri==0
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t

ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
srcxml=etree.parse(str(WORK/'document.xml'))
equations={}
for p in srcxml.findall('.//w:body/w:p',ns):
 txt=''.join(p.xpath('.//w:t/text() | .//m:t/text()',namespaces=ns))
 match=re.search(r'\((\d+)\)$',txt)
 if match and p.find('.//m:oMath',ns) is not None:equations[int(match[1])]=p
def equation(num):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(9)
 source=equations[num]
 for el in source:
  if etree.QName(el).localname=='pPr':continue
  copy=deepcopy(el)
  # The source tab stops belong to its original page layout. Use a shared equation then its number.
  for tab in copy.xpath('.//w:tab',namespaces=ns):tab.getparent().remove(tab)
  p._p.append(copy)
 # Preserve all mathematical elements; only the source paragraph layout is replaced.

def read_csv(path):return list(csv.DictReader((ROOT/path).open(encoding='utf-8-sig')))
models=['DLinear','固定等权','恒定输入路由','本文方法']
datasets=[('abilene','Abilene'),('geant','GÉANT')]
fig3=read_csv('abilene_forecasting/paper_figures/source_data_v24/fig5_relative_improvement_8seeds.csv')
cal=read_csv('paper_figures/fig4_calibration_trajectories_v39/calibration_source_8seeds_v39.csv')
f2=read_csv('abilene_forecasting/paper_figures/source_data_v25/fig3_lightts_improvement_seed42_44.csv')
weights={ds:[np.loadtxt(ROOT/f'abilene_forecasting/results/{ds}/dlinear_scale_seed{s}/router_weights.csv',delimiter=',') for s in range(42,50)] for ds,_ in datasets}
audit={'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'equations':[],'figures':[],'tables':[],'point_counts':{}}

def glossary():
 table(['符号或名称','可以这样理解'],[
 ['L与H','过去看几步、未来猜几步；本文96和24'],['C','有效OD列数；144或524'],['n与N','某个窗口编号、训练样本总数'],['t与h与c','时间位置、未来位置、OD编号'],['Nte','测试窗口数；813或451'],['X与Y与Ŷ','历史输入、真实未来、模型预测'],['v与u与x','原流量、对数流量、标准化流量'],['S与s','尺度集合{1,2,4}、其中一个尺度'],['T与R','趋势、原数减趋势后的余项'],['W与b与θ','线性系数、偏移量、模型参数总称'],['μ与σ','平均数与标准差，需看下标辨别用途'],['ℓ与d','最后一小时的平均状态、平均变化强度'],['z与a与α','四个特征、三个分数、三个融合份额'],['MSE与MAE与RMSE','三把误差尺子；都越小越好'],['DLinear','分解线性模型，本论文的专家主干'],['MLP与ReLU与Softmax','小网络、去掉负激活、把分数变成份额'],['Adam','按错误调整参数的优化器'],['Abilene与GÉANT','两套网络流量数据的名称']],[1.55,5.35],11)

def source_table(num):
 b=[x for x in blocks if x['type']=='tbl'][num-1]
 rows=b['rows'];widths={1:[.8,.65,.8,1.9,1.85,.9],2:[.85,1.65,.8,1.8,1.8],3:[.85,1.65,.8,1.8,1.8]}[num]
 table(rows[0],rows[1:],widths,10.5)
 audit['tables'].append(num)

def explain_table(num):
 rows=[b for b in blocks if b['type']=='tbl'][num-1]['rows'][1:]
 for ds,model,params,mse,mae in rows:
  mm,ms=mse.split('±');am,ass=mae.split('±')
  para(f'{ds}的{model}：有{params}个可训练参数；MSE平均{mm}，种子间标准差{ms}；MAE平均{am}，种子间标准差{ass}。这一行只概括该数据集上这个模型的重复实验。')

def fig2data():
 for ds,label in datasets:
  heading(label+'热图中的全部数字',3)
  rows=[]
  for metric in ['MSE','MAE','RMSE']:
   vals=[float(next(r['relative_improvement_pct'] for r in f2 if r['dataset']==ds and r['seed']==str(s) and r['metric']==metric)) for s in range(42,45)]
   rows.append([metric]+[f'{x:.1f}%' for x in vals]+[f'{np.mean(vals):.1f}%'])
  table(['指标','seed42','seed43','seed44','均值'],rows,[.9,1.5,1.5,1.5,1.5],11)

def fig3means():
 rows=[]
 for ds,label in datasets:
  for base in ['DLinear','固定等权']:
   vals=[np.mean([float(r['relative_improvement_pct']) for r in fig3 if r['dataset']==ds and r['baseline']==base and r['metric']==met]) for met in ['MSE','MAE','RMSE']]
   rows.append([label,base]+[f'{x:.2f}%' for x in vals])
 table(['数据集','比较基线','MSE','MAE','RMSE'],rows,[1.2,1.5,1.4,1.4,1.4],11)

def fig3points():
 for ds,label in datasets:
  for base in ['DLinear','固定等权']:
   heading(label+'相对'+base,2)
   rows=[];cols=[]
   for met in ['MSE','MAE','RMSE']:
    part=sorted([r for r in fig3 if r['dataset']==ds and r['baseline']==base and r['metric']==met],key=lambda r:int(r['seed']))
    assert [int(r['seed']) for r in part]==list(range(42,50))
    vals=np.array([float(r['relative_improvement_pct']) for r in part]);cols.append(vals)
   for i,s in enumerate(range(42,50)):rows.append([s]+[f'{v[i]:.2f}' for v in cols])
   rows.append(['均值']+[f'{v.mean():.2f}' for v in cols])
   ci=[2.364624251*v.std(ddof=1)/np.sqrt(8) for v in cols]
   rows.append(['95%区间']+[f'[{v.mean()-c:.2f}, {v.mean()+c:.2f}]' for v,c in zip(cols,ci)])
   table(['种子或统计量','MSE改进%','MAE改进%','RMSE改进%'],rows,[1.3,1.85,1.85,1.9],11)
 audit['point_counts']['figure3']=96

def fig5means():
 rows=[]
 for ds,label in datasets:
  mean=np.array([r.mean(axis=0) for r in weights[ds]]).mean(axis=0)
  rows.append([label]+[f'{x:.3f} 即{x*100:.1f}%' for x in mean])
 table(['数据集','尺度1','尺度2','尺度4'],rows,[1.2,1.9,1.9,1.9],11)

def calibration():
 for ds,label in datasets:
  heading('图4'+('a' if ds=='abilene' else 'b')+' '+label+'的40个分组',2)
  rows=[]
  for b in range(40):
   group=[r for r in cal if r['dataset']==ds and r['bin']==str(b)]
   counts=set(int(r['n_elements']) for r in group);true=np.array([float(r['true_mean']) for r in group])
   assert len(counts)==1 and np.allclose(true,true[0],atol=1e-12,rtol=0) and len(group)==32
   item=[b+1,counts.pop(),f'{true.mean():.3f}']
   for model in models:
    val=np.array([float(r['predicted_mean']) for r in group if r['model']==model])
    assert len(val)==8
    item.append(f'{val.mean():.3f}\n±{val.std(ddof=1):.3f}')
   rows.append(item)
  table(['组号','每种子\n点数','真实\n均值','DLinear\n均值±SD','固定等权\n均值±SD','恒定路由\n均值±SD','本文方法\n均值±SD'],rows,[.45,.9,.7,1.2,1.2,1.2,1.25],10)
  total=sum(int(r[1]) for r in rows)
  expected=(813*24*144 if ds=='abilene' else 451*24*524)
  assert total==expected
  para(f'各组每种子的点数相加为{total:,}，等于'+('813×24×144' if ds=='abilene' else '451×24×524')+'。8个种子各自使用这些预测位置，再计算均值和标准差。SD是标准差的缩写。')
 audit['point_counts']['figure4_calibration']=80

def trajectories():
 configs=[('c','abilene',46,278),('d','geant',296,397),('e','abilene',46,618),('f','geant',296,169)]
 for panel,ds,ch,window in configs:
  label=dict(datasets)[ds];heading(f'图4{panel} {label} OD{ch+1} 窗口{window+1}',2)
  arrays={};truth=None
  for model,directory in [('DLinear','dlinear'),('固定等权','dlinear_scale_static'),('本文方法','dlinear_scale')]:
   with np.load(ROOT/f'abilene_forecasting/results/{ds}/{directory}_seed42/predictions.npz') as z:
    arrays[model]=z['prediction'][window,:,ch].copy()
    yy=z['truth'][window,:,ch].copy()
    if truth is None:truth=yy
    else:assert np.array_equal(truth,yy)
  with np.load(ROOT/f'abilene_forecasting/results/constant_router_ablation/{ds}/dlinear_scale_constant_router_seed42/predictions.npz') as z:
   assert np.array_equal(truth,z['truth'][window,:,ch]);arrays['恒定输入路由']=z['prediction'][window,:,ch].copy()
  rows=[[h+1,f'{truth[h]:.4f}']+[f'{arrays[m][h]:.4f}' for m in models] for h in range(24)]
  table(['未来小时','真实','DLinear','固定等权','恒定路由','本文方法'],rows,[.7,1.24,1.24,1.24,1.24,1.24],10.5)
  errs={m:float(np.mean((arrays[m]-truth)**2)) for m in models}
  para('这24步的MSE：'+'；'.join(f'{m} {errs[m]:.6f}' for m in models)+'。')
  best=min(errs,key=errs.get);para('本例中24步MSE最低的是'+best+'。')
  audit.setdefault('trajectory_mse',{})[panel]=errs
 audit['point_counts']['figure4_trajectory_values']=4*24*5

def weight_tables():
 for ds,label in datasets:
  heading(label+'的权重标记',2)
  means=np.array([r.mean(axis=0) for r in weights[ds]]);pooled=np.vstack(weights[ds])
  rows=[[s]+[f'{x:.3f}' for x in vals] for s,vals in zip(range(42,50),means)]
  for name,vals in [('均值 黑菱形',means.mean(axis=0)),('标准差 误差线半长',means.std(axis=0,ddof=1)),('汇总中位数 短横线',np.median(pooled,axis=0))]:
   rows.append([name]+[f'{x:.3f}' for x in vals])
  table(['种子或统计量','尺度1 α1','尺度2 α2','尺度4 α4'],rows,[2.1,1.6,1.6,1.6],11)
  para(f'小提琴汇总8×{weights[ds][0].shape[0]}={len(pooled)}组窗口权重；每组有3个数，对应三个尺度。黑菱形和中位数可能不同，因为它们的算法不同。')
 audit['point_counts']['figure5_seed_mean_points']=48

def sourceindex():
 table(['原文对象','原文页码','本说明书对应位置'],[
 ['研究问题与总体流程','1至2','第一至二课'],['图1','2','第五课图1'],['式1至6','3','第四课式1至6'],['式7至13','4','第四课式7至13'],['式14至16与表1','5','第四课及第六课表1'],['表2与图2','6','第五课图2及第六课表2'],['表3与图3','7','第六课表3及第五课图3'],['图4','8','第五课图4及附录B C'],['图5及运行开销','9','第五课图5及第七课'],['高需求分析与结语','10','第七至八课'],['参考文献','10','第八课参考文献角色']],[2,1,3.9],11)

handlers={
 'GLOSSARY':glossary,'FIG2DATA':fig2data,'FIG3MEANS':fig3means,'FIG3POINTS':fig3points,
 'FIG5MEANS':fig5means,'CALIBRATION':calibration,'TRAJECTORIES':trajectories,'WEIGHTS':weight_tables,'SOURCEINDEX':sourceindex,
 'TABLE2EXPLAIN':lambda:explain_table(2),'TABLE3EXPLAIN':lambda:explain_table(3),
 'TIMING':lambda:table(['模型','Abilene每批ms','GÉANT每批ms'],[['DLinear','0.77','0.70'],['固定等权多尺度','2.38','2.10'],['本文方法','2.65','2.61']],[2.5,2.2,2.2],11),
 'HIGHDEMAND':lambda:table(['数据集','高需求MAE','平均低估幅度','两指标相对\n全局权重降幅'],[['Abilene','0.147±0.013','0.137±0.017','17.22% / 20.49%'],['GÉANT','0.632±0.060','0.577±0.080','3.59% / 6.30%']],[1.1,1.8,1.8,2.2],11)
}
figfiles={1:'image1.png',2:'image10.png',3:'image15.png',4:'image20.png',5:'image16.png'}
for line in (WORK/'guide.md').read_text(encoding='utf-8').splitlines():
 if not line:continue
 line=line.replace('꜀','c')
 if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):
  p=heading(line[3:],1);p.paragraph_format.page_break_before=line[3:].startswith('附录')
 elif line.startswith('### '):heading(line[4:],2)
 elif line.startswith('{{EQ:'):
  n=int(line[5:-2]);equation(n);audit['equations'].append(n)
 elif line.startswith('{{FIG:'):
  doc.paragraphs[-1].paragraph_format.keep_with_next=True
  n=int(line[6:-2]);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.add_run().add_picture(str(WORK/figfiles[n]),width=Inches(6.9))
  p.paragraph_format.space_after=Pt(9);audit['figures'].append(n)
 elif line.startswith('{{TABLE:'):source_table(int(line[8:-2]))
 elif line.startswith('{{'):handlers[line[2:-2]]()
 else:para(line)
assert audit['equations']==list(range(1,17));assert audit['figures']==list(range(1,6));assert audit['tables']==[1,2,3]
doc.save(OUT)
# Mathematical content must exactly match the source, including hats, radicals and limits.
with ZipFile(OUT) as z:generated=etree.fromstring(z.read('word/document.xml'))
orig_math=[etree.tostring(e,method='c14n') for n in range(1,17) for e in equations[n].xpath('.//m:oMath',namespaces=ns)]
new_math=[etree.tostring(e,method='c14n') for e in generated.xpath('.//m:oMath',namespaces=ns)]
# C14N namespace declarations may differ; compare all tag, text and attribute content instead.
def signature(e):return (e.tag,e.text,tuple(sorted(e.attrib.items())),tuple(signature(x) for x in e))
assert [signature(e) for n in range(1,17) for e in equations[n].xpath('.//m:oMath',namespaces=ns)]==[signature(e) for e in generated.xpath('.//m:oMath',namespaces=ns)]
audit['native_math_exactly_preserved']=True
assert '{{' not in ''.join(generated.xpath('.//w:t/text()',namespaces=ns))
audit['output']=str(OUT);audit['output_sha256']=hashlib.sha256(OUT.read_bytes()).hexdigest()
(WORK/'content_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved',OUT.name.encode('ascii',errors='backslashreplace').decode(), 'paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'characters',sum(len(p.text) for p in doc.paragraphs))
