"""Replace Figure 3 with the paired global-weight comparison in v33."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.26v33.docx"
TARGET = ROOT / "word_output/2026.09.26v34.docx"
FIGURE = ROOT / "paper_figures/fig3_global_vs_adaptive_paired.png"


def replace_text(paragraph, old, new):
    assert old in paragraph.text, old
    updated = paragraph.text.replace(old, new)
    for run in paragraph.runs:
        run.text = ""
    paragraph.runs[0].text = updated


doc = Document(SOURCE)
p = doc.paragraphs
assert len(p) == 114 and len(doc.inline_shapes) == 5 and len(doc.tables) == 3
assert p[88].text.startswith("图3 Abilene与GÉANT上本文方法相对两类基线")
assert len(p[87]._p.xpath(".//w:drawing")) == 1

replace_text(
    p[79],
    "DLinear和固定等权模型的内部比较放在表3和图3中。",
    "DLinear和固定等权模型的内部比较见表3。",
)
replace_text(
    p[84],
    "Abilene优势较小，GÉANT区间跨越零，不能据此断言跨初始化稳定优于全局权重。",
    "Abilene优势较小，GÉANT区间跨越零，不能据此断言跨初始化稳定优于全局权重。图3展示这一更直接对照的逐种子结果。",
)
p[86].text = (
    "图3以随机种子为配对单位，给出本文方法相对全局可学习固定权重模型的MSE降低率。"
    "每个圆点代表一个随机种子，正值表示样本级路由的MSE更低；下方菱形和误差线分别表示"
    "8个种子的平均降低率及基于t分布的95%置信区间。Abilene有6/8个种子改善，平均降低1.80%，"
    "置信区间为[0.01%,3.58%]；GÉANT有7/8个种子改善，平均降低4.31%，"
    "但种子43的MSE升高13.22%，使置信区间[−2.51%,11.13%]跨越零。"
    "因此，该图支持部分设置下的额外收益，不能证明对不同初始化均稳定有效。"
)
for run in list(p[87].runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
p[87].add_run().add_picture(str(FIGURE), width=Inches(6.25))
p[88].text = "图3 样本级路由相对全局可学习固定权重模型的八随机种子配对MSE降低率"
p[89].text = (
    "固定等权模型只能检验等权融合是否足够；全局可学习固定权重模型进一步控制了“权重不等”这一因素。"
    "图3中的GÉANT种子43与多数种子方向相反，提醒我们平均改进不足以说明稳定性。"
    "为检查改善是否集中于少数OD通道，图4进一步统计每条通道在8个种子下相对全局权重模型的MSE胜出次数。"
)

doc.save(TARGET)
check = Document(TARGET)
assert len(check.paragraphs) == 114 and len(check.inline_shapes) == 5
assert "图3 样本级路由相对全局可学习固定权重模型" in check.paragraphs[88].text
assert "图3汇总随机种子42～49下的相对改进结果" not in "\n".join(x.text for x in check.paragraphs)
assert "H=12" not in "\n".join(x.text for x in check.paragraphs)
print(TARGET)
