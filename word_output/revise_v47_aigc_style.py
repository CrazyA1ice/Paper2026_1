from pathlib import Path
from zipfile import ZipFile
from difflib import SequenceMatcher
from collections import Counter
import hashlib
import json
import re
import subprocess
import sys

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / '2026.09.27v46.docx'
TARGET = ROOT / '2026.09.28v47.docx'
SCAN_SCRIPT = Path(r'C:\Users\Quant\.codex\skills\aigc-reduce\scripts\aigc_scan.py')
PROSE_INDICES = [3,12,13,14,15,17,19,20,22,25,29,30,34,36,37,39,42,44,45,47,50,51,53,55,57,59,61,62,64,66,68,69,71,72,75,77,79,82,84,86,87,90,93,95,98,99,100,102]

# Each operation has one exact source span and an explicit editorial reason.
EDITS = [
    (12, '后一时段的OD流量估计可为网络管理系统预先调整资源提供依据。', '', '删去与前一句网络管理用途重复的说明。'),
    (13, '网络流量预测方法已由传统统计模型扩展至机器学习和深度学习模型。', '传统统计模型之外，机器学习和深度学习模型也已用于网络流量预测。', '调整背景句的语序，保留方法发展范围。'),
    (14, '模型结构开销也不可忽略。', '多尺度信息的利用还需考虑模型结构开销。', '把笼统强调改成与多尺度设计直接相关的约束。'),
    (15, '不同输入窗口的流量状态存在差异：某些窗口以短时波动为主，另一些窗口则表现为更平缓的趋势。', '短时波动与平缓趋势在不同输入窗口中的占比并不相同。', '合并对称列举，保留窗口状态差异。'),
    (17, '本文采用样本级自适应多尺度结构处理上述问题。', '针对输入窗口之间的状态差异，本文采用样本级自适应多尺度结构。', '明确方法对应的具体问题，减少泛指。'),
    (30, '预处理后的输入样本仍为原始时间分辨率。为使专家分别观察细粒度变化与较平滑的趋势，下一步从同一历史窗口构造不同尺度的表示。', '预处理不改变输入样本的时间分辨率。随后从同一历史窗口构造不同尺度的表示，供各专家分别观察细粒度变化与较平滑的趋势。', '以处理步骤衔接两节，调整长句语序。'),
    (36, '三者分别对应原始、中等和平滑程度更高的时间表示。', '', '删去与随后三个尺度的具体解释重复的概述。'),
    (37, '本文并不将{1,2,4}作为经过穷举得到的最优尺度组合，而是把它作为兼顾细粒度信息、趋势信息和模型规模的一组固定设计。', '尺度组合{1,2,4}是一组兼顾细粒度信息、趋势信息和模型规模的固定设计，未经穷举搜索，本文不将其视为最优尺度组合。', '改写成设计事实与适用边界，保留未经最优搜索的限定。'),
    (45, '采用DLinear作为尺度专家，有助于区分预测主干和尺度融合各自的影响。', '预测主干和尺度融合的影响需分别考察，因此采用DLinear作为尺度专家。', '先说明选择依据，再交代设计。'),
    (51, '路由器仅计算三个尺度的融合分数，无需再次预测完整的OD流量，因此本文未引入循环网络、注意力网络等额外时序编码器。', '路由器的任务是计算三个尺度的融合分数，无需再次预测完整的OD流量。本文未引入循环网络、注意力网络等额外时序编码器。', '拆分任务定义与结构选择，避免原因与方法堆叠。'),
    (61, '；两类固定权重对照均不能单独证明逐窗口权重分配的作用。', '。', '保留首层权重无效的结构说明；重复的结论限定在第2.4节和结论中保留。'),
    (62, '多尺度序列决定模型从哪些时间分辨率读取历史信息，尺度专家完成各自的预测，路由网络只负责组合结果。', '', '删去与前一句三个计算环节重复的解释。'),
    (75, 'RMSE是MSE开方后的结果，用于补充呈现误差幅值。', 'RMSE为MSE的平方根，补充描述误差幅值。', '压缩指标解释，保留定义和用途。'),
    (79, '按（LightTS误差−本文方法误差）/LightTS误差计算每次实验的相对降低率后，Abilene上的', '每次实验的相对降低率按（LightTS误差−本文方法误差）/LightTS误差计算。Abilene上的', '把计算定义与数值报告分为两个完整句子。'),
    (87, '半小提琴展示8次结果的分布；黑色菱形和横向误差线', '半小提琴展示8次结果的分布。黑色菱形和横向误差线', '拆开分布和推断量的读图说明。'),
    (90, '由于恒定输入路由的首层输入权重不起作用，且各模型分别训练，因此，这些对照显示', '恒定输入路由的首层输入权重不起作用，各模型也分别训练。这些对照显示', '去掉层叠连接词，保留两项比较局限。'),
    (93, '24个预测步及有效OD通道：按真实值分位点划为最多40个区间，四种模型的曲线', '24个预测步及有效OD通道。真实值按分位点划为最多40个区间，四种模型的曲线', '把统计范围与分组方式分开说明。'),
    (93, '再选择波动程度接近中位数和第90百分位的窗口，未依据模型误差挑选。', '再选择波动程度接近中位数和第90百分位的窗口。上述选择未依据模型误差。', '单独交代案例选择依据，保留避免按误差挑选的说明。'),
    (95, '在全部测试窗口及种子上的汇总分布；彩色散点', '在全部测试窗口及种子上的汇总分布。彩色散点', '拆分小提琴与种子均值的读图定义。'),
    (98, '虽然总体均值偏离1/3，但GÉANT三个尺度的', '总体均值偏离1/3。GÉANT三个尺度的', '分别陈述总体均值和方差来源，保留后文证据边界。'),
    (99, '分别有15、406、392个窗口，本文方法相对', '分别有15、406、392个窗口。本文方法相对', '将样本数量与误差比较分开，降低单句信息负荷。'),
    (100, '因此，本文方法参数量较小，但三个尺度专家及路由计算仍增加了推理开销；上述结果不代表端到端部署时延。', '本文方法参数量较小，三个尺度专家及路由计算仍增加了推理开销。端到端部署时延还包括数据读取与传输，不在上述计时范围内。', '用明确的计时范围替代泛化否定句，保留推理开销和部署边界。'),
]


def extract_prose(document):
    return '\n\n'.join(document.paragraphs[i].text for i in PROSE_INDICES)


def scan(label, text):
    path = ROOT / f'{label}_aigc_scan_input.txt'
    path.write_text(text, encoding='utf-8')
    result = subprocess.run([sys.executable, str(SCAN_SCRIPT), str(path), '--json'], capture_output=True, text=True, encoding='utf-8', check=True)
    (ROOT / f'{label}_aigc_scan.json').write_text(result.stdout, encoding='utf-8')
    return json.loads(result.stdout)


original = Document(SOURCE)
edited = Document(SOURCE)
reasons = {}
for index, old, new, reason in EDITS:
    p = edited.paragraphs[index]
    assert not p._p.xpath('.//m:oMath | .//w:drawing'), index
    assert p.text.count(old) == 1, (index, old)
    assert p.runs[0].text.count(old) == 1, index
    p.runs[0].text = p.runs[0].text.replace(old, new)
    reasons.setdefault(index, []).append(reason)

# Keep short opening fragments from being stranded at the page bottom.
edited.paragraphs[36].paragraph_format.keep_together = True
edited.paragraphs[46].paragraph_format.keep_with_next = True
edited.paragraphs[47].paragraph_format.keep_together = True
edited.paragraphs[69].paragraph_format.widow_control = True

edited.save(TARGET)
verified = Document(TARGET)
assert len(original.paragraphs) == len(verified.paragraphs)
numbers = lambda text: Counter(re.findall(r'\d+(?:\.\d+)?', text))
citations = lambda text: re.findall(r'\[\d+(?:[—～\-,]\d+)*\]', text)
for i, (a, b) in enumerate(zip(original.paragraphs, verified.paragraphs)):
    if i not in reasons:
        assert a.text == b.text, i
    assert numbers(a.text) == numbers(b.text), i
    assert citations(a.text) == citations(b.text), i
    assert re.findall(r'[图表]\d+', a.text) == re.findall(r'[图表]\d+', b.text), i
    assert re.findall(r'“[^”]*”', a.text) == re.findall(r'“[^”]*”', b.text), i

terms = ['DLinear','LightTS','Softmax','MSE','MAE','RMSE','Abilene','GÉANT','Adam','NVIDIA GTX 1060','CESNET TS-Zoo','全局可学习权重','恒定输入路由','固定等权']
all_original = '\n'.join(p.text for p in original.paragraphs)
all_edited = '\n'.join(p.text for p in verified.paragraphs)
for term in terms:
    assert all_original.count(term) == all_edited.count(term), term
assert len(original.tables) == len(verified.tables) == 3
for a, b in zip(original.tables, verified.tables):
    assert [[c.text for c in row.cells] for row in a.rows] == [[c.text for c in row.cells] for row in b.rows]
assert len(original.inline_shapes) == len(verified.inline_shapes) == 5
ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with ZipFile(SOURCE) as a, ZipFile(TARGET) as b:
    media = lambda z: {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
    assert media(a) == media(b)
    math = lambda z: [etree.tostring(x) for x in etree.fromstring(z.read('word/document.xml')).xpath('//m:oMath', namespaces=ns)]
    assert math(a) == math(b)
    for name in ['word/footnotes.xml', 'word/endnotes.xml']:
        if name in a.namelist():
            assert a.read(name) == b.read(name), name

before_text = extract_prose(original)
after_text = extract_prose(verified)
before_scan = scan('v46', before_text)
after_scan = scan('v47', after_text)
diff = SequenceMatcher(None, before_text, after_text, autojunk=False)
changed_chars = sum(max(j-i, l-k) for op,i,j,k,l in diff.get_opcodes() if op != 'equal')
ratio = changed_chars / max(len(before_text), len(after_text)) * 100
coverage = len(reasons) / len(PROSE_INDICES) * 100

lines = [
    '# V46 → V47 模板化语言修订与对比报告', '',
    '## 1 结果与范围', '',
    f'- 输入：{SOURCE.name}；输出：{TARGET.name}。',
    f'- 以中文摘要与正文中的48个段落为检查范围，修改其中{len(reasons)}个段落，共{len(EDITS)}项精确字符串替换；摘要、英文摘要未改。',
    f'- 段落覆盖率：{coverage:.2f}%。覆盖率表示发生修改的段落比例，并非字面修改率或AI检测率。',
    f'- 逐字符差异率：{ratio:.2f}%。按SequenceMatcher对齐文本，将各差异片段的新旧字符长度较大值求和，再除以两版扫描文本较长者；不将整段计为修改。',
    f'- 扫描文本字符数（含段落分隔）：{len(before_text)} → {len(after_text)}。',
    '- 主要操作：压缩重复解释，调整语序，拆分承载多项信息的长句，并用具体范围说明替代泛化否定句。',
    '- 本轮未为达到技能中的40%目标强行替换术语或改写已清楚的段落。保护学术语体、数据和证据边界优先于修改率。',
    '- 本报告未调用知网、万方或其他商业检测平台，不能报告或保证实际AIGC比例下降。下列扫描指标只是本地规则命中与文本统计。', '',
    '## 2 本地扫描前后对比', '',
    '扫描范围：中文摘要与中文正文；不包含英文摘要、标题、图表标题、表格单元格、参考文献、独立公式行。正文中的公式指代和引用保留，扫描器可能将其误记为编号或结构特征。', '',
    '| 指标 | V46 | V47 | 解读 |', '|---|---:|---:|---|',
]
metrics = [
    ('模板句式规则命中数', 'template_phrases', 'count', '规则匹配不证明AI作者身份；原命中为计时边界中的“不代表”。'),
    ('被动语态标记数', 'passive_voice', 'count', '方法中的合规被动表达保留。'),
    ('平均句长', 'burstiness', 'avg_len', '单位为扫描器统计的字符；小数、英文等会影响切句。'),
    ('句长变异系数', 'burstiness', 'cv', '句式统计量，不是AIGC概率。'),
    ('相近段长组数', 'para_symmetry', 'symmetrical_runs', '长度相近在学术文本中很常见，不凭此扩写或删减内容。'),
    ('编号规则命中数', 'nested_numbers', 'count', '正文的公式指代按原稿保留。'),
    ('冒号并列规则命中数', 'colon_lists', 'count', '仅对影响阅读的长句作结构调整。'),
    ('口语及网络用语命中数', 'colloquial', 'count', '前后均应无口语化新增。'),
    ('超限破折号段落数', 'dash_density', 'over_limit_paras', '前后均无超限。'),
]
for label, group, key, comment in metrics:
    a, b = before_scan[group][key], after_scan[group][key]
    if isinstance(a, list): a, b = len(a), len(b)
    lines.append(f'| {label} | {a} | {b} | {comment} |')
lines += ['', '## 3 保护内容核验', '',
    '- 所有段落中的数值及出现次数、引用编号、图表编号、引号内原文一致。',
    '- 14个关键术语的文本与出现次数一致。未将“融合”“优化”“有效”等正常术语按词表机械替换。',
    '- 3张表的全部单元格文本一致，5幅内嵌图的媒体文件哈希一致。',
    '- 全部原生公式XML一致，脚注与尾注文件一致；参考文献及图表标题未改。',
    '- GÉANT置信区间跨零、种子43退步、典型窗口例外、权重方差以种子间差异为主等结论限定均保留。',
    '- 第1.5节重复的证据限定删去一处；第2.4节及结论仍保留完整限定，未削弱不利结果。',
    '- 排版调整：尺度说明段落保持完整；第1.4节标题与首段保持衔接，避免只有一行正文落在页末。', '',
    '## 4 逐段原文与修改稿', '',
]
for index, reasons_list in reasons.items():
    lines += [f'### 段落{index}（Word文档从0计数）', '', '**修改理由：** ' + '；'.join(reasons_list), '', '**原文：**', '', original.paragraphs[index].text, '', '**修改后：**', '', verified.paragraphs[index].text, '']
lines += ['## 5 提交前核验', '', '- 需由作者复核语义与术语；若使用目标期刊认可的检测平台，可将实际报告交回，继续针对具体标记段落处理。', '- 不能依据本地规则把某段判定为AI生成，也不能以本轮字面修改率推算商业平台的检测分数。', '']
(ROOT/'2026.09.28v47_语言修订对比报告.md').write_text('\n'.join(lines), encoding='utf-8')
print(json.dumps({'changed_paragraphs':len(reasons),'operations':len(EDITS),'paragraph_coverage_pct':round(coverage,2),'character_diff_pct':round(ratio,2),'before':before_scan,'after':after_scan},ensure_ascii=False,indent=2))
print(TARGET)
