"""Add the constant-input router ablation to the existing v37 manuscript."""

from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "2026.09.27v37.docx"
TARGET = ROOT / "2026.09.27v38.docx"


def replace_paragraph(document, start, old, new):
    matches = [p for p in document.paragraphs if p.text.startswith(start)]
    assert len(matches) == 1, (start, len(matches))
    paragraph = matches[0]
    assert old in paragraph.text, start
    updated = paragraph.text.replace(old, new, 1)
    runs = paragraph.runs
    first_index = next(index for index, run in enumerate(runs) if run.text)
    runs[first_index].text = updated
    for index, run in enumerate(runs):
        if index != first_index:
            run.text = ""
    assert paragraph.text == updated


def set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    first = paragraph.runs[0]
    first.text = value
    for run in paragraph.runs[1:]:
        run.text = ""


document = Document(SOURCE)

replace_paragraph(
    document,
    "摘要：",
    "8个随机种子的核心消融进一步加入可学习的全局固定权重对照；本文方法相对该对照的逐种子平均MSE改进率分别为1.80%和4.31%，分别有6次和7次取得更低误差。GÉANT的配对95%置信区间跨越零，",
    "8个随机种子的核心消融加入全局可学习权重和恒定输入路由两类固定权重对照；本文方法相对两者在Abilene上的逐种子平均MSE改进率均约为1.80%，在GÉANT上分别为4.31%和4.97%。GÉANT的两项配对95%置信区间均跨越零，",
)
replace_paragraph(
    document,
    "Abstract:",
    "An eight-seed ablation adds a learned global-weight baseline. The mean seed-wise MSE reductions of the adaptive router over this baseline are 1.80% and 4.31%, with improvements in six and seven of eight seeds, respectively. The paired 95% confidence interval for GÉANT includes zero, and most of its pooled router-weight variance arises between seeds.",
    "An eight-seed ablation adds two fixed-weight controls: learned global weights and a router receiving constant input. The adaptive model reduces mean seed-wise MSE over both controls by about 1.80% on Abilene, and by 4.31% and 4.97%, respectively, on GÉANT. Both paired 95% confidence intervals on GÉANT include zero, and most of its pooled router-weight variance arises between seeds.",
)
replace_paragraph(
    document,
    "本文采用样本级自适应多尺度结构",
    "固定等权与全局可学习固定权重模型、多个随机种子",
    "固定等权、全局可学习固定权重与恒定输入路由模型、多个随机种子",
)
replace_paragraph(
    document,
    "图1给出模型的整体流程。",
    "固定等权与全局可学习权重对照保留相同的尺度和专家",
    "固定等权、全局可学习权重与恒定输入路由对照保留相同的尺度和专家",
)
replace_paragraph(
    document,
    "式（13）利用上述权重",
    "为了区分多尺度结构、非等权分配与样本相关分配的作用，设置两种共享相同专家的对照：固定等权模型将三个尺度恒定设为1/3；全局可学习权重模型以三个可训练参数经Softmax得到一组常权重，对所有窗口共用。后者提供对输入相关权重更直接的性能对照，但路由器另有少量可训练参数。",
    "为了区分多尺度结构、非等权分配与样本相关分配的作用，设置三种共享相同专家的对照：固定等权模型将三个尺度恒定设为1/3；全局可学习权重模型以三个可训练参数经Softmax得到一组常权重；恒定输入路由模型保留完整路由网络，但训练和推理时都向路由器输入全零窗口，使全部样本共用一组可学习权重。三种模型的尺度专家均接收真实历史窗口。",
)
replace_paragraph(
    document,
    "参数训练使用Adam优化器",
    "样本级路由模型为8339个参数。路由网络相对固定等权模型增加131个参数，相对全局权重模型增加128个参数。各多尺度模型共享专家结构，全局权重模型因此提供较接近的对照；两者仍有128个参数的差异，不能将全部误差差异直接归因于逐窗口加权。",
    "样本级路由模型和恒定输入路由模型均为8339个标称参数。路由网络相对固定等权模型增加131个参数，相对全局权重模型增加128个参数。恒定输入路由保留相同网络结构，但首层输入权重因输入恒为零而不参与有效计算；两类固定权重对照均不能单独证明逐窗口权重分配的作用。",
)
replace_paragraph(
    document,
    "实验按外部基线",
    "固定等权、全局可学习常权重与样本级路由模型共享三个尺度专家",
    "固定等权、全局可学习常权重、恒定输入路由与样本级路由模型共享三个尺度专家",
)
replace_paragraph(
    document,
    "所有对比模型使用相同的数据划分",
    "DLinear、固定等权多尺度、全局可学习权重和本文方法最多训练30轮",
    "DLinear、固定等权多尺度、全局可学习权重、恒定输入路由和本文方法最多训练30轮",
)
replace_paragraph(
    document,
    "外部基线比较使用随机种子",
    "比较DLinear、固定等权多尺度、全局可学习权重和本文方法。全局权重由等权初始化，与三个专家一同训练",
    "比较DLinear、固定等权多尺度、全局可学习权重、恒定输入路由和本文方法。全局权重由等权初始化，恒定输入路由对所有窗口使用全零路由输入；两者均与三个专家一同训练",
)
replace_paragraph(
    document,
    "实验设置DLinear、LightTS",
    "实验设置DLinear、LightTS、固定等权多尺度和全局可学习权重四类对照。DLinear用于比较单一尺度与多尺度结构；LightTS提供独立轻量模型基线。后两种多尺度对照与本文方法保留相同的尺度集合{1,2,4}和DLinear专家：固定等权模型始终按1/3融合，全局权重模型学习一组对所有窗口共用的Softmax权重。由此，全局权重模型与本文方法的比较用于评估输入相关权重是否带来额外的预测收益。",
    "实验设置DLinear、LightTS、固定等权多尺度、全局可学习权重和恒定输入路由五类对照。DLinear用于比较单一尺度与多尺度结构；LightTS提供独立轻量模型基线。后三种多尺度对照与本文方法保留相同的尺度集合{1,2,4}和DLinear专家，分别使用等权、全局可学习权重和恒定输入路由权重。后两者均对全部窗口使用一组固定权重，可检验完整模型相对常权重方案的误差变化。",
)
replace_paragraph(
    document,
    "表3中的核心消融覆盖",
    "表3中的核心消融覆盖随机种子42～49，共重复8次。相对单尺度DLinear，本文方法在Abilene和GÉANT上的MSE的逐种子平均改进率分别为2.77%和2.35%；相对固定等权模型，MSE的逐种子平均改进率分别为1.65%和4.54%。新增的全局可学习权重模型允许尺度权重不等，但对全部窗口保持恒定。与该模型逐种子配对比较，本文方法的平均MSE改进率分别为1.80%和4.31%，分别在6/8和7/8个种子上胜出；t分布近似的95%置信区间分别为[0.01%,3.58%]和[−2.51%,11.13%]。Abilene优势较小，GÉANT区间跨越零，不能据此断言跨初始化稳定优于全局权重。",
    "表3汇总随机种子42～49的8次核心消融。本文方法相对DLinear与固定等权模型的逐种子平均MSE改进率，Abilene分别为2.77%和1.65%，GÉANT分别为2.35%和4.54%。全局可学习权重与恒定输入路由均对所有窗口使用一组可学习权重，后者与本文方法同为8339个标称参数。相对这两种对照，本文方法在Abilene上的改进均约1.80%，在GÉANT上分别为4.31%和4.97%，各有6/8和7/8个种子胜出。相对恒定输入路由的95%置信区间分别为[0.49%,3.11%]和[−1.62%,11.56%]，相对全局权重则为[0.01%,3.58%]和[−2.51%,11.13%]；GÉANT两项区间均跨零。",
)

tables = [table for table in document.tables if table.rows[0].cells[0].text == "数据集"
          and table.rows[0].cells[1].text == "模型"
          and len(table.rows) == 9]
assert len(tables) == 1, len(tables)
table = tables[0]
for dataset in ("Abilene", "GÉANT"):
    reference = next(row for row in table.rows if row.cells[0].text == dataset
                     and row.cells[1].text == "全局可学习权重")
    following = next(row for row in table.rows if row.cells[0].text == dataset
                     and row.cells[1].text == "本文方法")
    copied = deepcopy(reference._tr)
    following._tr.addprevious(copied)
    new_row = table.rows[list(table._tbl.tr_lst).index(copied)]
    values = (dataset, "恒定输入路由", "8339", "0.192±0.002", "0.181±0.006") if dataset == "Abilene" else (
        dataset, "恒定输入路由", "8339", "0.449±0.012", "0.379±0.007")
    for cell, value in zip(new_row.cells, values):
        set_cell_text(cell, value)

for cell in table.rows[-1].cells:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.keep_with_next = False

replace_paragraph(
    document,
    "固定等权模型只能检验等权融合",
    "固定等权模型只能检验等权融合是否足够；全局可学习权重模型进一步控制了“权重不等”这一因素。表3显示，样本级路由在Abilene的6个种子、GÉANT的7个种子上取得较低MSE，但GÉANT种子43的MSE比全局权重模型高13.22%，使跨种子95%置信区间跨越零。图4进一步结合全测试点的平均预测偏差与按真实值波动选取的预测轨迹，观察整体趋势和具体误差。",
    "固定等权模型检验等权融合，全局可学习权重及恒定输入路由模型提供两种常权重对照。表3中本文方法相对两者的平均MSE均较低，但GÉANT的配对区间均跨越零；其中种子43相对恒定输入路由退步12.78%。由于恒定输入路由的首层输入权重不起作用，且各模型分别训练，这一比较仍不足以证明每个窗口的尺度权重分配本身有效。图4进一步结合全测试点的平均预测偏差与按真实值波动选取的预测轨迹，观察整体趋势和具体误差。",
)
replace_paragraph(
    document,
    "本文面向骨干网络OD流量",
    "加入全局可学习权重对照后，样本级路由的逐种子平均MSE改进率分别为1.80%和4.31%，但GÉANT的95%置信区间跨越零。",
    "相对全局可学习权重与恒定输入路由两类常权重对照，本文方法在Abilene上的逐种子平均MSE改进均约1.80%；GÉANT上的对应改进分别为4.31%和4.97%，但两项95%置信区间均跨越零。",
)

assert len(table.rows) == 11
document.save(TARGET)
print(TARGET)
