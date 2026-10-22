from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from collections import Counter
import hashlib, json, re
from docx import Document
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / '2026.09.29v52.docx'
TARGET = ROOT / '2026.09.29v53.docx'
doc = Document(SOURCE)
original = Document(SOURCE)
p = list(doc.paragraphs)
changes = []

def replace(i, old, new, reason):
    assert p[i].text.count(old) == 1
    runs = [r for r in p[i].runs if old in r.text]
    if len(runs) == 1:
        runs[0].text = runs[0].text.replace(old, new, 1)
    else:
        assert p[i].text == old
        assert not p[i]._p.xpath('.//m:oMath | .//w:drawing')
        assert p[i].runs
        p[i].runs[0].text = new
        for run in p[i].runs[1:]:
            run.text = ''
    changes.append((reason, old, new))

def new_para(text, template):
    elem = deepcopy(template._p)
    for child in list(elem):
        if child.tag != '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr':
            elem.remove(child)
    para = Paragraph(elem, template._parent)
    run = para.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return para

def insert_before(anchor, text, template):
    para = new_para(text, template)
    anchor._p.addprevious(para._p)
    return para

def insert_after(anchor, text, template):
    para = new_para(text, template)
    anchor._p.addnext(para._p)
    return para

# One application sentence in the introduction; original argument stays intact.
replace(12, '本文因而预测后续OD流量。',
        '本文因而预测后续OD流量。面向小时级骨干网需求分析，预测结果可作为运维和容量规划的输入，本文不直接生成资源分配动作。',
        '引言补充应用对象和预测模块边界')
replace(21, '1.1 问题定义与数据预处理', '1.1 预测任务与数据预处理', '明确任务小节名称')
replace(33, '1.2 多尺度序列构造', '1.2 多尺度构造与DLinear专家', '合并尺度构造与专家小节')
changes.append(('合并原1.2和1.3', p[38].text, '原1.3下全部正文及公式接续于1.2，不删除技术说明'))
p[38]._p.getparent().remove(p[38]._p)
replace(46, '1.4 样本级自适应尺度加权', '1.3 样本级自适应尺度加权', '更新小节编号')
replace(58, '1.5 训练目标与模型复杂度', '1.4 模型训练与预测输出', '将训练输出与运行开销分别组织')

time_span = ('在第2.1节所用的1 h聚合数据上，96步历史输入对应过去96 h，24步预测对应未来24 h。'
             '该设置适用于小时级需求变化分析；若采用逐小时滚动调用，每次仍预测未来24 h，预测跨度与更新周期应分别理解。')
insert_after(p[30], time_span, p[30])
changes.append(('1.1新增时间跨度说明', '无', time_span))

# Keep inline math intact and relocate only the control definitions.
control_start = '为了区分多尺度结构、非等权分配与样本相关分配的作用'
control = p[57].text[p[57].text.index(control_start):]
replace(57, control, '', '将常权重对照定义原文移至2.2')
# The shared-backbone explanation is moved as a whole original paragraph.
p[71]._p.addnext(p[45]._p)
control_para = insert_after(p[45], control, p[71])
changes.append(('统一主干理由移至2.2', original.paragraphs[45].text, original.paragraphs[45].text))

# Retain training sentence; move parameter counts without changing their wording.
training = '训练采用Adam优化器更新参数，并根据验证集MSE选择最终使用的模型。'
assert p[61].text.startswith(training)
parameter_text = p[61].text[len(training):]
replace(61, parameter_text, '', '参数量说明原文移至2.6，训练规则留在1.4')
inference = ('推理时，将历史OD窗口按训练阶段确定的规则预处理，再经多尺度专家与路由融合得到未来多步预测。'
             '这些输出描述OD流量需求，不是链路路由或带宽分配方案；具体管理动作需要由外部模块结合网络条件确定。')
insert_after(p[62], inference, p[62])
changes.append(('1.4新增预测输出接口说明', '无', inference))

replace(76, '2.3 两数据集外部基线比较', '2.3 总体预测性能与基线比较', '明确总体预测比较的章节作用')
replace(94, '2.5 路由权重分析', '2.5 预测行为与尺度权重分析', '图4和图5归入同一行为分析小节')
p[91]._p.addprevious(p[94]._p)
changes.append(('2.5起点前移至图4之前', '图4在2.4内，2.5从图5说明开始', '图4校准与案例、图5权重及分组分析统一归入2.5'))
heading = insert_before(p[100], '2.6 运行开销与网络管理应用讨论', p[94])
insert_after(heading, parameter_text, p[100])
changes.append(('新增2.6并移入原计时段', '计时段位于路由权重分析末尾', '参数量与原计时段原文归入2.6'))

applications = [
    ('应用时，可由运维模块定期读取小时级历史OD序列，调用预测模型生成未来一天的需求估计，'
     '辅助观察高需求时段与容量变化。这里给出的是使用流程，现有实验仍为离线预测评价，尚未验证在线管理效果。'),
    ('图4显示各模型在高真实值区间仍有低估，本文方法的分组平均偏差相对较小。'
     '对于依赖预测的容量评估，这提示需要关注高需求时段的低估风险；分组偏差改善尚不能直接解释为拥塞率下降或调度收益增加。'),
    ('实际使用还需将预测恢复到原始流量量纲，并结合拓扑、路由和链路容量，才能由OD需求进一步分析链路负载。'
     '当前误差指标在标准化空间计算，前向计时也不包含完整部署流程，因此不能据此认定端到端运行或资源调度效果。')
]
anchor = p[100]
for text in applications:
    anchor = insert_after(anchor, text, p[100])
    changes.append(('2.6新增应用讨论', '无', text))

# Split the existing conclusion at complete-sentence boundaries, adding one modest application sentence.
old_conclusion = p[102].text
split1 = old_conclusion.index('按真实值分组的校准结果表明')
split2 = old_conclusion.index('路由器是否稳定匹配窗口尺度需求')
part1, part2, part3 = old_conclusion[:split1], old_conclusion[split1:split2], old_conclusion[split2:]
app_end = '现有结果可为小时级骨干网需求分析提供预测参考，实际运维和资源配置收益仍需另行验证。'
replace(102, old_conclusion, part1, '结论按结果、应用意义、证据边界分段')
c2 = insert_after(p[102], part2 + app_end, p[102])
insert_after(c2, part3, p[102])
changes.append(('结论新增应用落点', '无', app_end))

doc.save(TARGET)
verified = Document(TARGET)
ns = {'m':'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(SOURCE) as a, ZipFile(TARGET) as b:
    xa, xb = etree.fromstring(a.read('word/document.xml')), etree.fromstring(b.read('word/document.xml'))
    objects = lambda x, path: [etree.tostring(e) for e in x.xpath(path, namespaces=ns)]
    assert objects(xa, '//m:oMath') == objects(xb, '//m:oMath'), 'math'
    assert objects(xa, '//w:drawing') == objects(xb, '//w:drawing'), 'drawings'
    media = lambda z: {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
    assert media(a) == media(b), 'media'
    assert a.namelist() == b.namelist()
    for n in a.namelist():
        if n != 'word/document.xml': assert a.read(n) == b.read(n), n
assert len(original.tables) == len(verified.tables) == 3
for a,b in zip(original.tables, verified.tables): assert a._tbl.xml == b._tbl.xml
assert len(original.inline_shapes) == len(verified.inline_shapes) == 5

# All original data paragraphs and all original unchanged paragraphs survive with exact XML.
changed_indices = {12,21,33,38,46,57,58,61,76,94,102}
final_xml = [x._p.xml for x in verified.paragraphs]
for i,x in enumerate(original.paragraphs):
    if i not in changed_indices: assert x._p.xml in final_xml, ('unchanged paragraph',i)
for text in [control, parameter_text, part1, part2, part3]:
    assert any(text in x.text for x in verified.paragraphs), ('relocation',text)
before = '\n'.join(x.text for x in original.paragraphs)
after = '\n'.join(x.text for x in verified.paragraphs)
assert re.findall(r'\[\d+\]',before) == re.findall(r'\[\d+\]',after), 'citation order'
assert Counter(re.findall(r'“[^”]*”',before)) == Counter(re.findall(r'“[^”]*”',after)), 'quotes'
old_numbers = Counter(re.findall(r'\d+(?:\.\d+)?',before))
new_numbers = Counter(re.findall(r'\d+(?:\.\d+)?',after))
# Heading renumbering intentionally changes subsection number tokens only.
old_heads = ['1.1 问题定义与数据预处理','1.2 多尺度序列构造','1.3 DLinear尺度专家',
             '1.4 样本级自适应尺度加权','1.5 训练目标与模型复杂度','2.3 两数据集外部基线比较','2.5 路由权重分析']
new_heads = ['1.1 预测任务与数据预处理','1.2 多尺度构造与DLinear专家',
             '1.3 样本级自适应尺度加权','1.4 模型训练与预测输出','2.3 总体预测性能与基线比较',
             '2.5 预测行为与尺度权重分析','2.6 运行开销与网络管理应用讨论']
added = time_span + inference + ''.join(applications) + app_end
expected = old_numbers.copy()
expected.subtract(Counter(re.findall(r'\d+(?:\.\d+)?',''.join(old_heads))))
expected.update(Counter(re.findall(r'\d+(?:\.\d+)?',''.join(new_heads)+added)))
assert +expected == new_numbers, 'all original numeric tokens and deliberate additions'

lines = ['# V52至V53结构与应用说明修改对比报告','','## 范围与结果','',
'- 来源：2026.09.29v52.docx；另存：2026.09.29v53.docx。按用户要求进行小范围结构调整和应用说明补充，未做全文语言重写。',
'- 合并相关方法小节、集中对照定义、调整图4所属小节，新增运行开销与应用讨论；已有技术说明通过移位保留。',
'- 中英文摘要、标题、参考文献及实验结果段落保持原文。图表内容、尺寸、顺序、标题及全部公式与编号保持不变；正文增补会造成自然分页变化。',
f'- 原正文段落数{len(original.paragraphs)}，新段落数{len(verified.paragraphs)}。仅删除被合并的一个小节标题，正文信息不删减。',
f'- 全部段落文本字符数（不含内嵌公式及表格文本）：{len(before)} → {len(after)}。', '',
'## 修改后的章节结构','',
'- 1.1 预测任务与数据预处理',
'- 1.2 多尺度构造与DLinear专家（合并原1.2、1.3）',
'- 1.3 样本级自适应尺度加权',
'- 1.4 模型训练与预测输出',
'- 2.1 数据集与实验设置',
'- 2.2 对比模型与评价指标（移入统一主干理由和常权重对照定义）',
'- 2.3 总体预测性能与基线比较',
'- 2.4 八随机种子核心消融与稳定性',
'- 2.5 预测行为与尺度权重分析（含图4、图5）',
'- 2.6 运行开销与网络管理应用讨论',
'- 3 结论（原内容分为三段）', '',
'## 逐项对比','']
for i,(reason,old,new) in enumerate(changes,1):
    lines += [f'### {i}. {reason}','','**修改前**','',old or '原位置仅保留相邻说明。','','**修改后或新位置**','',new or '此处移出，原文完整保留在新位置。','']
lines += ['## 保护核验','',
'- 3张表XML完全一致；5幅图的媒体哈希及绘图XML完全一致，图中文字、数据、尺寸和顺序未变。',
'- 全部41个数学对象XML及出现顺序完全一致；16个编号公式和所有内嵌数学对象保持原样。',
'- 非改写段落XML逐一匹配，包括所有实验结果、统计解释、计时数据、图表标题与参考文献。',
'- 参数量、常权重对照定义及原结论分段内容在新文档中逐段精确匹配。引用标记及出现顺序、直接引语均一致。',
'- 数值核验仅允许章节编号变化、新增小时换算及正文指代，全部原有实验数字保留。',
'- DOCX中除主文档XML外其他组成文件逐字节一致。', '',
'## 应用说明的证据边界','',
'- 新增小时跨度来自既有1 h/步、L=96、H=24设置，不是新增实验结果。逐小时滚动调用表述为建议方式。',
'- 运维和容量规划流程属于用途讨论，没有声称完成部署、降低拥塞率、提升吞吐或获得调度收益。',
'- 原始量纲恢复、拓扑/路由/容量信息以及端到端运行属于应用所需条件，不描述为已验证功能。',
'- 本轮未新增实验，未对外上传检测；没有新的AIGC复测结果。', '',
'## 版面复核','', '最终渲染检查完成后补充。','']
(ROOT/'2026.09.29v53_结构修改对比报告.md').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps({'target':str(TARGET),'changes':len(changes),'old_paragraphs':len(original.paragraphs),
                  'new_paragraphs':len(verified.paragraphs),'chars_before':len(before),'chars_after':len(after),
                  'protected_checks':'passed'},ensure_ascii=False))
