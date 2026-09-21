from __future__ import annotations
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.21v18.docx"
TARGET = ROOT / "word_output" / "2026.09.21v19.docx"

doc = Document(SOURCE)

def find(prefix):
    xs = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(xs) != 1:
        raise RuntimeError(f"{prefix!r}: {len(xs)} matches")
    return xs[0]

def set_text(p, text):
    first = p.runs[0] if p.runs else p.add_run()
    first.text = text
    for r in p.runs[1:]:
        r.text = ""

def repl(prefix, text):
    p = find(prefix)
    set_text(p, text)

replacements = {
"摘要：": "摘要：骨干网络源—目的节点对流量具有明显的时变性和多时间尺度特征，不同时间尺度在不同流量状态下的作用并不相同。针对固定尺度模型难以灵活利用这些信息的问题，提出一种样本级自适应多尺度预测方法。方法采用尺度1、2和4的非重叠平均池化构造多分辨率序列，各尺度分别由DLinear专家进行预测；同时从输入窗口提取均值、标准差、末时刻均值和平均绝对一阶差分，经轻量路由网络得到Softmax尺度权重，并据此融合各尺度预测结果。实验在Abilene和GÉANT两个骨干网流量数据集上进行。3个共同随机种子下，所提方法的MSE相对DLinear分别降低4.06%和4.43%，相对LightTS分别降低29.89%和46.83%；8个随机种子的消融实验中，相对固定等权融合分别降低1.67%和4.68%。结果显示，在仅增加少量参数的情况下，自适应尺度加权可以更充分地利用不同时间尺度的信息，从而改善多变量网络流量预测精度。",

"Abstract:": "Abstract: Backbone origin-destination (OD) traffic varies over time and contains information at multiple temporal scales. The contribution of each scale also changes with the current traffic state, which limits fixed-scale forecasting models. To address this issue, a sample-wise adaptive multiscale forecasting method is developed. Non-overlapping average pooling at scales 1, 2, and 4 is used to form multiresolution sequences, and a separate DLinear expert produces a forecast at each scale. Meanwhile, the mean, standard deviation, latest-step mean, and mean absolute first difference of the input window are fed into a lightweight routing network. The resulting Softmax weights are then used to combine the forecasts from the three scales. Experiments are conducted on the Abilene and GÉANT backbone traffic datasets. With three common random seeds, the method reduces MSE by 4.06% and 4.43% relative to DLinear, and by 29.89% and 46.83% relative to LightTS. In the eight-seed ablation experiment, the MSE reductions over fixed equal weighting are 1.67% and 4.68%, respectively. The results indicate that adaptive scale weighting can make better use of complementary temporal information while introducing only a small number of additional parameters.",

"随着云计算、视频业务、数据中心互联及智能网络应用的持续发展": "随着云计算、视频业务、数据中心互联及智能网络应用的发展，骨干网络中的业务类型不断增多，流量规模也持续扩大。源—目的节点对（origin-destination，OD）流量随时间变化明显，同时伴有突发和周期性波动。若能够较准确地预测后一时段的OD流量，网络管理系统就可以提前掌握业务变化，为容量规划、拥塞预警、资源调度和流量工程提供参考。因此，网络流量预测的关键之一，是如何从历史序列中提取有效的时间变化特征。",

"针对网络流量预测问题，相关研究已由传统统计模型逐步发展到机器学习和深度学习方法": "网络流量预测的研究经历了从统计方法到机器学习、深度学习方法的发展。Lohrasbinasab等[1]对相关方法进行了综述，认为数据驱动模型更适合处理复杂的流量变化。Qin等[2]将图神经网络用于通信网络流量预测，同时考虑空间关联和时间依赖；韦烜等[10]研究了大型IP网络流量矩阵的分析与预测，表明OD流量之间的关联和时间演化都会影响预测结果。对于网络流量中的周期变化，唐文杰等[11]提出CycleLLH模型，将周期性信息引入预测过程。由这些研究可以看出，时间依赖、变量关联和周期变化都是网络流量建模中需要考虑的因素。",

"除拓扑关联和周期规律外，网络流量在不同时间尺度上还具有不同的变化特征": "网络流量在不同时间尺度上呈现的特征并不完全相同。细粒度序列保留了较多短时波动，经过平滑后的粗粒度序列则更容易反映整体趋势。Zhou等[3]采用深度回声状态网络开展多尺度网络流量预测，通过不同时间分辨率提取信息。DLinear[4]使用时间序列分解和线性映射完成预测，结构较为简洁；TimeMixer[5]利用多尺度混合机制处理不同分辨率的时间序列；LightTS[6]采用采样导向的MLP结构进行快速多变量时间序列预测。这些方法说明，多尺度信息并不一定需要依赖复杂模型才能发挥作用，如何在轻量结构中组织和利用这些信息仍有研究空间。",

"然而，在实际骨干网络中，不同时间段的流量水平、波动幅度和短期变化状态并不相同": "实际骨干网络的流量状态会不断变化，同一个时间尺度并不一定适合所有输入窗口。固定尺度或固定融合比例虽然实现简单，但在流量水平、波动幅度发生变化时，可能无法充分发挥多尺度信息的作用。基于这一考虑，本文在尺度1、2和4上构造多分辨率序列，以DLinear作为各尺度预测专家，并利用输入窗口的统计特征生成样本级尺度权重。所提方法在Abilene和GÉANT两个公开骨干网数据集上进行验证，实验包括基线模型比较、固定等权消融、多随机种子重复实验和路由权重分析。",

"本文提出的自适应多尺度网络流量预测方法整体结构如图1所示": "本文方法的整体结构如图1所示。原始OD流量完成变换和标准化后，在尺度1、2和4上分别构造输入序列，各尺度序列送入对应的DLinear专家得到预测结果。与此同时，路由网络从当前输入窗口的统计特征中计算尺度权重，再对3个专家的输出进行加权。这样处理后，较细时间尺度中的局部变化和较粗时间尺度中的趋势信息可以同时参与最终预测。",

"其中μctr和σctr仅由训练集计算": "其中μctr和σctr由训练集计算，验证集和测试集沿用同一组参数。原始流量在不同OD维度上的数值差异较大，因此先进行log(1+x)变换，再按变量标准化。前者用于压缩数值范围，后者使不同OD变量处于较接近的尺度。数据按时间顺序以6∶2∶2划分为训练集、验证集和测试集，滑动窗口分别在各集合内部构造。",

"为同时保留短时变化和较长时间范围内的趋势信息，本文选取尺度1、2和4构造多分辨率序列": "本文选取尺度1、2和4构造多分辨率序列，对应的输入长度分别为96、48和24。尺度1直接保留原始时间分辨率；尺度2采用2点平均，减弱部分短时扰动；尺度4进一步平滑局部变化，使较慢的变化趋势更加突出。不同尺度均采用非重叠平均池化完成，因此这一过程本身不引入新的可训练参数。",

"在各时间尺度上，本文采用DLinear作为基础预测专家": "各时间尺度均采用独立的DLinear专家。DLinear通过移动平均将序列分为趋势项和余项，再分别进行线性映射[4]。由于其结构较简单，用作多尺度框架中的基础预测器不会明显增加模型复杂度。同一尺度内，各OD变量共享时间映射参数；不同尺度之间则使用独立参数，以分别学习对应时间分辨率下的变化规律。",

"多尺度专家可以从不同时间分辨率提取信息，但如果直接采用固定等权融合": "得到3个尺度的预测结果后，还需要确定各尺度在当前样本中的占比。固定等权融合对所有输入窗口都使用相同权重，无法反映不同时段流量状态的变化。本文从原始输入Xn中提取四维统计向量zn=[μn,σn,ℓn,dn]，其中",

"μn、σn、ℓn和dn分别描述输入窗口的总体水平、离散程度、最近状态和平均变化强度": "μn、σn、ℓn和dn分别表示输入窗口的总体水平、离散程度、最近状态和平均变化强度。这4个统计量用于概括当前窗口的基本状态，并输入隐藏维数为16的两层感知器，得到3个尺度对应的分数。路由网络只使用低维统计特征，因此增加的计算量较小。",

"经过Softmax归一化后，每个样本得到一组和为1的尺度权重": "经Softmax归一化后，每个样本得到一组和为1的尺度权重，该组权重在同一时间窗口内由全部OD变量共享。这样既可以让尺度比例随输入样本变化，也可以避免路由器参数量随OD变量数增长。固定等权消融模型将3个尺度权重统一设为1/3，其尺度集合、专家结构和训练过程均保持不变。",

"模型采用Adam优化，并以验证集MSE作为模型选择依据": "模型采用Adam进行优化，并根据验证集MSE选择参数。DLinear、固定等权多尺度模型和本文方法的可训练参数量分别为4656、8208和8339，其中路由器相对固定等权多尺度模型只增加131个参数。因此，本文方法的参数增量主要来自轻量路由网络，整体模型规模变化较小。",

"为验证所提方法在不同骨干网络流量场景下的预测效果": "实验选用CESNET TS-Zoo[7]中的Abilene[8]和GÉANT[9] 1 h聚合流量矩阵。Abilene保留144维有效OD流量，时间序列相对较长；GÉANT去除恒定通道后保留524维，变量数更多。两个数据集在OD维度和流量分布上存在差异。实验均按时间顺序以6∶2∶2划分数据，输入长度L=96，预测长度H=24，预处理方式与第1.1节一致，具体设置见表1。",

"为保证比较过程一致，所有模型采用相同的数据划分、输入长度和评价流程": "各模型使用相同的数据划分、输入长度和评价指标，批量大小均为16，并根据验证集MSE选择最优模型。DLinear、固定等权多尺度模型和本文方法最多训练30轮，LightTS最多训练50轮。外部基线比较采用随机种子42—44；DLinear、固定等权多尺度模型和本文方法另外使用42—49共8个随机种子重复实验。",

"实验选取DLinear、LightTS和固定等权多尺度模型作为对比方法": "对比模型包括DLinear、LightTS和固定等权多尺度模型。DLinear作为单尺度基线，LightTS作为轻量多变量时间序列预测模型。固定等权多尺度模型与本文方法使用完全相同的尺度集合{1,2,4}和DLinear专家，区别只在于其3个尺度权重始终为1/3，因此可用于观察自适应加权本身带来的变化。",

"上述指标均在标准化空间中计算": "MSE、MAE和RMSE均在标准化空间中计算，数值越小表示预测误差越低。MSE对较大的预测偏差更敏感，MAE反映平均绝对误差，RMSE则给出均方误差开方后的结果。正文主要使用MSE和MAE比较模型，表2报告3个共同随机种子的均值±标准差，表3报告8个随机种子的均值±标准差。",

"表2给出了随机种子42—44下各模型的预测结果": "表2列出了随机种子42—44下的预测结果。本文方法在Abilene和GÉANT上的MSE分别为0.186和0.434。与DLinear相比，MSE分别降低4.06%和4.43%；与LightTS相比，分别降低29.89%和46.83%。两个数据集上的MAE也低于对应基线。由此可以看出，在相同输入长度和预测长度下，多尺度输入与样本级权重结合后，整体预测误差有所下降。",

"为进一步观察不同随机初始化下的结果，图2给出了3个共同随机种子的逐次比较": "图2给出了3个共同随机种子的逐次结果。Abilene上，本文方法3次实验的MSE和MAE均低于DLinear和LightTS。GÉANT上，相对LightTS同样保持较低误差；与DLinear比较时，个别随机种子的结果存在波动，但3次实验的平均值仍然更低。结合表2可知，本文方法在两个数据集上的平均误差均有所下降，在GÉANT上相对LightTS的差距更明显。",

"在基线比较的基础上，为进一步考察多尺度结构和动态加权的作用": "表3将DLinear、固定等权多尺度模型和本文方法的重复次数扩展到8个随机种子。本文方法相对DLinear在Abilene和GÉANT上的平均MSE分别降低2.77%和2.37%。固定等权模型已经使用相同的多尺度专家，在此基础上，本文方法的MSE仍分别降低1.67%和4.68%。这表明误差下降并非只来自多尺度输入，尺度权重随样本变化也产生了作用。",

"图3进一步给出了固定等权模型与本文方法在相同随机种子下的配对结果": "图3比较了固定等权模型与本文方法在相同随机种子下的MSE。Abilene的8个随机种子中，有6个随机种子下本文方法的MSE更低；GÉANT中为7个。两个数据集的平均ΔMSE均为正，与表3中的平均结果一致。",

"在确认自适应加权能够改善预测结果后，进一步分析路由器对三个时间尺度的实际分配情况": "图4给出了随机种子42—49下全部测试窗口的尺度权重分布，用于观察路由器在测试阶段对不同时间尺度的分配情况。固定等权模型的3个尺度权重始终为1/3，而本文方法的权重随输入窗口变化，因此可以从分布中看到不同尺度的总体占比。",

"从图4可以看出，两个数据集上的尺度权重均没有集中在固定的1/3附近": "从图4可以看出，两个数据集的3个尺度并未得到相同权重。Abilene中α1、α2和α4的平均值分别为0.257、0.319和0.424，GÉANT中分别为0.283、0.318和0.399，尺度4的平均权重均最高，尺度1最低。在当前24步预测任务下，较粗时间尺度中的平滑趋势信息占比较大，但尺度1和尺度2仍保留一定权重，用于补充较细粒度的变化信息。",

"综合表2、表3和图2—图4可以看出": "结合表2、表3和图2—图4，本文方法在Abilene和GÉANT上的平均预测误差均低于相应对比模型。固定等权消融中，自适应权重在相同多尺度专家基础上继续降低了MSE；图4中3个尺度的权重分布也存在明显差异。两项结果对应表明，路由器并非简单地对多尺度结果做平均，而是在不同输入窗口下改变各尺度的参与程度。",

"本文提出了一种样本级自适应多尺度网络流量预测方法": "本文研究了骨干网络OD流量的多尺度预测问题，并提出样本级自适应多尺度网络流量预测方法。模型在尺度1、2和4上构造多分辨率序列，由DLinear专家分别预测，再利用轻量路由网络根据输入窗口状态确定尺度权重。Abilene和GÉANT上的实验中，本文方法相对DLinear、LightTS和固定等权多尺度模型均取得了更低的平均预测误差。路由权重结果显示，不同时间尺度在预测中的占比并不相同。该方法仅增加少量参数即可实现多尺度预测结果的动态融合，可用于骨干网络OD流量预测。"
}

for prefix, text in replacements.items():
    repl(prefix, text)

full = "\n".join(p.text for p in doc.paragraphs)

# 数据与术语保护
for token in [
    "源—目的节点对", "DLinear", "Softmax", "Abilene", "GÉANT", "LightTS",
    "CycleLLH", "TimeMixer", "CESNET TS-Zoo",
    "4.06%", "4.43%", "29.89%", "46.83%", "1.67%", "4.68%",
    "0.186", "0.434", "2.77%", "2.37%",
    "0.257", "0.319", "0.424", "0.283", "0.318", "0.399",
    "4656", "8208", "8339", "131"
]:
    assert token in full, token

assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 4
for caption in [
    "图1 自适应多尺度网络流量预测方法框架",
    "图2 Abilene与GÉANT上三随机种子外部基线配对结果",
    "图3 固定等权与自适应权重的八随机种子配对MSE及配对差值分布",
    "图4 八随机种子下样本级自适应路由权重分布",
]:
    assert caption in full

doc.save(TARGET)
print(TARGET)
print("tables", len(doc.tables), "figures", len(doc.inline_shapes))
print("abstract_chars", len(find("摘要：").text.replace("摘要：", "")))
