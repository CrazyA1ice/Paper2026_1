from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.21v14.docx"
TARGET = ROOT / "word_output" / "2026.09.21v15.docx"
FIGURE = (
    ROOT / "abilene_forecasting" / "paper_figures" / "exports" / "color"
    / "fig2_external_baseline_paired_v15_color.png"
)


def find_one(document: Document, startswith: str):
    matches = [p for p in document.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph starting with {startswith!r}, found {len(matches)}"
        )
    return matches[0]


def set_paragraph_text(paragraph, text: str) -> None:
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_figure2(document: Document) -> None:
    if len(document.inline_shapes) != 5:
        raise RuntimeError(f"Expected 5 inline figures, found {len(document.inline_shapes)}")
    shape = document.inline_shapes[1]
    blip = shape._inline.xpath(".//a:blip")[0]
    embed = blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    image_part = document.part.related_parts[embed]
    image_part._blob = FIGURE.read_bytes()

    with Image.open(FIGURE) as image:
        px_w, px_h = image.size
    shape.width = Inches(5.55)
    shape.height = Inches(5.55 * px_h / px_w)

    doc_pr = shape._inline.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("name", "图2 三随机种子外部基线配对结果")
        doc_pr[0].set(
            "descr",
            "Abilene和GÉANT上DLinear、LightTS与本文方法在seed42至44下的MSE和MAE配对结果；"
            "方形、圆形和三角形分别表示seed42、43和44，实线连接三种子均值。"
        )


def main() -> None:
    if TARGET.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {TARGET}")
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    doc = Document(SOURCE)
    replace_figure2(doc)

    p23 = find_one(doc, "逐随机种子结果进一步表明")
    set_paragraph_text(
        p23,
        "逐随机种子结果进一步见图2。图2分别给出Abilene和GÉANT上的MSE、MAE，"
        "左侧配对比较DLinear与本文方法，右侧配对比较LightTS与本文方法；"
        "方形、圆形和三角形分别对应随机种子42、43和44，细虚线连接同一随机种子的结果，"
        "较粗实线连接3个随机种子的均值。Abilene上，本文方法相对DLinear的MSE和MAE在随机种子42—44下均为3/3次降低，"
        "由3个共同随机种子均值计算的降幅分别为4.06%和10.52%；相对LightTS时均为3/3次降低，"
        "MSE和MAE降幅分别为29.89%和17.65%。GÉANT上，本文方法相对DLinear的MSE和MAE均为2/3次降低，"
        "seed43出现退化，均值降幅分别为4.43%和3.46%；相对LightTS时仍为3/3次降低，"
        "MSE和MAE降幅分别为46.83%和29.29%。因此，两数据集的平均结果均支持自适应多尺度融合降低预测误差，"
        "但GÉANT相对DLinear的逐种子结果存在更明显波动。本文方法仅含8339个可训练参数，明显少于当前适配下的LightTS；"
        "参数量仅反映模型规模，不等同于训练或推理速度。外部基线结果用于说明当前统一实验协议下的相对表现，"
        "不据此推断本文方法在其他任务上普遍优于LightTS。"
    )

    caption = find_one(doc, "图2 本文方法相对DLinear和LightTS的三随机种子误差降低率")
    set_paragraph_text(
        caption,
        "图2 Abilene与GÉANT上三随机种子外部基线配对结果"
    )

    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "图2 Abilene与GÉANT上三随机种子外部基线配对结果" in full_text
    assert "MSE和MAE降幅分别为46.83%和29.29%" in full_text
    assert "seed43出现退化" in full_text
    assert "代表性窗口预测结果" not in full_text
    assert len(doc.inline_shapes) == 5
    assert len(doc.tables) == 3

    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
