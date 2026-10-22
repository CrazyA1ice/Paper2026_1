"""Read-only audit of the H=24 results used by the manuscript."""
from pathlib import Path
from zipfile import ZipFile
from collections import defaultdict
import csv
import hashlib
import json
import re

import numpy as np
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
RESULTS = ROOT/'abilene_forecasting/results'
MANUSCRIPT = ROOT/'word_output/2026.10.02v55.docx'
models = {'DLinear':'dlinear','LightTS':'lightts','固定等权多尺度':'dlinear_scale_static',
          '全局可学习权重':'dlinear_scale_global','恒定输入路由':'dlinear_scale_constant_router',
          '本文方法':'dlinear_scale'}
labels = {'Abilene':'abilene','GÉANT':'geant'}
metrics = {}
checks = []
issues = []
configurations = defaultdict(set)
details = {'tables':[],'pairs':[],'weights':[],'trajectory_cases':[]}

def rows(path):
    with open(path,encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))

def check(label,condition,detail=''):
    checks.append({'check':label,'passed':bool(condition),'detail':detail})
    if not condition: issues.append(label)

def directory(ds,model,seed):
    group = ('global_weight_ablation' if model=='dlinear_scale_global' else
             'constant_router_ablation' if model=='dlinear_scale_constant_router' else '')
    return RESULTS/group/ds/f'{model}_seed{seed}'

def paired(values):
    a=np.asarray(values,float)
    critical = 4.302652730 if len(a)==3 else 2.364624251
    half=critical*a.std(ddof=1)/np.sqrt(len(a))
    return dict(mean=float(a.mean()),ci95=[float(a.mean()-half),float(a.mean()+half)],
                wins=int((a>0).sum()))

for ds in labels.values():
    with np.load(ROOT/f'abilene_forecasting/data/processed/{ds}_1hour.npz') as data:
        test=data['test']
        expected_truth=np.stack([test[i+96:i+120] for i in range(len(test)-119)])
        details.setdefault('dataset_shapes',{})[ds]={k:list(data[k].shape) for k in ['train','val','test']}
        check(f'{ds}: training normalization',np.max(np.abs(data['train'].mean(0)))<1e-5
              and np.max(np.abs(data['train'].std(0)-1))<1e-5)
        split_rows={k:len(data[k]) for k in ['train','val','test']}
        split_windows={k:len(data[k])-119 for k in ['train','val','test']}
    for label,model in models.items():
        for seed in range(42,45 if model=='lightts' else 50):
            run=directory(ds,model,seed)
            record=json.loads((run/'metrics.json').read_text(encoding='utf-8'))
            metrics[ds,model,seed]=record
            check(f'{ds}/{model}/{seed}: task configuration',
                  record['input_len']==96 and record['pred_len']==24 and record['batch_size']==16
                  and record['split_rows']==split_rows and record['split_windows']==split_windows)
            configurations[model].add((record['learning_rate'],record['epochs_max'],record['patience']))
            with np.load(run/'predictions.npz') as data:
                prediction,truth=data['prediction'],data['truth']
            check(f'{ds}/{model}/{seed}: saved truth and shape',
                  prediction.shape==expected_truth.shape and np.array_equal(truth,expected_truth))
            error=prediction.astype(np.float64)-truth
            calculated={'mse':float(np.mean(error**2)),'mae':float(np.mean(np.abs(error)))}
            calculated['rmse']=calculated['mse']**.5
            check(f'{ds}/{model}/{seed}: prediction metrics',
                  all(np.isclose(calculated[k],record['test_normalized'][k],rtol=2e-6,atol=1e-7)
                      for k in calculated),str(calculated))
            if seed==42 and model in ('dlinear','dlinear_scale_static','dlinear_scale','dlinear_scale_constant_router'):
                if model=='dlinear':
                    sd=truth.std(axis=(0,1)); eligible=np.flatnonzero(sd>0)
                    channel=int(eligible[np.argmin(abs(sd[eligible]-np.quantile(sd[eligible],.75)))])
                    variation=np.abs(np.diff(truth[:,:,channel],axis=1)).mean(axis=1)
                    selected=[int(np.argmin(abs(variation-np.quantile(variation,q)))) for q in (.5,.9)]
                    cases={w:{} for w in selected}
                for w in selected: cases[w][label]=float(np.mean(error[w,:,channel]**2))
            del prediction,truth,error
    for w,values in cases.items():
        details['trajectory_cases'].append(dict(dataset=ds,channel_one_based=channel+1,
            window_one_based=w+1,mse=values,adaptive_best=values['本文方法']==min(values.values())))

doc=Document(MANUSCRIPT)
for number,seeds in [(2,range(42,45)),(3,range(42,50))]:
    for row in doc.tables[number-1].rows[1:]:
        ds,label,count,mse,mae=[c.text for c in row.cells]
        ds,model=labels[ds],models[label]
        records=[metrics[ds,model,s] for s in seeds]
        expected=[]
        for key in ['mse','mae']:
            a=np.array([r['test_normalized'][key] for r in records])
            expected.append(f'{a.mean():.3f}±{a.std(ddof=1):.3f}')
        check(f'table {number}/{ds}/{model}',[mse,mae]==expected
              and all(r['parameters']==int(count) for r in records))
        details['tables'].append(dict(table=number,dataset=ds,model=model,parameters=int(count),
                                      mse=expected[0],mae=expected[1]))

for ds in labels.values():
    for baseline in ['dlinear','dlinear_scale_static','dlinear_scale_global','dlinear_scale_constant_router','lightts']:
        seeds=range(42,45 if baseline=='lightts' else 50)
        for metric in ['mse','mae','rmse']:
            values=[100*(1-metrics[ds,'dlinear_scale',s]['test_normalized'][metric]/
                          metrics[ds,baseline,s]['test_normalized'][metric]) for s in seeds]
            details['pairs'].append(dict(dataset=ds,baseline=baseline,metric=metric,
                                         seeds=list(seeds),**paired(values)))
    weights=np.stack([np.loadtxt(directory(ds,'dlinear_scale',s)/'router_weights.csv',delimiter=',')
                      for s in range(42,50)])
    check(f'{ds}: routing weights sum to one',np.allclose(weights.sum(-1),1,atol=1e-6)
          and np.all(weights>=0) and weights.shape[1]==len(expected_truth) if ds=='geant'
          else np.allclose(weights.sum(-1),1,atol=1e-6) and weights.shape[1]==813)
    mean=weights.mean((0,1)); within=weights.var(1).mean(0); between=weights.mean(1).var(0)
    details['weights'].append(dict(dataset=ds,mean=mean.tolist(),
        between_seed_variance_share=(between/(between+within)).tolist()))

for file,baseline_field in [
 ('abilene_forecasting/paper_figures/source_data_v25/fig3_lightts_improvement_seed42_44.csv','lightts'),
 ('abilene_forecasting/paper_figures/source_data_v24/fig5_relative_improvement_8seeds.csv',None)]:
    for row in rows(ROOT/file):
        ds,seed,key=row['dataset'],int(row['seed']),row['metric'].lower()
        baseline=baseline_field or {'DLinear':'dlinear','固定等权':'dlinear_scale_static'}[row['baseline']]
        expected=100*(1-metrics[ds,'dlinear_scale',seed]['test_normalized'][key]/
                       metrics[ds,baseline,seed]['test_normalized'][key])
        check(f'{Path(file).name}/{ds}/{seed}/{baseline}/{key}',
              np.isclose(expected,float(row['relative_improvement_pct']),atol=1e-7))

calibration=rows(ROOT/'paper_figures/fig4_calibration_trajectories_v39/calibration_source_8seeds_v39.csv')
grouped=defaultdict(list)
for row in calibration: grouped[row['dataset'],row['model'],int(row['seed'])].append(row)
cal_models={'DLinear':'dlinear','固定等权':'dlinear_scale_static','恒定输入路由':'dlinear_scale_constant_router','本文方法':'dlinear_scale'}
for (ds,label,seed),part in grouped.items():
    count=sum(int(r['n_elements']) for r in part)
    expected_count=813*24*144 if ds=='abilene' else 451*24*524
    check(f'figure4 calibration counts/{ds}/{label}/{seed}',count==expected_count)
    for key in ['mse','mae']:
        value=sum(float(r[key])*int(r['n_elements']) for r in part)/count
        check(f'figure4 calibration metric/{ds}/{label}/{seed}/{key}',
              np.isclose(value,metrics[ds,cal_models[label],seed]['test_normalized'][key],rtol=2e-6,atol=1e-7))

details['training_configurations']={k:[list(v) for v in sorted(values)] for k,values in configurations.items()}
text='\n'.join(p.text for p in doc.paragraphs)
check('H=12/H=48 absent from manuscript','H=12' not in text and 'H=48' not in text)
citations=[]
for p in doc.paragraphs[:112]:
    for n in re.findall(r'\[(\d+)\]',p.text):
        if int(n) not in citations: citations.append(int(n))
check('references numbered by first occurrence',citations==list(range(1,12)))
check('numbered equations sequential',
      [int(m) for p in doc.paragraphs for m in re.findall(r'^\s*\((\d+)\)\s*$',p.text)]==list(range(1,17)))
details['unreported_assignment_audit']=rows(RESULTS/'router_window_assignment_audit/summary.csv')
details['latency_sources']={folder:rows(RESULTS/folder/'inference_benchmark_summary.csv')
    for folder in ['inference_benchmark','inference_benchmark_repeat','inference_benchmark_cuda_events']}
details['volatility_groups']=rows(RESULTS/'global_weight_ablation/analysis/volatility_summary.csv')
report=dict(manuscript=str(MANUSCRIPT),runs_checked=len(metrics),checks=checks,
            failed_checks=issues,details=details)
(OUT/'experiment_data_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('Runs',len(metrics),'checks',len(checks),'failed',len(issues))
print(json.dumps(issues,ensure_ascii=False))
print(json.dumps({k:details[k] for k in ['trajectory_cases','training_configurations','weights']},ensure_ascii=False))
