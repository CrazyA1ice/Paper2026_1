"""Make focused evidence and reporting corrections to the v32 manuscript."""

from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "2026.09.25v32.docx"
TARGET = ROOT / "2026.09.26v33.docx"


def replace(paragraph, old, new):
    first = paragraph.runs[0]
    if old not in first.text:
        raise ValueError(f"Expected text not found: {old}")
    first.text = first.text.replace(old, new)


doc = Document(SOURCE)
p = doc.paragraphs
assert len(p) == 113 and len(doc.inline_shapes) == 5 and len(doc.tables) == 3

replace(p[3], "在Abilene和GÉANT数据集上，3个共同随机种子下", "在输入长度96、预测长度24的Abilene和GÉANT实验中，3个共同随机种子下")
replace(p[3], "4.06%和4.43%", "4.07%和4.54%")
replace(p[3], "29.89%和46.83%", "29.80%和46.89%")
replace(p[9], "On Abilene and GÉANT, the method reduces MSE", "With input length 96 and forecast length 24, the method reduces MSE on Abilene and GÉANT")
replace(p[9], "4.06% and 4.43%", "4.07% and 4.54%")
replace(p[9], "29.89% and 46.83%", "29.80% and 46.89%")

replace(p[17], "固定等权多尺度模型、多个随机种子重复实验", "固定等权与全局可学习固定权重模型、多个随机种子重复实验")
replace(p[19], "固定等权消融模型保留相同的尺度和专家，因此能够单独比较样本级权重带来的变化。", "固定等权与全局可学习权重对照保留相同的尺度和专家，用于分别比较等权、全局非等权与输入相关融合的结果。")
replace(p[45], "单尺度预测、固定等权多尺度和样本级自适应多尺度三种设置", "单尺度预测、固定等权、全局可学习固定权重和样本级路由四种设置")
replace(p[50], "供路由器判断当前样本更适合依赖哪些时间尺度，而不直接参与未来流量数值的预测。", "供路由器计算输入相关的尺度权重，而不直接参与未来流量数值的预测。")
replace(p[52], "路由器只需判断尺度的相对重要性，不需要再次完成完整的OD流量预测", "路由器仅计算三个尺度的融合分数，不需要再次完成完整的OD流量预测")
replace(p[57], "后者与本文方法的差别集中在权重是否依赖输入窗口。", "后者提供对输入相关权重更直接的性能对照，但路由器另有少量可训练参数。")
replace(p[61], "各多尺度模型共享专家结构，因此全局权重对照可用于检验输入相关融合的额外作用。", "各多尺度模型共享专家结构，全局权重模型因此提供较接近的对照；两者仍有128个参数的差异，不能将全部误差差异直接归因于逐窗口加权。")
replace(p[62], "单尺度DLinear、固定等权多尺度模型和本文方法", "单尺度DLinear、固定等权多尺度、全局可学习固定权重模型和本文方法")
replace(p[71], "由此，全局权重模型与本文方法的比较更直接地检验输入相关权重是否带来收益。", "由此，全局权重模型与本文方法的比较用于评估输入相关权重是否带来额外的预测收益。")
replace(p[75], "RMSE是MSE开方后的结果，可与误差幅值直接对照。正文以MSE和MAE作为主要比较指标，RMSE用于补充。", "RMSE是MSE开方后的结果，用于补充呈现误差幅值。正文以MSE和MAE作为主要比较指标。")
replace(p[75], "同时给出平均水平和重复实验间的波动。", "同时给出平均水平和重复实验间的波动。文中相对改进率均先按共同随机种子逐一计算，再取种子平均；因此不直接用表中四舍五入后的均值计算百分比。")
replace(p[77], "相对DLinear分别降低4.06%和4.43%，相对LightTS分别降低29.89%和46.83%", "逐种子计算后相对DLinear分别降低4.07%和4.54%，相对LightTS分别降低29.80%和46.89%")
replace(p[84], "平均MSE分别降低2.77%和2.37%", "MSE的逐种子平均改进率分别为2.77%和2.35%")
replace(p[84], "平均MSE分别降低1.67%和4.68%", "MSE的逐种子平均改进率分别为1.65%和4.54%")
replace(p[100], "在Abilene和GÉANT上，本文方法", "在预测长度24的Abilene和GÉANT任务上，本文方法")

runtime = (
    "模型规模之外，还需考察前向开销。对8个随机种子的预测长度24检查点，"
    "在NVIDIA GTX 1060上以批量16对预加载的测试张量计时，本文方法在Abilene和GÉANT上的"
    "平均每批前向耗时分别为2.65 ms和2.61 ms，固定等权模型分别为2.38 ms和2.10 ms，"
    "单尺度DLinear分别为0.77 ms和0.70 ms。计时仅覆盖模型前向计算，不含数据读取与传输。"
    "因此，本文方法参数量较小，但三个尺度专家及路由计算仍增加了推理开销；上述结果不代表端到端部署时延。"
)
new = p[93].insert_paragraph_before(runtime)
new_ppr = deepcopy(p[92]._p.pPr)
if new_ppr is not None:
    if new._p.pPr is not None:
        new._p.remove(new._p.pPr)
    new._p.insert(0, new_ppr)

doc.save(TARGET)
check = Document(TARGET)
assert len(check.paragraphs) == 114
assert len(check.inline_shapes) == 5 and len(check.tables) == 3
text = "\n".join(paragraph.text for paragraph in check.paragraphs)
assert "29.89%" not in text and "4.43%" not in check.paragraphs[3].text
assert "H=12" not in text and "H=48" not in text
assert "NVIDIA GTX 1060" in text
print(TARGET)
