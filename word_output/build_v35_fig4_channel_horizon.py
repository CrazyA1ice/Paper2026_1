"""Restore v33 Figure 3 and replace only Figure 4 with readable H=24 evidence."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.26v33.docx"
TARGET = ROOT / "word_output/2026.09.26v35.docx"
FIGURE = ROOT / "paper_figures/fig4_channel_horizon_v35.png"


def set_text(paragraph, text):
    assert paragraph.runs
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


doc = Document(SOURCE)
p = doc.paragraphs
assert len(p) == 114 and len(doc.inline_shapes) == 5
assert p[88].text == "图3 Abilene与GÉANT上本文方法相对两类基线的八随机种子多指标相对改进分布"
assert p[91].text.startswith("图4 样本级路由相对全局可学习权重模型")
assert len(p[90]._p.xpath(".//w:drawing")) == 1

set_text(
    p[89],
    "固定等权模型只能检验等权融合是否足够；全局可学习权重模型进一步控制了“权重不等”这一因素。"
    "表3显示，样本级路由在Abilene的6个种子、GÉANT的7个种子上取得较低MSE，"
    "但GÉANT种子43的MSE比全局权重模型高13.22%，使跨种子95%置信区间跨越零。"
    "图4进一步从OD通道和预测步两个维度检查这一差异的覆盖范围与边界。",
)
for run in list(p[90].runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
p[90].add_run().add_picture(str(FIGURE), width=Inches(5.85))
set_text(p[91], "图4 样本级路由相对全局可学习权重模型的通道覆盖与预测步误差变化")
set_text(
    p[92],
    "图4上排统计每条OD通道在8个共同随机种子中取得较低MSE的次数：Abilene的144条通道中有124条"
    "（86.1%）、GÉANT的524条通道中有410条（78.2%）在至少5个种子上改善；"
    "GÉANT仍有27条通道在8个种子中均未改善。下排逐种子计算相对全局权重模型的MSE降低率，"
    "浅线为单个种子，深线为8种子的逐步均值。Abilene在第1步平均降低28.70%，第24步转为升高19.36%；"
    "GÉANT第1步平均升高5.68%，第24步降低10.99%。两数据集分别有15/24和22/24个预测步的"
    "跨种子平均降低率为正。相邻预测步并非独立重复，图中曲线用于定位收益与失效区间，"
    "不能据此扩大显著性结论，也不能证明路由器按窗口识别尺度需求。",
)

conclusion = p[101]
old = "通道层面，分别有86.1%和78.2%的OD通道在至少5个种子上改善。"
assert old in conclusion.text
set_text(
    conclusion,
    conclusion.text.replace(
        old,
        old + "预测步分解显示改善方向随位置变化，尤其Abilene的后段预测步出现误差升高。",
    ),
)

doc.save(TARGET)
check = Document(TARGET)
assert check.paragraphs[88].text == p[88].text
assert len(check.paragraphs) == 114 and len(check.inline_shapes) == 5
assert "H=12" not in "\n".join(x.text for x in check.paragraphs)
print(TARGET)
