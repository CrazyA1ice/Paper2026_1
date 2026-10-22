"""Apply the global-weight control and router analysis to the v27 manuscript."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.23v27.docx"
TARGET = ROOT / "word_output" / "2026.09.23v28.docx"
FIGURE = ROOT / "paper_figures" / "router_seed_variation_v28.png"


def replace_paragraph(paragraph, expected, new_text):
    assert expected in paragraph.text, (expected, paragraph.text[:100])
    assert paragraph.runs
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_cell(cell, value):
    paragraph = cell.paragraphs[0]
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


def main():
    document = Document(SOURCE)
    paragraphs = document.paragraphs
    updates = {
        3: ("摘要：", "摘要：骨干网络源—目的节点对流量同时包含短时起伏与较慢趋势，单一时间分辨率难以兼顾。本文采用尺度1、2和4的DLinear专家，并由输入窗口统计量驱动的轻量路由网络计算Softmax融合权重。在Abilene和GÉANT数据集上，3个共同随机种子下本文方法相对DLinear的MSE分别降低4.06%和4.43%，相对LightTS分别降低29.89%和46.83%。8个随机种子的核心消融进一步加入可学习的全局固定权重对照；本文方法相对该对照的逐种子平均MSE改进率分别为1.80%和4.31%，分别有6次和7次取得更低误差。GÉANT的配对95%置信区间跨越零，且路由权重变化主要来自随机种子之间，因此当前结果支持平均误差改善，但尚不能证明路由器稳定识别每个窗口的最佳尺度。"),
        9: ("Abstract:", "Abstract: Backbone origin-destination traffic contains both short-term fluctuations and slower trends. We combine DLinear experts at scales 1, 2, and 4 using Softmax weights generated from input-window statistics. On Abilene and GÉANT, the method reduces MSE relative to DLinear by 4.06% and 4.43%, and relative to LightTS by 29.89% and 46.83%, over three common random seeds. An eight-seed ablation adds a learned global-weight baseline. The mean seed-wise MSE reductions of the adaptive router over this baseline are 1.80% and 4.31%, with improvements in six and seven of eight seeds, respectively. The paired 95% confidence interval for GÉANT includes zero, and most of its pooled router-weight variance arises between seeds. Thus, the observed average error reduction does not yet establish reliable window-specific scale selection."),
        57: ("为了区分", "式（13）利用上述权重对3个DLinear专家的预测结果进行加权求和，得到最终预测Ŷn。该组权重由输入窗口计算，并由同一窗口内的全部OD变量共享。为了区分多尺度结构、非等权分配与样本相关分配的作用，设置两种共享相同专家的对照：固定等权模型将三个尺度恒定设为1/3；全局可学习权重模型以三个可训练参数经Softmax得到一组常权重，对所有窗口共用。后者与本文方法的差别集中在权重是否依赖输入窗口。"),
        61: ("参数训练", "参数训练使用Adam优化器，并以验证集MSE选择最终模型。单尺度DLinear包含4656个可训练参数；三个尺度专家的固定等权模型包含8208个参数；增加三个全局权重参数后为8211个参数；样本级路由模型为8339个参数。路由网络相对固定等权模型增加131个参数，相对全局权重模型增加128个参数。各多尺度模型共享专家结构，因此全局权重对照可用于检验输入相关融合的额外作用。"),
        64: ("实验按", "实验按外部基线、核心消融和路由行为三个层次组织。DLinear和LightTS提供相同预测任务下的外部误差参照；固定等权、全局可学习常权重与样本级路由模型共享三个尺度专家，依次检验多尺度结构、非等权融合及输入相关权重的作用；八个共同随机种子的配对结果和逐种子权重分析用于评估结果波动。"),
        68: ("所有对比", "所有对比模型使用相同的数据划分、输入长度、预测长度、评价指标和批量大小，batch size统一为16。实验不为各模型单独搜索最优超参数，而是在统一预测任务下控制数据与训练条件。DLinear、固定等权多尺度、全局可学习权重和本文方法最多训练30轮，LightTS最多训练50轮；均根据验证集MSE选择测试所用参数，测试集仅用于最终统计。"),
        69: ("外部基线", "外部基线比较使用随机种子42～44共3个共同随机种子。核心消融使用随机种子42～49共8个共同随机种子，比较DLinear、固定等权多尺度、全局可学习权重和本文方法。全局权重由等权初始化，与三个专家一同训练，并按相同验证集准则选择模型；随机种子作为配对统计单位。"),
        71: ("实验设置3类", "实验设置DLinear、LightTS、固定等权多尺度和全局可学习权重四类对照。DLinear用于比较单一尺度与多尺度结构；LightTS提供独立轻量模型基线。后两种多尺度对照与本文方法保留相同的尺度集合{1,2,4}和DLinear专家：固定等权模型始终按1/3融合，全局权重模型学习一组对所有窗口共用的Softmax权重。由此，全局权重模型与本文方法的比较更直接地检验输入相关权重是否带来收益。"),
        82: ("固定等权消融", "两个数据集上的相对降幅并不相同。GÉANT包含的OD变量数高于Abilene，同时本文方法相对LightTS的MSE降幅也更大，但现有实验不能把这种差异直接归因于变量规模。固定等权与全局可学习权重对照均采用相同的多尺度专家结构，适合进一步检验融合方式本身的作用。"),
        84: ("表3中的", "表3中的核心消融覆盖随机种子42～49，共重复8次。相对单尺度DLinear，本文方法在Abilene和GÉANT上的平均MSE分别降低2.77%和2.37%；相对固定等权模型，平均MSE分别降低1.67%和4.68%。新增的全局可学习权重模型允许尺度权重不等，但对全部窗口保持恒定。与该模型逐种子配对比较，本文方法的平均MSE改进率分别为1.80%和4.31%，分别在6/8和7/8个种子上胜出；t分布近似的95%置信区间分别为[0.01%,3.58%]和[−2.51%,11.13%]。Abilene优势较小，GÉANT区间跨越零，不能据此断言跨初始化稳定优于全局权重。"),
        85: ("表3", "表3 八随机种子核心消融结果（均值±标准差）"),
        89: ("固定等权", "固定等权模型只能检验等权融合是否足够；一组经训练得到的全局常权重可进一步控制“权重不等”这一因素。表3显示本文方法在两数据集上的平均MSE低于全局权重模型，但种子间优势并不稳定，尤其GÉANT存在一次明显退化。图5仍展示相对DLinear和固定等权模型的多指标分布；全局权重的逐种子配对结果见表3及本节统计。"),
        91: ("图6汇总", "图6分别显示随机种子42～49下三个尺度的平均路由权重。每个圆点对应一个种子在全部测试窗口上的均值，误差线为同一种子内测试窗口的标准差，虚线表示1/3。这样可直接区分模型初始化造成的权重差异与同一模型随窗口变化的幅度。"),
        93: ("图6", "图6 各随机种子路由权重均值及测试窗口内标准差"),
        94: ("Abilene中", "Abilene三个尺度的汇总平均权重为0.257、0.319和0.424，GÉANT为0.283、0.318和0.399。虽然总体均值偏离1/3，但GÉANT三个尺度的汇总权重方差约96.2%、95.8%和97.7%来自种子之间；同一种子内随测试窗口的变化较小。Abilene的对应比例为48.7%、86.1%和80.6%。因此，合并所有种子的权重分布不能单独证明路由器有效识别每个窗口的尺度需求。"),
        95: ("表2、表3", "进一步以输入窗口平均绝对一阶差分衡量波动强度，并用训练窗口的三分位数阈值划分测试窗口。Abilene低、中、高波动组分别有15、406、392个窗口，本文方法相对全局权重模型的逐种子平均MSE改进率分别为4.43%、1.90%、1.69%；GÉANT在训练阈值下无低波动测试窗口，中、高组分别有76、375个窗口，改进率为5.85%和4.06%。Abilene高波动组的尺度1平均权重高于中波动组，但误差优势并未更强；GÉANT中、高组的权重近似相同。分组结果尚不支持“波动越强，路由收益越大”的解释。"),
        97: ("本文面向", "本文面向骨干网络OD流量中的短时波动和较慢趋势，以尺度1、2和4的DLinear专家构建多尺度预测，并通过轻量路由网络计算输入相关融合权重。在Abilene和GÉANT上，本文方法相对DLinear、LightTS和固定等权模型取得较低平均误差。加入全局可学习权重对照后，样本级路由的逐种子平均MSE改进率分别为1.80%和4.31%，但GÉANT的95%置信区间跨越零。逐种子权重及波动分组分析也未能充分证明路由器稳定地匹配每个窗口的尺度需求。后续需在更多数据和预测设置下检验这种输入相关融合的稳定收益。"),
    }
    for index, (expected, new_text) in updates.items():
        replace_paragraph(paragraphs[index], expected, new_text)

    table = document.tables[2]
    assert len(table.rows) == 7
    for source_index, cells in [
        (2, ["Abilene", "全局可学习权重", "8211", "0.192±0.002", "0.190±0.006"]),
        (6, ["GÉANT", "全局可学习权重", "8211", "0.446±0.014", "0.377±0.009"]),
    ]:
        original_row = table.rows[source_index]
        new_row = deepcopy(original_row._tr)
        original_row._tr.addnext(new_row)
        inserted_row = next(row for row in table.rows if row._tr is new_row)
        for cell, value in zip(inserted_row.cells, cells):
            replace_cell(cell, value)
    assert len(table.rows) == 9
    assert [row.cells[1].text for row in table.rows] == [
        "模型", "DLinear", "固定等权多尺度", "全局可学习权重", "本文方法",
        "DLinear", "固定等权多尺度", "全局可学习权重", "本文方法",
    ]

    figure_paragraph = paragraphs[92]
    assert len(figure_paragraph.runs) == 1
    assert figure_paragraph.runs[0]._element.xpath(".//w:drawing")
    figure_paragraph.runs[0]._element.getparent().remove(figure_paragraph.runs[0]._element)
    figure_paragraph.add_run().add_picture(str(FIGURE), width=Inches(5.4))
    figure_paragraph.paragraph_format.keep_with_next = True
    paragraphs[93].paragraph_format.keep_with_next = False
    paragraphs[94].paragraph_format.keep_together = True
    # Keep the method diagram's caption directly below its image.
    paragraphs[20]._p.addnext(paragraphs[22]._p)
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
