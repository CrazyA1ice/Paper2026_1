from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.20v5.docx"
TARGET = ROOT / "word_output" / "2026.09.21v6.docx"


def find_one(document: Document, startswith: str):
    matches = [p for p in document.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {startswith!r}, found {len(matches)}")
    return matches[0]


def set_paragraph_text(paragraph, text: str) -> None:
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def clone_paragraph(template, text: str):
    clone = deepcopy(template._p)
    for child in list(clone):
        if child.tag.endswith("}r"):
            clone.remove(child)
    run = OxmlElement("w:r")
    if template.runs and template.runs[0]._r.rPr is not None:
        run.append(deepcopy(template.runs[0]._r.rPr))
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    clone.append(run)
    return clone


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first.text = text
    for run in paragraph.runs[1:]:
        run.text = ""
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def build_table4(document: Document):
    table = document.tables[2]
    clone = deepcopy(table._tbl)
    values = [
        ["数据集", "模型", "参数量", "平均批延迟/ms", "吞吐量/(样本·s⁻¹)"],
        ["Abilene", "DLinear", "4656", "16.86±1.13", "951.9±56.7"],
        ["Abilene", "固定等权多尺度", "8208", "31.24±0.40", "510.9±6.2"],
        ["Abilene", "自适应多尺度", "8339", "40.65±0.19", "392.4±1.6"],
        ["GÉANT", "DLinear", "4656", "56.71±0.40", "274.5±1.9"],
        ["GÉANT", "固定等权多尺度", "8208", "101.53±2.69", "153.4±3.7"],
        ["GÉANT", "自适应多尺度", "8339", "131.35±1.25", "118.5±1.1"],
    ]
    from docx.table import Table

    wrapped = Table(clone, document)
    for row, row_values in zip(wrapped.rows, values):
        for cell, value in zip(row.cells, row_values):
            set_cell_text(cell, value)
    return clone


def main() -> None:
    if TARGET.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {TARGET}")
    document = Document(SOURCE)

    old_heading = find_one(document, "2.5 路由行为与适用范围")
    heading_template = old_heading
    body_template = find_one(document, "为从总体统计视角分析样本级路由器")
    caption_template = find_one(document, "表3 八随机种子核心模型结果")

    inserted = [
        clone_paragraph(heading_template, "2.5 参数规模与推理效率"),
        clone_paragraph(
            body_template,
            "为补充仅以参数量表征轻量性的局限，本文加载随机种子42—49的已训练checkpoint进行推理效率测试，不再训练模型。测试平台为Intel Core i7-8750H CPU，PyTorch固定为单线程，批量大小为16；每个checkpoint先预热30批，再对完整测试集重复计时10次。计时仅覆盖预先载入内存的测试张量在模型中的前向传播，不包含磁盘读取、数据预处理和训练时间。表4报告8个checkpoint的均值±样本标准差。",
        ),
        clone_paragraph(caption_template, "表4 核心模型参数量与CPU单线程推理效率（批量大小16，n=8）"),
        build_table4(document),
        clone_paragraph(
            body_template,
            "由表4可见，自适应多尺度模型相对固定等权模型仅增加131个参数，占固定等权模型参数量的1.60%；其平均批延迟在Abilene和GÉANT上分别增加30.10%和29.36%，表明样本级路由虽具有很小的参数开销，但仍会产生额外计算。在本测试协议下，本文模型在两数据集上的吞吐量分别为392.4和118.5样本/s，低于DLinear但保持可量化的推理能力。因此，本文所称“轻量”主要指参数规模受控，不将该结果外推为端到端部署速度或跨硬件的效率优势。",
        ),
    ]
    anchor = old_heading._p
    for element in inserted:
        anchor.addprevious(element)

    set_paragraph_text(old_heading, "2.6 路由行为与适用范围")

    synthesis = find_one(document, "综合表2、表3和图2—图4可见")
    set_paragraph_text(
        synthesis,
        synthesis.text.replace("综合表2、表3和图2—图4可见", "综合表2—表4和图2—图4可见"),
    )
    conclusion = find_one(document, "在Abilene和GÉANT上的实验表明")
    set_paragraph_text(
        conclusion,
        "在Abilene和GÉANT上的实验表明，样本级自适应多尺度加权可在较小参数增量下改善固定等权融合，并降低相对DLinear和LightTS的平均预测误差；CPU单线程测试进一步表明，本文模型在批量大小16时分别达到392.4和118.5样本/s，但其推理速度低于DLinear，故轻量性应限定为参数规模受控。GÉANT跨种子波动较大。后续将扩展更多网络、预测长度及拓扑约束，并在目标部署硬件上进一步检验稳定性与端到端效率。",
    )

    document.save(TARGET)

    check = Document(TARGET)
    full_text = "\n".join(p.text for p in check.paragraphs)
    assert "2.5 参数规模与推理效率" in full_text
    assert "2.6 路由行为与适用范围" in full_text
    assert "表4 核心模型参数量与CPU单线程推理效率" in full_text
    assert "392.4和118.5样本/s" in full_text
    assert len(check.tables) == 4
    assert len(check.inline_shapes) == 4
    assert len(re.findall(r"(?m)^\[\d+\]", full_text.split("参考文献", 1)[1])) == 11
    print(TARGET)


if __name__ == "__main__":
    main()
