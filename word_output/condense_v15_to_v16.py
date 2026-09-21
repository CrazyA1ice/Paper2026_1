from __future__ import annotations

from pathlib import Path
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"word_output"/"2026.09.21v15.docx"
TARGET=ROOT/"word_output"/"2026.09.21v16.docx"

doc=Document(SOURCE)

def find(prefix):
    xs=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(xs)!=1:
        raise RuntimeError(f"{prefix!r}: {len(xs)} matches")
    return xs[0]

def repl(prefix,text):
    p=find(prefix)
    first=p.runs[0] if p.runs else p.add_run()
    first.text=text
    for r in p.runs[1:]:
        r.text=""

def remove_para(p):
    el=p._p
    el.getparent().remove(el)
    p._p=p._element=None

replacements={
"摘要：":"摘要：针对骨干网络OD流量的多时间尺度变化，提出一种样本级自适应多尺度预测方法。采用尺度1、2和4的非重叠平均池化构造多分辨率序列，由DLinear尺度专家分别预测；基于窗口均值、标准差、末时刻均值和平均绝对一阶差分生成Softmax尺度权重，实现动态融合。在Abilene和GÉANT数据集上，与DLinear、LightTS及固定等权多尺度模型比较。共同3个随机种子下，本文方法MSE较DLinear分别降低4.06%和4.43%，较LightTS分别降低29.89%和46.83%；8个随机种子的核心消融中，较固定等权融合分别降低1.67%和4.68%。结果表明，样本级动态尺度加权能够以较小参数增量有效利用多时间尺度信息。",
"Abstract:":"Abstract: A sample-wise adaptive multiscale forecasting method is proposed for backbone origin-destination (OD) traffic. Multiresolution sequences at scales 1, 2, and 4 are constructed by non-overlapping average pooling and predicted by DLinear experts. Window mean, standard deviation, latest-step mean, and mean absolute first difference are used to generate Softmax scale weights for dynamic fusion. Experiments on Abilene and GÉANT compare the method with DLinear, LightTS, and fixed equal-weight multiscale fusion. Under three common random seeds, MSE is reduced by 4.06% and 4.43% versus DLinear and by 29.89% and 46.83% versus LightTS. In the eight-seed core ablation, MSE is reduced by 1.67% and 4.68% versus fixed equal weighting. The results show that sample-wise scale weighting effectively exploits multiscale temporal information with a small parameter increase.",
"随着云计算":"随着云计算、视频业务和数据中心互联的发展，骨干网络流量呈现明显的时变性和多时间尺度特征。OD流量预测可为容量规划、拥塞预警和流量工程提供依据。现有方法已由统计模型发展到机器学习和深度学习方法[1]，并进一步利用图神经网络等结构建模网络关联[2,10]。对于仅依赖历史OD序列且强调结构简洁的场景，仍有必要研究参数规模受控的时序预测方法。",
"网络流量在不同时间尺度上包含互补信息":"不同时间尺度包含互补流量信息：细粒度序列保留局部变化，粗粒度序列突出趋势。已有研究通过多尺度回声状态网络[3]、TimeMixer[5]和CycleLLH[11]提升多尺度建模能力；DLinear[4]和LightTS[6]则说明简洁结构也能获得较好的时间序列预测性能。这为在轻量结构中进一步研究自适应尺度融合提供了基础。",
"现有研究分别从拓扑关联":"现有方法多侧重拓扑关联、多尺度表征或轻量序列建模，而对“如何依据当前输入状态动态调整不同时间尺度的重要性”关注不足。本文采用尺度1、2和4构造多分辨率序列，以DLinear作为尺度专家，并由窗口统计特征驱动轻量路由器生成Softmax权重，实现样本级动态融合；在Abilene和GÉANT上通过外部基线、固定等权消融和8随机种子实验验证其有效性。",
"本文方法由数据预处理":"本文方法由多尺度序列构造、DLinear尺度专家和样本级自适应融合组成。输入窗口在尺度1、2和4上形成多分辨率序列，各尺度专家独立预测，路由器根据窗口统计特征生成Softmax权重并融合结果，整体流程见图1。",
"其中μctr和σctr只由训练集计算":"其中μctr和σctr仅由训练集计算，并用于验证集和测试集；数据按时间顺序以6∶2∶2划分，各集合内部独立构造滑动窗口。",
"三个尺度的输入长度依次为96、48和24":"三个尺度的输入长度依次为96、48和24。尺度1保留细粒度变化，尺度2和4逐步平滑局部扰动；平均池化不引入额外可训练参数。",
"式中，趋势项线性层和余项线性层分别具有独立的时间映射权重与偏置参数":"同一尺度内各变量共享时间映射参数，不同尺度使用独立参数。",
"固定平均默认三个尺度对每个窗口同等重要":"固定等权无法反映不同输入窗口的状态差异。本文从原始输入Xn提取四维统计向量zn=[μn,σn,ℓn,dn]，其中",
"μn描述窗口总体水平":"μn、σn、ℓn和dn分别描述窗口总体水平、离散程度、最近状态和平均变化强度。四维向量经隐藏维数为16的两层感知器得到三个尺度分数：",
"式中，两层感知器各自包含权重矩阵和偏置":"Softmax将三个尺度分数转换为非负且和为1的权重：",
"每个样本单独生成一组权重":"每个样本生成一组共享于全部OD变量的尺度权重。固定等权消融将αn,s设为1/3，其余结构与训练过程保持一致。",
"式中，N为训练样本数":"模型采用Adam优化并按验证集MSE选择参数。DLinear、固定等权多尺度和本文模型分别含4656、8208和8339个可训练参数，样本级路由器仅增加131个参数。",
"实验选用CESNET TS-Zoo":"实验采用CESNET TS-Zoo[7]中的Abilene[8]和GÉANT[9] 1 h聚合流量矩阵。Abilene保留144维有效OD流量，GÉANT去除恒定通道后保留524维。两数据集均按时间顺序以6∶2∶2划分，输入长度L=96、预测长度H=24，并采用第1.1节的预处理方式，具体规模见表1。",
"所有模型均通过统一训练接口完成":"所有模型采用统一数据划分和训练接口，批量大小为16，并按验证集MSE选择最优参数。DLinear、固定等权多尺度和本文模型最多训练30轮，LightTS最多训练50轮。外部基线比较使用共同随机种子42—44，核心模型进一步使用42—49共8个随机种子。",
"对比模型包括单尺度DLinear":"对比模型包括DLinear、LightTS、固定等权多尺度模型和本文方法。固定等权多尺度与本文方法采用相同尺度集合{1,2,4}和DLinear专家，仅将尺度权重固定为1/3，用于检验自适应加权的贡献。",
"采用均方误差（mean squared error":"采用均方误差（MSE）、平均绝对误差（MAE）和均方根误差（RMSE）评价预测性能。MSE按式（14）计算，MAE和RMSE分别为",
"上述指标均在标准化空间中计算":"上述指标均在标准化空间计算，数值越小表示误差越低。正文主要报告MSE和MAE；表2采用3个共同随机种子的均值±标准差，表3采用8个随机种子的均值±标准差。",
"为保证外部模型比较的重复次数一致":"表2给出了随机种子42—44的外部基线结果。本文方法在Abilene和GÉANT上的MSE分别较DLinear降低4.06%和4.43%，较LightTS降低29.89%和46.83%；MAE也均低于两类基线。",
"逐随机种子结果进一步见图2":"逐随机种子结果见图2。Abilene上，本文方法相对DLinear的MSE和MAE均为3/3次降低，相对LightTS也均为3/3次降低；GÉANT相对DLinear为2/3次降低，相对LightTS为3/3次降低。结合表2可见，本文方法在两个数据集上均取得更低的平均误差，其中相对LightTS的优势更为明显。",
"表3进一步使用随机种子42—49比较DLinear":"表3使用随机种子42—49比较DLinear、固定等权多尺度和本文方法。相较DLinear，本文方法在Abilene和GÉANT上的MSE分别降低2.77%和2.37%；相较固定等权多尺度分别降低1.67%和4.68%，说明自适应加权能够进一步改善多尺度融合。",
"同种子配对结果及配对差值分布见图3":"同种子配对结果见图3。Abilene中本文方法在6/8个随机种子上取得更低MSE，GÉANT为7/8；两数据集的平均ΔMSE均为正，与表3的均值结果一致。",
"为从总体统计视角分析样本级路由器":"图4汇总了随机种子42—49下全部测试窗口的三尺度路由权重，用于观察整体尺度偏好。",
"从图4可以看出":"图4显示，两数据集的权重均未固定在1/3。Abilene的α1、α2和α4平均为0.257、0.319和0.424，GÉANT分别为0.283、0.318和0.399，均表现为尺度4权重最高、尺度1最低。",
"综合表2、表3和图2—图5可见":"综合表2、表3和图2—图4，本文方法在两个骨干网数据集上均取得更低的平均预测误差，自适应权重相对固定等权进一步改善结果；图4的权重分布表明模型能够根据输入样本调整不同尺度的贡献。",
"两数据集实验表明":"两数据集实验表明，样本级自适应多尺度加权能够以较小参数增量改善固定等权融合，并降低相对DLinear和LightTS的平均预测误差。路由权重分析表明模型能够依据输入状态动态调整尺度组合。后续将进一步研究更多网络、预测长度和拓扑信息。"
}

for prefix,text in replacements.items():
    repl(prefix,text)

# Remove redundant Figure 5 and its local explanatory text.
repl("2.5 路由行为与适用范围","2.5 路由权重分析")
intro=find("为进一步观察权重在具体测试样本上的动态变化")
caption=find("图5 两数据集样本级路由权重变化")
analysis=find("图5中三个尺度权重均随测试窗口发生变化")
pic_el=caption._p.getprevious()
if pic_el is None or ("graphic" not in pic_el.xml and "drawing" not in pic_el.xml):
    raise RuntimeError("Expected Figure 5 drawing immediately before its caption")
pic_el.getparent().remove(pic_el)
remove_para(intro)
remove_para(caption)
remove_para(analysis)

# Final guardrails.
full_text="\n".join(p.text for p in doc.paragraphs)
assert "图5 两数据集样本级路由权重变化" not in full_text
assert "2.5 路由权重分析" in full_text
assert "图2 Abilene与GÉANT上三随机种子外部基线配对结果" in full_text
assert "图3 固定等权与自适应权重的八随机种子配对MSE及配对差值分布" in full_text
assert "图4 八随机种子下样本级自适应路由权重分布" in full_text
assert len(doc.inline_shapes)==4
assert len(doc.tables)==3
assert 300 <= len(find("摘要：").text.replace("摘要：","")) <= 400
assert len(find("两数据集实验表明").text) <= 150

doc.save(TARGET)
print(TARGET)
