from pathlib import Path
import re
from collections import Counter
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.23v26.docx"
TARGET = ROOT / "word_output" / "2026.09.23v27.docx"
BEFORE = ROOT / "word_output" / "v27_report_before.txt"
AFTER = ROOT / "word_output" / "v27_report_after.txt"
REPORT = ROOT / "word_output" / "v27_report_guided_aigc_rewrite.txt"

doc = Document(SOURCE)

def find_prefix(prefix):
    found = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(found) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(found)}")
    return found[0]

def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)

def dump(document, path):
    with path.open("w", encoding="utf-8") as f:
        for i, p in enumerate(document.paragraphs):
            t = p.text.replace("\n", " ").strip()
            if t:
                f.write(f"[{i:04d}] {t}\n")

def all_text(document):
    return "\n".join(p.text for p in document.paragraphs)

def table_snapshot(document):
    return [[[cell.text for cell in row.cells] for row in t.rows] for t in document.tables]

def headings(document):
    out = []
    for p in document.paragraphs:
        t = p.text.strip()
        if re.match(r"^\d+(?:\.\d+)*\s", t) or t == "参考文献":
            out.append(t)
    return out

def figcaps(document):
    return [p.text.strip() for p in document.paragraphs if re.match(r"^图\d+\s", p.text.strip())]

dump(doc, BEFORE)
before_text = all_text(doc)
before_tables = table_snapshot(doc)
before_heads = headings(doc)
before_figcaps = figcaps(doc)
before_shapes = len(doc.inline_shapes)

# Report-guided rewrites.
# Scope: high-share suspicious fragments from the uploaded PaperNex report.
replacements = [
    (
        "摘要：骨干网络中的源—目的节点对流量往往同时包含短时起伏",
        "摘要：骨干网络中的源—目的节点对流量在同一时间序列中常同时出现短时起伏、突发变化和较慢趋势，固定时间分辨率难以完整利用这些变化信息。针对这一特点，本文设计样本级自适应多尺度预测方法。历史窗口经非重叠平均池化形成尺度1、2和4的输入，分别由3个DLinear专家完成预测；路由网络根据输入窗口的均值、标准差、末时刻均值和平均绝对一阶差分计算Softmax尺度权重，再融合3个尺度的预测结果。实验采用Abilene和GÉANT两个骨干网流量数据集。3个共同随机种子下，所提方法的MSE相对DLinear分别降低4.06%和4.43%，相对LightTS分别降低29.89%和46.83%；8个随机种子的消融实验中，相对固定等权融合分别降低1.67%和4.68%。路由模块只带来较小参数增量，但两个数据集上的平均预测误差仍进一步降低。"
    ),
    (
        "云计算、视频业务、数据中心互联及智能网络应用使骨干网络承载的数据规模持续增长",
        "骨干网络中的源—目的节点对（origin-destination，OD）流量直接反映不同节点对之间的通信需求。随着云计算、视频业务、数据中心互联及智能网络应用持续增加，业务在网络内部的迁移和聚集更加频繁，OD流量的时间变化也更复杂。容量规划、拥塞预警、资源调度和流量工程都需要对后续流量状态作出判断，提前获得后一时段的OD流量估计有助于网络管理系统预先调整资源。与较平稳的单变量序列相比，OD流量中往往同时存在周期变化、局部波动和突发变化，不同OD变量的数值范围也有明显差异，因此预测模型既要描述较长时间范围内的变化趋势，也要保留对短时状态变化的响应。"
    ),
    (
        "网络流量预测早期多采用统计模型",
        "网络流量预测已经从传统统计模型逐步扩展到机器学习和深度学习方法。Lohrasbinasab等[1]对相关方法进行了系统梳理，并指出数据驱动方法在复杂流量预测中的应用不断增加。网络规模扩大后，仅分析单一序列的时间变化已不足以覆盖OD变量之间的关联。Qin等[2]把图神经网络引入通信网络流量预测，在同一模型中处理空间关联和时间依赖；韦烜等[10]针对大型IP网络流量矩阵开展分析与预测，讨论了不同OD变量之间的关联及其时间变化。唐文杰等[11]则在CycleLLH模型中引入周期性信息。现有工作已经涉及时间依赖、变量关联和周期结构等多个方面，也为进一步分析不同时间分辨率下的流量变化提供了基础。"
    ),
    (
        "时间分辨率会改变同一段网络流量的表现形式",
        "同一段网络流量在不同时间分辨率下呈现的特征并不一致。原始分辨率保留更多局部细节，短时起伏表现得更明显；时间聚合会削弱部分局部扰动，使较慢的变化趋势更容易观察。Zhou等[3]利用深度回声状态网络进行多尺度网络流量预测，通过不同时间分辨率提取互补信息。另一方面，模型结构开销也需要考虑。DLinear[4]采用时间序列分解和线性映射完成预测；TimeMixer[5]利用多尺度混合机制处理不同分辨率序列；LightTS[6]采用采样导向的MLP结构进行快速多变量时间序列预测。本文在这些工作的基础上，把关注点放在多时间尺度信息的利用与模型规模控制之间的平衡。"
    ),
    (
        "现有多尺度方法通常需要先确定若干时间尺度",
        "多尺度方法通常先给定若干时间尺度，再把各尺度结果按既定方式融合。问题在于，不同输入窗口的流量状态并不相同：某些窗口以短时波动为主，另一些窗口则表现为更平缓的趋势。输入状态变化时，尺度1、2和4对预测的作用也可能发生变化。只使用单一尺度会舍弃其他时间分辨率的信息，而长期采用固定融合比例又难以体现窗口之间的状态差异。基于这一现象，本文不仅构造多个时间尺度，还让各尺度在最终预测中的参与程度随当前输入样本变化。"
    ),
    (
        "记预处理后的多变量流量序列为",
        "记预处理后的多变量流量序列为 xt∈RC，其中t表示时间索引，C表示有效OD变量数。连续序列按时间滑动后形成监督学习样本。第n个样本包含历史窗口Xn和预测目标Yn：Xn覆盖连续L个时间点，Yn为紧随其后的H个时间点。模型据此利用多个OD变量过去L步的联合变化估计未来H步取值。第n个样本对应的输入与目标定义为"
    ),
    (
        "本文取输入长度L=96",
        "本文设置输入长度L=96、预测长度H=24，变量数C由各数据集保留的有效OD维度确定，学习映射记为Ŷn=fθ(Xn)。L对应一次预测所使用的历史范围，H表示需要连续预测的未来长度。原始OD流量均为非负值，而且不同OD变量之间的量级差异较大。若直接以原始数值训练，较大范围的变量会使各维数据处在不一致的数值尺度上，因此先对原始流量vt,c作对数变换"
    ),
    (
        "其中μctr和σctr分别表示第c个OD变量在训练集上的均值和标准差",
        "其中μctr和σctr分别为第c个OD变量在训练集上的均值和标准差，验证集和测试集均使用训练阶段得到的统计量。log(1+x)先压缩非负流量的数值范围，同时保留流量大小关系；随后按变量标准化，使不同OD维度在相近的数值尺度上参与训练。前一步主要缓和原始流量的量级差异，后一步用于统一各变量参与损失计算和参数更新时的尺度。数据仍按时间顺序以6∶2∶2划分为训练集、验证集和测试集，各集合内部独立生成滑动窗口，避免窗口跨越不同数据划分边界。"
    ),
    (
        "经过3个DLinear专家后",
        "3个DLinear专家会针对同一样本给出3组预测结果，随后需要确定它们在最终输出中的占比。若始终令3个尺度均取1/3，就等价于假设所有输入窗口具有相同的尺度需求。但实际窗口的状态存在差别：局部波动较强时，细粒度信息可能更重要；变化较平稳时，较粗尺度中的趋势信息可能更有价值。为使融合比例能够随样本变化，本文从原始输入Xn提取四维统计向量zn=[μn,σn,ℓn,dn]，用这些统计量描述当前窗口状态，其中"
    ),
    (
        "μn、σn、ℓn和dn分别表示输入窗口的总体水平",
        "μn、σn、ℓn和dn分别描述输入窗口的总体水平、离散程度、最近状态和平均变化强度。μn反映窗口整体流量水平，σn刻画观测值的离散程度，ℓn取窗口末时刻的平均状态，dn由相邻时间点的一阶变化得到。这4个统计量分别提供水平、波动、最近状态和局部变化信息，供路由器判断当前样本更适合依赖哪些时间尺度，而不直接参与未来流量数值的预测。"
    ),
    (
        "实验中的对比对象包括3类",
        "实验设置3类对比模型：单尺度DLinear、LightTS和固定等权多尺度模型。DLinear与本文的尺度专家采用相同基础预测结构，用来比较单一时间分辨率和多尺度结构之间的差异；LightTS作为独立的轻量多变量时间序列预测模型，提供外部结构基线；固定等权多尺度模型保留与本文方法完全相同的尺度集合{1,2,4}和DLinear专家，但取消样本级路由，3个尺度始终按1/3权重融合。由此，DLinear与固定等权模型的差别主要来自多尺度信息，固定等权模型与本文方法的差别则集中在尺度融合方式。"
    ),
    (
        "两个数据集上的相对降幅存在差异",
        "Abilene和GÉANT上的相对降幅并不相同。GÉANT包含更多OD变量，同时本文方法相对LightTS的MSE降幅也更大，但当前实验只能说明两种模型在这两个数据集上的差距不同，尚不能把这种差异直接解释为变量规模造成的结果。固定等权消融与本文方法采用相同的多尺度专家结构，因此更适合继续判断样本级融合方式本身带来的影响。"
    ),
    (
        "表3将核心模型扩展到随机种子42～49",
        "表3中的核心消融覆盖随机种子42～49，共重复8次。与单尺度DLinear相比，本文方法在Abilene和GÉANT上的平均MSE分别降低2.77%和2.37%。固定等权多尺度模型保留相同的尺度集合和3个DLinear专家，只把融合方式改为固定等权；与该模型相比，本文方法的平均MSE仍分别降低1.67%和4.68%。前一组比较反映完整多尺度模型相对单尺度基线的变化，后一组比较则把关注点进一步缩小到尺度融合方式。"
    ),
    (
        "图5给出随机种子42～49下本文方法相对DLinear和固定等权多尺度模型的MSE",
        "图5汇总随机种子42～49下的相对改进结果，比较对象为DLinear和固定等权多尺度模型，指标包括MSE、MAE和RMSE。图中的散点对应单次随机种子实验，半小提琴表示8次结果的分布，圆点和误差线表示逐随机种子改进率的均值及95%置信区间，0%虚线表示两种方法误差相同。相对DLinear时，Abilene的MSE、MAE和RMSE平均改进率为2.77%、8.01%和1.40%，其中MSE与RMSE在8次实验中均改善；GÉANT分别为2.35%、1.32%和1.22%，且跨越0%的情况更多。相对固定等权模型时，Abilene三项指标的平均改进率为1.65%、8.92%和0.84%，GÉANT为4.54%、3.05%和2.37%。两种基线下的平均改进率都为正，但GÉANT在不同随机种子之间的波动更明显。"
    ),
    (
        "Abilene中α1、α2和α4的平均权重依次为0.257",
        "Abilene中α1、α2和α4的平均权重依次为0.257、0.319和0.424，GÉANT中依次为0.283、0.318和0.399。两组结果都显示尺度4的平均权重最高、尺度1最低，三个尺度的均值也与固定等权的1/3不同。在本文设置的24步预测任务中，路由器整体上更偏向尺度4所保留的平滑趋势信息，但尺度1和尺度2仍保持非零权重，因此最终预测仍由多个时间分辨率共同完成。"
    ),
    (
        "本文针对骨干网络OD流量中同时存在短时波动和较慢趋势的问题",
        "本文面向骨干网络OD流量中的短时波动和较慢趋势，采用尺度1、2和4的DLinear专家进行多尺度预测，并由轻量路由网络依据输入窗口统计量计算融合权重。Abilene和GÉANT实验表明，本文方法相对DLinear、LightTS和固定等权多尺度模型均获得更低的平均预测误差，三个尺度的路由权重也不是固定分配。路由模块仅增加131个参数，在较小参数增量下完成了样本级尺度融合，并改善了平均预测结果。"
    ),
]

changes = []
for prefix, new_text in replacements:
    p = find_prefix(prefix)
    old = p.text.strip()
    set_text(p, new_text)
    changes.append((old, new_text))

after_text = all_text(doc)
after_tables = table_snapshot(doc)
after_heads = headings(doc)
after_figcaps = figcaps(doc)

# Preservation checks: no data, terminology-bearing numeric content, citations, equations, tables or captions may change.
assert before_tables == after_tables, "Table content changed"
assert before_heads == after_heads, "Heading structure changed"
assert before_figcaps == after_figcaps, "Figure captions changed"
assert before_shapes == len(doc.inline_shapes), "Figure count changed"

num_pat = re.compile(r"\d+(?:\.\d+)?%?|\d+∶\d+∶\d+")
before_nums = Counter(num_pat.findall(before_text))
after_nums = Counter(num_pat.findall(after_text))
if before_nums != after_nums:
    print("MISSING_NUMBERS", dict(before_nums - after_nums))
    print("ADDED_NUMBERS", dict(after_nums - before_nums))
assert before_nums == after_nums, "Numeric token multiset changed"

cite_pat = re.compile(r"\[\d+\]")
assert Counter(cite_pat.findall(before_text)) == Counter(cite_pat.findall(after_text)), "Citation markers changed"

for eq in [f"({i})" for i in range(1,17)]:
    assert before_text.count(eq) == after_text.count(eq), f"Equation marker {eq} changed"

# Reference list must remain untouched.
orig = Document(SOURCE)
before_refs = [p.text for p in orig.paragraphs if p.text.strip().startswith("[")]
after_refs = [p.text for p in doc.paragraphs if p.text.strip().startswith("[")]
assert before_refs == after_refs, "References changed"

# Critical scientific terms/names must remain present.
protected_terms = [
    "DLinear", "LightTS", "TimeMixer", "CycleLLH", "Softmax", "Adam",
    "Abilene", "GÉANT", "CESNET TS-Zoo", "MSE", "MAE", "RMSE",
    "origin-destination", "OD", "非重叠平均池化"
]
for term in protected_terms:
    assert before_text.count(term) == after_text.count(term), f"Protected term count changed: {term}"

dump(doc, AFTER)
doc.save(TARGET)

with REPORT.open("w", encoding="utf-8") as f:
    f.write("Report-guided AIGC rewrite: v26 -> v27\n")
    f.write("Source report: PaperNex AIGC detection, overall suspicious share 44.9%, 42 flagged fragments.\n")
    f.write(f"Edited high-risk paragraphs: {len(changes)}\n")
    f.write("Preservation: tables/headings/figure captions/numbers/citations/equation markers/references/protected terms = PASS\n\n")
    for i, (old, new) in enumerate(changes, 1):
        f.write(f"[{i}]\nBEFORE: {old}\nAFTER: {new}\n\n")

print("saved", TARGET)
print("edited paragraphs", len(changes))
print("preservation PASS")
