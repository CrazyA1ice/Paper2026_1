from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output" / "2026.09.20v3.docx"
TARGET = ROOT / "word_output" / "2026.09.20v4.docx"
FIGURE_DIR = ROOT / "abilene_forecasting" / "paper_figures" / "exports" / "color"
FIGURE_2 = FIGURE_DIR / "fig2_core_weight_pairing_color.png"
FIGURE_3 = FIGURE_DIR / "fig3_router_weight_dynamics_color.png"


def set_image(document: Document, shape, image_path: Path, width_inches: float,
              pixel_size: tuple[int, int]) -> None:
    blip = shape._inline.xpath(".//a:blip")[0]
    embed = blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    image_part = document.part.related_parts[embed]
    image_part._blob = image_path.read_bytes()
    shape.width = Inches(width_inches)
    shape.height = Inches(width_inches * pixel_size[1] / pixel_size[0])


def replace_caption(document: Document, old: str, new: str) -> None:
    matches = [p for p in document.paragraphs if p.text == old]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one caption {old!r}, found {len(matches)}")
    paragraph = matches[0]
    if not paragraph.runs:
        paragraph.add_run(new)
        return
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""


def main() -> None:
    if TARGET.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {TARGET}")
    document = Document(SOURCE)
    if len(document.inline_shapes) != 3:
        raise RuntimeError(f"Expected 3 inline figures, found {len(document.inline_shapes)}")

    set_image(document, document.inline_shapes[1], FIGURE_2, 6.30, (3938, 1634))
    set_image(document, document.inline_shapes[2], FIGURE_3, 6.30, (4005, 1747))

    replace_caption(
        document,
        "图2 固定等权与自适应权重的八随机种子配对MSE",
        "图2 固定等权与自适应权重的八随机种子配对MSE（细线连接同一随机种子，菱形及误差棒表示均值±标准差，n=8）",
    )
    replace_caption(
        document,
        "图3 两数据集样本级路由权重变化（随机种子42）",
        "图3 两数据集样本级路由权重变化（随机种子42，灰色点线为等权基准1/3）",
    )

    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
