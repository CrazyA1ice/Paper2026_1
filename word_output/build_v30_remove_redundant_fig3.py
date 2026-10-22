"""Remove the redundant paired-seed figure while retaining its key result."""

from pathlib import Path
import re

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.25v29.docx"
PREVIEW = ROOT / "word_output/2026.09.25v28_fig6_realdata_preview.docx"
TARGET = ROOT / "word_output/2026.09.25v30.docx"


def replace_text(paragraph, new_text):
    assert paragraph.runs
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


doc = Document(SOURCE)
original = Document(PREVIEW)
paragraphs = doc.paragraphs
assert len(doc.inline_shapes) == 5

remove = [
    p for p in paragraphs
    if p.text.startswith("图3逐随机种子展示本文方法")
    or p.text.startswith("图3 样本级路由相对全局可学习权重模型")
]
remove += [p for p in paragraphs if any("fig3_global_vs_adaptive_paired" in r.text for r in p.runs)]
# The picture's filename is stored in the DOCX relationship, not as run text.
table = doc.tables[2]
image_element = table._tbl.getnext().getnext()
assert len(image_element.xpath(".//w:drawing")) == 1
assert len(remove) == 2
for paragraph in remove:
    paragraph._p.getparent().remove(paragraph._p)
image_element.getparent().remove(image_element)

paragraphs = doc.paragraphs
core = next(p for p in paragraphs if p.text.startswith("表3中的核心消融覆盖"))
replace_text(core, original.paragraphs[84].text)
explanation = next(p for p in paragraphs if p.text.startswith("固定等权模型只能检验"))
replace_text(
    explanation,
    "固定等权模型只能检验等权融合是否足够；全局可学习权重模型进一步控制了“权重不等”这一因素。表3显示，样本级路由在Abilene的6个种子、GÉANT的7个种子上取得较低MSE，但GÉANT种子43的MSE比全局权重模型高13.22%，使跨种子95%置信区间跨越零。因此，现有结果不能证明样本级路由在不同初始化下稳定优于全局固定权重。图3另汇总相对DLinear和固定等权模型的多指标结果。",
)

for paragraph in paragraphs:
    if re.search(r"图[45]", paragraph.text):
        replace_text(paragraph, re.sub(r"图([45])", lambda m: "图" + {"4": "3", "5": "4"}[m[1]], paragraph.text))

assert len(doc.inline_shapes) == 4
doc.save(TARGET)
print(TARGET)
