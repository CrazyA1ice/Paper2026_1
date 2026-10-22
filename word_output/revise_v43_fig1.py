"""Replace the first manuscript figure with the selected simplified design."""

from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.27v42.docx"
TARGET = ROOT / "word_output/2026.09.27v43.docx"
FIGURE = ROOT / "paper_figures/fig1_redesign_options/A_refined.png"

document = Document(SOURCE)
figure_paragraph = document.paragraphs[31]
assert figure_paragraph._p.xpath(".//w:drawing")
assert document.paragraphs[32].text == "图1 自适应多尺度网络流量预测方法框架"
blip = figure_paragraph._p.xpath(".//a:blip")[0]
relationship_id = blip.get(
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)
assert document.part.rels[relationship_id].target_ref == "media/image1.png"

image_bytes = FIGURE.read_bytes()
with Image.open(BytesIO(image_bytes)) as image:
    width_px, height_px = image.size
document.part.related_parts[relationship_id]._blob = image_bytes
shape = document.inline_shapes[0]
shape.height = round(shape.width * height_px / width_px)
document.save(TARGET)

check = Document(TARGET)
assert len(check.inline_shapes) == 5 and len(check.tables) == 3
assert check.part.related_parts[relationship_id].blob == image_bytes
assert check.inline_shapes[0].height == shape.height
assert [p.text for p in check.paragraphs] == [p.text for p in document.paragraphs]
print(TARGET)
