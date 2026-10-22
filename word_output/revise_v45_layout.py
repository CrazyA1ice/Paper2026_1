"""Keep the training-objective heading and its equation together."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.27v44.docx"
TARGET = ROOT / "word_output/2026.09.27v45.docx"

document = Document(SOURCE)
paragraphs = document.paragraphs
assert paragraphs[58].text == "1.5 训练目标与模型复杂度"
assert paragraphs[59].text.startswith("模型训练以均方误差作为优化目标。")
assert paragraphs[60].text.endswith("(14)")
paragraphs[58].paragraph_format.keep_with_next = True
paragraphs[59].paragraph_format.keep_together = True
paragraphs[59].paragraph_format.keep_with_next = True
document.save(TARGET)

check = Document(TARGET)
assert [p.text for p in check.paragraphs] == [p.text for p in paragraphs]
assert len(check.inline_shapes) == 5 and len(check.tables) == 3
print(TARGET)
