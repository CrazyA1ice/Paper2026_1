from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"C:\Users\Quant\Desktop\Paper2026_1")
SOURCE = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_定稿.md"
TEMPLATE = ROOT / "word_output" / "template_inspection" / "论文格式模板_converted.docx"
OUTPUT = ROOT / "word_output" / "基于自适应多尺度加权的多变量网络流量预测方法_图1重构版.docx"
FIG_DIR = ROOT / "abilene_forecasting" / "paper_figures" / "exports"


def set_run_font(run, size=10.5, bold=False, east_asia="宋体", latin="Times New Roman"):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trpr.append(cant_split)


def set_cell_margins(cell, top=60, start=70, bottom=60, end=70):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")


def clear_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def style_document(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    style_names = {s.name for s in document.styles}
    h1 = document.styles["Heading 1"] if "Heading 1" in style_names else document.styles.add_style("Heading 1", WD_STYLE_TYPE.PARAGRAPH)
    h1.font.name = "Times New Roman"
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = None
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.first_line_indent = Cm(0)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(2)
    h1.paragraph_format.keep_with_next = True

    h2 = document.styles["Heading 2"] if "Heading 2" in style_names else document.styles.add_style("Heading 2", WD_STYLE_TYPE.PARAGRAPH)
    h2.font.name = "Times New Roman"
    h2.font.size = Pt(10.5)
    h2.font.bold = True
    h2.font.color.rgb = None
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.space_before = Pt(4)
    h2.paragraph_format.space_after = Pt(1)
    h2.paragraph_format.keep_with_next = True

    if "Equation" not in [s.name for s in document.styles]:
        eq = document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        eq = document.styles["Equation"]
    eq.font.name = "Cambria Math"
    eq.font.size = Pt(10.5)
    eq._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    eq.paragraph_format.first_line_indent = Cm(0)
    eq.paragraph_format.space_before = Pt(2)
    eq.paragraph_format.space_after = Pt(2)
    eq.paragraph_format.keep_together = True


def add_label_paragraph(document, text, label, size=10.5, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = document.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r1 = p.add_run(label)
    set_run_font(r1, size=size, bold=True, east_asia="黑体")
    r2 = p.add_run(text[len(label):])
    set_run_font(r2, size=size)
    return p


def add_body_paragraph(document, text, reference_mode=False):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if reference_mode and re.match(r"^\[\d+\]", text):
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_run_font(r, size=9)
    else:
        p.paragraph_format.first_line_indent = Cm(0.74)
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = 1
    table.autofit = True
    set_table_borders(table)
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        prevent_row_split(row)
        if i == 0:
            set_repeat_table_header(row)
        for j, value in enumerate(row_data):
            cell = row.cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "E7E6E6")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = i < len(rows) - 1
            r = p.add_run(value)
            set_run_font(r, size=8.5, bold=(i == 0), east_asia="黑体" if i == 0 else "宋体")
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = 1.0


def add_figure(document, filename, caption):
    widths = {
        "fig1_method_schematic.png": 16.2,
        "fig2_main_ablation.png": 16.0,
        "fig3_weight_ablation.png": 16.0,
        "fig4_frequency_sensitivity.png": 16.0,
        "fig5_representative_forecast.png": 12.8,
        "fig6_router_weight_distribution.png": 16.0,
    }
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(FIG_DIR / filename), width=Cm(widths[filename]))
    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Cm(0)
    cp.paragraph_format.line_spacing = 1.0
    cp.paragraph_format.space_after = Pt(3)
    cp.paragraph_format.keep_together = True
    r = cp.add_run(caption)
    set_run_font(r, size=9)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        vals = [x.strip() for x in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", x) for x in vals):
            rows.append(vals)
        i += 1
    return rows, i


def build():
    document = Document(str(TEMPLATE))
    clear_body(document)
    style_document(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    title_count = 0
    reference_mode = False
    equation_counter = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        text = raw.strip()
        if not text:
            i += 1
            continue
        if text.startswith("{{FIG:"):
            m = re.match(r"\{\{FIG:([^|]+)\|(.+)\}\}", text)
            if not m:
                raise ValueError(f"Bad figure marker: {text}")
            add_figure(document, m.group(1), m.group(2))
            i += 1
            continue
        if text.startswith("$$") and text.endswith("$$"):
            equation_counter += 1
            eq_text = text[2:-2].strip().replace("\\quad", "    ").replace("\\;", " ")
            p = document.add_paragraph(style="Equation")
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run("\t")
            r = p.add_run(eq_text)
            set_run_font(r, size=10.5, east_asia="Cambria Math", latin="Cambria Math")
            p.add_run("\t")
            number_run = p.add_run(f"({equation_counter})")
            set_run_font(number_run, size=10.5, east_asia="宋体", latin="Times New Roman")
            i += 1
            continue
        if text.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(document, rows)
            continue
        if text.startswith("# "):
            title_count += 1
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text[2:])
            set_run_font(r, size=14, bold=True, east_asia="黑体")
            i += 1
            continue
        if text.startswith("## "):
            heading = text[3:]
            p = document.add_paragraph(heading, style="Heading 1")
            if heading == "参考文献":
                reference_mode = True
            i += 1
            continue
        if text.startswith("### "):
            document.add_paragraph(text[4:], style="Heading 2")
            i += 1
            continue
        if re.match(r"^表\d+\s", text):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(text)
            set_run_font(r, size=9, bold=True, east_asia="黑体")
            i += 1
            continue
        if text.startswith("作者：") or text.startswith("（") or (text.startswith("[") and text.endswith("]") and not reference_mode):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.35
            r = p.add_run(text.replace("作者：", ""))
            set_run_font(r, size=10.5)
            i += 1
            continue
        matched_label = False
        for label in ("摘要：", "关键词：", "中图分类号：", "Abstract:", "Key words:"):
            if text.startswith(label):
                add_label_paragraph(document, text, label, indent=False if label in ("关键词：", "中图分类号：", "Key words:") else True)
                matched_label = True
                break
        if matched_label:
            i += 1
            continue
        add_body_paragraph(document, text, reference_mode=reference_mode)
        i += 1

    core = document.core_properties
    core.title = "基于自适应多尺度加权的多变量网络流量预测方法"
    core.subject = "多变量网络流量预测"
    core.keywords = "网络流量预测；多变量时间序列；多尺度学习；自适应加权；DLinear；频域分析"
    core.comments = "根据用户提供论文模板生成；作者与单位信息保留待填写占位符。"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
