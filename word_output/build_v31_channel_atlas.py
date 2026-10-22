"""Insert the OD-channel evidence figure into the manuscript."""

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.25v30.docx"
TARGET = ROOT / "word_output/2026.09.25v31.docx"
FIGURE = ROOT / "paper_figures/dense_options/option_D_channel_atlas.png"


def replace_text(paragraph, value):
    assert paragraph.runs
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


def clone_after(anchor, template):
    element = deepcopy(template._p)
    anchor.addnext(element)
    return Paragraph(element, template._parent)


doc = Document(SOURCE)
paragraphs = doc.paragraphs
assert paragraphs[89].text.startswith("固定等权模型只能检验")
assert paragraphs[90].text.startswith("2.5")
assert paragraphs[93].text.startswith("图4 ")
assert len(doc.inline_shapes) == 4

for paragraph in paragraphs:
    if re.search(r"图4", paragraph.text):
        replace_text(paragraph, paragraph.text.replace("图4", "图5"))

replace_text(
    paragraphs[89],
    "固定等权模型只能检验等权融合是否足够；全局可学习权重模型进一步控制了“权重不等”这一因素。表3显示，样本级路由在Abilene的6个种子、GÉANT的7个种子上取得较低MSE，但GÉANT种子43的MSE比全局权重模型高13.22%，使跨种子95%置信区间跨越零。为检查收益是否集中在少数OD通道，图4进一步统计每条通道在8个种子下相对全局权重模型的MSE胜出次数。",
)

image = clone_after(paragraphs[89]._p, paragraphs[87])
for run in list(image.runs):
    if run._element.xpath(".//w:drawing"):
        run._element.getparent().remove(run._element)
image.add_run().add_picture(str(FIGURE), width=Inches(5.4))
image.paragraph_format.keep_with_next = True

caption = clone_after(image._p, paragraphs[88])
replace_text(caption, "图4 样本级路由相对全局可学习权重模型的OD通道改善覆盖情况")
caption.paragraph_format.keep_with_next = False

interpretation = clone_after(caption._p, paragraphs[89])
replace_text(
    interpretation,
    "图4上排给出每条OD通道在8个随机种子中取得较低MSE的次数，下排按全局模型的通道MSE由低到高排列，并显示各通道的逐种子比较。Abilene的144条通道中有124条（86.1%）、GÉANT的524条通道中有410条（78.2%）在至少5个种子上改善。GÉANT仍有27条通道在8个种子中均未改善，种子43的负向结果也较集中。通道层面的覆盖范围为方法提供了补充证据，但不能消除初始化不稳定，也不能直接证明路由器按测试窗口识别尺度需求。",
)
interpretation.paragraph_format.keep_with_next = False
interpretation.paragraph_format.keep_together = True

conclusion = next(p for p in doc.paragraphs if p.text.startswith("本文面向骨干网络OD流量中的短时波动"))
replace_text(
    conclusion,
    conclusion.text.replace(
        "但GÉANT的95%置信区间跨越零。",
        "但GÉANT的95%置信区间跨越零。通道层面，分别有86.1%和78.2%的OD通道在至少5个种子上改善。",
    ),
)

assert len(doc.inline_shapes) == 5
doc.save(TARGET)
print(TARGET)
