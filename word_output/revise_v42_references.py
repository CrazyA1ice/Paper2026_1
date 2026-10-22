"""Number references in the order they first appear in the v41 manuscript."""

import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "word_output/2026.09.27v41.docx"
TARGET = ROOT / "word_output/2026.09.27v42.docx"
OLD_ORDER = (1, 2, 10, 11, 3, 4, 5, 6, 7, 8, 9)
NEW_NUMBER = {old: new for new, old in enumerate(OLD_ORDER, 1)}
CITATION = re.compile(r"\[(\d+)\]")

document = Document(SOURCE)
paragraphs = document.paragraphs
heading = next(p for p in paragraphs if p.text == "参考文献")
heading_index = paragraphs.index(heading)
assert heading_index == 103

first_seen = []
for paragraph in paragraphs[:heading_index]:
    original = paragraph.text
    for match in CITATION.finditer(original):
        number = int(match.group(1))
        assert number in NEW_NUMBER
        if number not in first_seen:
            first_seen.append(number)
    for run in paragraph.runs:
        updated = CITATION.sub(
            lambda m: f"[{NEW_NUMBER[int(m.group(1))]}]", run.text
        )
        if updated != run.text:
            run.text = updated
    assert paragraph.text == CITATION.sub(
        lambda m: f"[{NEW_NUMBER[int(m.group(1))]}]", original
    ), "Citation split across runs"
assert tuple(first_seen) == OLD_ORDER

reference_paragraphs = paragraphs[heading_index + 1 :]
assert len(reference_paragraphs) == len(OLD_ORDER)
references = {}
for paragraph in reference_paragraphs:
    match = re.match(r"^\[(\d+)\]", paragraph.text)
    assert match, paragraph.text
    references[int(match.group(1))] = paragraph
assert set(references) == set(OLD_ORDER)

anchor = heading._p
for old_number in OLD_ORDER:
    paragraph = references[old_number]
    anchor.addnext(paragraph._p)
    anchor = paragraph._p
    for run in paragraph.runs:
        if re.match(r"^\[\d+\]", run.text):
            run.text = re.sub(
                r"^\[\d+\]", f"[{NEW_NUMBER[old_number]}]", run.text, count=1
            )
            break
    else:
        raise AssertionError(f"Reference {old_number} label split across runs")

document.save(TARGET)
check = Document(TARGET)
check_paragraphs = check.paragraphs
check_heading_index = next(
    i for i, paragraph in enumerate(check_paragraphs) if paragraph.text == "参考文献"
)
assert check_heading_index == heading_index
seen = []
for paragraph in check_paragraphs[:check_heading_index]:
    for match in CITATION.finditer(paragraph.text):
        number = int(match.group(1))
        assert 1 <= number <= len(OLD_ORDER)
        if number not in seen:
            seen.append(number)
assert seen == list(range(1, len(OLD_ORDER) + 1))
assert [
    int(re.match(r"^\[(\d+)\]", paragraph.text).group(1))
    for paragraph in check_paragraphs[check_heading_index + 1 :]
] == list(range(1, len(OLD_ORDER) + 1))
for new_number, old_number in enumerate(OLD_ORDER, 1):
    original = references[old_number].text
    current = check_paragraphs[check_heading_index + new_number].text
    assert current == original
assert len(check.inline_shapes) == len(document.inline_shapes) == 5
assert len(check.tables) == len(document.tables) == 3
print(TARGET)
