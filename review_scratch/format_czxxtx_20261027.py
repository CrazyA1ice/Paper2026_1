"""Apply the supplied journal format while preserving scientific content."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import re, json, hashlib, shutil
from lxml import etree
from PIL import Image
from docx import Document
from docx.shared import Pt, Twips, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'paper_template'/'长江信息通信'
WORK=BASE/'格式核对工作_20261027'
WORK.mkdir(exist_ok=True)
SOURCE=BASE/'自适应多尺度加权的多变量网络流量预测_czxxtx20261026_结构精简稿.docx'
OUTPUT=BASE/'自适应多尺度加权的多变量网络流量预测_czxxtx20261027_格式核对修订稿.docx'
BACKUP=BASE/'his'/'自适应多尺度加权的多变量网络流量预测_czxxtx20261026_结构精简稿_格式修改前备份.docx'
ART=WORK/'独立灰度图_600dpi'
ART.mkdir(exist_ok=True)

def digest(b):return hashlib.sha256(b).hexdigest()
sourcebytes=SOURCE.read_bytes()
if BACKUP.exists():
    assert BACKUP.read_bytes()==sourcebytes, 'Backup already exists with different content; do not overwrite.'
else:BACKUP.write_bytes(sourcebytes)
doc=Document(BytesIO(sourcebytes))
old=list(doc.paragraphs)
original_body='\n'.join(p.text for p in old if p.style.name=='JournalBody')
oldmath=[''.join(p._p.xpath('.//m:t/text()')) for p in old if p._p.xpath('.//m:oMath')]
oldtables=[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
ops=[]

def fonts(rpr,ea='宋体',latin='Times New Roman'):
    f=rpr.find(qn('w:rFonts'))
    if f is None:f=OxmlElement('w:rFonts');rpr.insert(0,f)
    for a in list(f.attrib):del f.attrib[a]
    for a,v in [('ascii',latin),('hAnsi',latin),('eastAsia',ea),('cs',latin)]:f.set(qn('w:'+a),v)

def group_number(m):
    s=m.group()
    if '.' in s:
        a,b=s.split('.')
        a='\u00a0'.join([a[max(0,len(a)-j-3):len(a)-j] for j in range(0,len(a),3)][::-1])
        return a+'.'+'\u00a0'.join(b[j:j+3] for j in range(0,len(b),3))
    return '\u00a0'.join([s[max(0,len(s)-j-3):len(s)-j] for j in range(0,len(s),3)][::-1])

def grouped(text):
    pat=re.compile(r'(?<![\w.])\d{4,}(?![\w.年])|(?<![\w.])\d+\.\d{4,}(?![\w.])')
    # Chinese is a word character in Python, so use explicit Latin/digit boundaries.
    pat=re.compile(r'(?<![A-Za-z0-9_.])(?:\d{4,}(?![0-9.年])|\d+\.\d{4,})(?![A-Za-z0-9_.])')
    return pat.sub(group_number,text)

varpat=re.compile(r'(?<![A-Za-z0-9_])(?:[xXyYTRWbCLHNSPqmsijthcnvua]|[μσαΩℓ])(?:[₀₁₂₃₄₅₆₇₈₉ₙₜₛₑᵗʳ⁻⁸]*)(?![A-Za-z0-9_])')
submap=str.maketrans({'₀':'0','₁':'1','₂':'2','₃':'3','₄':'4','₅':'5','₆':'6','₇':'7','₈':'8','₉':'9','ₙ':'n','ₜ':'t','ₛ':'s','ₑ':'e'})
supmap=str.maketrans({'ᵗ':'t','ʳ':'r','⁻':'−','⁸':'8'})

def add_text(p,text,variables=False):
    matches=[]
    if variables:
        for m in re.finditer(r'\[\d+(?:[-,，]\d+)*\]',text):matches.append((m.start(),m.end(),'cite',m.group()))
    if variables:
        for m in varpat.finditer(text):
            if any(m.start()<e and m.end()>s for s,e,_,_ in matches):continue
            if m.group()=='h' and re.search(r'\d\s*$',text[:m.start()]):continue
            if m.group() in 'abcd' and text[max(0,m.start()-1):m.start()] in ['（','('] and text[m.end():m.end()+1] in ['）',')']:continue
            matches.append((m.start(),m.end(),'var',m.group()))
    pos=0
    for start,end,kind,token in sorted(matches):
        if start>pos:p.add_run(text[pos:start])
        if kind=='cite':p.add_run(token).font.superscript=True
        else:
            for bit in re.findall(r'[₀₁₂₃₄₅₆₇₈₉ₙₜₛₑ]+|[ᵗʳ⁻⁸]+|[^₀₁₂₃₄₅₆₇₈₉ₙₜₛₑᵗʳ⁻⁸]+',token):
                if all(ch in '₀₁₂₃₄₅₆₇₈₉ₙₜₛₑ' for ch in bit):
                    r=p.add_run(bit.translate(submap));r.font.subscript=True;r.italic=True
                elif all(ch in 'ᵗʳ⁻⁸' for ch in bit):
                    r=p.add_run(bit.translate(supmap));r.font.superscript=True;r.italic=bool(re.search('[ᵗʳ]',bit))
                else:r=p.add_run(bit);r.italic=True
        pos=end
    if pos<len(text):p.add_run(text[pos:])

def rewrite(p,text,variables=False):
    before=p.text
    p.clear();add_text(p,text,variables)
    if before!=text:ops.append({'before':before,'after':text})

replacements={
 3:'摘要：针对骨干网小时级源—目的（origin-destination，OD）流量的多尺度变化，提出自适应多尺度加权预测方法。采用1、2、4尺度的独立分解线性模型（decomposition-linear，DLinear）生成候选预测，利用四维窗口统计特征计算融合权重，并在OD变量间共享时间映射。Abilene和GÉANT的24 h预测均方误差（mean squared error，MSE）较LightTS平均降低25.77%和43.74%，参数量为其6.41%和2.17%；Abilene高需求平均绝对误差（mean absolute error，MAE）与平均低估误差（average underestimation error，AUE）较DLinear降低13.96%和16.07%；GÉANT一周预测MSE降低23.16%。该方法以低参数量实现多尺度流量预测。',
 5:'中图分类号：TP393    文献标识码：A',
 8:'(School of Information Engineering, Anhui Finance & Trade Vocational College, Hefei 231200, China)',
 9:'Abstract: An adaptive multiscale weighting method was developed to model multiscale variations in hourly origin-destination (OD) traffic in backbone networks. Three independent decomposition-linear (DLinear) experts generated candidate forecasts at scales of 1, 2, and 4. Four window statistics determined the fusion weights, while temporal mappings were shared across OD variables. The statistics consisted of the window mean, standard deviation, mean of the last time step, and mean absolute first-order difference. The three scales covered the same historical interval, and each expert produced forecasts at a common output horizon. For the 24-hour task, 96-hour inputs were used, and the adaptive model contained 8 339 parameters on both datasets. The method was evaluated on the Abilene and GÉANT datasets using eight common random seeds. Relative to LightTS, the mean paired reductions in 24-hour mean squared error (MSE) were 25.77% and 43.74%, with parameter counts equivalent to 6.41% and 2.17% of the respective LightTS configurations. On Abilene, high-demand mean absolute error (MAE) and average underestimation error (AUE) decreased by 13.96% and 16.07% relative to DLinear. For one-week forecasting on GÉANT, MSE decreased by 23.16% relative to DLinear. The method combines multiscale forecasts with a small parameter budget.',
 12:old[12].text.replace('源—目的节点对（origin-destination，OD）','源—目的（origin-destination，OD）节点对'),
 13:old[13].text.replace('分解线性模型（DLinear）','分解线性模型（decomposition-linear，DLinear）'),
 16:old[16].text.replace('经两层网络生成尺度权重','经两层多层感知机（multilayer perceptron，MLP）生成尺度权重')+' 图1展示24 h基础任务，一周任务的输出长度为168步。',
 18:'图1 自适应多尺度加权网络流量预测模型',
 57:old[57].text+' 散点表示单次配对，菱形及误差线表示均值±样本标准差；（a）、（b）分别为Abilene和GÉANT。',
 61:'图2 24 h预测误差的配对降幅与参数量比例',
 69:old[69].text,
 71:old[71].text+' 散点表示单次重复，圆形和方形分别表示MAE和AUE，误差线表示样本标准差；（a）、（b）分别为Abilene和GÉANT。',
 73:'图3 高需求目标点的MAE与AUE',
 76:old[76].text+' 图4中（a）、（c）分别为MSE和MAE，（b）、（d）分别为对应指标的配对降幅；散点表示单次重复，菱形及误差线表示均值±样本标准差。',
 79:'图4 GÉANT一周预测误差与配对降幅',
 87:old[87].text.replace('QIN Liang, GU Huaxi, WEI Wenting','QIN L, GU H, WEI W'),
 90:old[90].text.replace('ZHOU Jian, HAN Taotao, XIAO Fu','ZHOU J, HAN T, XIAO F'),
 91:old[91].text.replace('ZENG Ailing, CHEN Muxi, ZHANG Lei','ZENG A, CHEN M, ZHANG L'),
 92:'[7] WANG S, WU H, SHI X, et al. TimeMixer: decomposable multiscale mixing for time series forecasting[C/OL]//The Twelfth International Conference on Learning Representations, May 7-11, 2024, Vienna, Austria. 2024[2026-10-21]. https://openreview.net/forum?id=7oLshfEIC2.',
 93:old[93].text.replace('LIU Xinhe, WANG Wenmin','LIU X, WANG W'),
 94:'[9] KUREŠ M, KOUMAR J, HYNEK K. CESNET TS-Zoo: a library for reproducible analysis of network traffic time series[C]//2025 21st International Conference on Network and Service Management, October 27-31, 2025, Bologna, Italy. Piscataway: IEEE, 2025: 1-5. DOI: 10.23919/CNSM67658.2025.11297513.',
 95:old[95].text.replace('ZHANG Yin','ZHANG Y'),
 97:old[97].text.replace('[CP/OL]','[EB/OL]'),
 98:'作者简介：柯亚文（1998—），男（汉族），安徽池州人，硕士研究生学历、硕士学位，安徽财贸职业学院信息工程学院，助教，主要研究方向：计算机软件。联系电话：18983096676；电子邮箱：keyawen@afc.edu.cn。'
}
for i,p in enumerate(old):
    if p.style.name=='JournalBody':rewrite(p,grouped(replacements.get(i,p.text)),True)
    elif i in replacements:rewrite(p,replacements[i])
    if p.style.name in ['JournalHeading1','JournalHeading2']:
        t=p.text
        t=re.sub(r'^(\d+)(?=[.\u3000 ])',lambda m:str(int(m.group())+1),t)
        rewrite(p,t)

# Chinese and English bibliography entries retain one reference number.
translations={
88:'WEI X, LIU Z, LI Q, et al. Research on analysis and prediction of traffic matrix for large-scale IP network[J]. Systems Engineering and Electronics, 2024, 46(6): 2164-2173. DOI: 10.12305/j.issn.1001-506X.2024.06.35.',
89:'TANG W J, XIAO Y L, KONG X Y, et al. CycleLLH: a new network traffic prediction model based on cycle integration[J]. Chinese Journal of Computers, 2024, 47(12): 2867-2888. DOI: 10.11897/SP.J.1016.2024.02867.'
}
for i,text in translations.items():
    p=old[i].insert_paragraph_before() # Move it immediately after its Chinese entry.
    old[i]._p.addnext(p._p);p.style='JournalReference';add_text(p,text)
    ops.append({'kind':'English reference counterpart','ref':i-85,'after':text})

# Keep both displays immediately after their shared introduction, before interpretation.
old[61]._p.addnext(old[59]._p)

# Distilled page geometry. The template has one column, not a two-column article layout.
for s in doc.sections:
    s.page_width=Twips(11906);s.page_height=Twips(16838)
    s.left_margin=Twips(1134);s.right_margin=Twips(1134)
    s.top_margin=Twips(851);s.bottom_margin=Twips(851)
    s.header_distance=Twips(851);s.footer_distance=Twips(992)
    cols=s._sectPr.find(qn('w:cols'))
    if cols is not None:
        cols.set(qn('w:num'),'1')
        for ch in list(cols):cols.remove(ch)

sizes={'Title':14,'JournalAuthor':10.5,'JournalAffiliation':10.5,'JournalMetadata':10.5,
       'JournalAbstract':10.5,'JournalBody':10.5,'JournalHeading1':14,'JournalHeading2':10.5,
       'JournalCaption':10.5,'JournalReference':10.5,'JournalReferenceHeading':14,
       'JournalEquation':10,'JournalTable':9,'JournalFigure':10.5}
for name,size in sizes.items():
    st=doc.styles[name];st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    fonts(st.element.get_or_add_rPr())
    st.font.bold=name in ['Title','JournalHeading1','JournalHeading2','JournalReferenceHeading']
    pf=st.paragraph_format;pf.line_spacing=1.5;pf.space_before=Pt(0);pf.space_after=Pt(0)
    pf.first_line_indent=Pt(21 if name=='JournalBody' else 0)
    pf.left_indent=Pt(0);pf.right_indent=Pt(0)
    pf.widow_control=True
    if name in ['Title','JournalAuthor','JournalAffiliation','JournalCaption','JournalFigure','JournalTable']:pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif name in ['JournalHeading1','JournalHeading2','JournalReferenceHeading']:pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
    else:pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if name in ['JournalHeading1','JournalReferenceHeading']:pf.space_before=Pt(5);pf.space_after=Pt(3)
    if name=='JournalHeading2':pf.space_before=Pt(3);pf.space_after=Pt(2)
    if name=='Title':pf.space_after=Pt(6)
    pf.keep_with_next=name in ['Title','JournalHeading1','JournalHeading2','JournalReferenceHeading']
    for e in st.element.xpath('./w:pPr/w:pBdr'):e.getparent().remove(e)

for p in doc.paragraphs:
    name=p.style.name
    for e in p._p.xpath('.//w:lastRenderedPageBreak'):e.getparent().remove(e)
    for r in p.runs:
        fonts(r._r.get_or_add_rPr());r.font.color.rgb=RGBColor(0,0,0)
        if name in sizes:r.font.size=Pt(sizes[name])
    if name in ['JournalBody','JournalAbstract','JournalAffiliation','JournalAuthor','JournalMetadata','JournalReference']:
        # Remove stale local spacing and indents that would override the actual template settings.
        pf=p.paragraph_format;pf.line_spacing=1.5;pf.space_before=Pt(0);pf.space_after=Pt(0)
        pf.first_line_indent=Pt(21 if name=='JournalBody' else 0);pf.left_indent=Pt(0);pf.right_indent=Pt(0)
        pf.widow_control=True
        pf.keep_with_next=False
    if name=='JournalReference':
        p.paragraph_format.left_indent=Pt(18);p.paragraph_format.first_line_indent=Pt(-18)
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.superscript=False
        if p.text.startswith(('WEI X,','TANG W J,')):
            p.paragraph_format.first_line_indent=Pt(0)
    if name=='JournalCaption':
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=p.text.startswith('表')
        p.paragraph_format.space_after=Pt(3)
    if name in ['JournalHeading1','JournalHeading2']:p.paragraph_format.keep_with_next=True
    if name=='JournalEquation':
        pf=p.paragraph_format;pf.line_spacing=1.05;pf.space_before=Pt(3);pf.space_after=Pt(3)
        pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.keep_together=True
        pf.tab_stops.clear_all();pf.tab_stops.add_tab_stop(Twips(4819),WD_TAB_ALIGNMENT.CENTER);pf.tab_stops.add_tab_stop(Twips(9638),WD_TAB_ALIGNMENT.RIGHT)
        for rp in p._p.xpath('.//m:r/w:rPr'):
            fonts(rp)
            for tag in ['w:sz','w:szCs']:
                z=rp.find(qn(tag))
                if z is None:z=OxmlElement(tag);rp.append(z)
                z.set(qn('w:val'),'20')
        for mr in p._p.xpath('.//m:r'):
            if mr.find(qn('w:rPr')) is None:
                rp=OxmlElement('w:rPr');fonts(rp);z=OxmlElement('w:sz');z.set(qn('w:val'),'20');rp.append(z)
                mpr=mr.find(qn('m:rPr'));mr.insert(1 if mpr is not None else 0,rp)
        for t in p._p.xpath('.//m:t'):
            if t.text:t.text=re.sub(r'(?<!\d)8339(?!\d)','8\u00a0339',t.text)

for idx in [88,89]:old[idx].paragraph_format.keep_with_next=True

# Bold labels only, leaving abstracts in the third person and the conclusions in present tense.
for idx,label in [(3,'摘要：'),(4,'关键词：'),(5,'中图分类号：'),(9,'Abstract:'),(10,'Key words:')]:
    p=old[idx]
    if p.runs and p.runs[0].text.startswith(label) and p.runs[0].text!=label:
        r=p.runs[0];tail=r.text[len(label):];r.text=label;r.bold=True
        after=p.add_run(tail);r._r.addnext(after._r)
        fonts(after._r.get_or_add_rPr());after.font.size=Pt(10.5)

# True editable three-line tables; no gridlines, no images or shaded cells.
for ti,t in enumerate(doc.tables):
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[[1.9,1.7,1.9,5.7,5.8],[2.5,5.2,3.2,3.2,2.9],[2.5,7.2,3.25,3.25],[9.5,3.25,3.25]][ti]
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    pr=t._tbl.tblPr
    for e in list(pr):
        if e.tag in [qn('w:tblBorders'),qn('w:shd')]:pr.remove(e)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil')
        if edge in ['top','bottom']:b.set(qn('w:sz'),'8');b.set(qn('w:color'),'000000')
        borders.append(b)
    pr.append(borders)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for e in list(trpr):
            if e.tag==qn('w:trHeight'):trpr.remove(e)
        split=OxmlElement('w:cantSplit');trpr.append(split)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            for e in list(cp):
                if e.tag in [qn('w:tcBorders'),qn('w:shd')]:cp.remove(e)
            cb=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single' if ri==0 and edge=='bottom' else 'nil')
                if ri==0 and edge=='bottom':b.set(qn('w:sz'),'4');b.set(qn('w:color'),'000000')
                cb.append(b)
            # Table borders supply the outer rules; cell borders must not suppress them.
            if ri==0:
                cb.find(qn('w:top')).set(qn('w:val'),'single');cb.find(qn('w:top')).set(qn('w:sz'),'8')
            if ri==len(t.rows)-1:
                cb.find(qn('w:bottom')).set(qn('w:val'),'single');cb.find(qn('w:bottom')).set(qn('w:sz'),'8')
            cp.append(cb)
            mar=cp.find(qn('w:tcMar'))
            if mar is None:mar=OxmlElement('w:tcMar');cp.append(mar)
            for side,val in [('top',55),('bottom',55),('left',60),('right',60)]:
                z=mar.find(qn('w:'+side))
                if z is None:z=OxmlElement('w:'+side);mar.append(z)
                z.set(qn('w:w'),str(val));z.set(qn('w:type'),'dxa')
            for p in c.paragraphs:
                txt=grouped(p.text);p.clear();p.style='JournalTable';p.add_run(txt)
                pf=p.paragraph_format;pf.line_spacing=1.1;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.first_line_indent=Pt(0)
                pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:fonts(r._r.get_or_add_rPr());r.font.size=Pt(9);r.bold=ri==0 or '本文方法' in row.cells[1 if ti!=3 else 0].text

# Export exactly the supplied figures in grayscale at their native pixel count.
# This is a data-preserving color conversion, not a regenerated scientific image.
media_changes={};image_audit=[]
with ZipFile(BytesIO(sourcebytes)) as zin:
    for num,idx in enumerate([17,60,72,78],1):
        p=old[idx];rid=p._p.xpath('.//a:blip/@r:embed')[0];part=doc.part.related_parts[rid]
        name=str(part.partname).lstrip('/');original=zin.read(name)
        im=Image.open(BytesIO(original)).convert('RGB');gray=im.convert('L')
        buf=BytesIO();gray.save(buf,format='PNG',dpi=(600,600));media_changes[name]=buf.getvalue()
        gray.save(ART/f'图{num}.tif',compression='tiff_lzw',dpi=(600,600))
        gray.save(ART/f'fig{num}.png',dpi=(600,600))
        inner=[e for e in p._p.xpath('.//wp:inline') if e.xpath('./a:graphic/a:graphicData/pic:pic')]
        assert len(inner)>=1
        picture=deepcopy(inner[0]);p.clear();p.add_run()._r.append(OxmlElement('w:drawing'))
        p.runs[0]._r[-1].append(picture)
        p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
        p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(0)
        ext=picture.find(qn('wp:extent'));w=im.width*1524;h=im.height*1524
        ext.set('cx',str(w));ext.set('cy',str(h))
        for shapeext in picture.xpath('./a:graphic/a:graphicData/pic:pic/pic:spPr/a:xfrm/a:ext'):
            shapeext.set('cx',str(w));shapeext.set('cy',str(h))
        dpi=im.width/(w/914400)
        dp=picture.find(qn('wp:docPr'))
        if dp is not None:dp.set('name',f'图{num}');dp.set('descr',f'图{num} 灰度图')
        image_audit.append({'figure':num,'package_part':name,'pixels':list(im.size),'display_width_cm':w/360000,'effective_dpi':dpi,'mode':'L','tiff_dpi':600,'original_sha256':digest(original),'gray_sha256':digest(buf.getvalue()),'native_word_shapes':False})

# Ask Word to refresh fields on opening; do not invent a cached page count.
settings=doc.settings.element
up=settings.find(qn('w:updateFields'))
if up is None:up=OxmlElement('w:updateFields');settings.append(up)
up.set(qn('w:val'),'true')

buf=BytesIO();doc.save(buf)
# Keep every package part not deliberately edited byte-for-byte identical.
edits={}
with ZipFile(BytesIO(buf.getvalue())) as generated:
    for n in ['word/document.xml','word/styles.xml','word/settings.xml']:edits[n]=generated.read(n)
edits.update(media_changes)
with ZipFile(BytesIO(sourcebytes)) as zin,ZipFile(OUTPUT,'w',ZIP_DEFLATED) as zout:
    for info in zin.infolist():zout.writestr(info,edits.get(info.filename,zin.read(info.filename)))

new=Document(OUTPUT)
newmath=[''.join(p._p.xpath('.//m:t/text()')) for p in new.paragraphs if p._p.xpath('.//m:oMath')]
normalize=lambda s:re.sub(r'\s+','',s)
assert list(map(normalize,oldmath))==list(map(normalize,newmath)),'Equation notation changed'
newtables=[[[c.text for c in r.cells] for r in t.rows] for t in new.tables]
assert [[[normalize(c) for c in r] for r in t] for t in oldtables]==[[[normalize(c) for c in r] for r in t] for t in newtables],'Table data changed'
oldstats=re.findall(r'\d+\.\d+%', original_body)
newbody='\n'.join(p.text for p in new.paragraphs if p.style.name=='JournalBody')
# Every original percentage must still occur in the revised body with identical multiplicity.
from collections import Counter
assert Counter(oldstats)==Counter(re.findall(r'\d+\.\d+%',newbody))
assert len(new.element.xpath('//m:oMath'))==12
assert len(new.inline_shapes)==4
refs=[int(re.match(r'\[(\d+)\]',p.text).group(1)) for p in new.paragraphs if re.match(r'\[\d+\]',p.text)]
assert refs==list(range(1,13))
bodyrefs=[int(v) for p in new.paragraphs if p.style.name=='JournalBody' for v in re.findall(r'\[(\d+)\]',p.text)]
assert list(dict.fromkeys(bodyrefs))==list(range(1,13))
same=[]
with ZipFile(BytesIO(sourcebytes)) as z1,ZipFile(OUTPUT) as z2:
    assert z1.namelist()==z2.namelist()
    for name in z1.namelist():
        if name not in edits:
            assert z1.read(name)==z2.read(name),name
            same.append(name)
audit={'source_sha256':digest(sourcebytes),'backup':str(BACKUP),'output':str(OUTPUT),'output_sha256':digest(OUTPUT.read_bytes()),'title_chars':len(new.paragraphs[0].text),'abstract_nonspace_chars':len(normalize(new.paragraphs[3].text.removeprefix('摘要：'))),'abstract_Chinese_chars':len(re.findall('[\u4e00-\u9fff]',new.paragraphs[3].text)),'english_abstract_words':len(re.findall(r"[A-Za-zÉé]+(?:[-'][A-Za-z]+)*|\d+(?:\.\d+)?%?",new.paragraphs[9].text.removeprefix('Abstract:'))),'equation_groups':12,'equations_semantically_identical':True,'tables':len(new.tables),'all_table_data_identical':True,'figures':4,'image_audit':image_audit,'reference_numbers':refs,'first_citation_order':list(dict.fromkeys(bodyrefs)),'preserved_package_parts':same,'modified_package_parts':list(edits),'changes':ops,'native_render_verified':False,'native_render_blocker':'Bundled runtime has no LibreOffice; Microsoft Word COM activation fails with 0x80070520. HTML review is not native Word pagination.'}
(WORK/'format_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:audit[k] for k in ['output','title_chars','abstract_nonspace_chars','abstract_Chinese_chars','english_abstract_words','equation_groups','tables','figures','all_table_data_identical','equations_semantically_identical']},ensure_ascii=False,indent=2))
