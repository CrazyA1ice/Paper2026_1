from pathlib import Path
from zipfile import ZipFile
import csv, hashlib, importlib.util, json, os, re, statistics, sys
from lxml import etree
from docx import Document

sys.stdout.reconfigure(encoding='utf8')
QA=Path(__file__).resolve().parent
ROOT=QA.parents[1]
manifest=json.loads((QA/'manifest.json').read_text(encoding='utf8'))
target=Path(manifest['output'])
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with ZipFile(target) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    media=[z.read(n) for n in z.namelist() if n.startswith('word/media/')]
    footertext=''.join(''.join(etree.fromstring(z.read(n)).xpath('.//w:t/text()',namespaces=NS)) for n in z.namelist() if re.fullmatch(r'word/footer\d+\.xml',n))
assert '编辑部填写' not in footertext
assert '第一作者简介' not in footertext
assert '作者简介' in footertext
maths=xml.xpath('.//m:oMath',namespaces=NS)
assert [hashlib.sha256(etree.tostring(m,method='c14n')).hexdigest() for m in maths]==manifest['math_c14n_sha256']
assert len(maths)==17
assert len(xml.xpath('.//w:drawing',namespaces=NS))==2
assert len(xml.xpath('.//w:tbl',namespaces=NS))==4
assert hashlib.sha256(Path(manifest['source']).read_bytes()).hexdigest()==manifest['source_sha256']
assert hashlib.sha256(Path(manifest['original']).read_bytes()).hexdigest()==manifest['original_sha256']
doc=Document(target)
assert doc.paragraphs[0].style.name=='Title'
assert len(doc.paragraphs[0].text)<=20
abstract=next(p.text.removeprefix('摘要：') for p in doc.paragraphs if p.text.startswith('摘要：'))
assert len(abstract)<=200
for t,wanted in zip(doc.tables,manifest['tables']):
    rows=[[c.text for c in r.cells] for r in t.rows]
    assert rows==[wanted['headers']]+wanted['rows']

evidence=json.loads((ROOT/'review_scratch/rejection_revision_20261020/verified_evidence.json').read_text(encoding='utf8'))
calc={}
for ds in ['abilene','geant']:
    main=evidence['h24'][ds]
    prop=main['dlinear_scale']['mse_by_seed']
    calc[ds]={}
    for model in ['dlinear','dlinear_scale_static','dlinear_scale_global','dlinear_scale_constant_router']:
        baseline=main[model]['mse_by_seed']
        calc[ds][model]=statistics.mean((b-p)/b*100 for b,p in zip(baseline,prop))
    assert all(main['dlinear_scale'][m]['mean'] < main[k][m]['mean'] for k in ['dlinear','dlinear_scale_static','dlinear_scale_global','dlinear_scale_constant_router'] for m in ['mse','mae'])
    for model in ['dlinear','lightts']:
        if model=='lightts':
            baseline=[json.loads((ROOT/f'abilene_forecasting/results/{ds}/lightts_seed{s}/metrics.json').read_text(encoding='utf8'))['test_normalized']['mse'] for s in [42,43,44]]
        else:
            baseline=main['dlinear']['mse_by_seed'][:3]
        calc[ds]['external_'+model]=statistics.mean((b-p)/b*100 for b,p in zip(baseline,prop[:3]))
expected={'abilene':{'dlinear':2.77,'dlinear_scale_static':1.65,'dlinear_scale_global':1.80,'dlinear_scale_constant_router':1.80,'external_dlinear':4.07,'external_lightts':29.80},
          'geant':{'dlinear':2.35,'dlinear_scale_static':4.54,'dlinear_scale_global':4.31,'dlinear_scale_constant_router':4.97,'external_dlinear':4.54,'external_lightts':46.89}}
for ds,models in expected.items():
    for model,v in models.items(): assert round(calc[ds][model],2)==v,(ds,model,calc[ds][model])

def readcsv(path):
    with open(path,encoding='utf8',newline='') as f:return list(csv.DictReader(f))
seasonal=readcsv(ROOT/'abilene_forecasting/results/supplement_20261020/seasonal_by_seed.csv')
seasonal_calc={}
seasonal_names={'seasonal_24':'日周期朴素','seasonal_168':'周周期朴素','dlinear_scale':'本文方法'}
for ds,label in [('abilene','Abilene'),('geant','GÉANT')]:
    seasonal_calc[ds]={}
    for model,name in seasonal_names.items():
        values=[r for r in seasonal if r['dataset']==ds and r['model']==model]
        assert len(values)==(8 if model=='dlinear_scale' else 1)
        row=next(row for row in manifest['tables'][3]['rows'] if row[:2]==[label,name])
        for metric,col in [('mse',3),('mae',4)]:
            nums=[float(r[metric]) for r in values]
            expected=f'{statistics.mean(nums):.3f}'
            if len(nums)==8:expected+=f'±{statistics.stdev(nums):.3f}'
            assert row[col]==expected,(ds,model,metric,row[col],expected)
    prop=[float(r['mse']) for r in seasonal if r['dataset']==ds and r['model']=='dlinear_scale']
    for model in ['seasonal_24','seasonal_168']:
        base=float(next(r['mse'] for r in seasonal if r['dataset']==ds and r['model']==model))
        seasonal_calc[ds][model]=statistics.mean((base-p)/base*100 for p in prop)
for ds,vals in {'abilene':[44.93,47.80],'geant':[31.77,29.69]}.items():
    assert [round(seasonal_calc[ds][m],2) for m in ['seasonal_24','seasonal_168']]==vals
alltext='\n'.join(p.text for p in doc.paragraphs)
assert not any(s in alltext for s in ['首次提出','显著优于','所有条件','普遍优势','已实现实时部署','−1.17%','12 h预测','泛化能力与研究局限'])
assert all(s in alltext for s in ['完整训练历史','标准化后的空间','共同随机种子','均值±样本标准差'])
assert 2*24*sum(96//s+1 for s in [1,2,4])+131==8339
for ntrain,nval,ntest,windows in [(2793,931,932,[2674,812,813]),(1709,570,570,[1590,451,451])]:
    assert [n-96-24+1 for n in [ntrain,nval,ntest]]==windows
for f in manifest['figures']:
    assert f['effective_dpi']>=300 and f['width_cm']<14
    assert hashlib.sha256(Path(f['source']).read_bytes()).hexdigest()==hashlib.sha256(Path(f['file']).read_bytes()).hexdigest()
    assert Path(f['file']).read_bytes() in media

references=[p.text for p in doc.paragraphs if re.match(r'^\[\d+\]',p.text)]
assert len(references)==11
assert all(references[i].startswith(f'[{i+1}]') for i in range(11))
source=Document(manifest['source'])
assert references==[p.text for p in source.paragraphs if re.match(r'^\[\d+\]',p.text)]
checks={'native_equations_unchanged':17,'figures_unchanged':2,'tables_checked':4,'references_preserved':11,
        'source_files_unchanged':True,'parameter_count_verified':8339,'title_length':len(doc.paragraphs[0].text),
        'abstract_length':len(abstract),'main_comparisons':'mean errors across all stated seeds','computed_improvements':calc,
        'seasonal_improvements':seasonal_calc,
        'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest()}
(QA/'structural_numeric_qa.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(checks,ensure_ascii=False))

if '--render' in sys.argv:
    renderer=Path(r'C:\Users\Quant\.codex\plugins\cache\openai-primary-runtime\documents\26.1015.12040\skills\documents\render_docx.py')
    spec=importlib.util.spec_from_file_location('document_renderer',renderer)
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module)
    native=Path(r'C:\Users\Quant\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler')
    pdftoppm=next(native.rglob('pdftoppm.exe'))
    os.environ['PATH']=str(pdftoppm.parent)+os.pathsep+os.environ['PATH']
    pdf=QA/'render/review.pdf'
    module.convert_to_pdf=lambda *args,**kwargs:(str(pdf),'Native Microsoft Word readonly PDF export; packaged LibreOffice unavailable on Windows')
    pages=module.rasterize(str(target),str(QA/'render'),150,False,False)
    print('Rendered',len(pages),'pages')
