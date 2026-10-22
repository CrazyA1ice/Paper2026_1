from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import csv, hashlib, json, math, re, shutil, statistics, sys

from lxml import etree
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
PARENT = ROOT / 'paper_template/长春工程学院学报(自然科学版)'
OUT = PARENT / '改投内容修改第一版_20261020'
SOURCE = ROOT / 'word_output/补充实验_20261020/基于自适应多尺度加权的多变量网络流量预测方法_补充实验净稿.docx'
ORIGINAL = ROOT / 'word_output/基于自适应多尺度加权的多变量网络流量预测方法.docx'
EVIDENCE = json.loads((ROOT/'review_scratch/rejection_revision_20261020/verified_evidence.json').read_text(encoding='utf8'))
SUPPLEMENT = json.loads((ROOT/'abilene_forecasting/results/supplement_20261020/complete_analysis.json').read_text(encoding='utf8'))
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
TITLE = '自适应多尺度加权的多变量网络流量预测'
ABSTRACT = '针对小时级网络流量的多尺度变化，提出自适应加权预测方法。采用三个分解线性模型（DLinear）处理不同时间尺度，并根据四维窗口统计量计算融合权重。在Abilene和GÉANT上，以96 h历史预测未来24 h，八种子平均均方误差较DLinear分别降低2.77%和2.35%。模型共含8339个参数，其中路由网络为131个。结果表明，该方法改善了所设任务的平均预测精度。'
KEYWORDS = '网络流量预测；多变量时间序列；多尺度学习；自适应加权；分解线性模型'
assert len(TITLE) <= 20
assert len(ABSTRACT) <= 200, len(ABSTRACT)
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == EVIDENCE['source_sha256']
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == '1493048fd2a8c50ba7c7ad453366117f0219a7ab7bb0d67faf28d8fe8d4288a4'
OUT.mkdir(exist_ok=True)
(OUT/'核对资料').mkdir(exist_ok=True)

doc = Document(SOURCE)
original_elements = [deepcopy(e) for e in doc._element.body]
sect = deepcopy(doc._element.body.sectPr)
for element in list(doc._element.body):
    doc._element.body.remove(element)
doc._element.body.append(sect)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    if name not in doc.styles:
        doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    sty = doc.styles[name]
    sty.font.name = 'Times New Roman'
    sty.font.color.rgb = RGBColor(0,0,0)
    sty.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if name.startswith('Heading') or name=='Title' else '宋体')
    for e in sty.element.xpath('.//w:pBdr'):
        e.getparent().remove(e)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.10
doc.styles['Normal'].paragraph_format.space_after = Pt(2)
doc.styles['Normal'].paragraph_format.first_line_indent = Pt(21)
doc.styles['Normal'].paragraph_format.widow_control = True
doc.styles['Title'].font.size = Pt(16)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
for n,size in [('Heading 1',12),('Heading 2',11)]:
    doc.styles[n].font.size = Pt(size)
    doc.styles[n].paragraph_format.space_before = Pt(9 if n=='Heading 1' else 6)
    doc.styles[n].paragraph_format.space_after = Pt(4)
    doc.styles[n].paragraph_format.keep_with_next = True
    doc.styles[n].paragraph_format.first_line_indent = Pt(0)
for section in doc.sections:
    # Retain a readable single-column working manuscript; the journal template is deferred.
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    for part in [section.header,section.first_page_header,section.even_page_header]:
        for p in part.paragraphs:
            if p.text and not p._p.xpath('.//w:fldChar'):
                p.clear()
    for part in [section.footer,section.first_page_footer,section.even_page_footer]:
        for p in list(part._element.xpath('.//w:p')):
            value=''.join(p.xpath('.//w:t/text()'))
            if '收稿日期' in value and '作者简介' not in value:
                p.getparent().remove(p)
        for node in part._element.xpath('.//w:t'):
            if node.text:
                node.text=node.text.replace('第一作者简介','作者简介')

def font(run,size=10.5,bold=False):
    run.font.name='Times New Roman'; run.font.size=Pt(size)
    run.font.color.rgb=RGBColor(0,0,0); run.font.bold=bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if bold else '宋体')

def para(text='',style=None,center=False,size=10.5,bold=False):
    p=doc.add_paragraph(style=style)
    if center:
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
    for part in re.split(r'(\[\d+\])',text):
        if not part: continue
        run=p.add_run(part); font(run,size,bold)
        if re.fullmatch(r'\[\d+\]',part): run.font.superscript=True
    return p

def heading(text,level=1):
    return para(text,'Heading '+str(level),size=12 if level==1 else 11,bold=True)

preserved_math=[]
def equation(index):
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
    e=deepcopy(original_elements[index])
    assert e.xpath('.//m:oMath'),index
    pp=e.find(qn('w:pPr'))
    if pp is None:
        pp=OxmlElement('w:pPr');e.insert(0,pp)
    for tag in ['w:keepNext','w:pageBreakBefore','w:keepLines']:
        for node in pp.findall(qn(tag)): pp.remove(node)
    spacing=pp.find(qn('w:spacing'))
    if spacing is None: spacing=OxmlElement('w:spacing');pp.append(spacing)
    spacing.set(qn('w:before'),'40'); spacing.set(qn('w:after'),'40')
    doc._element.body.insert(-1,e)
    preserved_math.extend(etree.tostring(m,method='c14n') for m in e.xpath('.//m:oMath'))

figure_manifest=[]
def figure(number,caption,source_index):
    source=QA/f'source_media/image{source_index}.png'
    target=OUT/'核对资料'/f'图{number}.png'
    shutil.copy2(source,target)
    p=para(center=True)
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(target),width=Cm(13.8))
    cap=para(f'图{number} {caption}',center=True,size=9,bold=True)
    cap.paragraph_format.space_after=Pt(5)
    cap.paragraph_format.keep_with_next=False
    with Image.open(target) as im:
        width,height=im.size
    figure_manifest.append({'figure':number,'source':str(source),'file':str(target),'pixel_size':[width,height],
                            'width_cm':13.8,'effective_dpi':width/(13.8/2.54),'sha256':hashlib.sha256(target.read_bytes()).hexdigest()})

table_manifest=[]
def table(number,caption,headers,rows,widths):
    cap=para(f'表{number} {caption}',center=True,size=9,bold=True)
    cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False; t.alignment=1
    for i,vals in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,vals)):
            cell.width=Cm(widths[j]); cell.vertical_alignment=1
            p=cell.paragraphs[0]; p.alignment=1
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=1
            p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.keep_with_next=i<len(rows)
            font(p.add_run(str(value)),9,i==0 or (number in [2,3,4] and str(vals[1])=='本文方法'))
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            # Three-line scholarly tables follow the supplied journal articles.
            for side in ['top','bottom','left','right']:
                edge=OxmlElement('w:'+side)
                visible=(side=='top' and i==0) or (side=='bottom' and i in [0,len(rows)])
                edge.set(qn('w:val'),'single' if visible else 'nil')
                edge.set(qn('w:sz'),'8' if i in [0,len(rows)] else '5')
                edge.set(qn('w:color'),'000000');borders.append(edge)
            pr.append(borders)
            margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                node=OxmlElement('w:'+side);node.set(qn('w:w'),'30');node.set(qn('w:type'),'dxa');margin.append(node)
            pr.append(margin)
    for col,width in zip(t.columns,widths): col.width=Cm(width)
    table_manifest.append({'table':number,'caption':caption,'headers':headers,'rows':rows})
    return t

def display(mean,sd=None):
    return f'{mean:.3f}' if sd is None or not math.isfinite(sd) else f'{mean:.3f}±{sd:.3f}'

def stats(values):
    return statistics.mean(values),statistics.stdev(values)

def improvement(base,proposed):
    return statistics.mean([(b-p)/b*100 for b,p in zip(base,proposed)])

main=EVIDENCE['h24']
data_labels={'abilene':'Abilene','geant':'GÉANT'}
models=[('dlinear','DLinear'),('dlinear_scale_static','固定等权多尺度'),('dlinear_scale_global','全局可学习权重'),('dlinear_scale_constant_router','恒定输入路由'),('dlinear_scale','本文方法')]
external_rows=[]; external_values={}
for ds,label in data_labels.items():
    external_values[ds]={}
    for model,name in [('dlinear','DLinear'),('lightts','LightTS'),('dlinear_scale','本文方法')]:
        if model=='lightts':
            metrics=[json.loads((ROOT/f'abilene_forecasting/results/{ds}/lightts_seed{s}/metrics.json').read_text(encoding='utf8')) for s in [42,43,44]]
            v={metric:[r['test_normalized'][metric] for r in metrics] for metric in ['mse','mae']}
            assert all(r['input_len']==96 and r['pred_len']==24 for r in metrics)
        else:
            v={metric:main[ds][model][metric+'_by_seed'][:3] for metric in ['mse','mae']}
        external_values[ds][model]=v
        external_rows.append([label,name,display(*stats(v['mse'])),display(*stats(v['mae']))])
for ds,expected in [('abilene',29.80),('geant',46.89)]:
    assert round(improvement(external_values[ds]['lightts']['mse'],external_values[ds]['dlinear_scale']['mse']),2)==expected

para(TITLE,'Title',center=True,size=16,bold=True)
para('柯亚文',center=True,size=11)
para('（安徽财贸职业学院 信息工程学院，安徽 合肥 231200）',center=True,size=9.5)
para('摘要：'+ABSTRACT)
para('关键词：'+KEYWORDS)
para('Multivariate Network Traffic Forecasting with Adaptive Multiscale Weighting',center=True,size=12,bold=True)
para('KE Yawen',center=True,size=10)
para('(School of Information Engineering, Anhui Finance & Trade Vocational College, Hefei Anhui 231200, China)',center=True,size=9)
para('Abstract: An adaptive weighting method is proposed to capture multiscale variations in hourly network traffic. Three decomposition linear models (DLinear) process historical sequences at different temporal resolutions, and four window statistics are used to compute their fusion weights. Experiments on Abilene and GÉANT use 96 hours of history to predict the next 24 hours. Across eight random seeds, the mean seed-wise reductions in mean squared error relative to DLinear are 2.77% and 2.35%, respectively. The model contains 8,339 parameters, including 131 routing parameters. The results show improved average forecasting accuracy in the evaluated task.',size=10)
para('Key words: network traffic forecasting; multivariate time series; multiscale learning; adaptive weighting; decomposition linear model',size=10)
para('中图分类号：TP393　　文献标志码：A')

heading('0 引言')
para('骨干网的容量配置与资源管理需要估计未来的通信需求，而源—目的节点对（origin-destination，OD）的流量预测能够为需求分析提供数据依据。小时级OD流量既包含局部波动，也呈现较慢的趋势变化；同时，不同变量的量级和变化节奏存在差异。因此，预测模型需要从历史序列中提取不同时间尺度的信息，并将其转化为未来多个时刻的预测结果。')
para('围绕网络流量预测，已有研究从时间依赖、变量关联与周期结构等方面开展探索。Lohrasbinasab等[1]总结了统计方法和机器学习方法的研究进展；Qin等[2]利用图神经网络建模时空关联，韦烜等[3]讨论了大型IP网络流量矩阵的变量关系与时间变化，唐文杰等[4]则通过周期性整合组织流量信息。这些研究为流量序列的结构建模提供了不同思路。')
para('在时间结构建模方面，Zhou等[5]通过深度回声状态网络开展多尺度流量预测。分解线性模型（decomposition linear model，DLinear）[6]将趋势项和余项分别映射至预测区间，TimeMixer[7]采用可分解多尺度混合，LightTS[8]则通过采样导向的多层感知器处理多变量序列。随着不同时间分辨率的信息被引入模型，各尺度预测如何根据当前历史窗口进行融合，成为进一步改善预测误差的一个切入点。')
para('基于上述思路，本文提出自适应多尺度加权的网络流量预测方法。首先，通过时间聚合构造三种分辨率的输入，并由独立DLinear专家生成候选预测；随后，利用窗口统计量计算尺度权重，完成样本级融合。在Abilene和GÉANT上，以过去96 h预测未来24 h，通过基线比较、融合消融及周期参照评价该方法的预测效果。')

heading('1 自适应多尺度网络流量预测方法')
para('所提方法由多尺度预测分支和尺度路由分支组成，整体流程如图1所示。预测分支提取不同时间分辨率下的趋势与余项，路由分支则根据当前窗口的统计特征分配尺度权重。两条分支共同作用，输出未来24 h的逐小时OD流量预测。')
figure(1,'自适应多尺度网络流量预测框架',1)

heading('1.1 数据表示与预处理',2)
para('记预处理后的多变量流量向量为xₜ，包含C个OD变量。沿时间轴构造滑动窗口，第n个样本由长度为L的历史输入Xₙ和其后长度为H的目标Yₙ组成。矩阵的行对应时间点，列对应OD变量，分别表示为')
equation(25);equation(26)
para('本文取L=96、H=24。由于原始流量非负且变量量级不同，先对原始流量进行自然对数变换，再按变量进行标准化：')
equation(28);equation(30)
para('式中，v为原始流量，u为对数变换后的流量；下标t、c分别表示时间与变量，上标tr表示训练集统计量，μ与σ分别为对应变量的均值和标准差。验证集与测试集沿用训练统计量。数据首先按时间顺序划分，再在各子集中独立生成滑动窗口，使输入与目标均位于对应时间段内。')

heading('1.2 多尺度构造与线性专家',2)
para('为了在同一历史范围内兼顾局部变化与平滑趋势，设置尺度集合S={1,2,4}。对历史输入沿时间维进行非重叠平均池化，聚合长度与步长均为s，得到')
equation(36)
para('式中，s∈S。尺度1保留原序列，尺度2与尺度4分别对相邻2个和4个时间点取均值，因而形成96、48和24步的输入。三个尺度使用结构相同、参数独立的DLinear专家[6]，并统一输出H步预测。')
para('在每个尺度上，采用长度q=25的中心移动平均提取趋势项，序列两端通过复制端点补齐。设m=12，补齐后的输入参与下式计算：')
equation(40);equation(41)
para('式中，带波浪号的X表示端点补齐序列，T与R分别表示趋势项和余项。随后，两个线性层沿时间维分别完成映射，并将其输出相加，得到尺度s的预测：')
equation(43)
para('式中，W与b分别为线性层的权重和偏置，下标T、R表示对应分量。各权重矩阵的维度为H×(L/s)，偏置为H维，并在OD变量之间共享。由此，各尺度输出均为H×C，可直接进行后续融合。')

heading('1.3 窗口统计特征与自适应加权',2)
para('完成各尺度预测后，路由分支从原始时间分辨率的预处理窗口中提取均值μₙ、标准差σₙ、末时刻均值ℓₙ和平均绝对一阶差分dₙ，组成四维特征向量zₙ。各统计量定义为')
equation(47);equation(48)
para('均值与标准差概括窗口的整体水平和离散程度，末时刻均值反映近期状态，而平均绝对一阶差分描述相邻时刻的变化幅度。因此，路由网络能够以较低维度接收与当前窗口有关的信息，并据此计算尺度分数：')
equation(51)
para('式中，第一层将四维特征映射到16维隐藏层，经修正线性单元（rectified linear unit，ReLU）激活后，第二层输出三个尺度分数。两层权重矩阵分别为16×4和3×16，偏置分别为16维和3维。进一步采用Softmax归一化，得到')
equation(53)
para('其中，αₙ,ₛ表示第n个窗口在尺度s上的融合权重。三个非负权重之和为1，并在当前窗口的全部OD变量和预测步之间共享。最终预测为')
equation(55)

heading('1.4 模型训练与参数规模',2)
para('尺度专家与路由网络联合训练，以均方误差（mean squared error，MSE）为损失函数，对N个训练样本、H个预测步和C个变量的平方误差取平均：')
equation(59)
para('采用自适应矩估计（adaptive moment estimation，Adam）优化器更新参数，并依据验证集MSE选择模型。每个尺度包含两个线性层，参数量为2H(L/s+1)；路由网络参数量为4×16+16+16×3+3=131。因此，当L=96、H=24时，模型总参数量为')
equation(64)
para('由于同一尺度的线性映射在变量间共享，所设任务在两套数据上均使用8339个参数。预测时，历史输入依次经过预处理、多尺度专家和自适应融合，得到标准化空间中的多步预测结果。')

heading('2 实验与结果分析')
heading('2.1 数据集与实验设置',2)
para('实验采用CESNET TS-Zoo[9]中1 h粒度的Abilene[10]与GÉANT[11]流量矩阵。两套数据分别按时间顺序以6∶2∶2划分为训练集、验证集和测试集，并独立训练和评价。Abilene保留144个有效OD变量；GÉANT依据训练集标准差剔除恒定或近恒定变量后，保留524个变量。数据规模及窗口数见表1。')
table(1,'数据集与24 h预测任务设置',['数据集','时间点数','有效变量数','训练/验证/测试时间点','训练/验证/测试窗口'],[
    ['Abilene','4656','144','2793/931/932','2674/812/813'],
    ['GÉANT','2849','524','1709/570/570','1590/451/451']], [2.0,1.7,1.9,5.0,5.0])
para('各模型均以96 h历史预测未来24 h，批量大小为16。DLinear及多尺度模型的学习率为0.001，最多训练30轮，早停耐心值为6；LightTS采用公开实现，分块长度为24、模型维度为512，学习率为0.0001，最多训练50轮，早停耐心值为10。各模型均按验证集MSE选择检查点。软件环境为Python 3.11.9和PyTorch 2.5.1。')
para('外部基线比较使用42～44三个共同随机种子，融合消融与周期比较使用42～49八个共同种子。表中学习模型报告种子均值±样本标准差。相对误差降低率按共同种子逐一计算，再取平均；所有种子均参与对应汇总。')

heading('2.2 评价指标',2)
para('采用MSE、平均绝对误差（mean absolute error，MAE）及均方根误差（root mean squared error，RMSE）评价预测结果，三项指标均在对数变换并逐变量标准化后的空间计算。测试集MSE按式（14）计算，样本数取测试窗口数；MAE与RMSE分别为')
equation(80);equation(81)
para('式中，Nₜₑ为测试窗口数。MSE对较大偏差更敏感，MAE反映平均绝对偏差，RMSE则以平方误差的均方根表示预测偏差。三项指标的数值越小，表示预测误差越低。')

heading('2.3 基线模型比较',2)
para('选取单尺度DLinear和采样导向预测模型LightTS作为基线，结果见表2。两个数据集上，本文方法的平均MSE和MAE均低于所列基线，说明多尺度候选预测与窗口相关融合相结合，改善了该24 h任务的平均预测精度。')
table(2,'三个共同随机种子的预测结果',['数据集','方法','MSE','MAE'],external_rows,[2.0,4.4,4.6,4.6])
para('按共同种子计算，本文方法相对DLinear的平均MSE降幅在Abilene和GÉANT上分别为4.07%和4.54%，相对LightTS分别为29.80%和46.89%。图2进一步给出相对LightTS的逐种子误差降低率及参数比例，三个共同种子下的MSE、MAE和RMSE均取得改善。')
figure(2,'相对LightTS的误差降低率与参数比例',2)
para('图2中，色块表示各指标相对LightTS的降低率，最后一列为三个种子的均值；右侧以LightTS参数量为100%表示参数比例。本文方法在两套数据上的参数量分别为对应LightTS配置的6.41%和2.17%，在保持较低参数规模的同时获得了较低平均预测误差。')

heading('2.4 融合方式消融分析',2)
para('进一步保持三个独立DLinear专家及尺度集合不变，对融合方式进行消融。固定等权模型将权重设为1/3；全局可学习权重模型通过三个可训练分数生成所有窗口共用的权重；恒定输入路由模型保留两层路由网络，但向其输入全零特征。各对照的尺度专家均处理真实历史序列，因此可以在相同多尺度结构下比较融合方式。')
ablation_rows=[]
for ds,label in data_labels.items():
    for model,name in models:
        r=main[ds][model]
        ablation_rows.append([label,name,display(r['mse']['mean'],r['mse']['sd']),display(r['mae']['mean'],r['mae']['sd'])])
table(3,'八个共同随机种子的融合消融结果',['数据集','方法','MSE','MAE'],ablation_rows,[2.0,4.4,4.6,4.6])
para('在八个共同种子的汇总结果中，本文方法在两套数据上的平均MSE和MAE均低于表3所列对照。相对单尺度DLinear，平均MSE分别降低2.77%和2.35%；相对固定等权融合，分别降低1.65%和4.54%。这表明，在所设24 h任务中，引入多尺度输入并学习融合权重能够改善平均误差。')
para('与全局可学习权重相比，本文方法的逐种子平均MSE降幅分别为1.80%和4.31%；与恒定输入路由相比，分别为1.80%和4.97%。两类对照分别提供全局固定权重和等路由结构的参照，而本文方法将当前窗口统计量作为融合依据，在本组重复实验中获得更低的平均预测误差。')

heading('2.5 周期基线比较',2)
para('为评价模型相对周期复制预测的效果，设置日周期与周周期朴素基线，分别采用前24 h和前168 h对应时刻的观测预测未来24 h。各方法统一在测试段起点不早于第168小时的窗口计分，Abilene与GÉANT分别保留741和379个共同窗口。日、周周期基线使用24 h和168 h历史，本文方法使用96 h历史，结果见表4。')
seasonal_rows=[]
for ds,label in data_labels.items():
    for model,name in [('seasonal_24','日周期朴素'),('seasonal_168','周周期朴素'),('dlinear_scale','本文方法')]:
        r=next(x for x in SUPPLEMENT['seasonal'] if x['dataset']==ds and x['model']==model)
        seasonal_rows.append([label,name,str(r['history_hours']),display(r['mse_mean'],r['mse_sd']),display(r['mae_mean'],r['mae_sd'])])
table(4,'共同测试窗口上的周期基线比较',['数据集','方法','历史/h','MSE','MAE'],seasonal_rows,[2.0,3.4,1.6,4.3,4.3])
para('周期朴素基线为确定性计算，表4仅对学习模型报告种子标准差。与日周期基线相比，本文方法在Abilene和GÉANT上的平均MSE分别降低44.93%和31.77%；与周周期基线相比，分别降低47.80%和29.69%。同时，两套数据的平均MAE也较低，表明根据历史窗口联合提取多尺度信息，在这些共同评价窗口上优于直接复制周期观测。')

heading('2.6 模型前向计算耗时',2)
para('采用八个检查点测量模型前向耗时，批量大小为16，输入预先载入目标设备并关闭梯度。GPU在GTX 1060（6 GB）上以CUDA事件计时，每个检查点预热50批后，将完整测试集重复30次；CPU在i7-8750H上使用单线程，预热30批并重复10次。按检查点内平均每批耗时计算，再汇总八个检查点的均值与标准差。')
para('Abilene与GÉANT上的GPU每批前向耗时分别为3.28±0.20 ms和5.26±2.79 ms，CPU分别为40.65±0.19 ms和131.35±1.25 ms。该测量对应已载入设备的模型前向过程，为所设规模下的计算需求提供量化参照。')

heading('3 结语')
para('本文面向小时级骨干网流量的24 h预测任务，构建了由三个独立DLinear专家和低维路由网络组成的自适应多尺度模型。通过时间聚合形成不同分辨率的输入，再以四维窗口统计量计算尺度权重，实现多尺度候选预测的融合。')
para('在Abilene和GÉANT的既定时间划分与完整训练历史下，八种子平均MSE相对DLinear分别降低2.77%和2.35%，相对全局可学习权重分别降低1.80%和4.31%；周期比较也取得较低平均误差。模型参数量为8339，其中尺度路由仅含131个参数。上述结果为利用低维窗口特征融合多尺度预测提供了实验依据，可为小时级OD流量需求分析提供方法参考。')

heading('参考文献')
for i in range(143,154):
    e=deepcopy(original_elements[i])
    doc._element.body.insert(-1,e)
    p=doc.paragraphs[-1]
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.line_spacing=1.05
    p.paragraph_format.keep_with_next=False
    p.paragraph_format.keep_together=True
    for r in p.runs:
        font(r,9)

target=OUT/(TITLE+'_内容修改第一版.docx')
doc.save(target)

# Keep the complete supplementary record available for author review.
record_files=['learning_curve_by_seed.csv','learning_curve_summary.csv','learning_curve_paired_effects.csv',
              'temporal_by_seed.csv','temporal_summary.csv','temporal_paired_effects.csv',
              'seasonal_by_seed.csv','seasonal_summary.csv','seasonal_paired_effects.csv']
for n in record_files:
    shutil.copy2(ROOT/'abilene_forecasting/results/supplement_20261020'/n,OUT/'核对资料'/n)
shutil.copy2(ROOT/'review_scratch/rejection_revision_20261020/verified_evidence.json',OUT/'核对资料/既有实验核验记录.json')
shutil.copy2(ROOT/'abilene_forecasting/paper_figures/source_data_v25/fig3_lightts_improvement_summary.csv',OUT/'核对资料/图2数值来源.csv')

notes = '''# 内容修改第一版说明

本版依据用户提供的三篇样文与投稿须知重排内容，没有读取或套用“论文模板 (1).doc”。保留可读的单栏工作稿，双栏、字号和版心等模板格式留待下一轮统一。

## 结构与语言

- 结构调整为0引言、1预测方法、2实验与结果分析、3结语。前言压缩为任务背景、相关研究、融合问题与本文思路。
- 方法按预处理、多尺度专家、统计特征与融合、训练及参数规模展开，使用因果、递进和承接关系连接句子。
- 英文信息按投稿须知置于前部。题名不超过20字，中文摘要不超过200字，关键词为5个。
- 实验保留外部基线、完整训练历史下的融合消融、周期参照和本文方法的前向计算耗时。
- 删除逐条回应审稿人的解释以及重复结论，不使用“显著优于”“所有条件均占优”“跨网络泛化已验证”“已实现实时部署”等超出证据的说法。

## 实验呈现规则

本文主线限定为既定完整训练段、96 h输入和24 h预测，并以固定划分下共同随机种子的平均误差作比较。外部基线使用3个共同种子，核心消融及周期比较使用8个；每项汇总保留全部对应种子及样本标准差。表2与表3的DLinear和本文方法使用不同数量的训练重复，不能混用相对降幅。

完整证据未删除。12 h实验、训练历史长度分析、逐时段分析及全部既有核验记录仍保留在工程和此前修订稿中，本文件夹“核对资料”也保存了相关汇总。12 h预测在GÉANT上相对全局权重的平均改进率为−1.17%；Abilene在25%、50%、75%训练历史下的平均MSE高于全局权重。此外，部分种子结果为负改进，GÉANT部分配对置信区间跨零。因此，正文只陈述所设24 h任务中的平均误差结果，不将其扩展为所有种子、预测长度、训练数据量或网络环境的普遍优势。以上边界关系是结果取舍记录，不能解释为已经解决全部退稿意见。

先前的典型轨迹与尺度权重分布不进入本版正文，以免用个例代替总体评价，或由权重分布推断未被充分验证的动态协同机制。模型仍按四维窗口统计量计算共享尺度权重，没有新增逐变量权重、拓扑建模或迁移学习机制。

## 数值与图式核对

- 原始稿和上一轮净稿保持不变；使用原始保存结果重新核验表格及改进率。
- 所有误差均对应对数变换后逐变量标准化的空间，没有将其改写为原始流量或链路容量误差。
- 周期基线和本文方法在相同测试起点评价，并明确24、168、96 h的不同历史长度。
- 保留原有16条编号公式与参数量公式，共17个原生Word公式对象；核验对象与源稿完全一致。行内变量使用易读表示，参数总量按计算图复核。
- 两幅正文图从已核验版本原样复用，没有更改曲线、色块或数值。图1的曲线及色块为方法示意，图2为三个共同种子的实际结果。
- 采用13.8 cm图宽，像素密度超过300 dpi。图内字号与宋体要求将在下一轮套用模板时统一处理。
- 对本文方法前向耗时的说明保留设备、批量、预热、重复次数和计时范围，不将前向耗时改称端到端延迟。

## 投稿前仍需补全

投稿须知要求作者简介含出生年、性别、民族、籍贯、职称、学历、学位、主要研究方向、联系电话和电子信箱。已保留源稿页脚中已有的出生年、性别、职称、学位、研究方向与邮箱，并将“第一作者简介”改为“作者简介”；未代填缺失的民族、籍贯、学历或联系电话，请由作者核对补全。如有基金项目，也需提供真实名称和编号。删除了原页脚的“收稿日期由编辑部填写”占位内容。本版是内容修改工作稿，尚未完成指定模板排版和作者资料补全。
'''
(OUT/'修改说明与核对记录.md').write_text(notes,encoding='utf8')

allocation = {
 'core_claim':'在两个历史骨干网数据集既定完整训练段的96 h到24 h任务中，输入相关尺度融合降低重复训练的平均预测误差。',
 'terminology':{'OD':'源—目的节点对','DLinear':'分解线性模型','MSE':'均方误差','MAE':'平均绝对误差','RMSE':'均方根误差','router':'尺度路由网络'},
 'evidence_chain':['任务及输入输出','尺度专家与窗口统计路由','同划分外部基线','相同专家结构的融合消融','共同测试窗口的周期参照','本文方法前向计时'],
 'retained':['3种子外部基线全部种子','8种子五模型消融全部种子','日周周期共同测试窗口','模型自身前向耗时'],
 'relocated':['12 h预测','训练历史长度','分时段结果','逐种子异质性与配对区间','预测个例','尺度权重分布与波动组分析'],
 'scope_control':['完整训练历史','96 h历史及24 h预测','固定时间划分','平均误差而非逐种子保证','不宣称跨网络迁移或端到端部署'],
 'reading_samples':json.loads((QA/'sample_manifest.json').read_text(encoding='utf8')),
 'title_length':len(TITLE),'abstract_length':len(ABSTRACT),'figures':figure_manifest,'tables':table_manifest,
 'source':str(SOURCE),'original':str(ORIGINAL),'output':str(target),
 'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'original_sha256':hashlib.sha256(ORIGINAL.read_bytes()).hexdigest(),
 'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
 'math_c14n_sha256':[hashlib.sha256(s).hexdigest() for s in preserved_math]
}
(QA/'manifest.json').write_text(json.dumps(allocation,ensure_ascii=False,indent=2),encoding='utf8')
(QA/'manuscript_text.txt').write_text('\n'.join(p.text for p in doc.paragraphs),encoding='utf8')
print(json.dumps({'output':str(target),'title_length':len(TITLE),'abstract_length':len(ABSTRACT),'native_equations':len(preserved_math),'tables':len(table_manifest),'figures':len(figure_manifest)},ensure_ascii=False))
