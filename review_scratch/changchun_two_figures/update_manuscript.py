from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib,json,re,sys
from lxml import etree
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).resolve().parents[2]
QA=Path(__file__).resolve().parent
OUT=ROOT/'paper_template/长春工程学院学报(自然科学版)/改投内容修改第一版_20261020'
BASE=OUT/'自适应多尺度加权的多变量网络流量预测_内容修改第一版.docx'
TARGET=OUT/'自适应多尺度加权的多变量网络流量预测_补充两图版.docx'
assert hashlib.sha256(BASE.read_bytes()).hexdigest()=='e17bedc4a3ad444ae387ec607a4a922b0a21adb81c1d7f1fdea9de26b0105d14'
audit=json.loads((QA/'source_audit.json').read_text(encoding='utf8'))
doc=Document(BASE)

def font(run,size=10.5,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if bold else '宋体')

def find(prefix):
    matches=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches)==1,(prefix,len(matches))
    return matches[0]

def insert(anchor,text,style=None,center=False,size=10.5,bold=False):
    p=anchor.insert_paragraph_before(text,style=style)
    if center:
        p.alignment=1;p.paragraph_format.first_line_indent=Pt(0)
    for run in p.runs:font(run,size,bold)
    return p

def figure(anchor,number,caption,path):
    p=insert(anchor,'',center=True)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(path),width=Cm(13.8))
    caption=insert(anchor,f'图{number} {caption}',center=True,size=9,bold=True)
    caption.paragraph_format.keep_with_next=False
    caption.paragraph_format.space_after=Pt(4)

anchor=find('2.5 周期基线比较')
insert(anchor,'为观察多尺度融合中的权重分配，图3汇总八个种子在全部测试窗口上的尺度权重。小提琴表示窗口权重的汇总分布，彩色点表示各种子的窗口平均权重；黑色菱形和误差线分别表示这些均值的平均值与样本标准差，短横线为汇总分布中位数，虚线为等权1/3。')
figure(anchor,3,'尺度权重分布与种子均值',Path(audit['figure3_file']))
insert(anchor,'Abilene在尺度1、2、4上的平均权重分别为0.257、0.319、0.424，GÉANT分别为0.283、0.318、0.399。两套数据均为尺度4分配了较高的平均权重，同时保留尺度1与尺度2的预测贡献，形成不同时间分辨率信息的组合。')

anchor=find('2.6 模型前向计算耗时')
newheading=insert(anchor,'2.6 高需求预测点分析',style='Heading 2',size=11,bold=True)
newheading.paragraph_format.keep_with_next=True
insert(anchor,'进一步以训练集各OD变量的第90百分位为阈值，将真实值超过该阈值的测试预测点定义为高需求点。分别计算这些点的MAE，以及真实值减预测值取正值部分后在全部高需求点上的平均值，即平均低估误差。各指标先在单个种子内计算，再汇总八个共同种子，结果见图4。')
figure(anchor,4,'高需求预测点的误差比较',Path(audit['figure4_file']))
insert(anchor,'图4中，圆点和方点分别表示MAE与平均低估误差的种子均值，横向误差线为种子间样本标准差。两项指标均沿用前述标准化坐标。本文方法在Abilene上的两项均值分别为0.147、0.137，在GÉANT上分别为0.632、0.577，均低于图中所列对照。相对全局可学习权重，Abilene上MAE与平均低估误差的逐种子平均降幅分别为17.22%和20.49%，GÉANT上分别为3.59%和6.30%。上述结果补充反映了该24 h任务中高需求预测点的平均误差表现。')
anchor.clear();font(anchor.add_run('2.7 模型前向计算耗时'),11,True)
find('外部基线比较使用').text='外部基线比较使用42～44三个共同随机种子，融合消融、周期比较、权重及高需求分析使用42～49八个共同种子。表中学习模型报告种子均值±样本标准差。相对误差降低率按共同种子逐一计算，再取平均；所有种子均参与对应汇总。'
for run in find('外部基线比较使用').runs:font(run)
doc.save(TARGET)

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
def xml(path):
    with ZipFile(path) as z:return etree.fromstring(z.read('word/document.xml'))
base_xml,new_xml=xml(BASE),xml(TARGET)
assert len(new_xml.xpath('.//w:drawing',namespaces=NS))==4
assert len(new_xml.xpath('.//w:tbl',namespaces=NS))==4
assert [etree.tostring(e,method='c14n') for e in base_xml.xpath('.//m:oMath',namespaces=NS)]==[etree.tostring(e,method='c14n') for e in new_xml.xpath('.//m:oMath',namespaces=NS)]
assert len(new_xml.xpath('.//m:oMath',namespaces=NS))==17
assert [[c.text for row in t.rows for c in row.cells] for t in Document(BASE).tables]==[[c.text for row in t.rows for c in row.cells] for t in Document(TARGET).tables]
before_refs=[p.text for p in Document(BASE).paragraphs if re.match(r'^\[\d+\]',p.text)]
after_refs=[p.text for p in Document(TARGET).paragraphs if re.match(r'^\[\d+\]',p.text)]
assert before_refs==after_refs and len(after_refs)==11
with ZipFile(BASE) as old,ZipFile(TARGET) as new:
    for name in old.namelist():
        if name.startswith('word/media/'):
            assert old.read(name)==new.read(name)
assert hashlib.sha256(BASE.read_bytes()).hexdigest()=='e17bedc4a3ad444ae387ec607a4a922b0a21adb81c1d7f1fdea9de26b0105d14'
manifest={'base':str(BASE),'output':str(TARGET),'base_sha256':hashlib.sha256(BASE.read_bytes()).hexdigest(),
          'output_sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'figures':4,'tables':4,'equations':17,'references':11,
          'figure3':audit['figure3_file'],'figure4':audit['figure4_file'],'source_audit':str(QA/'source_audit.json')}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
notes='''# 补充两图说明

本轮在内容修改第一版的基础上增加两张图，另存为补充两图版，原版不变。原有图1、图2、四张表、17个可编辑公式及11条参考文献保持不变。期刊论文模板仍未读取或套用。

图3复用上一轮已核验文稿中的尺度权重分布图，图像未经裁剪或改数值。核对了两个数据集、八个种子和全部测试窗口的路由权重；图中文字说明区分窗口分布、单种子窗口均值及种子间标准差。该图用于描述尺度分配，不将权重差异直接解释为预测改善的因果机制。

图4使用已完成的高需求分析结果重绘，未新增训练。采用训练集第90百分位确定各OD变量的高需求阈值，保留五种模型和八个共同种子。每项指标先在种子内计算，再对种子汇总，误差线为样本标准差。正文陈述平均误差结果，未将其表述为每个种子均占优或统计显著。

原材料中的逐种子改进分布图和典型轨迹图也有分析价值，但前者展示训练重复差异，后者展示个例行为，与本轮新增两图的作用不同。没有通过删去负改进种子、删去较差轨迹或调整坐标数值来改造这些旧图。原始图和完整实验记录继续保留。

新增图表及相应CSV保存在“补充图表”子文件夹。图4同时保留PDF、SVG、600 dpi PNG及TIFF。正文已补齐图号、引用、图注和相应结果段落，并新增高需求分析小节。作者简介缺失项目及后续模板排版要求沿用上一版修改说明。
'''
(OUT/'补充两图说明.md').write_text(notes,encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False))
