from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib, json, re, shutil, sys
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).resolve().parents[2]
QA=Path(__file__).resolve().parent
PARENT=ROOT/'paper_template/长春工程学院学报(自然科学版)'
BASE=PARENT/'改投内容修改第一版_20261020/自适应多尺度加权的多变量网络流量预测_柱状图修订版.docx'
OUT=PARENT/'长预测核验与修订_20261021'
ASSETS=OUT/'数据与图表'
MAIN=OUT/'自适应多尺度加权的多变量网络流量预测_扩展预测修订版.docx'
SUPPLEMENT=OUT/'长预测时长扩展实验_完整结果与核验.docx'
BASE_HASH='eb2f06a71158bb9162efb38a8bad3943d4ad357ebe7e163c3aa8abfd5fe0ef1f'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(BASE)==BASE_HASH
audit=json.loads((QA/'data_verification.json').read_text(encoding='utf8'))
assert audit['records']==320 and audit['prediction_files_hash_verified']==320
assert json.loads((QA/'figure5_collision.json').read_text(encoding='utf8'))['verdict']=='PASS'
MODELS=['dlinear','dlinear_scale_static','dlinear_scale_global','dlinear_scale_constant_router','dlinear_scale']
LABELS=dict(zip(MODELS,['DLinear','固定等权','全局可学习权重','恒定输入路由','本文方法']))
summary=pd.read_csv(ASSETS/'summary.csv')
paired=pd.read_csv(ASSETS/'paired_reductions.csv')
protocol=json.loads((ASSETS/'protocol.json').read_text(encoding='utf8'))

def text(p,value,size=10.5,bold=False):
    p.clear();r=p.add_run(value)
    r.font.name='Times New Roman';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if bold else '宋体')
    return p

shutil.copy2(BASE,MAIN)
doc=Document(MAIN)
def find(prefix):
    items=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(items)==1,(prefix,len(items))
    return items[0]

abstract='针对小时级网络流量的多尺度变化，提出自适应加权预测方法。采用三个分解线性模型处理多尺度输入，并由四维窗口统计量计算融合权重。以96 h历史预测未来24 h时，八种子平均均方误差较DLinear在Abilene和GÉANT上分别降低2.77%和2.35%；在GÉANT的168 h任务中，均方误差和平均绝对误差分别降低23.16%和11.16%。24 h模型含8339个参数，其中路由参数为131个。'
assert len(abstract)<=200,len(abstract)
text(find('摘要：'),'摘要：'+abstract)
text(find('Abstract:'),'Abstract: An adaptive weighting method is proposed to capture multiscale variations in hourly network traffic. Three decomposition linear models process inputs at different temporal resolutions, and four window statistics determine their fusion weights. With 96 hours of history and a 24-hour forecast horizon, the mean seed-wise reductions in mean squared error relative to DLinear are 2.77% on Abilene and 2.35% on GÉANT across eight random seeds. For the 168-hour GÉANT task, the corresponding reductions in mean squared error and mean absolute error are 23.16% and 11.16%. The 24-hour model has 8,339 parameters, including 131 routing parameters.')
text(find('基于上述思路'),'基于上述思路，本文提出自适应多尺度加权的网络流量预测方法。首先，通过时间聚合构造三种分辨率的输入，并由独立DLinear专家生成候选预测；随后，利用窗口统计量计算尺度权重，完成样本级融合。在Abilene和GÉANT上，以过去96 h预测未来24 h，通过基线比较、融合消融及周期参照评价预测效果，并在GÉANT的一周预测任务中进一步考察多尺度融合的作用。')
text(find('所提方法由'),'所提方法由多尺度预测分支和尺度路由分支组成，图1以24 h预测任务为例展示整体流程。预测分支提取不同时间分辨率下的趋势与余项，路由分支则根据当前窗口的统计特征分配尺度权重。两条分支共同作用，输出未来H个小时的逐小时OD流量预测。')
text(find('本文取L=96'),'基础实验取L=96、H=24，扩展实验保持L=96并调整H。由于原始流量非负且变量量级不同，先对原始流量进行自然对数变换，再按变量进行标准化：')
text(find('由于同一尺度'),'由于同一尺度的线性映射在变量间共享，24 h任务在两套数据上均使用8339个参数。预测时，历史输入依次经过预处理、多尺度专家和自适应融合，得到标准化空间中的多步预测结果。')
text(find('各模型均以96 h历史'),'基础实验中，各模型均以96 h历史预测未来24 h，批量大小为16。DLinear及多尺度模型的学习率为0.001，最多训练30轮，早停耐心值为6；LightTS采用公开实现，分块长度为24、模型维度为512，学习率为0.0001，最多训练50轮，早停耐心值为10。各模型均按验证集MSE选择检查点。软件环境为Python 3.11.9和PyTorch 2.5.1。')
text(find('外部基线比较'),'外部基线比较使用42～44三个共同随机种子，融合消融、周期比较、权重、高需求及扩展预测分析使用42～49八个共同种子。学习模型报告种子均值±样本标准差。相对误差降低率按共同种子逐一计算，再取平均；所有种子均参与对应汇总。')
text(find('为评价模型相对周期'),'为评价模型相对周期复制预测的效果，设置日周期复制与周周期复制基线，分别用前24 h和前168 h相应时刻的已观测值预测未来24 h。各方法统一在测试段起点不早于第168小时的窗口计分，Abilene与GÉANT分别保留741和379个共同窗口。日、周周期基线使用24 h和168 h历史，本文方法使用96 h历史，结果见表4。')
p=find('周期朴素基线');text(p,p.text.replace('周期朴素基线','周期复制基线'))
for row in doc.tables[3].rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            if '周期朴素' in p.text:text(p,p.text.replace('周期朴素','周期复制'),9)

timing=find('2.7 模型前向计算耗时')
new_heading=timing.insert_paragraph_before(style=timing.style)
text(new_heading,'2.7 GÉANT一周预测对照',11,True)
new_setup=timing.insert_paragraph_before()
text(new_setup,'进一步保持小时采样、96 h输入、数据划分和训练设置不变，将五种模型的预测时长扩展为24、48、96和168 h，完整结果见补充材料。图5给出GÉANT的168 h任务，五种模型均在相同的307个测试窗口上评价，并覆盖全部预测小时与524个OD变量。该任务中，本文模型含57587个参数，其中路由网络仍为131个。')
picture=timing.insert_paragraph_before()
picture.paragraph_format.first_line_indent=Pt(0);picture.alignment=WD_ALIGN_PARAGRAPH.CENTER
picture.paragraph_format.keep_with_next=True
shape=picture.add_run().add_picture(str(ASSETS/'图5_GEANT一周预测误差.png'),width=Cm(13.8))
shape._inline.docPr.set('descr','GÉANT一周预测中五种模型的MSE与MAE；柱高和误差线分别为八个训练种子的均值与样本标准差。')
caption=timing.insert_paragraph_before()
caption._p.get_or_add_pPr().append(deepcopy(find('图4 高需求')._p.pPr.find(qn('w:spacing'))))
caption.paragraph_format.first_line_indent=Pt(0);caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
text(caption,'图5 GÉANT一周预测任务的MSE与MAE',9,True)
text(timing.insert_paragraph_before(),'图5中，柱高与误差线分别表示八个种子的均值与样本标准差，“全局权重”和“恒定路由”沿用图4的简称。本文方法的MSE为0.841±0.088，MAE为0.536±0.028，两项均值均低于所列对照。相对DLinear，逐种子平均降幅分别为23.16%和11.16%；相对全局可学习权重，分别为26.02%和12.34%。八个种子在这两项指标上均优于四种对照，说明该方法在GÉANT一周预测任务中取得了较低的预测误差。')
text(timing,'2.8 24 h任务的模型前向计算耗时',11,True)
p=find('采用八个检查点测量');text(p,p.text.replace('采用八个检查点','采用24 h任务的八个检查点',1))
text(find('本文面向小时级骨干网'),'本文面向小时级骨干网流量预测，构建了由三个独立DLinear专家和低维路由网络组成的自适应多尺度模型。通过时间聚合形成不同分辨率的输入，再以四维窗口统计量计算尺度权重，实现多尺度候选预测的融合。')
text(find('在Abilene和GÉANT的既定'),'在Abilene和GÉANT的24 h任务中，八种子平均MSE相对DLinear分别降低2.77%和2.35%，相对全局可学习权重分别降低1.80%和4.31%。在GÉANT的168 h任务中，MSE与MAE相对DLinear的逐种子平均降幅分别为23.16%和11.16%。24 h模型含8339个参数，168 h模型含57587个参数，两者的尺度路由均为131个参数。上述结果为利用低维窗口特征融合多尺度预测提供了实验依据，可为小时级OD流量需求分析提供方法参考。')
doc.save(MAIN)

# Copy the retained reference and reuse its typography and page geometry.
shutil.copy2(BASE,SUPPLEMENT)
si=Document(SUPPLEMENT)
body=si._element.body
for element in list(body):
    if element.tag!=qn('w:sectPr'):body.remove(element)
for section in si.sections:
    for hdrftr in [section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer]:
        for p in hdrftr.paragraphs:p.clear()

def para(value='',style=None,bold=False,size=10.5):
    return text(si.add_paragraph(style=style),value,size,bold)

def heading(value,level=1,newpage=False):
    p=para(value,'Heading '+str(level),True,12 if level==1 else 11)
    p.paragraph_format.page_break_before=newpage
    return p

def table(caption,headers,rows,widths):
    cp=para(caption,bold=True,size=9)
    cp.paragraph_format.first_line_indent=Pt(0);cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next=True
    t=si.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for c,value in zip(t.rows[0].cells,headers):c.text=value
    for values in rows:
        for c,value in zip(t.add_row().cells,values):c.text=str(value)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'55' if edge in ['top','bottom'] else '75');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if ri==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E6E9EC');tcpr.append(fill)
            for p in c.paragraphs:
                value=p.text;text(p,value,9.5,ri==0)
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.05
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==1 and len(headers)==4 else WD_ALIGN_PARAGRAPH.CENTER
    return t

para('长预测时长扩展实验完整结果','Title',True,16)
para('自适应多尺度加权的多变量网络流量预测  补充材料',size=10.5).paragraph_format.first_line_indent=Pt(0)
para('本材料汇总Abilene与GÉANT在24、48、96和168 h预测时长下的完整对照。共评价320组模型与种子配置，其中240组为本轮新训练，80组复用已有24 h训练记录并重新计分。正文展示GÉANT的168 h任务，其他条件及全部配对比较列于本材料。')
heading('S1 实验设置与评价口径')
para('保留1 h采样粒度、96 h历史输入及按时间顺序6∶2∶2的数据划分。自然对数变换及逐变量标准化沿用训练集统计量。五种模型为DLinear、固定等权、全局可学习权重、恒定输入路由和本文方法，定义与正文一致。每个模型与预测时长均独立训练，使用42～49八个共同随机种子；批量大小16，学习率0.001，最多30轮，早停耐心值6，按验证集MSE选择检查点。')
para('为比较不同预测时长，取能够容纳168 h目标的共同预测起点，Abilene为669个，GÉANT为307个。每个起点均统计对应时长的全部目标小时和保留变量，先计算单个种子的MSE与MAE，再汇总八种子的均值和样本标准差。训练种子是重复训练的汇总单位，滑动窗口中的重叠目标不作为独立实验重复。')
window_rows=[]
for ds,name in [('abilene','Abilene'),('geant','GÉANT')]:
    for h in [24,48,96,168]:
        counts=protocol['data_readiness'][ds]['windows'][str(h)]
        window_rows.append([name,h,counts['train'],counts['val'],counts['test'],protocol['data_readiness'][ds]['common_test_origins']])
table('表S1 各预测任务的样本窗口数',['数据集','时长/h','训练','验证','全部测试','共同起点'],window_rows,[2.5,2,2.4,2.4,2.4,2.7])
para('各时长全部有效测试起点的评分同时保留在配套数据的all字段中；共同起点评分记为common。正文表3的24 h基础实验使用全部有效起点，因此其数值与本材料的24 h共同起点评分对应不同评价集合。本文模型参数量为342H+131，其中H为预测小时数；H取24、48、96、168时，参数量分别为8339、16547、32963、57587，路由网络均为131个参数。')

for ds,name,number in [('abilene','Abilene',2),('geant','GÉANT',3)]:
    heading(f'S{number} {name}的完整预测结果',newpage=True)
    para(f'表S{number}列出四种预测时长下的五种模型结果，每项为八个种子的均值±样本标准差。所有指标均在对数变换与逐变量标准化后的空间计算；数值越小，预测误差越低。')
    rows=[]
    for h in [24,48,96,168]:
        group=summary[(summary.dataset==ds)&(summary.horizon==h)].set_index('model')
        for model in MODELS:
            row=group.loc[model]
            rows.append([h,LABELS[model],f'{row.common_mse_mean:.6f}±{row.common_mse_sd:.6f}',f'{row.common_mae_mean:.6f}±{row.common_mae_sd:.6f}'])
    table(f'表S{number} {name}的共同起点预测误差',['预测时长/h','模型','MSE','MAE'],rows,[2.1,4.0,4.2,4.1])
    para('各时长均保留五种模型及全部八个种子，逐种子结果见配套数据。',size=9.5)

heading('S4 配对降幅与时长变化',newpage=True)
para('对每个共同种子分别计算100×（对照误差－本文方法误差）／对照误差，再对八个种子取平均。表S4和表S5中的正值表示本文方法误差较低，负值表示误差较高；该均值与先求两组平均误差再计算比例的结果可能不同。')
for ds,name,number in [('abilene','Abilene',4),('geant','GÉANT',5)]:
    rows=[]
    for model in MODELS[:-1]:
        for metric in ['mse','mae']:
            group=paired[(paired.dataset==ds)&(paired.baseline==model)&(paired.metric==metric)].set_index('horizon')
            rows.append([LABELS[model],metric.upper(),*[f'{group.loc[h,"mean_reduction_pct"]:.2f}' for h in [24,48,96,168]]])
    table(f'表S{number} {name}的逐种子平均相对降幅  单位为百分比',['对照模型','指标','24 h','48 h','96 h','168 h'],rows,[4.0,1.6,2.2,2.2,2.2,2.2])
para('不同预测时长的相对误差变化并不一致。在两网络、四种对照和两项指标组成的16组比较中，没有一组相对降幅随24、48、96、168 h依次不减；其中8组的168 h降幅高于24 h，均来自GÉANT。GÉANT的168 h任务中，本文方法在八个种子的MSE与MAE上均优于四种对照。上述比较为描述性汇总，未进行趋势显著性检验。')
para('配套数据保留320条逐种子评分、40组均值与标准差、64组配对降幅，以及数据、代码和预测文件的校验信息。其中summary包含共同起点与全部有效起点两种评分；paired_reductions包含平均降幅、种子间样本标准差和改善种子数。训练记录、检查点和原始预测保存在本项目实验目录，可按数据集、预测时长、模型和种子逐一复核。',size=9.5)
si.save(SUPPLEMENT)

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with ZipFile(BASE) as old,ZipFile(MAIN) as new:
    a=etree.fromstring(old.read('word/document.xml'));b=etree.fromstring(new.read('word/document.xml'))
    assert [etree.tostring(x,method='c14n') for x in a.xpath('.//m:oMath',namespaces=NS)]==[etree.tostring(x,method='c14n') for x in b.xpath('.//m:oMath',namespaces=NS)]
    for name in old.namelist():
        if name.startswith('word/media/'):assert old.read(name)==new.read(name),name
    changed_parts=[name for name in old.namelist() if name in new.namelist() and old.read(name)!=new.read(name)]
original=Document(BASE);revised=Document(MAIN)
for i,(ot,nt) in enumerate(zip(original.tables,revised.tables)):
    assert [[c.text.replace('周期朴素','周期复制') for c in r.cells] for r in ot.rows]==[[c.text for c in r.cells] for r in nt.rows],i
refs=lambda d:[p.text for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)]
assert refs(original)==refs(revised) and len(refs(revised))==11
assert len(revised.inline_shapes)==5 and len(revised.tables)==4
assert len(b.xpath('.//m:oMath',namespaces=NS))==17
assert len(si.tables)==5
assert sha(BASE)==BASE_HASH
manifest={'base':str(BASE),'base_sha256':BASE_HASH,'main':str(MAIN),'supplement':str(SUPPLEMENT),'main_sha256':sha(MAIN),'supplement_sha256':sha(SUPPLEMENT),'figures':5,'tables':4,'equations_unchanged':17,'references_unchanged':11,'original_media_unchanged':4,'chinese_abstract_characters':len(abstract),'changed_package_parts':changed_parts,'data_records_verified':320,'supplement_summary_rows':40,'supplement_paired_comparisons':64,'figure5':str(ASSETS/'图5_GEANT一周预测误差.png')}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False))
