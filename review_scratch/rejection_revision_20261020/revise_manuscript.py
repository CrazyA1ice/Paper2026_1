from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
import json, re, hashlib, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
OUT = ROOT/'word_output/退稿意见修改_20261020'
OUT.mkdir(exist_ok=True)
SOURCE = ROOT/'word_output/基于自适应多尺度加权的多变量网络流量预测方法.docx'
EVIDENCE = json.loads((QA/'verified_evidence.json').read_text(encoding='utf8'))
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
W, M = NS['w'], NS['m']
def E(name, **attrs):
    prefix, local = name.split(':')
    n = etree.Element('{'+NS[prefix]+'}'+local)
    for k,v in attrs.items(): n.set('{'+NS[prefix]+'}'+k,str(v))
    return n
def text_of(el):
    return ''.join(el.xpath('.//w:t/text()|.//m:t/text()',namespaces=NS))
def color_runs(el, color):
    for r in el.xpath('.//w:r|.//m:r',namespaces=NS):
        rp = r.find('{'+W+'}rPr')
        if rp is None:
            rp = E('w:rPr'); r.insert(0,rp)
        c=rp.find('{'+W+'}color')
        if c is None: c=E('w:color'); rp.append(c)
        for a in list(c.attrib): del c.attrib[a]
        c.set('{'+W+'}val',color)

with ZipFile(SOURCE) as z:
    original_xml = z.read('word/document.xml')
base_root = etree.fromstring(original_xml)
base_body = base_root.find('w:body',NS)
originals = [deepcopy(e) for e in base_body]
changes=[]

def new_p(text, template=12, heading=False):
    p=E('w:p')
    pp=originals[template].find('w:pPr',NS)
    if pp is not None: p.append(deepcopy(pp))
    if heading:
        pp=p.find('w:pPr',NS)
        if pp is None: pp=E('w:pPr');p.insert(0,pp)
        if pp.find('w:keepNext',NS) is None: pp.append(E('w:keepNext'))
    run_text(p,text,heading=heading)
    if template == 79:
        pp=p.find('w:pPr',NS)
        if pp is None: pp=E('w:pPr');p.insert(0,pp)
        if pp.find('w:keepNext',NS) is None:pp.append(E('w:keepNext'))
        for rp in p.xpath('./w:r/w:rPr',namespaces=NS):rp.append(E('w:b'))
    return p

def run_text(p, text, heading=False):
    # Recreate only edited text runs. Equations and original images are retained separately.
    for s in re.split(r'(\[\d+\])',text):
        if not s: continue
        r=E('w:r');rp=E('w:rPr');r.append(rp)
        rf=E('w:rFonts',ascii='Times New Roman',hAnsi='Times New Roman',eastAsia='黑体' if heading else '宋体');rp.append(rf)
        rp.append(E('w:color',val='000000'));rp.append(E('w:sz',val='24' if heading else '21'));rp.append(E('w:szCs',val='24' if heading else '21'))
        if heading: rp.append(E('w:b'))
        if re.fullmatch(r'\[\d+\]',s): rp.append(E('w:vertAlign',val='superscript'))
        t=E('w:t');t.text=s;t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');r.append(t);p.append(r)

REPLACEMENTS = {
4: '摘要：针对小时级骨干网流量中短时波动与较慢趋势并存的问题，提出基于自适应多尺度加权的预测方法。模型以三个独立分解线性专家处理不同时间分辨率的历史序列，根据窗口的四项统计量计算尺度权重，并对各尺度预测加权融合。在Abilene和GÉANT上，使用96 h历史预测未来24 h；八随机种子实验中，相对单尺度DLinear的平均均方误差降幅分别为2.77%和2.35%。模型共8339个参数，其中路由网络含131个参数。补充的12 h预测中，相对全局可学习权重，Abilene平均改进15.97%，GÉANT平均改进率为−1.17%。结果表明，输入相关的尺度融合在部分设置下可降低平均误差，但优势依赖数据集和预测长度。现有结果限于两套历史骨干网数据上的离线评价，实际部署效果仍需验证。',
10: 'Abstract: An adaptive multiscale weighting method is proposed for hourly backbone network traffic forecasting, where short-term fluctuations coexist with slower trends. Three independent decomposition linear experts process historical sequences at different temporal resolutions. A routing network computes scale weights from four window statistics and combines the candidate forecasts. Experiments on Abilene and GÉANT use 96 hours of history to predict the next 24 hours. Across eight random seeds, the mean seed-wise reductions in mean squared error relative to single-scale DLinear are 2.77% and 2.35%, respectively. The model has 8,339 parameters, including 131 routing parameters. In additional 12-hour forecasting experiments, the mean improvement over learned global weights is 15.97% on Abilene and −1.17% on GÉANT. These results show that input-dependent scale fusion can reduce mean errors in some settings, but its benefit depends on the dataset and forecast horizon. The evidence is limited to offline evaluation on two historical backbone network datasets; deployment performance requires further validation.',
12: '骨干网容量规划需要估计源—目的节点对（origin-destination，OD）未来的通信需求。本文研究根据过去96 h的OD流量预测未来24 h逐小时需求的问题。不同OD变量的量级和变化规律存在差异，同一序列中又常同时出现短时波动与较慢趋势。因此，模型既要保留局部变化，也要利用时间聚合后的趋势信息，并控制多尺度处理带来的额外开销。',
13: '现有研究从时间依赖、变量关联和周期结构等角度处理网络流量预测问题。Lohrasbinasab等[1]综述了统计方法与机器学习方法的发展；Qin等[2]采用图神经网络建模空间关联与时间依赖；韦烜等[3]分析大型网际协议（Internet protocol，IP）网络流量矩阵的变量关联和时间变化；唐文杰等[4]在CycleLLH中引入周期性信息。这些研究说明，多变量流量预测需要根据任务选择时间结构与变量关系的建模方式。',
14: '多尺度方法通过不同时间分辨率描述流量变化。Zhou等[5]采用深度回声状态网络进行多尺度流量预测。在轻量预测方法中，分解线性模型（decomposition linear model，DLinear）[6]利用序列分解和线性映射生成预测，TimeMixer[7]通过多尺度混合组织时间信息，LightTS[8]采用采样导向的多层感知器（multilayer perceptron，MLP）。本文沿用DLinear作为尺度专家，将研究重点放在不同尺度预测的融合方式上。',
15: '在专家结构相同、参数相互独立的条件下，固定融合权重无法随当前窗口调整各尺度的贡献。本文考察以低维窗口统计量驱动尺度权重，能否在有限参数增量下获得优于固定融合的预测误差。这里的自适应针对尺度，所有OD变量在同一窗口内共用一组权重；模型不显式学习逐变量权重或网络拓扑关系。',
16: '为检验这一问题，本文设置尺度1、2和4的独立DLinear专家，并以四维统计特征驱动尺度路由网络。在Abilene和GÉANT上分别训练和测试，通过固定等权、全局可学习权重及恒定输入路由对照区分不同融合方式的作用，进一步比较12 h与24 h预测结果，并分析计算开销和适用边界。',
18: '图1给出模型的预测分支与路由分支。历史OD窗口经对数变换和标准化后，构造三个时间尺度的输入；各DLinear专家分别预测趋势项与余项，并将二者相加，得到相同预测长度的候选结果。路由网络从预处理窗口提取四项统计量，经Softmax生成尺度权重，最终对候选预测加权求和。图中曲线和色块仅作示意。',
21: '三个专家结构相同、参数独立。细尺度保留局部变化，粗尺度通过时间聚合平滑短时波动；尺度权重决定这些候选预测在最终输出中的相对贡献。',
33: '主实验使用过去96 h预测未来24 h。补充实验保持历史长度不变，将预测长度改为12 h，用于检验融合效果对预测长度的依赖。',
37: '非重叠平均池化不引入可训练参数。96步历史在尺度1、2和4上分别形成96、48和24步输入，粗尺度通过局部取均值平滑短时波动。',
38: '尺度集合{1,2,4}使各尺度输入长度均为整数，每个尺度配置一个独立专家；同一尺度内，每个时间点仅参与一次聚合。',
50: '尺度路由网络由两层全连接层组成，将四维统计向量映射到16维隐藏层，再输出三个尺度分数。隐藏层采用修正线性单元（rectified linear unit，ReLU）激活。该模块只负责尺度融合，不再增加循环网络或注意力编码器。尺度分数计算如下：',
61: '推理时，历史窗口沿用训练阶段的预处理规则，经多尺度专家和路由网络生成未来多步预测。本文的误差评价在标准化空间完成；用于实际需求分析前，还需恢复原始流量量纲。',
63: '实验评价总体误差、融合方式、预测长度、尺度权重和计算开销，并据此讨论适用范围。',
65: '本实验选用CESNET TS-Zoo[9]中的1 h粒度Abilene[10]和GÉANT[11]流量矩阵。主实验均使用96步历史预测未来24步，规模及窗口数见表1。Abilene包含144个有效OD变量；GÉANT仅依据训练集标准差剔除恒定或近恒定变量后保留524个变量。两套数据分别按时间顺序以6∶2∶2划分，并独立训练和测试，预处理见第1.1节。',
70: '各模型使用相同的数据划分、输入长度、主实验预测长度和评价指标，批量大小均为16，优化器均为Adam；超参数预先固定，未对各模型分别搜索。可训练融合参数与尺度专家共同优化，以验证集MSE选择检查点。DLinear及各多尺度模型的学习率为0.001，最多训练30轮，早停耐心值为6；LightTS的学习率为0.0001，最多50轮，早停耐心值为10。因此，表2反映这些固定配置下的比较，不代表各基线经充分调优后的最优性能。',
72: '外部基线比较采用42～44三个共同随机种子；核心消融采用42～49八个共同随机种子。补充的12 h预测仅比较全局可学习权重与本文方法，保持数据划分、历史长度、批量大小和优化设置不变，并按相同八个种子重新训练。两套数据的训练、验证、测试窗口数在12 h设置下分别为2686/824/825和1602/463/463。',
76: '所有指标在对数变换及逐变量标准化后的空间计算，数值越小表示误差越低，不能直接解释为原始流量或链路负载的误差。表2和表3分别报告三个、八个种子的均值±样本标准差。改进率按共同种子逐一计算，即对照与本文方法的误差之差除以对照误差，再取平均；不使用表内舍入值计算。八种子改进率的95%置信区间使用自由度7的t分布和种子间标准误。该区间仅描述固定数据划分下的训练随机性，不代表跨网络不确定性；重叠测试窗口不作为独立重复，本文不据此进行显著性检验。',
78: '表2报告三个共同随机种子的基线结果。在本实验配置下，本文方法在Abilene和GÉANT上的逐种子平均MSE相对DLinear分别降低4.07%和4.54%，相对LightTS分别降低29.80%和46.89%。图2补充各指标的逐种子改进率及参数比例。',
83: '两个数据集的三个共同种子中，本文方法相对LightTS的MSE、MAE和RMSE均降低。8339个参数分别约为LightTS参数量的6.41%和2.17%。这一结果说明本文方法在所设轻量基线下具有参数规模优势，不能据此推断其优于其他类型的流量预测模型。',
84: '以下使用相同的多尺度专家比较融合方式，以减小主干结构差异对结果解释的影响。',
85: '2.3 融合方式消融与预测长度影响',
86: '表3报告八个共同种子的消融结果。该设置比表2增加了训练重复次数，但未增加独立的流量观测数据。',
89: '相对全局可学习权重和恒定输入路由，本文方法在Abilene上的逐种子平均MSE改进率均约为1.80%，在GÉANT上分别为4.31%和4.97%。Abilene相对两类对照均有6/8个种子改善，GÉANT均为7/8。相对全局权重的95%置信区间分别为[0.01%,3.58%]和[−2.51%,11.13%]；相对恒定输入路由分别为[0.49%,3.11%]和[−1.62%,11.56%]。GÉANT的两项区间均跨零，因此其平均改进不能表述为稳定优势。',
93: '相对DLinear，Abilene的MSE、MAE和RMSE平均改进率为2.77%、8.01%和1.40%，GÉANT为2.35%、1.32%和1.22%。相对固定等权多尺度，两套数据的MSE平均改进率分别为1.65%和4.54%。GÉANT的种子间波动更大，部分种子出现负改进。',
94: '在24 h设置下，GÉANT种子43相对恒定输入路由的MSE增加12.77%，说明平均误差下降并不等于每次训练均改善。进一步保持历史长度96 h，将预测长度改为12 h，结果见表4。不同预测长度的测试目标与窗口数不同，因此仅在同一预测长度内比较两种融合方式。',
99: '图4中，四种模型在较高真实值区间均出现低估，本文方法的分组平均偏差较小。四例典型轨迹中，本文方法有三例的24步MSE低于所示对照，在Abilene典型波动窗口中表现较差。典型轨迹仅用于展示行为，整体比较仍以表3为准。',
103: 'Abilene三个尺度的汇总平均权重为0.257、0.319和0.424，GÉANT为0.283、0.318和0.399。GÉANT各尺度权重总方差约有96.2%、95.8%和97.7%来自种子间差异，Abilene对应比例为48.7%、86.1%和80.6%。因此，偏离等权只能说明模型形成了非等权分配，不能单独证明权重随窗口发生有效调整，尤其是GÉANT同一种子内的权重变化较小。',
104: '两个数据集分别按各自训练窗口的平均绝对一阶差分三分位阈值划分低、中、高波动组。Abilene测试窗口数为15、406、392，相对全局权重的平均MSE改进率为4.43%、1.90%、1.69%；GÉANT低波动组为空，中、高组为76、375个窗口，改进率为5.85%、4.06%。Abilene高波动组的尺度1权重较高，但误差改善未同步增大；GÉANT中、高组权重接近，误差改善却有差异。这些结果不支持将误差下降简单归因于更强的动态权重调整。',
105: '2.5 计算开销与网络管理应用',
106: '表5比较参数量与前向耗时。本文方法比固定等权多尺度增加131个路由参数，比全局可学习权重增加128个参数；但相对单尺度DLinear，参数量由4656增至8339，推理也需执行三个专家。因此，路由参数增量较小不能等同于整个模型的计算开销与单尺度相当。',
107: '采用种子42～49的八个检查点，在GTX 1060（6 GB）上以CUDA事件计时，在i7-8750H上进行单线程CPU计时。批量大小为16，输入张量预先载入目标设备并关闭梯度；GPU各检查点预热50批、完整测试集重复30次，CPU预热30批、重复10次。每次总耗时除以批次数，先在检查点内平均，再报告八个检查点的均值±标准差。表5不包含数据读取、设备传输、预处理和结果还原，不能作为在线端到端延迟。',
108: '在相同历史长度与预测长度下，144与524个变量均对应8339个参数，但CPU前向耗时明显增加，说明参数共享不能消除变量规模带来的计算增长。已有训练记录中，本文方法在Abilene和GÉANT上的完整训练评估流程耗时分别为52.47±22.48 s和32.73±34.48 s。该计时包含训练、验证、检查点读写及测试，早停轮次也不同，不能作为纯训练速度或扩展性基准。',
109: '面向小时级需求分析，模型可在新观测到达后，以最近96 h历史生成未来24 h的逐小时OD预测。实际使用前需依次逆标准化、逆对数变换并截断至非负，再结合已知拓扑和路由关系汇总链路需求。本文未执行这一步的在线部署，也未测量调度收益，因而当前实验仅评价离线预测能力。',
112: 'Abilene高需求预测点的改善在不同种子下较一致；GÉANT两项配对改进率的95%置信区间均跨零。上述结果只说明标准化空间内的高需求预测误差，尚不能证明原始流量量纲下的峰值保障或容量调度效果。',
114: '本文研究独立DLinear专家的多尺度融合，以四维窗口统计量生成样本级尺度权重。24 h预测的重复实验在两套数据上得到较低平均误差，但GÉANT的种子间波动较大，12 h设置下也未获得稳定改善。该方法以增加尺度专家计算为代价控制融合参数规模，现阶段适用于所测历史数据上的离线比较。更长时期数据、不同业务网络和在线端到端评价仍是检验其适用性的必要工作。',
}

def rewrite(root, idx, text):
    p=root.find('w:body',NS)[idx]
    assert not p.xpath('.//m:oMath|.//w:drawing',namespaces=NS),idx
    old=text_of(p)
    for child in list(p):
        if child.tag != '{'+W+'}pPr': p.remove(child)
    heading=bool(re.match(r'^\d+\.\d+ ',text))
    run_text(p,text,heading)
    changes.append({'original_body_index':idx,'action':'replace','before':old,'after':text})

root=deepcopy(base_root)
for idx, text in REPLACEMENTS.items(): rewrite(root,idx,text)
body=root.find('w:body',NS)
edited_elements=[body[i] for i in REPLACEMENTS]

def span_replace(p, old, new):
    # Keep native math, fields, and untouched runs byte-for-byte.
    ts=p.xpath('.//w:t',namespaces=NS)
    flat=''.join(t.text or '' for t in ts)
    assert flat.count(old)==1,(old,flat)
    start=flat.index(old); end=start+len(old); pos=0; added=False
    for t in ts:
        s=t.text or '';nxt=pos+len(s)
        if nxt>start and pos<end:
            left=s[:max(0,start-pos)];right=s[max(0,end-pos):] if end<nxt else ''
            t.text=left+(new if not added else '')+right
            added=True
        pos=nxt
    assert old not in ''.join(t.text or '' for t in ts)
    return p

span_replace(body[49], '反映窗口整体流量水平', '表示标准化窗口的平均水平')
span_replace(body[56], '输入窗口决定这一组权重，同一窗口内的全部OD变量共用该组取值。',
    '同一窗口内的全部OD变量和预测步共用这组尺度权重。变量信息通过共享线性映射和汇总统计量参与预测，模型不构造变量对之间的显式交互。该设计控制了参数规模，但会压缩各变量状态差异。')
for idx in [49,56]:
    changes.append({'original_body_index':idx,'action':'local_replace','before':text_of(originals[idx]),'after':text_of(body[idx])}); edited_elements.append(body[idx])

# Keep stable references to the source elements before insertions change their indices.
anchors={i:body[i] for i in range(len(body))}
for idx,tag in [(64,'keepNext'),(65,'keepLines'),(115,'keepNext')]:
    pp=anchors[idx].find('w:pPr',NS)
    if pp is None:pp=E('w:pPr');anchors[idx].insert(0,pp)
    if pp.find('w:'+tag,NS) is None:pp.append(E('w:'+tag))
inserted=[]
def add_after(anchor, nodes, label):
    if not isinstance(nodes,list): nodes=[nodes]
    at=body.index(anchor)+1
    for node in nodes:
        body.insert(at,node);at+=1;inserted.append(node)
    changes.append({'action':'insert','location':label,'after':'\n'.join(text_of(n) for n in nodes)})

def math_r(s):
    r=E('m:r');t=E('m:t');t.text=s;r.append(t);return r
def math_sup(base, sup):
    x=E('m:sSup');e=E('m:e');e.append(math_r(base));p=E('m:sup');p.append(math_r(sup));x.extend([e,p]);return x
def native_math_p(parts):
    p=new_p('',template=25)
    pp=p.find('w:pPr',NS)
    if pp is None:pp=E('w:pPr');p.insert(0,pp)
    jc=pp.find('w:jc',NS)
    if jc is None:jc=E('w:jc');pp.append(jc)
    jc.set('{'+W+'}val','center')
    mp=E('m:oMath')
    for part in parts:mp.append(math_r(part) if isinstance(part,str) else part)
    p.append(mp);return p

# Complexity is derived from the checked implementation, with a fixed moving-average width.
sum_node=E('m:nary');pr=E('m:naryPr');pr.append(E('m:chr',val='∑'));pr.append(E('m:limLoc',val='subSup'));pr.append(E('m:supHide',val='1'));sum_node.append(pr)
sub=E('m:sub');sub.append(math_r('s∈S'));sup=E('m:sup');term=E('m:e');term.append(math_r('2H(L/s+1)'));sum_node.extend([sub,sup,term])
param_formula=native_math_p(['P = ',sum_node,' + 131 = 8339'])
complexity=[new_p('1.5 参数规模与计算复杂度',57,True),
 new_p('设尺度数为K，路由隐藏层宽度为d。每个尺度的两个线性层合计包含2H(L/s+1)个参数，且在OD变量间共享；路由网络参数量为4×16+16+16×3+3=131。主实验取L=96、H=24、K=3，参数总量为：'),
 param_formula,
 new_p('对单个窗口，尺度池化及四项统计量的计算量随LC线性增长；各尺度移动平均的直接计算量与qC乘以各尺度输入长度之和成正比，线性预测与HC乘以该长度之和成正比，输出融合与KHC成正比。固定尺度集合及q=25时，主要计算量为O(CHL+CL)，路由MLP的计算量为O(4d+dK)。忽略固定算子的工作区，输入和候选预测的存储量为O(CL+KHC)。'),
 new_p('因此，在历史长度和预测长度固定时，参数量不随OD变量数增加，但计算量与中间张量仍随变量数增长。历史长度或预测长度增加也会提高参数量和计算量。若保留完整OD矩阵，变量数最多随节点数的平方增长，本文的参数共享并不意味着对网络节点数具有线性扩展性。')]
add_after(anchors[61],complexity,'1.5')

add_after(anchors[67],new_p('按1 h采样点数量折算，两套数据的总观测量约为194.0 d和118.7 d，训练部分约为116.4 d和71.2 d。相邻滑动窗口有重叠，表1窗口数不等于独立观测数，OD变量数也不等于独立网络数。上述数据量可用于所设模型的离线比较，但不足以覆盖长期季节变化及复杂业务变化；增加随机种子不能弥补这一覆盖不足。'),'2.1 表1后')

def make_table(headers, rows, widths):
    # Follow the manuscript's existing three-line table, rather than imposing a report theme.
    tpl=deepcopy(originals[80]);tblpr=tpl.find('w:tblPr',NS);grid=tpl.find('w:tblGrid',NS)
    for c in list(grid):grid.remove(c)
    for w in widths:grid.append(E('w:gridCol',w=w))
    samples=originals[80].findall('w:tr',NS)
    headcell=deepcopy(samples[0].find('w:tc',NS));bodycell=deepcopy(samples[1].find('w:tc',NS))
    for r in tpl.findall('w:tr',NS):tpl.remove(r)
    for ri,data in enumerate([headers]+rows):
        tr=E('w:tr');trpr=E('w:trPr');trpr.append(E('w:cantSplit'))
        if ri==0:trpr.append(E('w:tblHeader'))
        tr.append(trpr)
        for ci,txt in enumerate(data):
            tc=deepcopy(headcell if ri==0 else bodycell)
            for child in list(tc):
                if child.tag != '{'+W+'}tcPr':tc.remove(child)
            pr=tc.find('w:tcPr',NS);cw=pr.find('w:tcW',NS);cw.set('{'+W+'}w',str(widths[ci]))
            if ri==len(rows):
                borders=pr.find('w:tcBorders',NS)
                if borders is None:borders=E('w:tcBorders');pr.append(borders)
                b=borders.find('w:bottom',NS)
                if b is None:b=E('w:bottom');borders.append(b)
                b.set('{'+W+'}val','single');b.set('{'+W+'}sz','8');b.set('{'+W+'}color','000000')
            p=new_p(str(txt),template=12);pp=p.find('w:pPr',NS)
            if pp is None:pp=E('w:pPr');p.insert(0,pp)
            for tag in ['ind','spacing']:
                el=pp.find('w:'+tag,NS)
                if el is not None:pp.remove(el)
            pp.append(E('w:ind',firstLine='0'));pp.append(E('w:spacing',before='30',after='30',line='240',lineRule='auto'))
            if ri < len(rows):pp.append(E('w:keepNext'))
            jc=pp.find('w:jc',NS)
            if jc is None:jc=E('w:jc');pp.append(jc)
            jc.set('{'+W+'}val','center')
            for sz in p.xpath('.//w:sz|.//w:szCs',namespaces=NS):sz.set('{'+W+'}val','20')
            if ri==0:
                for rp in p.xpath('./w:r/w:rPr',namespaces=NS):rp.append(E('w:b'))
            tc.append(p);tr.append(tc)
        tpl.append(tr)
    return tpl

h12rows=[]
for ds,label in [('abilene','Abilene'),('geant','GÉANT')]:
    a=EVIDENCE['h12'][ds];g=a['dlinear_scale_global']['mse'];r=a['dlinear_scale']['mse'];comp=a['router_vs_global']
    h12rows.append([label,f"{g['mean']:.3f}±{g['sd']:.3f}",f"{r['mean']:.3f}±{r['sd']:.3f}",f"{comp['mean']:.2f}",f"[{comp['ci95'][0]:.2f},{comp['ci95'][1]:.2f}]",f"{comp['wins']}/8"])
add_after(anchors[94],[new_p('表4 12 h预测的融合方式比较',79),
    make_table(['数据集','全局权重MSE','本文方法MSE','改进率/%','95%区间/%','改善种子'],h12rows,[1050,1800,1800,1250,2100,1000]),
    new_p('表4中MSE为八种子的均值±标准差，改进率按种子配对计算。12 h预测时，本文方法与全局权重模型分别为4235和4107个参数。Abilene有7/8个种子改善；GÉANT仅3/8个，平均MSE反而略高。因此，24 h预测上的平均改进不能直接推广到其他预测长度。')],'2.3 表4')

timingrows=[]
for ds,label in [('abilene','Abilene'),('geant','GÉANT')]:
    for mod,title in [('dlinear','DLinear'),('dlinear_scale_static','固定等权多尺度'),('dlinear_scale','本文方法')]:
        vals=[]
        for bench in ['inference_benchmark_cuda_events','inference_benchmark_cpu1']:
            r=next(r for r in EVIDENCE['timing'][bench]['rows'] if r['dataset']==ds and r['model']==mod)
            vals.append(f"{float(r['latency_ms_mean']):.2f}±{float(r['latency_ms_sd_across_seeds']):.2f}")
        timingrows.append([label,title,r['parameters'],*vals])
add_after(anchors[107],[new_p('表5 参数量与每批前向耗时',79),
 make_table(['数据集','模型','参数量','GPU耗时/ms','CPU耗时/ms'],timingrows,[1100,2500,1200,2100,2100])],'2.5 表5')

add_after(anchors[112],[new_p('2.6 泛化能力与研究局限',57,True),
 new_p('两套骨干网数据的OD规模和流量分布不同，可用于检验同一建模方案在两个数据集上分别训练后的表现；这不等于训练于一个网络、直接迁移至另一网络的跨网络泛化。本文未开展跨网络迁移、业务类型分层或拓扑变化实验，也未评价在线更新和概念漂移适应。因而不能据现有结果推断其适用于接入网、数据中心或动态路由环境。'),
 new_p('当前证据还受观测时期较短、单一时间划分和有限基线范围约束。GÉANT的种子波动及12 h负改进进一步表明，低维汇总统计量未必能提供充分的状态区分。后续验证需要更长时期及不同业务网络的数据，在保持验证与测试时间段固定的条件下绘制训练数据量学习曲线，并进行滚动时间评价；实际应用还需测量原始量纲误差、单窗口端到端时延及链路容量决策收益。')],'2.6')

edited_elements.extend(inserted)

def write_doc(path, tree):
    with ZipFile(SOURCE) as zin, ZipFile(path,'w',ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True) if info.filename=='word/document.xml' else zin.read(info.filename)
            zout.writestr(info,data)

clean=OUT/'基于自适应多尺度加权的多变量网络流量预测方法_修改稿.docx'
marked=OUT/'基于自适应多尺度加权的多变量网络流量预测方法_标红稿.docx'
write_doc(clean,root)
for el in edited_elements:color_runs(el,'C00000')
write_doc(marked,root)

# Author-facing change record. This is not an appeal or a claim that missing experiments were done.
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
for name in ['Normal','Title','Heading 1','Heading 2']:
    sty=doc.styles[name];sty.font.name='Times New Roman';sty.font.color.rgb=RGBColor(0,0,0)
    sty.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if name!='Normal' else '宋体')
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(18)
doc.styles['Heading 1'].font.size=Pt(13)
doc.styles['Heading 2'].font.size=Pt(11)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.add_paragraph('退稿意见逐条修改说明',style='Title')
doc.add_paragraph('基于自适应多尺度加权的多变量网络流量预测方法\n2026年10月20日')
doc.add_paragraph('本轮以指定最终投稿版为基准，保留原文件，另存净稿和标红稿。修改围绕计算开销、数据与泛化证据、语言逻辑展开，并将既有12 h预测结果和CPU计时补入正文。该说明供作者核对及后续改投使用，不是申诉信。')
doc.add_paragraph('修订状态',style='Heading 1')
doc.add_paragraph('已完成正文修订和既有实验记录核验。训练数据覆盖不足、跨网络迁移和在线部署验证仍未解决，不能将本轮文字修改视为这三项实验已经完成。当前稿件可供进一步完善和改投评估，尚不能声称已充分回应全部实证质疑。')
doc.add_paragraph('一 关于计算开销与训练数据不足',style='Heading 1')
comment1='（1）本文方法的核心研究难点在于如何实现多尺度特征的有效提取、不同尺度与变量权重的动态协同，以及复杂网络环境下预测精度、实时性和泛化能力的平衡。多尺度特征提取、自适应权重计算以及多变量关联建模会增加模型的计算开销。当网络规模扩大、变量数量增加或预测时间窗口延长时，可能出现训练成本高、推理延迟增加等问题，影响模型在实时网络管理中的应用。实验中的训练数据明显不足，说服力偏弱。'
doc.add_paragraph('原意见',style='Heading 2');doc.add_paragraph(comment1)
doc.add_paragraph('修改及证据',style='Heading 2')
for t in [
 '第1.3节明确同一窗口的全部OD变量和预测步共用尺度权重，不宣称逐变量权重协同或显式变量交互。模型的贡献限定在输入相关的尺度融合。',
 '新增第1.5节，依据实际共享线性层推导8339个参数及131个路由参数，解释计算量、存储量随OD变量数和历史长度增长。若节点数增加使OD变量数按平方增长，计算量也不能称为对节点数线性扩展。',
 '第2.5节新增表5，采用已保存的CUDA事件与单线程CPU计时结果。GTX 1060上本文方法每批耗时在Abilene、GÉANT分别为3.28±0.20 ms、5.26±2.79 ms；CPU为40.65±0.19 ms、131.35±1.25 ms。表中同时列出DLinear及固定等权多尺度，明确多尺度模型需承担更高计算成本。',
 'GPU旧稿的2.65/2.61 ms来自另一轮计时，本次统一为记录了CUDA事件后端、预热50批和重复30次的结果。没有选取更有利的旧数字，也未将前向耗时称为在线延迟。',
 '第2.1节新增观测量和有效样本解释。训练数据折合约116.4 d和71.2 d；重叠窗口、OD变量和随机种子都不能视为新增独立网络或独立观测。第2.6节明确尚未绘制训练数据量学习曲线，数据覆盖不足仍是未完成项。',
 '训练记录的elapsed_seconds包含训练、验证、检查点读写和测试，且早停不同。第2.5节据实报告该流程耗时，同时说明它不是纯训练速度基准。']:
    doc.add_paragraph(t)
doc.add_paragraph('仍需完成',style='Heading 2')
doc.add_paragraph('增加更长时期及不同业务来源的数据；在验证、测试时间段固定时，比较不同训练数据量，并报告误差曲线。若要论证实时部署，还需在目标设备上测量批量为1的完整处理流程及峰值内存。')
doc.add_paragraph('二 关于泛化能力和实际应用验证',style='Heading 1')
comment2='（2）模型泛化能力和实际应用验证不足。 不同网络环境下的流量分布、业务构成和负载特征存在差异。如果模型主要基于单一数据集开展实验，其在不同网络拓扑、不同业务场景以及动态变化环境中的适应能力可能不足，实际部署效果仍需进一步验证。'
doc.add_paragraph('原意见',style='Heading 2');doc.add_paragraph(comment2)
doc.add_paragraph('修改及证据',style='Heading 2')
for t in [
 '第2.1节突出Abilene与GÉANT分别训练、分别测试的设计，改善双数据集证据的可见性；第2.6节区分数据集内时间外推与跨网络迁移，不把两个数据集上的结果当作跨拓扑迁移成功。',
 '第2.3节补入此前已完成、未写入投稿正文的12 h实验及表4。已从32份配置和保存预测重算MSE、MAE，并核对同一种子两模型测试目标完全一致。Abilene平均MSE改进15.97%，95%区间为[5.05%,26.88%]，7/8个种子改善；GÉANT平均改进率为−1.17%，区间为[−5.79%,3.46%]，仅3/8个种子改善。负结果同步写入中英文摘要和结语。',
 '第2.4节收紧权重机制解释。偏离等权不能单独证明有效动态适应，GÉANT权重方差主要来自种子差异，波动分层也不支持把误差改善直接归因于权重变化。',
 '第2.5及2.6节明确原始量纲、链路汇总、概念漂移、在线更新和实际调度收益尚未验证。容量规划表述改为潜在用途和部署前步骤。']:
    doc.add_paragraph(t)
doc.add_paragraph('仍需完成',style='Heading 2')
doc.add_paragraph('跨网络训练测试需要明确变量对齐和目标网络预处理协议；滚动时间评价需要将训练、验证和测试按先后顺序多次推进；实际应用评价需要真实拓扑、路由及容量信息。这些工作本轮没有执行，也未用拟议方案代替结果。')
doc.add_paragraph('三 关于语言表达和逻辑',style='Heading 1')
comment3='（3）文章中的语句表达、描述之间逻辑性较弱，可读性不强（是否由翻译软件引起？），建议进一步修改及完善。'
doc.add_paragraph('原意见',style='Heading 2');doc.add_paragraph(comment3)
doc.add_paragraph('修改及证据',style='Heading 2')
doc.add_paragraph('引言按应用问题、既有方法、具体问题与验证路线组织；方法部分压缩对候选预测和尺度长度的重复解释；结果部分区分三个种子与八个种子的比较口径；讨论与结语明确证据支持范围。同步修订英文摘要，保持与中文摘要的数据、条件和负结果一致。')
doc.add_paragraph('局部纠正包括将“共享同一结构”改为“结构相同、参数独立”，将窗口均值限定为标准化空间，将“6/8和7/8”明确对应到数据集，明确波动阈值由各自训练集计算，并将种子43的误差增加率由12.78%更正为12.77%。')
doc.add_paragraph('统一用语',style='Heading 2')
for t in ['尺度权重指三个时间尺度的融合系数；变量权重未在本方法中建模。',
          '样本级尺度路由指每个输入窗口生成一组权重；同一窗口内所有OD变量和预测步共用。',
          '标准化空间指标指对数变换和逐OD标准化后的误差；不能替代原始流量或链路负载误差。',
          '种子间波动指固定数据划分下训练随机性的变化；不代表新网络泛化的不确定性。']:
    doc.add_paragraph(t)
doc.add_paragraph('四 证据来源与修改范围',style='Heading 1')
doc.add_paragraph('本轮没有重新训练模型，也没有增加新网络或新观测。新增表4来自既有horizon12_global_vs_router实验，新增表5来自既有inference_benchmark_cuda_events与inference_benchmark_cpu1记录。主实验统计由metrics.json交叉核对，新预测长度还核对了保存的predictions.npz。相关实验文件保持不变。')
doc.add_paragraph('净稿与标红稿正文一致，标红以指定最终投稿版为基准，仅标记本轮替换及新增段落。原有41处原生公式、5幅图、3张表和11条参考文献保留；新增参数量原生公式和两张结果表。未加入未经核验的引用或声称已完成的部署实验。')
before_chars=sum(len(text_of(p)) for p in originals if p.tag=='{'+W+'}p')
after_chars=sum(len(text_of(p)) for p in body if p.tag=='{'+W+'}p')
doc.add_paragraph(f'为防止反复追加意见造成正文膨胀，本轮主要替换并压缩原段落，同时加入必要的复杂度与失败边界。按正文段落文本字符统计（不计表格），修改前{before_chars}字/字符，修改后{after_chars}字/字符。该统计不是Word字数统计，是否满足目标期刊篇幅需按改投要求确认。')
doc.add_paragraph('建议补充验证的最小方案',style='Heading 2')
for t in ['训练数据量验证。在每个数据集上固定验证与测试区间，从同一训练截止点向前截取25%、50%、75%、100%的训练观测，保持预处理仅使用对应训练段统计量。至少比较DLinear、全局权重与本文方法，保留八共同种子；结果尚未计算。',
          '时间稳健性验证。在不重叠的测试区间进行滚动起点评价，分别报告各时段误差。不得把重叠预测点当作独立重复进行推断；结果尚未计算。',
          '部署验证。以目标环境中的单个96 h窗口测量数据读取、预处理、推理、结果还原的完整延迟，记录目标硬件及峰值内存，再用真实路由映射评价链路需求；结果尚未计算。']:
    doc.add_paragraph(t)
report=OUT/'退稿意见逐条修改说明.docx'
for p in doc.paragraphs:
    for border in p._p.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(report)
(QA/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf8')
manifest={'source':str(SOURCE),'source_sha256':EVIDENCE['source_sha256'],
          'clean':str(clean),'marked':str(marked),'report':str(report),
          'before_paragraph_characters':before_chars,'after_paragraph_characters':after_chars,
          'replaced_paragraphs':len(REPLACEMENTS)+2,'new_elements':len(inserted)}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False,indent=2))
