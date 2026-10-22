from pathlib import Path
from zipfile import ZipFile
import hashlib,importlib.util,json,os,re,sys
from docx import Document
from lxml import etree
import pdfplumber

sys.stdout.reconfigure(encoding='utf8')
QA=Path(__file__).resolve().parent
manifest=json.loads((QA/'manifest.json').read_text(encoding='utf8'))
target=Path(manifest['output']);base=Path(manifest['base'])
assert hashlib.sha256(target.read_bytes()).hexdigest()==manifest['output_sha256']
assert hashlib.sha256(base.read_bytes()).hexdigest()==manifest['base_sha256']
doc=Document(target)
assert len(doc.inline_shapes)==4 and len(doc.tables)==4
text='\n'.join(p.text for p in doc.paragraphs)
assert '圆点和方点' not in text and '横向误差线' not in text
assert '柱高与误差线分别表示八个种子的均值和样本标准差' in text
assert '否则记为0，再对全部高需求点取平均' in text
assert [int(re.match(r'^图(\d+) ',p.text).group(1)) for p in doc.paragraphs if re.match(r'^图(\d+) ',p.text)]==[1,2,3,4]
with ZipFile(target) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    assert len(xml.xpath('.//m:oMath',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}))==17
    assert Path(manifest['figure4']).read_bytes() in [z.read(n) for n in z.namelist() if n.startswith('word/media/')]
renderer=Path(r'C:\Users\Quant\.codex\plugins\cache\openai-primary-runtime\documents\26.1020.11127\skills\documents\render_docx.py')
spec=importlib.util.spec_from_file_location('bar_document_renderer',renderer)
module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module)
native=Path(r'C:\Users\Quant\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler')
os.environ['PATH']=str(next(native.rglob('pdftoppm.exe')).parent)+os.pathsep+os.environ['PATH']
pdf=QA/'render/review.pdf'
module.convert_to_pdf=lambda *args,**kwargs:(str(pdf),'Native Microsoft Word readonly export')
pages=module.rasterize(str(target),str(QA/'render'),150,False,False)
stats=json.loads((QA/'word_statistics.json').read_text(encoding='utf-8-sig'))
assert len(pages)==stats['pages'] and stats['words']<=7500
texts=[];records=[]
with pdfplumber.open(pdf) as document:
    for i,page in enumerate(document.pages,1):
        page_text=page.extract_text() or '';texts.append(page_text)
        bad=[c for c in page.chars if c['text'].strip() and (c['x0']<-1 or c['top']<-1 or c['x1']>page.width+1 or c['bottom']>page.height+1)]
        assert not bad
        records.append({'page':i,'first':page_text[:110],'last':page_text[-150:]})
text='\n'.join(texts)
for n in range(1,5):assert re.search(r'图\s*'+str(n),text)
for value in ['17.22%','20.49%','3.59%','6.30%','0.147','0.137','0.632','0.577']:assert value in text,value
assert '作者简介' in text and '编辑部填写' not in text
(QA/'rendered_text.txt').write_text(text,encoding='utf8')
result={**manifest,'pages':stats['pages'],'word_count':stats['words'],'page_details':records,'out_of_page_text':0,'visual_review_pending':True}
(QA/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in result.items() if k!='page_details'},ensure_ascii=False))
