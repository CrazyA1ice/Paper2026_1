from pathlib import Path
from zipfile import ZipFile
import hashlib,json,re,shutil,sys
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Cm,Emu,Pt,RGBColor
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).resolve().parents[2]
QA=Path(__file__).resolve().parent
PREVIOUS=ROOT/'review_scratch/changchun_two_figures'
OUT=ROOT/'paper_template/长春工程学院学报(自然科学版)/改投内容修改第一版_20261020'
BASE=OUT/'自适应多尺度加权的多变量网络流量预测_补充两图版.docx'
TARGET=OUT/'自适应多尺度加权的多变量网络流量预测_柱状图修订版.docx'
BASE_HASH='b3560b3b5542146e4e27872fc6ab430d940c6fee05c8730dd714ba2b552d09b2'
assert hashlib.sha256(BASE.read_bytes()).hexdigest()==BASE_HASH
DATA=ROOT/'abilene_forecasting/results/high_demand_analysis'
per_seed=pd.read_csv(DATA/'per_seed.csv')
assert len(per_seed)==80 and not per_seed.duplicated(['dataset','model','seed']).any()
model_order=['dlinear','dlinear_scale_static','dlinear_scale_global','dlinear_scale_constant_router','dlinear_scale']
audited=[]
for ds in ['abilene','geant']:
    processed=ROOT/f'abilene_forecasting/data/processed/{ds}_1hour.npz'
    with np.load(processed) as raw:
        threshold=np.quantile(raw['train'].astype(np.float64),.9,axis=0)
        test=raw['test']
    reference=np.stack([test[i+96:i+120] for i in range(len(test)-119)])
    mask=reference>threshold[None,None,:]
    actual=reference[mask].astype(np.float64)
    for model in model_order:
        group=per_seed[(per_seed.dataset==ds)&(per_seed.model==model)].sort_values('seed')
        assert group.seed.tolist()==list(range(42,50))
        for _,row in group.iterrows():
            seed=int(row.seed)
            parent=ROOT/'abilene_forecasting/results'
            if model=='dlinear_scale_global':parent=parent/'global_weight_ablation'
            if model=='dlinear_scale_constant_router':parent=parent/'constant_router_ablation'
            path=parent/ds/f'{model}_seed{seed}'/'predictions.npz'
            with np.load(path) as raw:
                assert np.array_equal(raw['truth'],reference)
                prediction=raw['prediction']
            assert prediction.shape==reference.shape and np.isfinite(prediction).all()
            error=prediction[mask].astype(np.float64)-actual
            values={'mae':float(np.abs(error).mean()),'mean_shortfall':float(np.maximum(-error,0).mean())}
            assert int(mask.sum())==int(row.n_high_targets)
            for metric,value in values.items():assert np.isclose(value,row[metric],rtol=1e-12,atol=1e-14),(ds,model,seed,metric)
            audited.append({'dataset':ds,'model':model,'seed':seed,**values,'prediction_sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    print(f'{ds}: verified all 40 model-seed records against saved predictions',flush=True)

alternative=OUT/'图4备选'
plotted=pd.read_csv(alternative/'备选B_绘图数据.csv')
assert len(plotted)==20
for _,row in plotted.iterrows():
    values=per_seed[(per_seed.dataset==row.dataset)&(per_seed.model==row.model)][row.metric]
    assert len(values)==8
    assert np.isclose(values.mean(),row['mean'],rtol=1e-12)
    assert np.isclose(values.std(ddof=1),row.sample_sd,rtol=1e-12)
    assert int(row.n_seeds)==8
expected_pdf='6db549e85bc086978eafb19808cad4ddc7def7bd39eb9d6554ec63f54f973cbb'
assert hashlib.sha256((alternative/'备选B_高需求柱状图.pdf').read_bytes()).hexdigest()==expected_pdf
figure_prefix=OUT/'补充图表/图4_高需求预测误差_柱状图'
for suffix in ['.png','.pdf','.svg','.tiff']:
    shutil.copy2(alternative/('备选B_高需求柱状图'+suffix),figure_prefix.with_suffix(suffix))
per_seed.to_csv(figure_prefix.parent/'图4_柱状图_全部种子.csv',index=False,encoding='utf-8-sig')
plotted.to_csv(figure_prefix.parent/'图4_柱状图_绘图数据.csv',index=False,encoding='utf-8-sig')

doc=Document(BASE)
def find(prefix):
    matches=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches)==1,(prefix,len(matches))
    return matches[0]
def set_text(p,text,size=10.5,bold=False):
    p.clear();run=p.add_run(text)
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if bold else '宋体')

before=find('进一步以训练集各OD变量')
after=find('图4中，圆点和方点')
old_text=[before.text,after.text]
set_text(before,'为考察较高通信需求下的预测表现，以训练集各OD变量的第90百分位为阈值，将测试集中真实值超过对应阈值的预测点定义为高需求点。在这些点上计算MAE，并进一步统计平均低估误差。当预测值低于真实值时，记录二者的差额，否则记为0，再对全部高需求点取平均。各指标均在前述标准化空间计算，先在单个种子内计算，再对八个共同种子汇总，结果见图4。')
set_text(after,'图4中，柱高与误差线分别表示八个种子的均值和样本标准差；“全局权重”和“恒定路由”分别对应全局可学习权重和恒定输入路由。本文方法在Abilene上的MAE与平均低估误差分别为0.147和0.137，在GÉANT上分别为0.632和0.577，两项均值均低于图中所列对照。相对全局可学习权重，Abilene上两项指标的逐种子平均降幅分别为17.22%和20.49%，GÉANT上分别为3.59%和6.30%，表明本文方法在所设24 h任务的高需求点上取得了较低的平均预测误差。')
set_text(find('图4 高需求预测点'),'图4 高需求预测点的MAE与平均低估误差',9,True)
shape=doc.inline_shapes[3]
rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
doc.part.related_parts[rid]._blob=figure_prefix.with_suffix('.png').read_bytes()
with Image.open(figure_prefix.with_suffix('.png')) as im:
    shape.width=Cm(13.8);shape.height=Emu(round(int(shape.width)*im.height/im.width))
shape._inline.docPr.set('descr','两个网络上五种模型的高需求预测误差分组柱状图；柱高为八个种子均值，误差线为样本标准差。')
doc.save(TARGET)

ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with ZipFile(BASE) as z:old=etree.fromstring(z.read('word/document.xml'))
with ZipFile(TARGET) as z:new=etree.fromstring(z.read('word/document.xml'))
assert [etree.tostring(x,method='c14n') for x in old.xpath('.//m:oMath',namespaces=ns)]==[etree.tostring(x,method='c14n') for x in new.xpath('.//m:oMath',namespaces=ns)]
assert len(new.xpath('.//m:oMath',namespaces=ns))==17
assert [etree.tostring(x,method='c14n') for x in old.xpath('.//w:tbl',namespaces=ns)]==[etree.tostring(x,method='c14n') for x in new.xpath('.//w:tbl',namespaces=ns)]
assert len(doc.inline_shapes)==4 and len(doc.tables)==4
assert [p.text for p in Document(BASE).paragraphs if re.match(r'^\[\d+\]',p.text)]==[p.text for p in doc.paragraphs if re.match(r'^\[\d+\]',p.text)]
with ZipFile(BASE) as old_zip,ZipFile(TARGET) as new_zip:
    for media in old_zip.namelist():
        if media.startswith('word/media/') and media!=doc.part.related_parts[rid].partname[1:]:assert old_zip.read(media)==new_zip.read(media)
assert hashlib.sha256(BASE.read_bytes()).hexdigest()==BASE_HASH
manifest={'base':str(BASE),'output':str(TARGET),'base_sha256':BASE_HASH,'output_sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'figure4':str(figure_prefix.with_suffix('.png')),'figures':4,'tables':4,'equations':17,'references':11,'raw_records_verified':len(audited),'source_csv_sha256':hashlib.sha256((DATA/'per_seed.csv').read_bytes()).hexdigest(),'figure4_pdf_sha256':expected_pdf}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
(QA/'raw_data_verification.json').write_text(json.dumps(audited,ensure_ascii=False,indent=2),encoding='utf8')
(QA/'prose_revision.json').write_text(json.dumps({'classification':'necessary_support','terminology':{'全局权重':'全局可学习权重','恒定路由':'恒定输入路由','误差线':'八个种子间样本标准差'},'before_characters':sum(map(len,old_text)),'after_characters':len(before.text)+len(after.text),'changes':'Replace two existing paragraphs; define the metric, explain bar/SD encoding and retain the same bounded numerical results.'},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False))
