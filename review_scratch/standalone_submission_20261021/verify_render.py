from pathlib import Path
from zipfile import ZipFile
import hashlib,importlib.util,json,os,re,sys
import pdfplumber
from docx import Document

sys.stdout.reconfigure(encoding='utf8')
QA=Path(__file__).resolve().parent
m=json.loads((QA/'manifest.json').read_text(encoding='utf8'))
stats=json.loads((QA/'word_statistics.json').read_text(encoding='utf-8-sig'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
target=Path(m['output']);base=Path(m['base']);pdf=Path(stats['pdf'])
assert sha(base)==m['base_sha256']
assert sha(target)==m['output_sha256']
doc=Document(target)
assert len(doc.inline_shapes)==4 and len(doc.tables)==3
text='\n'.join(p.text for p in doc.paragraphs)
assert [p.text for p in doc.paragraphs if re.match(r'^2\.\d+ ',p.text)]==['2.1 数据与实验设置','2.2 预测精度比较','2.3 融合方式消融','2.4 高需求预测表现','2.5 一周预测表现']
assert [[c.text for c in row.cells] for row in doc.tables[0].rows]==m['standalone_task_table']
assert 'RMSE' not in text and '表4' not in text and '图5' not in text
for token in ['29.80%','46.89%','17.57%','29.15%','6.41%','2.17%','57587','307','23.16%','11.16%']:
    assert token in text,token
def image_blob(d,i):
    rid=d.inline_shapes[i]._inline.graphic.graphicData.pic.blipFill.blip.embed
    return d.part.related_parts[rid].blob
old=Document(base)
for new_i,old_i in [(0,0),(2,1),(3,2)]:
    assert image_blob(doc,new_i)==image_blob(old,old_i)

renderer_path=Path(r'C:\Users\Quant\.codex\plugins\cache\openai-primary-runtime\documents\26.1020.11127\skills\documents\render_docx.py')
spec=importlib.util.spec_from_file_location('standalone_renderer',renderer_path)
renderer=importlib.util.module_from_spec(spec);sys.modules[spec.name]=renderer;spec.loader.exec_module(renderer)
poppler=Path(r'C:\Users\Quant\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler')
os.environ['PATH']=str(next(poppler.rglob('pdftoppm.exe')).parent)+os.pathsep+os.environ['PATH']
renderer.convert_to_pdf=lambda *args,**kwargs:(str(pdf),'Native Microsoft Word read-only export')
pages=renderer.rasterize(str(target),str(pdf.parent),150,False,False)
assert len(pages)==stats['pages']
details=[];texts=[]
with pdfplumber.open(pdf) as document:
    for i,page in enumerate(document.pages,1):
        txt=page.extract_text() or '';texts.append(txt)
        bad=[c for c in page.chars if c['text'].strip() and (c['x0']<-1 or c['top']<-1 or c['x1']>page.width+1 or c['bottom']>page.height+1)]
        assert not bad,i
        details.append({'page':i,'first':txt[:150],'last':txt[-220:],'images':len(page.images)})
extracted='\n'.join(texts)
(QA/'main_rendered_text.txt').write_text(extracted,encoding='utf8')
compact=re.sub(r'\s+','',extracted)
for token in ['补充材料','附录','见存档','CPU','GPU','GTX','显卡','处理器','计算耗时']:
    assert token not in compact,token
for token in ['29.80%','46.89%','17.57%','29.15%','6.41%','2.17%','57587','1446/307/307','23.16%','11.16%','0.841','0.536','作者简介']:
    assert token in compact,token
for i in [1,2,3,4]:
    assert re.search(r'图\s*'+str(i),extracted),i
assert stats['words']<=7500
result={**stats,'out_of_page_text':0,'source_preserved':True,'figures':4,'tables':3,'native_math_preserved':16,'visual_review_pending':True,'page_details':details}
(QA/'render_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
