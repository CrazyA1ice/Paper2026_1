from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib,json,re,shutil,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).resolve().parents[2];QA=Path(__file__).resolve().parent
PARENT=ROOT/'paper_template/长春工程学院学报(自然科学版)'
BASE=PARENT/'数据处理精简修订_20261021/自适应多尺度加权的多变量网络流量预测_数据分析精简版.docx'
OUT=PARENT/'主稿独立投稿修订_20261021';ASSETS=OUT/'存档数据与图表'
TARGET=OUT/'自适应多尺度加权的多变量网络流量预测_独立投稿主稿.docx'
BASE_HASH='3a79dfd13db4bac005b8ffff1e12845f6e7846d84e2acaad7a6520db2f061086'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(BASE)==BASE_HASH
assert json.loads((QA/'figure2_collision.json').read_text(encoding='utf8'))['verdict']=='PASS'
shutil.copy2(BASE,TARGET)
doc=Document(TARGET)

def find(prefix):
    ps=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(ps)==1,(prefix,len(ps))
    return ps[0]

def text(p,value,size=10.5,bold=False):
    p.clear();r=p.add_run(value);r.font.name='Times New Roman';r.font.size=Pt(size)
    r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if bold else '宋体')
    return p

def remove(p):p._p.getparent().remove(p._p)

text(find('基于上述思路'),'基于上述思路，本文提出自适应多尺度加权的网络流量预测方法。首先，通过时间聚合构造三种分辨率的输入，并由独立DLinear专家生成候选预测；随后，利用窗口统计量计算尺度权重，完成样本级融合。在Abilene和GÉANT上，以过去96 h预测未来24 h，通过基线比较、融合消融与高需求分析评价预测效果，并在GÉANT上考察未来168 h的一周预测表现。')
text(find('基础实验取L=96'),'基础任务取L=96、H=24，一周任务取L=96、H=168。数据读取时，将缺失值和无穷值记为0，并将负值截为0。考虑到流量非负且变量量级不同，先进行自然对数变换，再按变量标准化：')
text(find('同一尺度的线性映射'),'同一尺度的线性映射在变量间共享，因此参数量不随OD变量数增加。当H=168时，模型含57587个参数，路由网络仍为131个参数。模型依次通过多尺度预测与自适应融合，输出标准化空间中的多步预测结果。')

text(find('实验采用CESNET TS-Zoo'),'实验采用CESNET TS-Zoo[9]中1 h粒度的Abilene[10]与GÉANT[11]流量矩阵，两套数据分别含4656和2849个时间点。按时间顺序以6∶2∶2划分训练集、验证集和测试集，并独立训练和评价。依据训练集标准差筛选有效变量后，Abilene保留144个OD变量，GÉANT保留524个OD变量。两套数据均开展24 h预测，GÉANT另设168 h预测任务，数据划分与样本窗口数见表1。')
text(find('表1 数据集与24 h预测任务设置'),'表1 数据集规模与预测任务设置',9,True)
t=doc.tables[0]
t._tbl.append(deepcopy(t.rows[-1]._tr))
values=[['数据集','有效变量数','预测时长/h','训练/验证/测试时间点','训练/验证/测试窗口'],['Abilene','144','24','2793/931/932','2674/812/813'],['GÉANT','524','24','1709/570/570','1590/451/451'],['GÉANT','524','168','1709/570/570','1446/307/307']]
widths=[2.1,2.1,2.0,4.9,4.7]
t.autofit=False
for col,width in zip(t.columns,widths):col.width=Cm(width)
for i,(row,row_values) in enumerate(zip(t.rows,values)):
    for c,value,width in zip(row.cells,row_values,widths):
        c.width=Cm(width)
        p=c.paragraphs[0];text(p,value,9,i==0)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for c in t.rows[-2].cells:
    c._tc.get_or_add_tcPr().find(qn('w:tcBorders')).find(qn('w:bottom')).set(qn('w:val'),'nil')

text(find('基线比较采用42～44'),'表2与图2的误差比较采用42～44三个共同随机种子，表3、图3和图4采用42～49八个共同种子。学习模型的误差报告均值与样本标准差；相对降幅在共同种子下逐一计算，再取平均。图中误差线均表示样本标准差。')
text(find('两个数据集上，本文方法的平均MSE'),'两个数据集上，本文方法的平均MSE和MAE均低于表2所列基线。图2进一步比较其相对LightTS的误差降幅与参数比例：MSE的逐种子平均降幅在Abilene和GÉANT上分别为29.80%和46.89%，MAE分别为17.57%和29.15%。本文模型均含8339个参数，而LightTS的两套配置分别含130162和384382个参数，对应参数比例为6.41%和2.17%。')
anchor=find('2.3 融合方式消融')
picture=anchor.insert_paragraph_before()
picture.paragraph_format.first_line_indent=Pt(0);picture.alignment=WD_ALIGN_PARAGRAPH.CENTER
picture.paragraph_format.keep_with_next=True
shape=picture.add_run().add_picture(str(ASSETS/'图2_LightTS误差降低率与参数比例.png'),width=Cm(13.8))
shape._inline.docPr.set('descr','三个面板分别显示本文方法相对LightTS的MSE降低率、MAE降低率与参数量比例；误差降低率使用三个共同训练种子，参数比例以LightTS为100%。')
caption=anchor.insert_paragraph_before()
caption.paragraph_format.first_line_indent=Pt(0);caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after=Pt(4)
text(caption,'图2 相对LightTS的预测误差降低率与参数比例',9,True)
text(anchor.insert_paragraph_before(),'图2(a)、(b)的柱高为三个共同种子的平均相对降幅，误差线为样本标准差；图2(c)以对应LightTS配置的参数量为100%，给出本文模型的参数比例。该对照表明，在所设24 h任务中，本文方法以较少的模型参数取得了较低的平均预测误差。')

p=find('表3中，本文方法');text(p,p.text.replace('尺度权重分布见补充材料S6。',''))
p=find('以训练集各OD变量的第90百分位');text(p,p.text.replace('结果见图2','结果见图3'))
text(find('图2 高需求预测点'),'图3 高需求预测点的MAE与平均低估误差（8次重复）',9,True)
p=find('图2中，本文方法在Abilene');text(p,p.text.replace('图2中','图3中',1))

text(find('保持小时采样、96 h输入'),'为考察一周范围内的预测表现，在GÉANT上设置H=168的任务，并保持小时采样、96 h输入、数据划分和预处理方式不变。DLinear、固定等权、全局可学习权重、恒定输入路由与本文方法均按168 h输出独立训练，沿用2.1节中DLinear及多尺度模型的训练设置。五种模型在相同的307个测试窗口上评价，覆盖全部168个预测小时与524个OD变量，结果见图4。')
text(find('图3 GÉANT一周预测任务'),'图4 GÉANT一周预测任务的MSE与MAE（8次重复）',9,True)
p=find('本文方法的MSE和MAE分别为0.841');text(p,p.text.replace('低于图中四种对照','低于图4中的四种对照'))
remove(find('2.6 模型计算耗时'))
remove(find('在GTX 1060'))

text(find('在Abilene和GÉANT的24 h任务中'),'24 h任务的基线比较和融合消融表明，本文方法在Abilene和GÉANT上均取得了较低的平均预测误差。相对LightTS，MSE的逐种子平均降幅分别为29.80%和46.89%，参数量分别为对应LightTS配置的6.41%和2.17%。在GÉANT的一周任务中，MSE与MAE相对DLinear的逐种子平均降幅分别为23.16%和11.16%。这些结果为利用低维窗口特征融合多尺度预测提供了实验依据，可为小时级OD流量需求分析提供方法参考。')
doc.save(TARGET)

ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with ZipFile(BASE) as old,ZipFile(TARGET) as new:
    a=etree.fromstring(old.read('word/document.xml'));b=etree.fromstring(new.read('word/document.xml'))
    assert [etree.tostring(x,method='c14n') for x in a.xpath('.//m:oMath',namespaces=ns)]==[etree.tostring(x,method='c14n') for x in b.xpath('.//m:oMath',namespaces=ns)]
    at=a.xpath('.//w:tbl',namespaces=ns);bt=b.xpath('.//w:tbl',namespaces=ns)
    assert [etree.tostring(x,method='c14n') for x in at[1:]]==[etree.tostring(x,method='c14n') for x in bt[1:]]
    allowed={'word/document.xml','word/_rels/document.xml.rels'}
    changed=[n for n in old.namelist() if old.read(n)!=new.read(n)]
    assert set(changed)<=allowed,changed
    for name in old.namelist():
        if name.startswith('word/media/'):assert old.read(name)==new.read(name)
    alltext='\n'.join(''.join(etree.fromstring(new.read(n)).itertext()) for n in new.namelist() if n.startswith('word/') and n.endswith('.xml'))
    for token in ['补充材料','附录','见存档','CPU','GPU','GTX','i7-8750H','显卡','处理器','计算耗时']:
        assert token not in alltext,token
assert len(doc.inline_shapes)==4 and len(doc.tables)==3
refs=lambda d:[p.text for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)]
assert refs(Document(BASE))==refs(doc) and len(refs(doc))==11
assert sha(BASE)==BASE_HASH
assert len(b.xpath('.//m:oMath',namespaces=ns))==16
assert [p.text for p in doc.paragraphs if re.match(r'^图\d+ ',p.text)]==['图1 自适应多尺度网络流量预测框架','图2 相对LightTS的预测误差降低率与参数比例','图3 高需求预测点的MAE与平均低估误差（8次重复）','图4 GÉANT一周预测任务的MSE与MAE（8次重复）']
manifest={'base':str(BASE),'base_sha256':BASE_HASH,'output':str(TARGET),'output_sha256':sha(TARGET),'figure2':str(ASSETS/'图2_LightTS误差降低率与参数比例.png'),'figures':4,'tables':3,'native_math_preserved':16,'references_preserved':11,'changed_package_parts':changed,'supplement_references':0,'hardware_mentions':0,'reported_tasks':['Abilene H24','GÉANT H24','GÉANT H168'],'standalone_task_table':values}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False))
