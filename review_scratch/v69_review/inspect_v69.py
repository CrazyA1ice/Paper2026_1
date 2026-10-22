from pathlib import Path
import zipfile, hashlib, json, difflib
from lxml import etree
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
SOURCE = ROOT / 'word_output/2026.10.04v69.docx'
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
def content(node):
    return ''.join(n.text or '' for n in node.iter() if n.tag in (f'{{{NS["w"]}}}t', f'{{{NS["m"]}}}t'))
with zipfile.ZipFile(SOURCE) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
    paragraphs = xml.xpath('./w:body/w:p', namespaces=NS)
    blocks=[]
    for i,p in enumerate(paragraphs):
        blocks.append(f'P{i}: {content(p)}')
        for j,m in enumerate(p.xpath('.//m:oMath',namespaces=NS)):
            (OUT/f'math-P{i}-{j}.xml').write_bytes(etree.tostring(m, pretty_print=True))
    for i,table in enumerate(xml.xpath('./w:body/w:tbl',namespaces=NS),1):
        blocks.append(f'\nTABLE {i}')
        for row in table.xpath('./w:tr',namespaces=NS):
            blocks.append(' | '.join(content(c) for c in row.xpath('./w:tc',namespaces=NS)))
    for name in z.namelist():
        if name.startswith('word/media/'):
            (OUT/Path(name).name).write_bytes(z.read(name))
    stats={'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
           'paragraph_count':len(paragraphs),'math_count':len(xml.xpath('.//m:oMath',namespaces=NS)),
           'tables':len(xml.xpath('./w:body/w:tbl',namespaces=NS)),
           'drawings':len(xml.xpath('.//w:drawing',namespaces=NS)),
           'sections':[dict(n.attrib) for n in xml.xpath('.//w:sectPr/*',namespaces=NS)]}
    (OUT/'source.xml').write_bytes(etree.tostring(xml,pretty_print=True))
    (OUT/'stats.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'manuscript.txt').write_text('\n'.join(blocks),encoding='utf-8')
old=Document(ROOT/'word_output/2026.10.04v68.docx')
new=Document(SOURCE)
diff='\n'.join(difflib.unified_diff([content(p._p) for p in old.paragraphs],
                                  [content(p._p) for p in new.paragraphs],fromfile='v68',tofile='v69',lineterm=''))
(OUT/'changes_from_v68.txt').write_text(diff,encoding='utf-8')
print(json.dumps(stats,ensure_ascii=False))
