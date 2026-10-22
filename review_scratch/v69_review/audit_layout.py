from pathlib import Path
import json, zipfile, io, re
from docx import Document
from PIL import Image
import numpy as np
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
a=json.loads((OUT/'numeric_audit.json').read_text(encoding='utf-8'))
print('Non-single spacing:')
for row in a['paragraph_format']:
    if row['line_spacing'] not in ('None','1.0'):
        print({k:v for k,v in row.items() if k!='runs'})
d=Document(ROOT/'word_output/2026.10.04v69.docx')
print('Style spacing:',[(n,str(d.styles[n].paragraph_format.line_spacing)) for n in ['Normal','Heading 1','Heading 2']])
with zipfile.ZipFile(ROOT/'word_output/2026.10.04v69.docx') as z:
    print('Images:',[(n,Image.open(io.BytesIO(z.read(n))).size) for n in z.namelist() if n.startswith('word/media/')])
print('Weights:')
for ds in ['abilene','geant']:
    means=np.stack([np.loadtxt(ROOT/'abilene_forecasting/results'/ds/f'dlinear_scale_seed{s}'/'router_weights.csv',delimiter=',').mean(axis=0) for s in range(42,50)])
    print(ds,means.mean(axis=0))
print('Han count:',sum(len(re.findall(r'[\u4e00-\u9fff]',p.text)) for p in d.paragraphs))
print('Page breaks:',[(i,p.text[:25]) for i,p in enumerate(d.paragraphs) if p.paragraph_format.page_break_before or 'w:type="page"' in p._p.xml])
print('Figure sizes cm:',[(s.width/360000,s.height/360000) for s in d.inline_shapes])
print('Source unchanged:',a['source_unchanged'])
