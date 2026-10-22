from pathlib import Path
import json, statistics, math, re, zipfile, hashlib
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
RESULTS=ROOT/'abilene_forecasting/results'
def metrics(ds,model,seed):
    prefix={'dlinear_scale_global':'global_weight_ablation',
            'dlinear_scale_constant_router':'constant_router_ablation'}.get(model,'')
    p=RESULTS/prefix/ds/f'{model}_seed{seed}'/'metrics.json'
    r=json.loads(p.read_text())
    assert (r['input_len'],r['pred_len'])==(96,24)
    return r
report={'tables':[],'paired':[]}
doc=Document(ROOT/'word_output/2026.10.04v69.docx')
for ti,seeds in [(1,range(42,45)),(2,range(42,50))]:
    for row in doc.tables[ti].rows[1:]:
        cells=[c.text for c in row.cells]
        ds={'Abilene':'abilene','GÉANT':'geant'}[cells[0]]
        model={'本文方法':'dlinear_scale','DLinear':'dlinear','LightTS':'lightts',
               '固定等权多尺度':'dlinear_scale_static','全局可学习权重':'dlinear_scale_global',
               '恒定输入路由':'dlinear_scale_constant_router'}[cells[1]]
        runs=[metrics(ds,model,s) for s in seeds]
        expected=[str(runs[0]['parameters'])]
        for metric in ['mse','mae']:
            values=[r['test_normalized'][metric] for r in runs]
            expected.append(f'{statistics.mean(values):.3f}±{statistics.stdev(values):.3f}')
        report['tables'].append({'table':ti+1,'dataset':ds,'model':model,'shown':cells[2:],
                                 'recomputed':expected,'match':cells[2:]==expected})
for n,seeds,models in [(3,range(42,45),['dlinear','lightts']),
                        (8,range(42,50),['dlinear','dlinear_scale_static','dlinear_scale_global','dlinear_scale_constant_router'])]:
    for ds in ['abilene','geant']:
        for model in models:
            for metric in ['mse','mae','rmse']:
                values=[100*(1-metrics(ds,'dlinear_scale',s)['test_normalized'][metric]/
                                   metrics(ds,model,s)['test_normalized'][metric]) for s in seeds]
                mean=statistics.mean(values)
                row={'n':n,'dataset':ds,'baseline':model,'metric':metric,'mean':mean,
                     'wins':sum(v>0 for v in values),'seed_values':values}
                if n==8:
                    half=2.3646242510102993*statistics.stdev(values)/math.sqrt(n)
                    row['ci95']=[mean-half,mean+half]
                report['paired'].append(row)
texts=[p.text for p in doc.paragraphs]
report['abstract_character_counts']={'all_nonspace':len(re.sub(r'\s','',texts[3].removeprefix('摘要：'))),
                                    'han':len(re.findall(r'[\u4e00-\u9fff]',texts[3]))}
report['conclusion_character_counts']={'all_nonspace':len(re.sub(r'\s','',texts[115])),
                                      'han':len(re.findall(r'[\u4e00-\u9fff]',texts[115]))}
report['reference_first_appearance']=[]
for text in texts[:116]:
    for number in re.findall(r'\[(\d+)\]',text):
        if int(number) not in report['reference_first_appearance']:
            report['reference_first_appearance'].append(int(number))
report['paragraph_format']=[{'i':i,'style':p.style.name,'text_prefix':p.text[:30],
   'alignment':str(p.alignment),'line_spacing':str(p.paragraph_format.line_spacing),
   'runs': [{'text':r.text[:30],'size':r.font.size.pt if r.font.size else None,'font':r.font.name,
             'superscript':r.font.superscript} for r in p.runs if r.text]} for i,p in enumerate(doc.paragraphs)]
with zipfile.ZipFile(ROOT/'word_output/2026.10.04v68.docx') as old, zipfile.ZipFile(ROOT/'word_output/2026.10.04v69.docx') as new:
    report['media_changed']=[n for n in new.namelist() if n.startswith('word/media/') and
                            (n not in old.namelist() or new.read(n)!=old.read(n))]
report['source_unchanged']=hashlib.sha256((ROOT/'word_output/2026.10.04v69.docx').read_bytes()).hexdigest()==json.loads((OUT/'stats.json').read_text())['source_sha256']
(OUT/'numeric_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('Tables all match:',all(r['match'] for r in report['tables']))
print('Counts:',report['abstract_character_counts'],report['conclusion_character_counts'])
print('Reference order:',report['reference_first_appearance'])
print('Changed media:',report['media_changed'])
for row in report['paired']:
    print(json.dumps({k:v for k,v in row.items() if k!='seed_values'},ensure_ascii=False))
