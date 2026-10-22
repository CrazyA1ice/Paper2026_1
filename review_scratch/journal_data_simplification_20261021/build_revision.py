from pathlib import Path
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import hashlib,json,re,shutil,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).resolve().parents[2]
QA=Path(__file__).resolve().parent
OLD=ROOT/'paper_template/长春工程学院学报(自然科学版)/长预测核验与修订_20261021'
BASE=OLD/'自适应多尺度加权的多变量网络流量预测_扩展预测修订版.docx'
SI_BASE=OLD/'长预测时长扩展实验_完整结果与核验.docx'
OUT=OLD.parent/'数据处理精简修订_20261021'
OUT.mkdir(exist_ok=True)
MAIN=OUT/'自适应多尺度加权的多变量网络流量预测_数据分析精简版.docx'
SI=OUT/'补充材料_完整结果与附加分析.docx'
REPORT=OUT/'样文对照与数据处理改进说明.md'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(BASE)=='76b2b907abd388a49503c7ad21e8a4b009848b033e65d59e5660f2033168c744'
assert sha(SI_BASE)=='57177770317afe921c081a98e4b51fd18158c9e9795d318b0909e0068a095a46'
assert json.loads((QA/'source_inspection.json').read_text(encoding='utf8'))['processed_data'][0]['all_finite']
shutil.copy2(BASE,MAIN);shutil.copy2(SI_BASE,SI)
source=Document(BASE);doc=Document(MAIN);si=Document(SI)

def find(d,prefix):
    values=[p for p in d.paragraphs if p.text.startswith(prefix)]
    assert len(values)==1,(prefix,len(values))
    return values[0]

def set_text(p,value,size=10.5,bold=False):
    p.clear();r=p.add_run(value)
    r.font.name='Times New Roman';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if bold else '宋体')
    return p

def remove(p):p._p.getparent().remove(p._p)

def clone_to_si(element):
    cloned=deepcopy(element)
    for blip in cloned.xpath('.//a:blip'):
        old_id=blip.get(qn('r:embed'))
        new_id,_=si.part.get_or_add_image(BytesIO(source.part.related_parts[old_id].blob))
        blip.set(qn('r:embed'),new_id)
    si._element.body.insert(len(si._element.body)-1,cloned)
    return cloned

def si_text(value,size=10.5,bold=False,style=None):
    return set_text(si.add_paragraph(style=style),value,size,bold)

def si_heading(value,new_page=False):
    p=si_text(value,12,True,'Heading 1');p.paragraph_format.page_break_before=new_page
    return p

def si_caption(value):
    p=si_text(value,9,True);p.paragraph_format.first_line_indent=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5)
    return p

def figure_paragraph(d,shape_index):
    return d.inline_shapes[shape_index]._inline.getparent().getparent().getparent()

def result_counts(d):
    inside=False;by_heading={};current=None
    for p in d.paragraphs:
        if p.text.startswith('2 实验'):inside=True
        if p.text.startswith('3 结语'):break
        if inside:
            if re.match(r'^2\.\d+ ',p.text):current=p.text;by_heading[current]=0
            elif current:by_heading[current]+=len(p.text)
    return by_heading

# Preserve relocated evidence before removing its main-text locations.
si_heading('S5 逐种子误差与参数比例',True)
si_text('图S1保留三个共同种子下相对LightTS的MSE、MAE和RMSE降幅，以及模型参数比例。RMSE为MSE的平方根，用作同一误差的另一种表达；其关系为')
rmse=clone_to_si(find(source,'\t\t(16)')._p)
for node in rmse.xpath('.//w:t'):
    if node.text:node.text=node.text.replace('(16)','(S1)')
clone_to_si(figure_paragraph(source,1))
si_caption('图S1 相对LightTS的逐种子误差降低率与参数比例')
si_text('色块为共同种子42～44的误差降低率，最后一列为三个种子的均值；右侧以LightTS参数量为100%表示模型参数比例。本文方法在Abilene与GÉANT上的参数量分别为对应LightTS配置的6.41%和2.17%。')

si_heading('S6 尺度权重分布')
si_text('图S2汇总八个种子在全部24 h任务测试窗口上的尺度权重。小提琴表示窗口权重的汇总分布，彩色点表示各种子的窗口平均权重；黑色菱形和误差线分别为这些均值的平均值与样本标准差，短横线为汇总分布中位数，虚线为等权1/3。')
clone_to_si(figure_paragraph(source,2))
si_caption('图S2 尺度权重分布与种子均值')
clone_to_si(find(source,'Abilene在尺度1')._p)

si_heading('S7 周期复制对照',True)
si_text(find(source,'为评价模型相对周期').text.replace('表4','表S6'))
si_caption('表S6 共同测试窗口上的周期复制预测误差')
clone_to_si(source.tables[3]._tbl)
period_text=find(source,'周期复制基线').text.replace('表4','表S6')
si_text(period_text)

si_heading('S8 计算耗时测量设置')
clone_to_si(find(source,'采用24 h任务的八个检查点')._p)
si_text('测量对象为24 h任务模型，不包含数据读取、预处理和设备传输时间。对同一检查点，先对所有重复批次的耗时取平均，再对八个检查点汇总均值和样本标准差；主稿报告对应结果。')
si_text('逐种子记录、完整时长结果及原始预测可按数据集、预测时长、模型和种子标识逐一对应。',9.5)

# Clarify an implemented processing policy without changing any observations.
p=find(doc,'基础实验取L=96')
set_text(p,'基础实验取L=96、H=24，扩展实验保持L=96并调整H。数据读取时，将缺失值和无穷值记为0，并将负值截为0。考虑到流量非负且变量量级不同，先进行自然对数变换，再按变量标准化：')
set_text(find(doc,'由于同一尺度'),'同一尺度的线性映射在变量间共享，因此参数量不随OD变量数增加。模型依次通过多尺度预测与自适应融合，输出标准化空间中的多步预测结果。')

set_text(find(doc,'2.1 数据集与实验设置'),'2.1 数据与实验设置',11,True)
set_text(find(doc,'外部基线比较'),'基线比较采用42～44三个共同随机种子，其余学习模型比较采用42～49八个共同种子。结果以均值±样本标准差表示；误差降幅在共同种子下分别计算相对降幅，再取平均。图中柱高和误差线分别表示均值和样本标准差。')
remove(find(doc,'2.2 评价指标'))
set_text(find(doc,'采用MSE、平均绝对误差'),'采用MSE和平均绝对误差（mean absolute error，MAE）评价预测结果，两项指标均在对数变换与逐变量标准化后的空间计算。测试集MSE按式（14）计算，N取相应测试窗口数；MAE为')
remove(find(doc,'\t\t(16)'))
set_text(find(doc,'式中，Nₜₑ为测试窗口数'),'式中，Nₜₑ为测试窗口数。MSE对较大偏差更敏感，MAE反映平均绝对偏差；两项指标越小，预测误差越低。')

set_text(find(doc,'2.3 基线模型比较'),'2.2 预测精度比较',11,True)
set_text(find(doc,'选取单尺度DLinear'),'选取单尺度DLinear和LightTS作为基线，在相同测试窗口上进行比较，结果见表2。')
set_text(find(doc,'表2 三个共同随机种子'),'表2 基线模型的24 h预测误差（3次重复）',9,True)
set_text(find(doc,'按共同种子计算'),'两个数据集上，本文方法的平均MSE和MAE均低于所列基线。相对LightTS，MSE的逐种子平均降幅在Abilene和GÉANT上分别为29.80%和46.89%，参数量分别为对应LightTS配置的6.41%和2.17%。逐种子结果与参数对照见补充材料S5。')
old_picture=figure_paragraph(doc,1)
old_picture.getparent().remove(old_picture)
remove(find(doc,'图2 相对LightTS'))
remove(find(doc,'图2中，色块'))

set_text(find(doc,'2.4 融合方式消融分析'),'2.3 融合方式消融',11,True)
p=find(doc,'进一步保持三个独立DLinear专家')
set_text(p,'保持三个独立DLinear专家及尺度集合不变，比较不同融合方式。固定等权模型将权重设为1/3；全局可学习权重模型（图中简称“全局权重”）以三个可训练分数生成所有窗口共用的权重；恒定输入路由模型（图中简称“恒定路由”）保留两层路由网络，并输入全零特征。三个对照的尺度专家均处理真实历史序列。')
set_text(find(doc,'表3 八个共同随机种子'),'表3 融合方式的24 h预测误差（8次重复）',9,True)
set_text(find(doc,'在八个共同种子的汇总结果中'),'表3中，本文方法在两套数据上的平均MSE和MAE均低于所列对照。相对DLinear，MSE的逐种子平均降幅分别为2.77%和2.35%；相对全局可学习权重，分别为1.80%和4.31%。结合固定等权与恒定输入路由的对照结果，依据当前窗口统计量生成尺度权重，在该24 h任务中取得了较低的平均预测误差。尺度权重分布见补充材料S6。')
remove(find(doc,'与全局可学习权重相比'))
remove(find(doc,'为观察多尺度融合中的权重分配'))
weight_picture=figure_paragraph(doc,1)
weight_picture.getparent().remove(weight_picture)
remove(find(doc,'图3 尺度权重分布'))
remove(find(doc,'Abilene在尺度1'))

start=find(doc,'2.5 周期基线比较')._p
stop=find(doc,'2.6 高需求预测点分析')._p
element=start
while element is not stop:
    next_element=element.getnext();element.getparent().remove(element);element=next_element

set_text(find(doc,'2.6 高需求预测点分析'),'2.4 高需求预测表现',11,True)
set_text(find(doc,'为考察较高通信需求'),'以训练集各OD变量的第90百分位为高需求阈值，在测试集中超过对应阈值的目标点上计算MAE与平均低估误差。后者将低估时的差额计入，其他情况记为0，再对全部高需求点取平均，结果见图2。')
set_text(find(doc,'图4 高需求预测点'),'图2 高需求预测点的MAE与平均低估误差（8次重复）',9,True)
set_text(find(doc,'图4中，柱高与误差线'),'图2中，本文方法在Abilene上的MAE与平均低估误差分别为0.147和0.137，在GÉANT上分别为0.632和0.577，两项均值均低于所列对照，体现了该24 h任务下较低的高需求预测误差。')

set_text(find(doc,'2.7 GÉANT一周预测对照'),'2.5 一周预测表现',11,True)
set_text(find(doc,'进一步保持小时采样'),'保持小时采样、96 h输入及其余实验设置不变，将预测时长扩展至24、48、96和168 h，两套数据的完整结果见补充材料S1～S4。图3展示GÉANT的168 h任务，五种模型均在相同的307个测试窗口上评价，覆盖全部预测小时与524个OD变量。')
set_text(find(doc,'图5 GÉANT一周预测任务'),'图3 GÉANT一周预测任务的MSE与MAE（8次重复）',9,True)
set_text(find(doc,'图5中，柱高与误差线'),'本文方法的MSE和MAE分别为0.841±0.088和0.536±0.028，低于图中四种对照。相对DLinear，两项指标的逐种子平均降幅分别为23.16%和11.16%，说明该方法在GÉANT的一周预测任务中取得了较低的平均预测误差。')

set_text(find(doc,'2.8 24 h任务的模型前向计算耗时'),'2.6 模型计算耗时',11,True)
set_text(find(doc,'采用24 h任务的八个检查点'),'在GTX 1060（6 GB）和i7-8750H单线程环境下，使用24 h任务的八个检查点测量模型前向耗时，批量大小为16。Abilene与GÉANT上的GPU每批耗时分别为3.28±0.20 ms和5.26±2.79 ms，CPU分别为40.65±0.19 ms和131.35±1.25 ms。详细计时设置见补充材料S8。')
remove(find(doc,'Abilene与GÉANT上的GPU每批前向耗时'))

# Retained figures and table values remain byte-identical.
doc.save(MAIN);si.save(SI)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with ZipFile(BASE) as old,ZipFile(MAIN) as new:
    oldxml=etree.fromstring(old.read('word/document.xml'));newxml=etree.fromstring(new.read('word/document.xml'))
    oldmath=[etree.tostring(x,method='c14n') for x in oldxml.xpath('.//m:oMath',namespaces=ns)]
    newmath=[etree.tostring(x,method='c14n') for x in newxml.xpath('.//m:oMath',namespaces=ns)]
    assert newmath==oldmath[:-1] and len(newmath)==16
    changed=[n for n in old.namelist() if old.read(n)!=new.read(n)]
    assert changed==['word/document.xml'],changed
    tables=lambda x:[etree.tostring(t,method='c14n') for t in x.xpath('.//w:tbl',namespaces=ns)]
    assert tables(newxml)==tables(oldxml)[:3]
with ZipFile(SI_BASE) as old,ZipFile(SI) as new:
    before=etree.fromstring(old.read('word/document.xml'));after=etree.fromstring(new.read('word/document.xml'))
    a=before.find(qn('w:body'));b=after.find(qn('w:body'))
    assert [etree.tostring(x,method='c14n') for x in list(a)[:-1]]==[etree.tostring(x,method='c14n') for x in list(b)[:len(a)-1]]
assert len(doc.tables)==3 and len(doc.inline_shapes)==3
assert len(si.tables)==6 and len(si.inline_shapes)==2
refs=lambda d:[p.text for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)]
assert refs(doc)==refs(source)
assert find(doc,'摘要：').text==find(source,'摘要：').text
before_counts=result_counts(source);after_counts=result_counts(doc)
manifest={'base':str(BASE),'base_sha256':sha(BASE),'si_base':str(SI_BASE),'si_base_sha256':sha(SI_BASE),'main':str(MAIN),'supplement':str(SI),'main_sha256':sha(MAIN),'supplement_sha256':sha(SI),'report':str(REPORT),'main_figures':3,'main_tables':3,'main_math_objects':16,'unchanged_references':11,'main_changed_parts':changed,'before_results_characters':sum(before_counts.values()),'after_results_characters':sum(after_counts.values()),'before_subsections':before_counts,'after_subsections':after_counts}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False))
