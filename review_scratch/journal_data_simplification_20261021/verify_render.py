from pathlib import Path
from zipfile import ZipFile
import hashlib,importlib.util,json,os,re,sys
import pdfplumber
from docx import Document
from lxml import etree
from PIL import Image,ImageChops

sys.stdout.reconfigure(encoding='utf8')
QA=Path(__file__).resolve().parent
ROOT=QA.parents[1]
m=json.loads((QA/'manifest.json').read_text(encoding='utf8'))
stats=json.loads((QA/'word_statistics.json').read_text(encoding='utf-8-sig'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(Path(m['base']))==m['base_sha256']
assert sha(Path(m['si_base']))==m['si_base_sha256']
main=Document(m['main']);si=Document(m['supplement']);old=Document(m['base'])
assert [p.text for p in main.paragraphs if re.match(r'^图\d+ ',p.text)]==['图1 自适应多尺度网络流量预测框架','图2 高需求预测点的MAE与平均低估误差（8次重复）','图3 GÉANT一周预测任务的MSE与MAE（8次重复）']
text='\n'.join(p.text for p in main.paragraphs)
assert 'RMSE' not in text and '图4' not in text and '图5' not in text and '表4' not in text
assert [p.text for p in main.paragraphs if re.match(r'^2\.\d+ ',p.text)]==['2.1 数据与实验设置','2.2 预测精度比较','2.3 融合方式消融','2.4 高需求预测表现','2.5 一周预测表现','2.6 模型计算耗时']
for new_index,old_index in enumerate([0,3,4]):
    def blob(d,i):
        rid=d.inline_shapes[i]._inline.graphic.graphicData.pic.blipFill.blip.embed
        return d.part.related_parts[rid].blob
    assert blob(main,new_index)==blob(old,old_index)
for record in json.loads((QA/'source_inspection.json').read_text(encoding='utf8'))['processed_data']:
    assert sha(ROOT/f'abilene_forecasting/data/processed/{record["dataset"]}_1hour.npz')==record['sha256']

renderer_path=Path(r'C:\Users\Quant\.codex\plugins\cache\openai-primary-runtime\documents\26.1020.11127\skills\documents\render_docx.py')
spec=importlib.util.spec_from_file_location('simplified_paper_renderer',renderer_path)
renderer=importlib.util.module_from_spec(spec);sys.modules[spec.name]=renderer;spec.loader.exec_module(renderer)
poppler=Path(r'C:\Users\Quant\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler')
os.environ['PATH']=str(next(poppler.rglob('pdftoppm.exe')).parent)+os.pathsep+os.environ['PATH']
results=[]
for item in stats:
    kind=item['kind'];target=Path(m[kind]);pdf=Path(item['pdf'])
    assert sha(target)==m[kind+'_sha256']
    renderer.convert_to_pdf=lambda *args,**kwargs:(str(pdf),'Native Microsoft Word read-only export')
    pages=renderer.rasterize(str(target),str(pdf.parent),150,False,False)
    assert len(pages)==item['pages']
    texts=[];detail=[]
    with pdfplumber.open(pdf) as document:
        for i,page in enumerate(document.pages,1):
            txt=page.extract_text() or '';texts.append(txt)
            bad=[c for c in page.chars if c['text'].strip() and (c['x0']<-1 or c['top']<-1 or c['x1']>page.width+1 or c['bottom']>page.height+1)]
            assert not bad,(kind,i)
            detail.append({'page':i,'first':txt[:120],'last':txt[-180:]})
    extracted='\n'.join(texts)
    (QA/f'{kind}_rendered_text.txt').write_text(extracted,encoding='utf8')
    for v in ['23.16%','11.16%','0.841','0.536'] if kind=='main' else ['-12.79','0.326210','0.840610','表S6','图S1','图S2']:
        assert v in re.sub(r'\s+','',extracted),v
    if kind=='main':
        assert '作者简介' in extracted and item['words']<=7500
        for i in [1,2,3]:assert re.search(r'图\s*'+str(i),extracted)
    else:
        assert '作者简介' not in extracted
    results.append({**item,'out_of_page_text':0,'page_details':detail,'visual_review_pending':True})

unchanged=[]
previous=ROOT/'review_scratch/long_horizon_followup_20261021'
for n in [1,2,3,4]:
    with Image.open(previous/f'render_supplement/page-{n}.png') as a,Image.open(QA/f'render_supplement/page-{n}.png') as b:
        same=a.size==b.size and ImageChops.difference(a.convert('RGB'),b.convert('RGB')).getbbox() is None
        unchanged.append({'page':n,'pixel_identical':same})
(QA/'render_verification.json').write_text(json.dumps({'documents':results,'supplement_original_pages':unchanged},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'documents':[{k:v for k,v in r.items() if k!='page_details'} for r in results],'supplement_original_pages':unchanged},ensure_ascii=False))
