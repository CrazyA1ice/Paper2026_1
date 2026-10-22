from pathlib import Path
from zipfile import ZipFile
import json,hashlib,sys,importlib.util,os
from lxml import etree
from docx import Document

sys.stdout.reconfigure(encoding='utf8')
QA=Path(__file__).resolve().parent
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
records=json.loads((QA/'reference_manifest.json').read_text(encoding='utf-8-sig'))
renderer_path=Path(r'C:\Users\Quant\.codex\plugins\cache\openai-primary-runtime\documents\26.1020.11127\skills\documents\render_docx.py')
spec=importlib.util.spec_from_file_location('template_renderer',renderer_path)
renderer=importlib.util.module_from_spec(spec);sys.modules[spec.name]=renderer;spec.loader.exec_module(renderer)
poppler=Path(r'C:\Users\Quant\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler')
os.environ['PATH']=str(next(poppler.rglob('pdftoppm.exe')).parent)+os.pathsep+os.environ['PATH']
all_records=[]
for item in records:
    p=Path(item['converted']);d=Document(p);paras=[]
    for i,para in enumerate(d.paragraphs):
        runs=[]
        for run in para.runs:
            if not run.text:continue
            runs.append({'text':run.text,'color':str(run.font.color.rgb),'font':run.font.name,'size':run.font.size.pt if run.font.size else None,'bold':run.bold,'italic':run.italic,'xml':etree.tostring(run._r,encoding='unicode')})
        paras.append({'index':i,'text':para.text,'style':para.style.name,'runs':runs,'pPr':etree.tostring(para._p.pPr,encoding='unicode') if para._p.pPr is not None else None})
    with ZipFile(p) as z:
        inventory=[{'part':n,'bytes':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()]
        xmltext={n:' | '.join(etree.fromstring(z.read(n)).xpath('.//w:t/text()|.//a:t/text()',namespaces=NS)) for n in z.namelist() if n.startswith('word/') and n.endswith('.xml') and ('document' in n or 'footnote' in n or 'endnote' in n or 'header' in n or 'footer' in n or 'comment' in n)}
        sections=[etree.tostring(x,encoding='unicode') for x in etree.fromstring(z.read('word/document.xml')).xpath('.//w:sectPr',namespaces=NS)]
    audit={**item,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'original_sha256':hashlib.sha256(Path(item['source']).read_bytes()).hexdigest(),'sections':sections,'paragraphs':paras,'tables':[[[c.text for c in r.cells] for r in t.rows] for t in d.tables],'all_text_parts':xmltext,'package_parts':inventory}
    (QA/(item['name']+'_inspection.json')).write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
    readable=[]
    for para in paras:
        readable.append(f"[{para['index']}] ({para['style']}) "+''.join(('[RED]'+r['text']+'[/RED]') if r['color'] in ['FF0000','C00000'] else r['text'] for r in para['runs']))
    (QA/(item['name']+'_text.txt')).write_text('\n'.join(readable),encoding='utf8')
    print(item['name']+'\n'+'\n'.join(readable))
    renderer.convert_to_pdf=lambda *args,**kwargs:(item['pdf'],'Native Word read-only export')
    images=renderer.rasterize(str(p),str(QA/(item['name']+'_render')),150,False,False)
    assert len(images)==item['pages']
    all_records.append({'name':item['name'],'pages':len(images),'sha256':audit['sha256'],'original_sha256':audit['original_sha256']})
(QA/'template_evidence.json').write_text(json.dumps(all_records,ensure_ascii=False,indent=2),encoding='utf8')
