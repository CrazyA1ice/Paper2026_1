from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
import hashlib,json,re,sys
from lxml import etree as E
from PIL import Image
from docx import Document

sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).resolve().parents[2];QA=Path(__file__).resolve().parent
PARENT=ROOT/'paper_template/长春工程学院学报(自然科学版)'
SOURCE=PARENT/'主稿独立投稿修订_20261021/自适应多尺度加权的多变量网络流量预测_独立投稿主稿.docx'
REFERENCE=QA/'论文模板.docx'
OUT=PARENT/'模板规范修订_20261021';ASSETS=OUT/'图表源文件'
TARGET=OUT/'自适应多尺度加权的多变量网络流量预测_期刊模板稿.docx'
old_manifest=json.loads((ROOT/'review_scratch/standalone_submission_20261021/manifest.json').read_text(encoding='utf8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(SOURCE)==old_manifest['output_sha256']
ref_evidence=json.loads((QA/'template_evidence.json').read_text(encoding='utf8'))[0]
assert sha(REFERENCE)==ref_evidence['sha256']
ref_manifest=json.loads((QA/'reference_manifest.json').read_text(encoding='utf-8-sig'))[0]
assert sha(Path(ref_manifest['source']))==ref_evidence['original_sha256']

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
V='urn:schemas-microsoft-com:vml'
O='urn:schemas-microsoft-com:office:office'
WP='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
PIC='http://schemas.openxmlformats.org/drawingml/2006/picture'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
NS={'w':W,'m':M,'r':R,'v':V,'wp':WP,'a':A}
def tag(name):
    prefix,local=name.split(':');return '{'+{'w':W,'m':M,'r':R,'v':V,'o':O,'wp':WP,'a':A,'pic':PIC}[prefix]+'}'+local
def el(name,attrs=None,text=None):
    x=E.Element(tag(name))
    for k,v in (attrs or {}).items():x.set(tag(k),str(v))
    if text is not None:x.text=text
    return x
def xml_bytes(root):return E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
def visible(p):return ''.join(p.xpath('.//w:t/text()',namespaces=NS))
with ZipFile(REFERENCE) as z:parts={n:z.read(n) for n in z.namelist()}
with ZipFile(SOURCE) as z:source_xml=E.fromstring(z.read('word/document.xml'))
root=E.fromstring(parts['word/document.xml']);body=root.find(tag('w:body'))
sections=[deepcopy(s) for s in root.xpath('.//w:sectPr',namespaces=NS)]
assert len(sections)==3
source_body=source_xml.find(tag('w:body'))
source_ps=source_body.findall(tag('w:p'))
def source_find(prefix):
    hits=[p for p in source_ps if visible(p).startswith(prefix)];assert len(hits)==1,(prefix,len(hits));return hits[0]
title=visible(source_ps[0]);assert len(title)<=20
cn_abstract='针对小时级网络流量预测，提出自适应多尺度加权方法。采用三个分解线性模型（DLinear）处理多尺度输入，以四维窗口统计量生成融合权重。在Abilene和GÉANT上，24 h预测的八种子平均均方误差较DLinear分别降低2.77%和2.35%；在GÉANT的168 h任务中，均方误差和平均绝对误差分别降低23.16%和11.16%。24 h模型含8339个参数，降低了所设任务的平均预测误差。'
assert len(cn_abstract)==198
en_abstract='An adaptive multiscale weighting method is proposed for hourly network traffic forecasting. Three decomposition linear models (DLinear) process multiscale inputs, and four window statistics determine their fusion weights. In 24-hour forecasting, the mean seed-wise reductions in mean squared error relative to DLinear are 2.77% on Abilene and 2.35% on GÉANT across eight random seeds. For the 168-hour GÉANT task, the corresponding reductions in mean squared error and mean absolute error are 23.16% and 11.16%. The 24-hour model has 8,339 parameters and reduces the average prediction error in the evaluated tasks.'
biography='作者简介：柯亚文（1998—），男（汉族），安徽池州人，助教，硕士研究生学历，硕士学位，主要研究方向：计算机软件。联系电话：12345678900；电子邮箱：keyawen@afc.edu.cn。'

styles=E.fromstring(parts['word/styles.xml'])
roles={
    'Title':(14,True,'center',0,8),
    'JournalAuthor':(10.5,False,'center',0,4),
    'JournalAffiliation':(10.5,False,'center',0,6),
    'JournalMetadata':(10.5,False,'left',0,2),
    'JournalAbstract':(10.5,False,'both',0,3),
    'JournalBody':(10.5,False,'both',21,0),
    'JournalHeading1':(12,True,'left',0,5),
    'JournalHeading2':(10.5,True,'left',0,3),
    'JournalCaption':(9,True,'center',0,0),
    'JournalReference':(10.5,False,'left',0,0),
    'JournalReferenceHeading':(10.5,True,'left',0,3),
    'JournalEquation':(10.5,False,'left',0,3),
    'JournalFigure':(10.5,False,'center',0,5),
    'JournalTable':(7.5,False,'center',0,0),
    'JournalBlank':(10.5,False,'left',0,0),
}
def rpr(size=10.5,bold=False,italic=False,sup=False):
    r=el('w:rPr');r.append(el('w:rFonts',{'w:ascii':'Times New Roman','w:hAnsi':'Times New Roman','w:eastAsia':'宋体','w:cs':'Times New Roman'}))
    if bold:r.append(el('w:b'))
    if italic:r.append(el('w:i'))
    r.append(el('w:color',{'w:val':'000000'}));r.append(el('w:sz',{'w:val':int(size*2)}));r.append(el('w:szCs',{'w:val':int(size*2)}))
    if sup:r.append(el('w:vertAlign',{'w:val':'superscript'}))
    return r
def ppr(role):
    size,bold,align,indent,after=roles[role]
    pp=el('w:pPr');pp.append(el('w:pStyle',{'w:val':role}));pp.append(el('w:widowControl'))
    pp.append(el('w:jc',{'w:val':align}));pp.append(el('w:spacing',{'w:before':0,'w:after':int(after*20),'w:line':360,'w:lineRule':'exact'}))
    pp.append(el('w:ind',{'w:left':0,'w:right':0,'w:firstLine':int(indent*20)}))
    pp.append(el('w:snapToGrid',{'w:val':0}))
    if role in ['Title','JournalHeading1','JournalHeading2','JournalReferenceHeading']:
        pp.append(el('w:keepNext'));pp.append(el('w:keepLines'))
    if role in ['JournalHeading1','JournalHeading2']:pp.append(el('w:outlineLvl',{'w:val':0 if role.endswith('1') else 1}))
    if role in ['JournalHeading1','JournalHeading2']:pp.find(tag('w:spacing')).set(tag('w:before'),'100' if role.endswith('1') else '60')
    return pp
for name in roles:
    old=styles.xpath('./w:style[@w:styleId="'+name+'"]',namespaces=NS)
    for n in old:styles.remove(n)
    style=el('w:style',{'w:type':'paragraph','w:styleId':name});style.append(el('w:name',{'w:val':name}));style.append(el('w:basedOn',{'w:val':'Normal'}));style.append(el('w:next',{'w:val':'JournalBody'}));style.append(ppr(name));style[-1].remove(style[-1].find(tag('w:pStyle')))
    style.append(rpr(roles[name][0],roles[name][1]));styles.append(style)
parts['word/styles.xml']=xml_bytes(styles)
for child in list(body):body.remove(child)

def run(p,text,size=10.5,bold=False,italic=False,sup=False):
    r=el('w:r');r.append(rpr(size,bold,italic,sup));t=el('w:t',text=text)
    if text.startswith(' ') or text.endswith(' '):t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    r.append(t);p.append(r);return r
var_pattern=r'(?<![A-Za-z])(?:[LCHNqrtcsPvuWRXYTbdzamnSxyij][\u2080-\u209fₙₜₑₛ]*|[μσαℓ][ₙₜₑₛ,]*)(?![A-Za-z])'
def para(text='',role='JournalBody',label=None,citations=True):
    p=el('w:p');p.append(ppr(role));size,bold=roles[role][:2]
    if label and text.startswith(label):run(p,label,size,True);text=text[len(label):]
    pattern=r'\[\d+\]' if citations else r'(?!)'
    for chunk in re.split('('+pattern+')',text):
        if not chunk:continue
        if citations and re.fullmatch(pattern,chunk):run(p,chunk,size,bold,sup=True)
        elif role=='JournalBody':
            cursor=0
            for match in re.finditer(var_pattern,chunk):
                run(p,chunk[cursor:match.start()],size,bold);value=match.group()
                panel_marker=match.start()>0 and chunk[match.start()-1] in '(（' and match.end()<len(chunk) and chunk[match.end()] in ')）'
                run(p,value,size,bold=bold if panel_marker else value[0] in 'WXRTYzbaxy' or bold,italic=not panel_marker);cursor=match.end()
            if cursor<len(chunk):run(p,chunk[cursor:],size,bold)
        else:run(p,chunk,size,bold)
    body.append(p);return p

para(title,'Title',citations=False)
para('收稿日期：','JournalMetadata',citations=False)
para(biography,'JournalMetadata',citations=False)
para('柯亚文','JournalAuthor',citations=False)
para(visible(source_ps[2]),'JournalAffiliation',citations=False)
para('摘要：'+cn_abstract,'JournalAbstract',label='摘要：',citations=False)
para(visible(source_ps[4]),'JournalAbstract',label='关键词：',citations=False)
p=para('','JournalMetadata',citations=False)
for text,bold in [('中图分类号：',True),('TP393    ',False),('文献标志码：',True),('A    ',False),('文章编号：',True)]:run(p,text,bold=bold)
p=para('doi：','JournalMetadata',label='doi：',citations=False);p.find(tag('w:pPr')).append(sections[0])

rels=E.fromstring(parts['word/_rels/document.xml.rels'])
existing_ids={r.get('Id') for r in rels};next_id=max(int(s[3:]) for s in existing_ids if s.startswith('rId') and s[3:].isdigit())+1
fig_manifest=json.loads((QA/'figure_manifest.json').read_text(encoding='utf8'))
figure_records=[]
def picture_p(number,width_cm=13.8):
    global next_id
    path=ASSETS/f'图{number}.png';rid=f'rId{next_id}';next_id+=1
    with Image.open(path) as im:pixel_width,pixel_height=im.size
    cx=round(width_cm*360000);cy=round(cx*pixel_height/pixel_width)
    part_name=f'word/media/journal_figure{number}.png';parts[part_name]=path.read_bytes()
    rel=E.SubElement(rels,'{'+REL+'}Relationship');rel.set('Id',rid);rel.set('Type',R+'/image');rel.set('Target',part_name[5:])
    p=el('w:p');pp=ppr('JournalFigure');pp.find(tag('w:spacing')).attrib.update({tag('w:before'):'0',tag('w:after'):'0',tag('w:line'):'240',tag('w:lineRule'):'auto'});p.append(pp)
    rr=el('w:r');drawing=el('w:drawing');inline=el('wp:inline');inline.append(el('wp:extent'));inline[-1].set('cx',str(cx));inline[-1].set('cy',str(cy))
    prop=el('wp:docPr');prop.set('id',str(100+number));prop.set('name',f'图{number}');prop.set('descr',fig_manifest['figures'][number-1]['evidence']);inline.append(prop)
    nf=el('wp:cNvGraphicFramePr');locks=el('a:graphicFrameLocks');locks.set('noChangeAspect','1');nf.append(locks);inline.append(nf)
    graphic=el('a:graphic');gd=el('a:graphicData');gd.set('uri',PIC);pic=el('pic:pic');nv=el('pic:nvPicPr');cn=el('pic:cNvPr');cn.set('id','0');cn.set('name',f'图{number}.png');nv.append(cn);nv.append(el('pic:cNvPicPr'));pic.append(nv)
    bf=el('pic:blipFill');bf.append(el('a:blip',{'r:embed':rid}));stretch=el('a:stretch');stretch.append(el('a:fillRect'));bf.append(stretch);pic.append(bf)
    sp=el('pic:spPr');xf=el('a:xfrm');off=el('a:off');off.set('x','0');off.set('y','0');ext=el('a:ext');ext.set('cx',str(cx));ext.set('cy',str(cy));xf.extend([off,ext]);sp.append(xf);geom=el('a:prstGeom');geom.set('prst','rect');geom.append(el('a:avLst'));sp.append(geom);pic.append(sp);gd.append(pic);graphic.append(gd);inline.append(graphic);drawing.append(inline);rr.append(drawing);p.append(rr)
    return p,cy/12700,pixel_width/(width_cm/2.54)
def figure(number,caption):
    ip,height_pt,dpi=picture_p(number)
    outer=para('','JournalFigure');outer.find(tag('w:pPr')).find(tag('w:spacing')).set(tag('w:lineRule'),'auto')
    rr=el('w:r');pict=el('w:pict');shape=el('v:shape');shape.set('id',f'JournalFigureBox{number}');shape.set('type','#_x0000_t202');shape.set('style',f'width:{13.8/2.54*72:.4f}pt;height:{height_pt+22:.4f}pt');shape.set('stroked','f');shape.set('filled','f')
    tb=el('v:textbox');tb.set('inset','0,0,0,0');content=el('w:txbxContent');content.append(ip)
    cp=el('w:p');cp.append(ppr('JournalCaption'));run(cp,caption,9,True);content.append(cp);tb.append(content);shape.append(tb);pict.append(shape);rr.append(pict);outer.append(rr)
    figure_records.append({'figure':number,'caption':caption,'box_height_pt':height_pt+22,'width_cm':13.8,'effective_dpi':dpi,'png_sha256':sha(ASSETS/f'图{number}.png')})

def equation(source_p):
    p=deepcopy(source_p);old=p.find(tag('w:pPr'))
    if old is not None:p.remove(old)
    pp=ppr('JournalEquation');tabs=el('w:tabs');tabs.append(el('w:tab',{'w:val':'center','w:pos':5040}));tabs.append(el('w:tab',{'w:val':'right','w:pos':10080}));pp.append(tabs);pp.find(tag('w:spacing')).set(tag('w:lineRule'),'atLeast');pp.find(tag('w:spacing')).set(tag('w:line'),'420');p.insert(0,pp)
    for rr in p.xpath('.//w:r',namespaces=NS):
        old=rr.find(tag('w:rPr'))
        if old is not None:rr.remove(old)
        rr.insert(0,rpr())
    for mr in p.xpath('.//m:r',namespaces=NS):
        props=mr.find(tag('w:rPr'))
        if props is None:
            props=el('w:rPr');mr.insert(1 if mr.find(tag('m:rPr')) is not None else 0,props)
        fonts=props.find(tag('w:rFonts'))
        if fonts is None:fonts=el('w:rFonts');props.insert(0,fonts)
        for key in ['w:ascii','w:hAnsi','w:cs']:fonts.set(tag(key),'Times New Roman')
        for key in ['w:sz','w:szCs']:
            size=props.find(tag(key))
            if size is None:size=el(key);props.append(size)
            size.set(tag('w:val'),'21')
    if not re.search(r'\(\d+\)',visible(p)):
        pp.find(tag('w:jc')).set(tag('w:val'),'center')
        math=p.xpath('.//m:oMath',namespaces=NS)[0];punct=el('m:r');punct.append(rpr());punct.append(el('m:t',text='.'));math.append(punct)
    # Times New Roman needs explicit upright styling for numerals and operators.
    for mr in list(p.xpath('.//m:r',namespaces=NS)):
        mt=mr.find(tag('m:t'))
        if mt is None or not mt.text:continue
        chunks=re.split(r'([0-9]+(?:\.[0-9]+)?|[=+\-−*/(),.;:×≥∈…]+|\s+)',mt.text)
        parent=mr.getparent();position=parent.index(mr)
        for chunk in filter(None,chunks):
            segment=deepcopy(mr);segment.find(tag('m:t')).text=chunk
            if re.fullmatch(r'[0-9=+\-−*/(),.;:×≥∈…\s]+',chunk):
                props=segment.find(tag('m:rPr'))
                if props is None:props=el('m:rPr');segment.insert(0,props)
                sty=props.find(tag('m:sty'))
                if sty is None:sty=el('m:sty');props.append(sty)
                sty.set(tag('m:val'),'p')
                wrp=segment.find(tag('w:rPr'))
                if wrp is None:wrp=rpr();segment.insert(1,wrp)
                for key in ['w:i','w:b']:
                    flag=wrp.find(tag(key))
                    if flag is None:flag=el(key);wrp.append(flag)
                    flag.set(tag('w:val'),'0')
            parent.insert(position,segment);position+=1
        parent.remove(mr)
    body.append(p)

table_index=0
def table(source_table):
    global table_index
    table_index+=1;t=deepcopy(source_table)
    for p in t.xpath('.//w:p',namespaces=NS):
        old=p.find(tag('w:pPr'))
        if old is not None:p.remove(old)
        pp=ppr('JournalTable');pp.find(tag('w:spacing')).set(tag('w:line'),'240');pp.find(tag('w:spacing')).set(tag('w:lineRule'),'auto');pp.find(tag('w:spacing')).set(tag('w:after'),'40');pp.find(tag('w:spacing')).set(tag('w:before'),'40');p.insert(0,pp)
        for rr in p.findall(tag('w:r')):
            previous=rr.find(tag('w:rPr'));b=previous.find(tag('w:b')) if previous is not None else None
            bold=b is not None and b.get(tag('w:val'),'1') not in ['0','false','off']
            if previous is not None:rr.remove(previous)
            rr.insert(0,rpr(7.5,bold))
    body.append(t)

special={
    '在时间结构建模方面':'在时间结构建模方面，Zhou等[5]通过深度回声状态网络开展多尺度流量预测。分解线性模型（decomposition linear model，DLinear）[6]将趋势项和余项分别映射至预测区间，TimeMixer[7]采用可分解多尺度混合，而轻量时序模型LightTS则以多层感知机（multilayer perceptron，MLP）建模时间依赖，并作为公开时序预测基准中的对照模型[8]。随着不同时间分辨率的信息被引入模型，各尺度预测如何根据当前历史窗口进行融合，成为进一步改善预测误差的一个切入点。',
    '[7] WANG':'[7] WANG Shiyu, WU Haixu, SHI Xiaoming, et al. TimeMixer: decomposable multiscale mixing for time series forecasting[C/OL]//The Twelfth International Conference on Learning Representations. 2024[2026-10-21]. https://openreview.net/forum?id=7oLshfEIC2.',
    '[8] ZHANG':'[8] WU Haixu, HU Tengge, LIU Yong, et al. TimesNet: temporal 2D-variation modeling for general time series analysis[C/OL]//The Eleventh International Conference on Learning Representations. 2023[2026-10-21]. https://openreview.net/forum?id=ju_Uqw384Oq.',
}
start=source_body.index(source_find('0 引言'))
children=list(source_body)[start:]
figure_num=0;skip_caption=False;in_refs=False
for i,child in enumerate(children):
    if child.tag==tag('w:sectPr'):continue
    if child.tag==tag('w:tbl'):table(child);continue
    assert child.tag==tag('w:p'),child.tag
    value=visible(child)
    if child.xpath('.//w:drawing',namespaces=NS):
        figure_num+=1;next_p=children[i+1];caption=visible(next_p);assert caption.startswith(f'图{figure_num} ')
        figure(figure_num,caption);skip_caption=True;continue
    if skip_caption:skip_caption=False;continue
    if child.xpath('.//m:oMath',namespaces=NS):equation(child);continue
    if value=='参考文献':
        para('','JournalBlank',citations=False);para('参考文献：','JournalReferenceHeading',citations=False);in_refs=True;continue
    for prefix,replacement in special.items():
        if value.startswith(prefix):value=replacement;break
    if not in_refs:
        value=value.replace('大型IP网络','大型网际协议（Internet Protocol，IP）网络')
        value=value.replace('进一步采用Softmax归一化，得到','进一步采用归一化指数函数（Softmax），得到')
        value=value.replace('模型总参数量为','模型总参数量P为')
        if value.startswith('式中，s∈S。'):value=value.replace('式中，s∈S。','式中，s∈S，i为聚合窗口内的时间索引。',1)
        if value.startswith('式中，第一层将四维特征'):value=value.replace('式中，第一层','式中，aₙ为尺度分数向量，W₁、W₂与b₁、b₂分别为两层网络的权重矩阵和偏置。第一层',1)
        if value.startswith('其中，αₙ,ₛ'):value=value.replace('上的融合权重。','上的融合权重，j为尺度集合中的求和索引。',1)
        if value.startswith('基础实验中'):value=value.replace('LightTS采用公开实现','LightTS采用公开基准中的实现[8]')
        if value.startswith('保持三个独立DLinear专家'):value+='各模型的预测误差见表3。'
    if re.match(r'^[0-3](?:\.\d+)? ',value):
        number,heading=value.split(' ',1);para(number+'\u3000'+heading,'JournalHeading2' if '.' in number else 'JournalHeading1',citations=False)
    elif re.match(r'^表\d+ ',value):
        pp=para(value,'JournalCaption',citations=False).find(tag('w:pPr'));pp.append(el('w:keepNext'))
    elif in_refs:para(value,'JournalReference',citations=False)
    else:para(value)
assert figure_num==4 and table_index==3
section_end=para('','JournalBlank',citations=False);section_end.find(tag('w:pPr')).append(sections[1])
en_title=para('Multivariate network traffic forecasting with adaptive multiscale weighting','Title',citations=False)
en_title.find('w:pPr/w:spacing',NS).set(tag('w:before'),'360')
para('KE Yawen','JournalAuthor',citations=False)
para('(School of Information Engineering, Anhui Finance & Trade Vocational College, Hefei Anhui 231200, China)','JournalAffiliation',citations=False)
para('','JournalBlank',citations=False)
para('Abstract: '+en_abstract,'JournalAbstract',label='Abstract:',citations=False)
para(visible(source_ps[9]),'JournalAbstract',label='Key words:',citations=False)
body.append(sections[2])

parts['word/document.xml']=xml_bytes(root);parts['word/_rels/document.xml.rels']=xml_bytes(rels)
ct=E.fromstring(parts['[Content_Types].xml'])
if not ct.xpath('./c:Default[@Extension="png"]',namespaces={'c':CT}):
    n=E.SubElement(ct,'{'+CT+'}Default');n.set('Extension','png');n.set('ContentType','image/png')
parts['[Content_Types].xml']=xml_bytes(ct)
with ZipFile(TARGET,'w',ZIP_DEFLATED) as z:
    for name,data in parts.items():z.writestr(name,data)

final=Document(TARGET);old=Document(SOURCE)
assert [[[c.text for c in row.cells] for row in t.rows] for t in old.tables]==[[[c.text for c in row.cells] for row in t.rows] for t in final.tables]
old_math=source_xml.xpath('.//m:oMath',namespaces=NS);new_math=root.xpath('.//m:oMath',namespaces=NS)
math_value=lambda m:''.join(m.xpath('.//m:t/text()',namespaces=NS))
assert len(old_math)==len(new_math)==16
for a,b in zip(old_math,new_math):assert math_value(b) in [math_value(a),math_value(a)+'.']
assert len(root.xpath('.//v:textbox',namespaces=NS))==4 and len(root.xpath('.//w:drawing',namespaces=NS))==4
assert [visible(p) for p in body.findall(tag('w:p')) if p.find('w:pPr/w:pStyle',NS).get(tag('w:val'))=='JournalHeading1']==['0\u3000引言','1\u3000自适应多尺度网络流量预测方法','2\u3000实验与结果分析','3\u3000结语']
assert re.findall(r'\((\d+)\)', '\n'.join(visible(p) for p in root.xpath('./w:body/w:p[.//m:oMath]',namespaces=NS)))==[str(i) for i in range(1,16)]
all_text='\n'.join(root.xpath('.//w:t/text()',namespaces=NS))
for bad in ['补充材料','附录','CPU','GPU','GTX','显卡','处理器','四号宋体','五号宋体','XXX','Charged active contour','2008BADC4B15','Less is more: fast multivariate']:
    assert bad not in all_text,bad
assert not root.xpath('.//w:color[@w:val="FF0000"]',namespaces=NS)
assert '12345678900' in all_text and '汉族' in all_text and '安徽池州人' in all_text
for index in range(1,12):
    assert len([p for p in final.paragraphs if p.text.startswith(f'[{index}] ')])==1,index
assert sha(SOURCE)==old_manifest['output_sha256'] and sha(REFERENCE)==ref_evidence['sha256']
with ZipFile(REFERENCE) as z:
    changed=[n for n in z.namelist() if z.read(n)!=parts[n]]
assert set(changed)<=set(['word/document.xml','word/styles.xml','word/_rels/document.xml.rels','[Content_Types].xml']),changed
manifest={'output':str(TARGET),'output_sha256':sha(TARGET),'source':str(SOURCE),'source_sha256':sha(SOURCE),'template_original':ref_manifest['source'],'template_original_sha256':ref_evidence['original_sha256'],'template_converted':str(REFERENCE),'template_converted_sha256':sha(REFERENCE),'native_math':16,'numbered_equations':15,'figures':figure_records,'tables':3,'references':11,'chinese_abstract_characters':len(cn_abstract),'english_block_position':'after references','template_package_parts_changed':changed,'editorial_values_intentionally_blank':['receipt date','article number','DOI'],'fund_project':'none, per user','user_author_metadata_completed':True}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
contract=QA/'artifact.md';contract.write_text(re.sub(r'SHA256 5419c896[a-f0-9]*', 'SHA256 '+sha(SOURCE),contract.read_text(encoding='utf8')),encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False))
