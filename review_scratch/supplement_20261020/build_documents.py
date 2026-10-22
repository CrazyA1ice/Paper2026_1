"""Apply the completed supplementary evidence to the previous verified revision."""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
RESULTS = ROOT / "abilene_forecasting/results/supplement_20261020"
OUT = ROOT / "word_output/补充实验_20261020"
BASE = ROOT / "word_output/退稿意见修改_20261020/基于自适应多尺度加权的多变量网络流量预测方法_修改稿.docx"
ORIGINAL = ROOT / "word_output/基于自适应多尺度加权的多变量网络流量预测方法.docx"
MODELS = [("dlinear", "DLinear"), ("dlinear_scale_global", "全局权重"), ("dlinear_scale", "本文方法")]
LABELS = {"abilene": "Abilene", "geant": "GÉANT"}
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def text(element):
    return "".join(t.text or "" for t in element.iter() if t.tag in [qn("w:t"), f"{{{M}}}t"])


def setfont(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体" if bold else "宋体")


def replace_paragraph(p, value, edited, heading=False):
    assert not p._p.xpath(".//m:oMath|.//w:drawing")
    old = p.text
    p.clear()
    setfont(p.add_run(value), 12 if heading else 10.5, heading)
    if heading:
        p.paragraph_format.keep_with_next = True
    edited.append(p._p)
    return {"before": old, "after": value}


def add_before(anchor, value, template, edited, heading=False, caption=False):
    p = anchor.insert_paragraph_before()
    if template._p.pPr is not None:
        p._p.insert(0, deepcopy(template._p.pPr))
    setfont(p.add_run(value), 12 if heading else 10.5, heading or caption)
    if heading or caption:
        p.paragraph_format.keep_with_next = True
    edited.append(p._p)
    return p


def table_before(doc, anchor, headers, rows, widths, edited, report=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    for cell, width in zip(table.columns, widths):
        cell.width = Inches(width)
    for row_index, values in enumerate([headers] + rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        trpr = table.rows[row_index]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for column, value in enumerate(values):
            cell = cells[column]
            cell.width = Inches(widths[column])
            cell.vertical_alignment = 1
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = 1
            p.paragraph_format.keep_with_next = False
            setfont(p.add_run(str(value)), 10 if not report else 9.5, row_index == 0)
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, twips in [("top", 35), ("bottom", 35), ("left", 60), ("right", 60)]:
                edge = OxmlElement("w:" + side)
                edge.set(qn("w:w"), str(twips)); edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            pr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "bottom", "left", "right"):
                edge = OxmlElement("w:" + side)
                visible = report or (side == "top" and row_index == 0) or (side == "bottom" and row_index in [0, len(rows)])
                edge.set(qn("w:val"), "single" if visible else "nil")
                edge.set(qn("w:sz"), "5" if report else "8")
                edge.set(qn("w:color"), "D9D9D9" if report else "000000")
                borders.append(edge)
            pr.append(borders)
            if report and row_index == 0:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E8EBED"); pr.append(shade)
    anchor._p.addprevious(table._tbl)
    edited.append(table._tbl)
    return table


def value(row, metric="mse", digits=3):
    mean, sd = row[metric + "_mean"], row[metric + "_sd"]
    return f"{mean:.{digits}f}" if sd is None or sd != sd else f"{mean:.{digits}f}±{sd:.{digits}f}"


def pick(rows, **filters):
    matches = [r for r in rows if all(r[k] == v for k, v in filters.items())]
    assert len(matches) == 1, filters
    return matches[0]


def main():
    result = json.loads((RESULTS / "complete_analysis.json").read_text(encoding="utf8"))
    prose = json.loads((QA / "reviewed_interpretation.json").read_text(encoding="utf8"))
    assert all(prose.values()), "All result interpretations must be reviewed before document creation"
    assert result["learning_conditions"] == 192 and result["new_trainings"] == 144
    assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == "e31bd839cfbea24a3a6916d22a4bbbf00f9030141add6e45a79688e9fba5103a"
    OUT.mkdir(exist_ok=True)
    doc = Document(BASE)
    edited, changes = [], []
    def find(starts):
        matches = [p for p in doc.paragraphs if p.text.startswith(starts)]
        assert len(matches) == 1, (starts, len(matches))
        return matches[0]
    normal = find("当前证据还受观测时期较短")
    heading_template = find("2.6 泛化能力")
    caption_template = find("表5 参数量")
    for starts, key in [("摘要：", "abstract_cn"), ("Abstract:", "abstract_en"),
                        ("当前证据还受观测时期较短", "limitations"),
                        ("本文研究独立DLinear专家", "conclusion")]:
        changes.append(replace_paragraph(find(starts), prose[key], edited))
    anchor = heading_template
    add = lambda s: add_before(anchor, s, normal, edited)
    caption = lambda s: add_before(anchor, s, caption_template, edited, caption=True)
    add_before(anchor, "2.6 周期基线与稳定性补充实验", heading_template, edited, heading=True)
    add("为检验模型相对简单周期预测的收益，增加24 h与168 h滞后的季节性朴素基线，即使用前一日或前一周对应时刻的观测预测未来24 h。共同测试起点从测试段第168小时开始，Abilene和GÉANT分别保留741和379个预测窗口，三个学习模型均在相同起点重新计分。周周期基线使用168 h历史，长于学习模型的96 h，因此该比较允许不同历史长度。所有输入仅来自预测起点之前，结果见表6。")
    caption("表6 共同测试起点上的周期基线比较")
    seasonal_rows = []
    for ds in LABELS:
        for model, label in [("seasonal_24", "日周期朴素"), ("seasonal_168", "周周期朴素")] + MODELS:
            r = pick(result["seasonal"], dataset=ds, model=model)
            seasonal_rows.append([LABELS[ds], label, r["history_hours"], value(r), value(r, "mae")])
    table_before(doc, anchor, ["数据集", "方法", "历史/h", "MSE", "MAE"], seasonal_rows,
                 [.83, 1.22, .68, 1.48, 1.48], edited)
    add("表6中学习模型报告八个种子的均值±样本标准差；朴素基线为确定性计算，不设种子标准差。指标沿用原完整训练集定义的标准化坐标。该表评价的窗口少于表3，二者不能跨表直接比较数值。" + prose["seasonal_discussion"])
    add("训练历史实验固定原验证、测试时间段及OD变量集合，从同一训练截止点向前保留最近25%、50%、75%和100%的训练观测。前三种比例分别重新拟合训练段的逐变量均值与标准差，再训练DLinear、全局权重和本文方法；训练配置与主实验一致，各用42～49八个共同种子。100%条件沿用主实验检查点，其余比例按相同配置重新训练，均以验证误差选择检查点。所有预测还原到原完整训练集定义的标准化坐标后计分，以保证不同训练比例之间的误差可比。结果见图6。")
    figp = add_before(anchor, "", caption_template, edited)
    figp.alignment = 1
    figp.paragraph_format.keep_with_next = True
    figp.add_run().add_picture(str(RESULTS / "figure6_training_history.png"), width=Inches(6.35))
    figure_caption = caption("图6 不同训练历史长度下的测试MSE")
    figure_caption.paragraph_format.keep_with_next = False
    add("图6中点为八种子均值，误差棒为种子间样本标准差。横轴为原训练段保留比例；增加比例同时增加观测量和较早时期的覆盖，不能把二者的作用完全分离。训练轮数上限和早停规则相同，但优化更新次数未匹配。" + prose["history_discussion"])
    add("进一步将原测试段从第96小时起的目标时刻均分为三个连续区间T1、T2和T3，保留24 h预测目标全部落在对应区间内的窗口，排除跨区间预测。Abilene三个区间分别有255、256和256个窗口，GÉANT均为135个。使用100%训练历史下的同一模型逐段评价，不在区间之间重新训练。各段目标时刻互不重叠，但段内滑动窗口仍重叠；该分析描述时间稳定性，不把窗口视为独立统计重复。")
    caption("表7 三个连续测试时段的MSE")
    temporal_rows = []
    for ds in LABELS:
        for block in ["T1", "T2", "T3"]:
            temporal_rows.append([LABELS[ds], block] + [value(pick(result["temporal"], dataset=ds, block=block, model=m)) for m, _ in MODELS])
    table_before(doc, anchor, ["数据集", "时段", "DLinear", "全局权重", "本文方法"], temporal_rows,
                 [.80, .55, 1.45, 1.45, 1.45], edited)
    add("表7为八种子均值±样本标准差，不同区间的目标分布不同，误差大小不能单独解释为模型随时间退化。" + prose["temporal_discussion"])
    changes.append(replace_paragraph(anchor, "2.7 泛化能力与研究局限", edited, heading=True))
    clean = OUT / "基于自适应多尺度加权的多变量网络流量预测方法_补充实验净稿.docx"
    marked = OUT / "基于自适应多尺度加权的多变量网络流量预测方法_补充实验标红稿.docx"
    doc.save(clean)
    for element in edited:
        for run in element.xpath(".//w:r|.//m:r"):
            pr = run.find(qn("w:rPr"))
            if pr is None:
                pr = OxmlElement("w:rPr"); run.insert(0, pr)
            color = pr.find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color"); pr.append(color)
            color.attrib.clear(); color.set(qn("w:val"), "C00000")
    doc.save(marked)

    report = Document()
    sec = report.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.85)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        sty = report.styles[name]
        sty.font.name = "Times New Roman"; sty.font.color.rgb = RGBColor(0, 0, 0)
        sty.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体" if name == "Normal" else "黑体")
        for border in sty.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    report.styles["Normal"].font.size = Pt(11)
    report.styles["Normal"].paragraph_format.line_spacing = 1.15
    report.styles["Normal"].paragraph_format.space_after = Pt(4)
    report.styles["Title"].font.size = Pt(18)
    report.styles["Heading 1"].font.size = Pt(13)
    report.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    report.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    report.styles["Heading 2"].font.size = Pt(11)
    report.add_paragraph("补充实验结果与修改记录", style="Title")
    report.add_paragraph("基于自适应多尺度加权的多变量网络流量预测方法\n2026年10月20日")
    report.add_paragraph("本轮完成日周期与周周期朴素基线、训练历史长度曲线及三个连续测试时段的评价。共核验192个模型条件，其中144次为新增训练，48个完整训练历史条件使用既有检查点重新推理并核对保存预测。结果保留全部预设比例和八个随机种子，未根据误差选择种子或省略不利条件。")
    report.add_paragraph(prose["report_summary"])
    report.add_paragraph("一 固定实验规则", style="Heading 1")
    for s in [
        "数据源及原60%/20%/20%时间划分不变；未新增独立网络或流量观测。所有学习模型以96 h预测未来24 h。训练历史由同一截止点向前截取，原验证、测试段保持不变。",
        "各截取段独立拟合训练均值和标准差；验证与测试只应用该变换。OD变量集合仍按原任务固定，不随比例重新筛选；截取段中标准差小于10的负8次方的变量保留，并将标准差置为1。预测统一还原至原完整训练集的标准化坐标再评价，以免不同训练子集的归一化尺度改变误差单位。",
        "模型为DLinear、全局可学习尺度权重、本文路由方法；种子42～49。批量16，Adam学习率0.001，最多30轮，早停耐心值6；只依验证MSE选择检查点。训练量不同也会改变优化更新次数，本实验未使其相等。",
        "实验规则及源数据、模型、训练代码校验值保存在protocol.json。本轮是在退稿后设计的补充分析，未预注册。部分互不重叠的训练任务并行运行；新增训练耗时含资源竞争，不用于速度对比。",
        "均值±标准差反映同一划分下八个训练种子的变化。配对改进率逐种子计算；配对结果文件中的95%区间为自由度7的t区间，仅描述训练随机性，未作多重比较校正，也不代表独立新网络上的不确定性。未对重叠测试窗口进行独立样本检验。",
    ]:
        report.add_paragraph(s)
    report.add_paragraph("二 日周期与周周期基线", style="Heading 1")
    report.add_paragraph("每个预测点直接取前24 h或前168 h的对应观测，不训练参数。为保证两种基线可评价且不读取预测起点后的观测，全部方法统一在测试段起点不早于168 h的窗口计分，Abilene为741窗，GÉANT为379窗。周前基线使用更长的168 h历史；它与96 h学习模型不是等历史预算比较。")
    p = report.add_paragraph()
    table_before(report, p, ["数据集", "方法", "历史/h", "MSE", "MAE"], seasonal_rows,
                 [.80, 1.20, .60, 1.45, 1.45], [], report=True)
    report.add_paragraph("确定性朴素基线只报告一个数值。学习模型为八种子均值±标准差。" + prose["seasonal_discussion"])
    report.add_paragraph("三 训练历史长度实验", style="Heading 1")
    history_rows = []
    for ds in LABELS:
        for percent in [25, 50, 75, 100]:
            r = pick(result["learning"], dataset=ds, train_percent=percent, model="dlinear")
            history_rows.append([LABELS[ds], f"{percent}%", f"{r['train_rows']}/{r['train_windows']}"] +
                                [value(pick(result["learning"], dataset=ds, train_percent=percent, model=m), digits=4) for m, _ in MODELS])
    p = report.add_paragraph()
    table_before(report, p, ["数据集", "比例", "小时/窗口", "DLinear MSE", "全局权重MSE", "本文MSE"], history_rows,
                 [.65, .40, 1.02, 1.15, 1.15, 1.15], [], report=True)
    report.add_paragraph("八种子均值±标准差。增加历史量同时纳入更早的观测，本实验不能完全区分数量效应与分布变化。" + prose["history_discussion"])
    report.add_paragraph("四 连续测试时段评价", style="Heading 1")
    report.add_paragraph("测试小时索引从0开始。Abilene目标区间依次为[96,374)、[374,653)、[653,932)，保留255、256、256窗；GÉANT依次为[96,254)、[254,412)、[412,570)，各135窗。排除跨区间的24 h预测；每个时段均使用完整训练历史的同一检查点，不逐段重新训练。该设计是固定模型的分时段诊断，不是滚动重训或跨网络迁移。")
    p = report.add_paragraph()
    table_before(report, p, ["数据集", "时段", "DLinear MSE", "全局权重MSE", "本文MSE"], temporal_rows,
                 [.72, .48, 1.4, 1.4, 1.4], [], report=True)
    report.add_paragraph(prose["temporal_discussion"])
    report.add_paragraph("五 论文修改及尚存边界", style="Heading 1")
    report.add_paragraph("在上一轮净稿基础上新增第2.6节、表6、图6及表7，原第2.6节顺延为2.7节；同步替换摘要、局限和结语中的相关表述。原稿与上一轮修订稿均保留。此次标红仅表示相对上一轮净稿的新增及替换段落；两份新稿正文完全相同。")
    report.add_paragraph(prose["limitations"])
    report.add_paragraph("这些实验增加了数据量敏感性和测试期内稳定性的证据，没有增加独立网络、延长原始观测时期，也没有完成业务分层、拓扑变化、概念漂移适应或真实部署。不能据此宣称审稿人关于数据覆盖和实际应用的全部问题已解决。")
    report.add_paragraph("六 结果文件与复核", style="Heading 1")
    report.add_paragraph("完整结果位于abilene_forecasting/results/supplement_20261020。learning_curve_by_seed.csv保留192个条件；seasonal_by_seed.csv和temporal_by_seed.csv保存周期共同起点及分时段结果；对应summary和paired_effects文件保留均值、标准差及逐种子配对改进率。每个条件保存模型来源、检查点校验值、训练与预处理记录，以及逐窗口误差。新训练条件另存最佳检查点。")
    report.add_paragraph("新增训练程序与汇总程序独立保存，原模型及原训练程序未修改。协议测试核查了训练子集标准化、季节性预测无未来读取及跨时段目标不重叠。192个条件汇总时检查指标有限、逐窗口平均值一致，以及完整训练条件与既有预测一致。")
    report_path = OUT / "补充实验结果与修改说明.docx"
    report.save(report_path)
    manifest = dict(original=str(ORIGINAL), base=str(BASE), clean=str(clean), marked=str(marked),
                    report=str(report_path), original_sha256=hashlib.sha256(ORIGINAL.read_bytes()).hexdigest(),
                    base_sha256=hashlib.sha256(BASE.read_bytes()).hexdigest(), changed_elements=len(edited))
    (QA / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf8")
    (QA / "changes.json").write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding="utf8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
