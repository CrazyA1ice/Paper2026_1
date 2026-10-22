from pathlib import Path
import json
from docx import Document
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[2]
QA=Path(__file__).resolve().parent
manifest=json.loads((QA/'manifest.json').read_text(encoding='utf8'))
result=json.loads((ROOT/'abilene_forecasting/results/supplement_20261020/complete_analysis.json').read_text(encoding='utf8'))
doc=Document(manifest['report'])
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(5)
for p in doc.paragraphs:
    for run in p.runs:
        run.text=run.text.replace('附表95%区间','配对结果文件中的95%区间')
        if '这一比较提供简单周期复制' in p.text:
            run.text=run.text.replace('仍按表3','仍按论文表3')
for row,(dataset,percent) in zip(doc.tables[1].rows[1:],[(d,p) for d in ['abilene','geant'] for p in [25,50,75,100]]):
    for cell,model in zip(row.cells[3:],['dlinear','dlinear_scale_global','dlinear_scale']):
        r=next(x for x in result['learning'] if x['dataset']==dataset and x['train_percent']==percent and x['model']==model)
        nodes=cell._tc.xpath('.//w:t')
        assert len(nodes)==1
        nodes[0].text=f"{r['mse_mean']:.4f}±{r['mse_sd']:.4f}"
doc.save(manifest['report'])
print('Report spacing and four-decimal history statistics updated')
